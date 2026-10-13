import sys, copy, openpyxl
from openpyxl.styles import Alignment, Font
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb.active
C1 = "PDF列（E・G・I）の空欄は「画面と同じ」の意味か（共通確認①）"
same = lambda r: f"No.{r}（{ws[f'C{r}'].value}）と同じ"
K = {
 4: C1 + "。技術解説書では、PDFの診療開始日に「hh時mm分」まで出す形式もあるが、PDFも画面と同じ形式（時分なし）で出力しているか。",
 5: C1 + "。参照元の「AbatementDateTime」は abatementDateTime の表記ゆれという理解でよいか。",
 6: C1,
 7: "本記載は現在の画面の挙動か、案件213（臨床情報詳細画面の傷病名をPDFと同様に出力）のリリース後の挙動か（共通確認③）。前置・後置修飾語の参照元「code.extention.valueCodeableConcept」は、技術解説書の Condition.code.extension:diseasePrefixModifier 等と同じという理解でよいか。",
 8: "画面に出す傷病名の条件（例：否定・取消の傷病名は画面に出さないか）を追記してほしい（共通確認②）。PDF列の空欄は「PDFには出力しない」の意味か、「画面と同じ」の意味か。",
 9: "「詳細画面は自施設で登録したもののみ」とあるが、他施設が登録した情報は詳細画面に一切表示されないという理解でよいか（画面とPDFで名称の出し方が違うことの影響を判断する前提のため）。",
 10: "加工の「未提供の場合”×”」の「未提供」は、未告知病名（参照元で\"1\"）を指すという理解でよいか。技術解説書2.1.0版では、令和8年12月末のリリースからPDFにも「患者共有済み」を出力するとあるが、PDFの対応予定を記載してほしい。",
 11: C1,
 12: "PDFの「JLAC10のコードセットのディスプレイ値」は、技術解説書の「FHIR識別文字列から変換したFHIR項目名称」と同じものか。案件221（JLAC11を優先して出力）のリリース後のPDFの出力内容を追記してほしい。あわせて、詳細画面でも同様の対応（コードの表示名での出力）ができるかの富士通への確認結果と、回答の見通し（時期）を共有してほしい。",
 13: "下限単位の省略（上限単位と同じ場合）は、案件143・144で対応予定と認識している。リリース済み（現在の挙動）か、リリース後の挙動を記載したものか（共通確認③）。PDFも同じ加工か（共通確認①）。",
 14: C1,
 15: "評価値”N”の非表示は、案件143・144で対応予定と認識している。リリース済み（現在の挙動）か（共通確認③）。PDFも同じ加工か（共通確認①）。",
 16: same(9),
 17: "傷病名は未告知病名タグ（uninformed）、感染症は患者共有タグ（pts）で判定しており、区分により判定の元が異なるが、仕様どおりという理解でよいか。PDFへの出力予定（令和8年12月末のリリース以降）を記載してほしい。",
 18: C1,
 19: "PDFの「JFAGY医薬品（YJコードベース／派生（成分）グループコードベース）」は、技術解説書の薬剤名のコード体系（YCM＋YJコード／剤形・規格・銘柄不明コード）と同じものを指すか。詳細画面でもコードの表示名で出力できるかの再確認の結果と、回答の見通し（時期）を共有してほしい。",
 20: C1,
 21: "【回答】臨床的状態がPDFに出力されない差は、現状のままとすることで差し支えない。【確認】あわせて、PDFに出すレコードの条件（clinicalStatus＝activeのみか、verificationStatus（未確認・確認済み・否定）で制御しているか）を明記してほしい（技術解説書の記載を実装に合わせるため）。画面に出すレコードの条件も追記してほしい（共通確認②）。",
 22: same(9),
 23: C1,
 24: "参照元にFHIRパスの記載がないため追記してほしい（薬剤アレルギーと同じ code.text・code.coding という理解でよいか）。詳細画面でもコードの表示名で出力できるかの再確認の結果と、回答の見通し（時期）を共有してほしい。",
 25: C1,
 26: "参照元にFHIRパスの記載がないため追記してほしい（薬剤アレルギーと同じ clinicalStatus.coding.code という理解でよいか）。PDFに出力されない差は、No.21と同様に現状のままで差し支えない。",
 27: same(9),
 28: same(12),
 29: C1,
 30: same(9),
 31: same(13) + "。備考のとおり、詳細画面の「最新〜」は最新の検査ではなく臨床情報に紐づく検査が出るとのことだが、画面の項目名を実態に合わせて見直す予定はあるか（急ぎではない）。",
 32: C1,
 33: same(15),
 34: "画面に出す検査の条件（ステータスが registered（結果未着）・cancelled（検査中止）のものも出るか）を追記してほしい（共通確認②）。PDFは技術解説書では preliminary・final・corrected のみ出力とあるが、実装も同じか。",
 37: same(13), 39: same(15), 43: same(13), 45: same(15),
 40: same(34), 46: same(34),
}
# header (K2:K3) styled like H2
h = ws['H2']; k = ws['K2']
k.value = "機構への確認事項（厚労省ACN案）"
k.font = copy.copy(h.font); k.fill = copy.copy(h.fill); k.border = copy.copy(h.border); k.alignment = copy.copy(h.alignment)
ws['K3'].border = copy.copy(ws['H3'].border); ws['K3'].fill = copy.copy(h.fill)
ws.merge_cells('K2:K3')
ws.column_dimensions['K'].width = 70
ws.column_dimensions['J'].width = 24
for r in range(4, 47):
    c = ws[f'K{r}']; ref = ws[f'I{r}']
    c.border = copy.copy(ref.border); c.font = copy.copy(ref.font)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    if r in K: c.value = K[r]
notes = ["【共通の確認事項】（右端の列から参照）",
 "共通確認①：PDF列（E・G・I）の空欄は「画面と同じ」の意味か。空欄の意味が分かるよう補記してほしい。",
 "共通確認②：臨床情報詳細画面に出すレコードの条件（出力制御）を区分ごとに追記してほしい（傷病名：clinicalStatus・verificationStatus、検査：status、アレルギー：clinicalStatus・verificationStatus）。否定された傷病名・中止した検査・現存しないアレルギーが画面に出ると、誤って読まれるおそれがあるため。",
 "共通確認③：本調査結果は現在の挙動か、案件143・144・213・221のリリース後の挙動か。リリース後の挙動を含む場合は、どの行が該当するかを示してほしい。"]
for i, t in enumerate(notes):
    c = ws.cell(row=48+i, column=11, value=t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if i == 0: c.font = Font(bold=True)
wb.save(dst)
print('written', len(K))
