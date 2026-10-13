# v2（K列＝機構への確認事項）に、L列「確認事項なしと判断した理由（ACN内部メモ）」を追加して v3 を作る。
# 対象：加工が「そのまま出力」以外で、K列にコメントがない行。
# 注意：L列はACN内部の判断メモ。機構へ送るときはL列を削除する。
import sys, copy, openpyxl
from openpyxl.styles import Alignment, Font
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb.active
SAME = "【同じ表示】桁数・参照元・加工が画面とPDFで結合されており（同じ処理）、同じ登録データからは同じ表示になる。"
HOSP = ("【画面に出ないが問題なし】詳細画面は自施設が登録した情報だけを出すため（J列の機構の説明）、登録医療機関名は自院の名前になり、表示する意味がない。"
        "他院の情報はPDFでだけ見るので、PDFに医療機関名（廃止なら【廃】）が出れば、どこが登録したかは判断できる。")
PREV = ("【PDFにだけある項目・問題なし】画面には前回・前々回の構造がない（J31）ため、比べる相手がない。"
        "PDFの前回・前々回は、最新（29〜34行）と同じ加工のため、PDFの中で最新と前回・前々回の表示はそろう。過去の推移はPDFで確認できる。")
L = {
 4: SAME + "形式は登録された日付の精度（年／年月／年月日）のまま出すため、精度が違っても誤解は生じない。",
 5: SAME,
 6: SAME + "「1」を「主」に置き換えるだけで、意味は変わらない。",
 8: SAME + "コード値（active等）を日本語（存続・寛解・軽快または治癒・不明）に置き換えるだけで、画面とPDFで同じ言葉になる。",
 9: HOSP,
 11: SAME + "形式（YYYY年MM月DD日HH時mm分）も同じ。",
 13: SAME + "上限と下限の単位が同じときに下限の単位を省くのは表示を短くするためで、値の意味は変わらない。",
 15: SAME + "評価値「N」（基準値内）を出さないのは画面もPDFも同じ。異常値（H・L等）は出るので、医療機関の判断には影響しない。",
 16: HOSP,
 17: SAME + "「1」を「〇」、それ以外を「×」に置き換えるだけ。",
 18: SAME,
 20: SAME + "症状は、コードがあればその表示名、なければ登録された文字を出す処理で、画面とPDFで同じ。薬剤名（19行）と違い、症状は画面でも表示名を使っている。",
 22: HOSP,
 23: SAME,
 25: SAME + "20行（薬剤アレルギーの症状）と同じ処理。",
 27: HOSP,
 30: HOSP,
}
for r in range(35, 47): L[r] = PREV
h = ws['K2']; c = ws['L2']
c.value = "確認事項なしと判断した理由\n（ACN内部メモ。機構へ送るときは削除）"
c.font = copy.copy(h.font); c.fill = copy.copy(h.fill); c.border = copy.copy(h.border)
c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
ws['L3'].border = copy.copy(ws['K3'].border); ws['L3'].fill = copy.copy(h.fill)
ws.merge_cells('L2:L3')
ws.column_dimensions['L'].width = 60
for r in range(4, 47):
    x = ws[f'L{r}']; ref = ws[f'K{r}']
    x.border = copy.copy(ref.border); x.font = copy.copy(ref.font)
    x.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    if r in L:
        assert ws[f'K{r}'].value in (None, ''), r   # K列にコメントがある行には書かない
        x.value = L[r]
note = ws.cell(row=48, column=12, value="【前提】L列の判断は、機構の調査結果（20261006）の記載どおりに動くことが前提（実機では確認していない）。"
       "K列下の共通確認①（どの時点の画面・PDFか）・②（出てくるレコードが同じか）の回答次第で、判断が変わることがある。")
note.alignment = Alignment(wrap_text=True, vertical='top'); note.font = Font(bold=True)
wb.save(dst)
print('written', len(L))
