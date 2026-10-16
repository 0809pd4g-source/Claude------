# ユーザーが作成したExcel（見出し「厚労省向けコメント」の列があるもの）に、厚労省向けコメントを書き込む。
# 使い方: python 20261014_厚労省向けコメント記入.py <入力xlsx> <出力xlsx>
# ・「厚労省向けコメント」の列は、1〜6行目の見出しから探す（列の位置が変わっても動く）
# ・書き込む行は、C列の項目名とB列の区分（結合セルを展開）で探す（行がずれても動く）
# ・既に何か書いてあるセルには書かない（ユーザーの記入を上書きしない）
import sys, copy, openpyxl
from openpyxl.styles import Alignment
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb.active

col = None
hdr = None
for r in range(1, 7):
    for c in ws[r]:
        if isinstance(c.value, str) and '厚労省向けコメント' in c.value:
            col = c.column; hdr = r
assert col, '見出し「厚労省向けコメント」が見つからない'

def cat_of(r):
    for m in ws.merged_cells.ranges:
        if m.min_col <= 2 <= m.max_col and m.min_row <= r <= m.max_row:
            return ws.cell(m.min_row, 2).value
    return ws.cell(r, 2).value

NAME = ("【画面とPDFで名称が違う】同じ{0}でも、詳細画面は医療機関が登録した名称、PDFは標準マスタの名称で表示されます。"
        "機構は10/5に「画面側の修正は3月リリースでは難しい」と説明しています。"
        "傷病名は案件213で画面をPDFにそろえる方針のため、ACNとしては他の区分も基本はそろえていただく（修正を求める）方針がよいと考えています。"
        "機構から対応できないと回答があった場合の扱い（求める時期・それまでの周知）を、別途ご相談させてください。")
ALLERGY = ("【ご確認いただきたいこと】アレルギーの「臨床的状態」（今もそのアレルギーがあるか：存続・治った等）が、詳細画面には表示され、PDFには表示されない差について、"
           "機構から「現状のまま残したい」との確認依頼が来ています。ACNとしては、現状のままで問題ないと考えますが、この方針で機構へ回答してよろしいでしょうか。\n"
           "【問題ないと考える理由】①PDFには「存続」のアレルギーだけが載る決まりのため、欄を作っても全て「存続」になり、欄がなくても読み違えは生じません（治ったアレルギーはPDFに載りません）。"
           "②詳細画面では状態が表示され、案件196（3月リリース予定）で「他の医療機関に表示されるのは存続のものだけ」と分かる注釈を画面に出す予定です。\n"
           "【前提（機構に確認中）】PDFが本当に「存続」で絞り込んでいるか（以前、別の条件で絞り込んでいるとの説明があったため）、PDF上にその旨の注記があるか。前提と異なる場合は改めてご相談します。")
TEXT = {
 ('傷病名', '傷病名'): "【12月までは表記が違う】詳細画面の傷病名は、案件213（12月実装予定）でPDFと同じ表記（修飾語・「の疑い」付き）にそろいます。"
                       "それまでの間は、画面は医療機関が登録した名称、PDFは修飾語・「の疑い」を付けた名称で表示され、表記が異なります。12月までの間、医療機関への周知が必要かをご判断いただきたい点です。",
 ('傷病名', '患者共有済み'): "【12月にそろう】PDFは、案件193（12月希望）で項目名が「告知状況（未告知／済）」から「患者共有済み（×／〇）」に変わり、画面と同じ表示になります。"
                           "なお、画面に「患者共有済み」の欄が付くのは案件196（3月希望）のため、画面とPDFが完全にそろうのは3月以降の見込みです。",
 ('感染症', '検査項目'): NAME.format('検査項目') + "なお、案件221（3月）でPDFはJLAC11コードを優先して表示しますが、画面は対象外のため、差が広がる見込みです。",
 ('検査', '検査項目'): "感染症の検査項目と同じです。",
 ('薬剤アレルギー', '薬剤名'): NAME.format('薬剤'),
 ('その他アレルギー', 'アレルギー物質'): NAME.format('アレルギー物質'),
 ('薬剤アレルギー', '臨床的状態'): ALLERGY,
 ('その他アレルギー', '臨床的状態'): "薬剤アレルギーの臨床的状態と同じです（同じ方針で機構へ回答してよいかをご確認ください）。",
 ('検査', '最新採取日時'): "【画面の「最新」が最新とは限らない】詳細画面には時系列（最新・前回・前々回）がなく、選んだ登録分の検査結果が表示されます。"
                          "そのため、画面の「最新〜」の欄に最新ではない結果が出たり、画面で見ている結果がPDFでは「前回」の欄に出たりする場合があります。"
                          "医師が古い結果を最新と誤認するおそれがあるため、ACNとしては、少なくとも画面の項目名を実態に合わせる修正を求めるのがよいと考えています。",
}
written = []
for r in range(hdr + 1, ws.max_row + 1):
    key = (cat_of(r), ws.cell(r, 3).value)
    if key in TEXT:
        c = ws.cell(r, col)
        if c.value not in (None, ''):
            print('既に記入あり・スキップ', r, key); continue
        c.value = TEXT[key]
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        written.append((r, key))
missing = set(TEXT) - {k for _, k in written}
print('書き込み', len(written), [w[0] for w in written]); print('見つからない', missing)
ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width or 0, 70)
wb.save(dst)
