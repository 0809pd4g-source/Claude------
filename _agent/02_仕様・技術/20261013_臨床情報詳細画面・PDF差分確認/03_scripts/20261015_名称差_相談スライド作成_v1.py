"""WBS確認会向け：詳細画面とPDFの名称の差（検査項目・薬剤名・アレルギー物質）の相談スライドを作る。
体裁は 06_会議/WBS確認会/20261009/02_output/20261008_血液型の登録表示抑制スライド_v1.pptx に合わせる。
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(sys.argv[1])
FONT = "Meiryo UI"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0xD9, 0xD9, 0xD9)
PINK = RGBColor(0xF2, 0xDC, 0xDB)
LIGHT = RGBColor(0xDE, 0xEB, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Emu(10689336)
prs.slide_height = Emu(7562088)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, paras, size=12, bold=False, color=BLACK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: str（改行区切り）または [(文字列, {bold,color,size})] のリスト"""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Cm(0.15))
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    items = paras.split("\n") if isinstance(paras, str) else paras
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = it if isinstance(it, list) else [(it, {})] if isinstance(it, str) else [it]
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
    return tb


def header(slide, title, sub):
    text(slide, 1.1, 0.8, 27.4, 1.4, title, size=22, bold=True)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(0.8), Cm(2.3), Cm(29.0), Cm(2.3))
    ln.line.color.rgb = NAVY
    ln.line.width = Pt(1.5)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(2.8), Cm(27.4), Cm(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    text(slide, 1.4, 2.8, 26.9, 1.1, sub, size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, col_w, rows, row_h=0.9, size=10.5, head_fill=NAVY, first_col_fill=GRAY, fills=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Cm(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Cm(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.15)
            c.margin_top = c.margin_bottom = Cm(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = it if isinstance(it, list) else [(it, {})]
                for s, opt in runs:
                    r = p.add_run()
                    r.text = s
                    r.font.name = FONT
                    r.font.size = Pt(opt.get("size", size))
                    r.font.bold = opt.get("bold", i == 0)
                    r.font.color.rgb = opt.get("color", WHITE if i == 0 else BLACK)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = head_fill
            elif fills and (i, j) in fills:
                c.fill.fore_color.rgb = fills[(i, j)]
            elif j == 0:
                c.fill.fore_color.rgb = first_col_fill
            else:
                c.fill.fore_color.rgb = WHITE
    return tbl


B = {"bold": True}
R = {"bold": True, "color": RED}

# ---- スライド1：相談の概要と差の一覧 ----
s = prs.slides.add_slide(BLANK)
header(s, "X. 臨床情報詳細画面とPDF帳票の名称の差に係る相談",
       "ご相談事項：検査項目・薬剤名・アレルギー物質の名称が、詳細画面とPDFで異なる件の対応方針")
text(s, 1.1, 4.2, 27.4, 2.6, [
    "・ToDo#1228（10/5開発課題検討会）を受け、DX機構から「PDFと詳細画面の差異に関する調査結果」（10/6）が共有された。",
    "・画面とPDFの両方に出る項目のうち、下表の3項目は、同じ登録データでも画面とPDFで表示される名称が異なる。",
    [("・富士通からは「品質担保の観点から3月リリースでの修正は困難」との回答", {}), ("（10/5開発課題検討会 DX機構）", {"size": 10})],
])
table(s, 1.1, 7.0, 27.4, [3.4, 3.0, 9.4, 11.6], [
    ["区分", "項目", "詳細画面（オン資端末）", "PDF帳票（他医療機関が取得）"],
    ["感染症・検査", "検査項目", "医療機関が登録した名称（code.text）",
     ["JLACコードがある場合：JLACマスタの名称（FHIR項目名称）", "上記以外：医療機関が登録した名称（code.text）"]],
    ["薬剤アレルギー等", "薬剤名", "医療機関が登録した名称（code.text）",
     ["JFAGY医薬品コード（ダミー以外）の場合：コードの名称（display）", "上記以外：医療機関が登録した名称（code.text）"]],
    ["その他アレルギー等", "アレルギー物質", "医療機関が登録した名称（code.text）",
     ["JFAGY食品・非食品非医薬品コード（ダミー以外）の場合：コードの名称（display）", "上記以外：医療機関が登録した名称（code.text）"]],
    [["（参考）", "傷病名"], "傷病名", ["医療機関が登録した名称（code.text）。疑い病名でも「の疑い」が付かない", "→ 12月リリース（案件213）でPDFと同様に出力"],
     ["前置修飾語＋傷病名＋後置修飾語＋「の疑い」（MEDISの名称）", "※100桁を超えた分は切り捨て"]],
], row_h=1.55, fills={(4, 0): PINK, (4, 1): PINK, (4, 2): PINK, (4, 3): PINK})
text(s, 1.1, 15.2, 27.4, 1.6, [
    "・詳細画面は、自施設が登録した臨床情報を確認するための画面（登録医療機関名は出力しない）。PDFは他の医療機関が取得する帳票。",
    "・3月リリース（案件221）で、PDF・XMLの検査項目はJLAC11を優先して出力するよう変わる（詳細画面は対象外）。",
], size=11)
text(s, 1.1, 19.4, 27.4, 0.8, "出典：DX機構「PDFと詳細画面の差異に関する調査結果」（2026/10/6）、開発課題検討会（2026/10/5）、案件台帳（令和8年度）", size=9)

# ---- スライド2：確認結果 ----
s = prs.slides.add_slide(BLANK)
header(s, "X. 臨床情報詳細画面とPDF帳票の名称の差に係る相談",
       "確認結果：FHIRの定義と、画面とPDFで名称がずれるリスク")
text(s, 1.1, 4.2, 27.4, 0.8, [[("① 詳細画面が使う名称（code.text）は必須か", B)]], size=12)
table(s, 1.1, 5.0, 27.4, [6.6, 2.8, 18.0], [
    ["プロファイル（JP-CLINS v1.13.0）", "code.text", "定義（抜粋）"],
    ["JP_Observation_LabResult_eCS", "1..1（必須）", "項目名。報告書などに記載する場合に使用する表示名。受信側はこの文字列を項目表示文字列として使用できる。"],
    ["JP_AllergyIntolerance_eCS", "1..1（必須）", "電子カルテシステム等で登録され表示されている文字列をcode.textに必ず設定すること。"],
    [["（参考）", "JP_Condition_eCS"], "1..1（必須）", "前置修飾語、後置修飾語も含めて電子カルテ上での表示（入力）文字列全体を設定する。"],
], row_h=1.0, size=10)
text(s, 1.1, 9.2, 27.4, 0.8, "→ 3項目とも必須のため、詳細画面の名称が空欄になることはない。差は「医療機関の電子カルテ上の名称」と「コードの標準名称」の違い。", size=11)

text(s, 1.1, 10.2, 27.4, 0.8, [[("② 仕様上、閲覧時に表示する名称はどちらか", B)]], size=12)
text(s, 1.1, 11.0, 27.4, 1.6, [
    "・技術解説書2.1.0版 表24では、FHIR項目名称を「臨床情報一覧やマイナポータル等の閲覧画面で利用者が直感的に理解しやすい形で表示するための検査項目名称」と定義。",
    "・JFAGY（薬剤名・アレルギー物質）も、ダミーコード以外はコードの名称を表示し、ダミーコードの場合だけtextを表示する考え方（技術解説書 表18）。",
    [("→ 閲覧時は標準名称（コードの名称）を表示するのが仕様の考え方であり、", {}), ("基本的にはPDFの表示が正しい。", R)],
], size=11)

text(s, 1.1, 13.4, 27.4, 0.8, [[("③ 画面とPDFで名称がずれるリスク", B)]], size=12)
table(s, 1.1, 14.2, 27.4, [6.0, 21.4], [
    ["リスク", "内容"],
    ["コードの誤りに気づけない",
     "自施設の詳細画面は登録した名称（text）しか出ないため、誤ったコードを付けて登録しても、他の医療機関のPDFに別の名称（コードの名称）が出ていることを、自施設の確認では見つけられない。"],
    ["同じ1件の名称が違って見える",
     "自施設が画面で確認した名称と、他の医療機関が受け取るPDFの名称が異なり、問合せや読み違いの原因になる。"],
    [["（参考）傷病名", "＝先に対応"],
     "画面に「の疑い」が付かず、病名の意味が変わってしまうため、12月リリース（案件213）で先に対応。3項目は名称の表現の差で意味は変わらないことが多いが、上記のリスクは残る。"],
], row_h=1.25, size=10, fills={(1, 1): PINK})

# ---- スライド3：対応案と相談事項 ----
s = prs.slides.add_slide(BLANK)
header(s, "X. 臨床情報詳細画面とPDF帳票の名称の差に係る相談",
       "対応案：費用対効果を踏まえた進め方")
table(s, 1.1, 4.2, 27.4, [5.2, 9.2, 7.0, 6.0], [
    ["案", "内容", "効果", "費用・時期"],
    ["A 画面をPDFに揃える", "詳細画面もPDFと同じ名称の出し方にする（傷病名の案件213と同じ考え方）",
     "差がなくなり、自施設の確認でコードの誤りに気づける", "改修費用が発生。3月リリースは困難（富士通）→ 時期は機構に照会中"],
    ["B 画面に標準名称を併記", "画面に登録した名称とコードの名称を並べて表示する",
     "コードの誤りに気づける。登録した名称も残る", "改修費用が発生（Aとの差は機構に要確認）"],
    ["C 改修せず周知", "差があることを技術解説書・運用マニュアル等に明記し、医療機関に周知する",
     "費用はかからないが、コードの誤りに気づけないリスクは残る", "費用なし。技術解説書の次回改訂で対応"],
], row_h=2.0, size=10.5, fills={(1, 0): LIGHT, (1, 1): LIGHT, (1, 2): LIGHT, (1, 3): LIGHT})
text(s, 1.1, 12.6, 27.4, 0.8, "※機構には、改修の可否と、改修できる場合の時期を照会中（ToDo#1228への確認事項）。費用の見込みは回答を受けて整理する。", size=10)

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(13.8), Cm(27.4), Cm(0.9))
bar.fill.solid()
bar.fill.fore_color.rgb = NAVY
bar.line.fill.background()
text(s, 1.4, 13.8, 26.9, 0.9, "ご相談事項", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.1, 14.9, 27.4, 3.6, [
    [("（1）方向性：", B), ("閲覧時は標準名称を表示するのが仕様の考え方のため、基本は案A（画面をPDFに揃える）で機構に改修を求めることでよいか。", {})],
    [("（2）進め方：", B), ("機構から改修の可否・時期・費用の回答を受けたうえで、費用対効果を見てA・B・Cを判断することでよいか。", {})],
    [("（3）当面の扱い：", B), ("改修までの間、差があることを医療機関に周知するか（案Cを暫定で行うか）。", {})],
], size=12)
text(s, 1.1, 19.4, 27.4, 0.8, "出典：JP-CLINS v1.13.0（StructureDefinition）、技術解説書2.1.0版（表18・表24）、DX機構 調査結果（2026/10/6）", size=9)

cp = prs.core_properties
for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
    setattr(cp, k, "")
prs.save(OUT)
print("saved", OUT)
