"""v7：表からXMLレイアウトの列を外し、差分の理由をFHIRパスの違い（画面＝code.text、PDF＝code.coding.display）に。3月リリース後の挙動の列を追加。参考をFHIRの定義（検査のdisplayは固定値）に差し替え。
v6：厚労省推奨の構成（背景・状況→ご相談事項→参考）と体裁に合わせて作り直す。
体裁の手本：20260902_血液型の共有について.pptx（テンプレートの「7_タイトルとコンテンツ」、紺のセクションラベル、•箇条書き12pt、薄ピンク見出しの表）。
手本ファイルを土台にし、元のスライドを除いて同じレイアウトでスライドを作る。
使い方：python 本ファイル <手本pptx> <出力pptx>
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.oxml.ns import qn
from lxml import etree

REF, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(REF)
base = prs.slides[3]
LAYOUT = base.slide_layout
LABEL = BODY = None
for sh in base.shapes:
    if sh.name == "四角形: 角を丸くする 3":
        LABEL = copy.deepcopy(sh._element)
    if sh.name == "コンテンツ プレースホルダー 10" and sh.is_placeholder:
        BODY = copy.deepcopy(sh._element)
assert LABEL is not None and BODY is not None

# 手本のスライドを除く
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
RED = "C00000"
_id = [100]


def _next_id():
    _id[0] += 1
    return _id[0]


def new_slide(title):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1 and ph.placeholder_format.idx == 1:  # 本文プレースホルダーは自前で置く
            ph._element.getparent().remove(ph._element)
    s.shapes.title.text = title
    return s


def _set_xfrm(el, x, y, w, h):
    off = el.find(".//" + qn("a:off"))
    ext = el.find(".//" + qn("a:ext"))
    off.set("x", str(int(Cm(x)))); off.set("y", str(int(Cm(y))))
    ext.set("cx", str(int(Cm(w)))); ext.set("cy", str(int(Cm(h))))


def label(s, text, x, y, w=4.57):
    el = copy.deepcopy(LABEL)
    el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id()))
    el.find(".//" + qn("p:cNvPr")).set("name", f"ラベル {text}")
    _set_xfrm(el, x, y, w, 0.69)
    el.find(".//" + qn("a:t")).text = text
    s.shapes._spTree.append(el)


def _run(p, text, bold=False, color=None, size=1200):
    r = etree.SubElement(p, qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"), lang="ja-JP", altLang="en-US", sz=str(size))
    if bold:
        rpr.set("b", "1")
    if color:
        sf = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(r, qn("a:t")).text = text


def body(s, paras, x, y, w, h, size=1200):
    """paras: [str | [(text, {bold, color})]]。•付きの箇条書き（手本と同じ字下げ）"""
    el = copy.deepcopy(BODY)
    el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id()))
    el.find(".//" + qn("p:cNvPr")).set("name", "本文")
    _set_xfrm(el, x, y, w, h)
    tx = el.find(qn("p:txBody"))
    for p in tx.findall(qn("a:p")):
        tx.remove(p)
    for para in paras:
        p = etree.SubElement(tx, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"), marL="285750", indent="-285750")
        sa = etree.SubElement(ppr, qn("a:spcAft"))
        etree.SubElement(sa, qn("a:spcPts"), val="300")
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial", panose="020B0604020202020204", pitchFamily="34", charset="0")
        etree.SubElement(ppr, qn("a:buChar"), char="•")
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            _run(p, t, o.get("bold", False), o.get("color"), size)
    s.shapes._spTree.append(el)


def _border(tcPr):
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w="6350", cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val="000000")
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def table(s, x, y, col_w, rows, row_h=0.6, size=900, head_fill="F2CEEF", fills=None):
    shp = s.shapes.add_table(len(rows), len(col_w), Cm(x), Cm(y), Cm(sum(col_w)), Cm(row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Cm(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Cm(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            tf = c.text_frame
            tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = it if isinstance(it, list) else [(it, {})]
                for t, o in runs:
                    r = p.add_run()
                    r.text = t
                    r.font.size = Pt(size / 100)
                    r.font.name = "Meiryo UI"
                    rpr = r._r.get_or_add_rPr()
                    ea = etree.SubElement(rpr, qn("a:ea"), typeface="Meiryo UI")
                    r.font.bold = o.get("bold", i == 0)
                    from pptx.dml.color import RGBColor
                    r.font.color.rgb = RGBColor.from_string(o.get("color", "000000"))
            c.margin_left = c.margin_right = Cm(0.1)
            c.margin_top = Cm(0.05)
            c.margin_bottom = Cm(0.03)
            from pptx.enum.text import MSO_ANCHOR
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tcPr = c._tc.get_or_add_tcPr()
            for child in list(tcPr):
                tcPr.remove(child)
            _border(tcPr)
            sf = etree.SubElement(tcPr, qn("a:solidFill"))
            fill = head_fill if i == 0 else (fills or {}).get((i, j), "FFFFFF")
            etree.SubElement(sf, qn("a:srgbClr"), val=fill)
    return tbl


B = {"bold": True}
R = {"bold": True, "color": RED}
TITLE = "X.臨床情報詳細画面とPDF帳票の名称の差について"
X0, W0 = 1.37, 24.8

# ---- スライド1：背景・状況 → ご相談事項 ----
s = new_slide(TITLE)
label(s, "背景・状況", X0, 2.3)
body(s, [
    [("案件196の検討で、詳細画面の傷病名に「疑い」が表示されないことが判明した。詳細画面は自施設の登録内容を確認する手段のため、", {}), ("12月リリース（案件213）で画面をPDFと同様に出力する", B), ("こととなった。", {})],
    "傷病名以外にも同様の差がないかをFHIRパスで確認し、DX機構から調査結果が共有された（10/6、ToDo#1228）。",
    [("検査項目・薬剤名・アレルギー物質は、", {}), ("詳細画面とPDFで参照するFHIRパスが異なるため、表示に差分が生じる", B), ("（下表）。", {})],
    "富士通からは「品質担保の観点から3月リリースでの修正は困難」との回答（10/5）。3月リリース後も、画面の名称の出し方は変わらない。",
], X0, 3.0, W0, 3.9)
table(s, 1.9, 7.0, [2.6, 4.6, 7.8, 4.6, 4.2], [
    ["項目", ["詳細画面（FHIRパス）", [("※code.textはFHIRの定義上、入力必須", {"bold": False})]], "PDF帳票（FHIRパス）", "表示の差分", "3月リリース後"],
    [["検査項目", "（感染症・検査）"], ["Observation.code.text", "（医療機関の項目名）"],
     ["Observation.code.coding（JLAC10）.display", "＝項目ごとの固定の文字列（例：TP、ALB）", "※JLAC10がない場合はcode.text"],
     [[("生じる", R)], "項目名が固定の文字列と同じ場合を除く"],
     ["画面：変更なし", "PDF：JLAC11併記時はJLAC11優先（案件221）"]],
    ["薬剤名", ["AllergyIntolerance.code.text", "（電子カルテ上の名称）"],
     ["AllergyIntolerance.code.coding（JFAGY）.display", "※ダミーコードの場合はcode.text"],
     [[("生じうる", R)], "電子カルテ上の名称がJFAGYの名称と異なる場合"],
     "変更なし"],
    ["アレルギー物質", ["AllergyIntolerance.code.text", "（電子カルテ上の名称）"],
     ["AllergyIntolerance.code.coding（JFAGY）.display", "※ダミーコードの場合はcode.text"],
     [[("生じうる", R)], "薬剤名と同じ"],
     "変更なし"],
    ["（参考）傷病名", ["Condition.code.text", "（修飾語を含む入力文字列）"],
     ["前置修飾語＋傷病名（MEDIS）＋後置修飾語＋「の疑い」（verificationStatus）"],
     [[("生じる", B)], "疑いフラグのみの場合「の疑い」が出ない"],
     ["12月（案件213）で", "画面をPDFと同様に出力"]],
], row_h=0.8)

label(s, "ご相談事項", X0, 13.4)
body(s, [
    [("表示に差分が生じるため、", {}), ("傷病名（案件213）と同様に、詳細画面の名称をPDFに揃える改修を機構に求める方向", B), ("でよろしいでしょうか。", {})],
    [("薬剤名・アレルギー物質は、電子カルテ上の名称がJFAGYの名称と同じ場合は差分が出ませんが、", {}), ("その場合でも念のため改修を求めますか。", B)],
    "改修の可否・時期・費用を機構に照会し、費用対効果を見て判断する進め方でよろしいでしょうか。あわせて、改修までの間、差分があることを医療機関に周知するかについてご意向を伺いたい。",
], X0, 14.1, W0, 3.9, size=1100)

# ---- スライド2：参考 → 機構への確認事項 ----
s = new_slide(TITLE)
label(s, "参考", X0, 2.3)
body(s, [[("詳細画面・PDFが参照するFHIRパスの定義", B), ("（JP-CLINS v1.13.0）", {})]], X0, 3.0, W0, 0.8)
table(s, 1.9, 3.8, [5.6, 3.6, 2.2, 12.4], [
    ["プロファイル", "FHIRパス", "多重度", "定義（抜粋）"],
    [["JP_Observation_LabResult_eCS", "（検査・感染症）"], "code.text", "1..1（必須）", "項目名。報告書などに記載する場合に使用する表示名。（注釈：多くの場合、coding.displayと同一になるが…）"],
    ["", ["code.coding", "（JLAC10）.display"], "1..1（必須）", ["項目ごとに固定値（fixedString）。例：総蛋白＝TP、アルブミン＝ALB、HbA1c＝HbA1c-NGSP", "（技術解説書の「FHIR識別文字列」）"]],
    [["JP_AllergyIntolerance_eCS", "（薬剤・その他アレルギー）"], "code.text", "1..1（必須）", "電子カルテシステム等で登録され表示されている文字列をcode.textに必ず設定すること。"],
    ["", ["code.coding", "（JFAGY）.display"], "1..1（必須）", "コードに対応する表示名（JFAGYのコード表の名称）。"],
    [["（参考）JP_Condition_eCS", "（傷病名）"], "code.text", "1..1（必須）", "前置修飾語、後置修飾語も含めて電子カルテ上での表示（入力）文字列全体を設定する。"],
], row_h=0.85)
body(s, [
    "詳細画面・PDFがどのFHIRパスを参照しているかは、DX機構の調査結果（10/6）の「参照元」欄による。",
    "code.textは医療機関が自由に設定する名称、code.coding.displayはコード表で決まった名称であり、両者が一致する保証はない。",
], X0, 9.4, W0, 1.9, size=1100)

label(s, "機構への確認事項", X0, 11.6)
body(s, [
    "PDFの検査項目名は、調査結果では「JLAC10のディスプレイ値」（＝固定の文字列。例：TP）だが、技術解説書の「FHIR識別文字列をFHIR項目名称に自動変換して出力」のとおり変換して表示しているか。",
    "案件221（3月、JLAC11優先）の後、PDFの検査項目名はJLAC11のdisplayになるか。JLAC11のみで登録された場合はcode.textのままか。",
    "詳細画面の名称をPDFに揃える改修の可否、可能な場合の時期と費用。",
], X0, 12.3, W0, 4.5, size=1100)

cp = prs.core_properties
for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
    setattr(cp, k, "")
pkg = prs.part.package
for rId, rel in list(pkg._rels.items()):
    if rel.reltype.endswith("/custom-properties"):
        pkg._rels.pop(rId)
prs.save(OUT)
print("saved")
