# 正式回答（20261006）の複製の右端（K列）に、機構への確認事項を書く（v2）。
# 観点：画面とPDFの両方に出る項目で、同じ情報が出るか。
# 注意：画面列とPDF列を横に結合したセル（例 D4:E4）は「画面とPDFで同じ」の意味。空欄と読まない。
import sys, copy, openpyxl
from openpyxl.styles import Alignment, Font
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb.active
NAME = ("画面は医療機関が登録した名称（code.text）、PDFはコードの表示名を出すため、同じ{0}でも画面とPDFで名称が異なる（同じ情報にならない）という理解でよいか。"
        "あわせて、10/5にご報告いただいた詳細画面側の対応可否（{1}）の確認結果と、回答の見通し（時期）を共有してほしい。")
K = {
 7: "本記載は、案件213（臨床情報詳細画面の傷病名をPDFと同様に出力する対応、12月実装予定）のリリース後の挙動という理解でよいか。"
    "そうであれば、リリースまでは画面はTEXT値のままのため、画面とPDFで傷病名の表記が異なる（例：PDFは修飾語・「の疑い」を付けた名称、画面は医療機関が登録した名称）という理解でよいか。"
    "また、リリース後の違いは、PDFだけ100桁を超えた分を切る点のみという理解でよいか（100桁を超える傷病名は、画面では全部出て、PDFでは途中で切れる）。",
 10: "画面とPDFで同じ表示（未提供の場合×、それ以外は〇）という理解でよいか。案件143（PDFのレイアウト変更）の要求事項では、PDFの傷病名は「告知状況」の列に「未告知／済」と表示するとあり、画面（患者共有済み：×／〇）と項目名・表示が異なるように読める。どちらが最新の仕様か、画面とPDFで同じ表示になるのはいつのリリースからかを示してほしい。",
 12: NAME.format("検査項目", "コードセットのディスプレイ値での出力") + "案件221（JLAC11を優先して出力）のリリース後は、PDFの名称がどう変わるかもあわせて示してほしい。",
 19: NAME.format("薬剤", "コードセットのディスプレイ値での出力"),
 21: "【回答】臨床的状態がPDFに出力されない差は、ご提案のとおり現状のままとすることで差し支えない。"
     "【確認】画面に出る薬剤アレルギーのうち、臨床的状態がactiveのものだけがPDFに出る（画面とPDFで出てくる件数が異なることがある）という理解でよいか。",
 24: NAME.format("アレルギー物質", "コードセットのディスプレイ値での出力"),
 26: "【回答】No.21（薬剤アレルギー）と同様、現状のままとすることで差し支えない。【確認】画面に出るその他アレルギーのうち、activeのものだけがPDFに出るという理解でよいか。",
 28: "No.12（感染症の検査項目）と同じ。",
 29: "備考（J31）のとおり、詳細画面の「最新〜」は最新の検査ではなく臨床情報に紐づく検査が出るとのこと。そうであれば、同じ検査項目でも、画面とPDFで「最新」として出る検査結果（採取日時・基準値・結果・評価値・ステータス）が異なることがある（例：同じ検査の新しい結果が後から登録された場合）という理解でよいか。どのような場合に食い違うかを示してほしい。",
 31: "No.29と同じ。", 32: "No.29と同じ。", 33: "No.29と同じ。", 34: "No.29と同じ。",
}
h = ws['H2']; k = ws['K2']
k.value = "機構への確認事項（厚労省ACN案）\n観点：画面とPDFで同じ情報が出るか"
k.font = copy.copy(h.font); k.fill = copy.copy(h.fill); k.border = copy.copy(h.border)
k.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
ws['K3'].border = copy.copy(ws['H3'].border); ws['K3'].fill = copy.copy(h.fill)
ws.merge_cells('K2:K3')
ws.column_dimensions['K'].width = 70
ws.column_dimensions['J'].width = 24
for r in range(4, 47):
    c = ws[f'K{r}']; ref = ws[f'I{r}']
    c.border = copy.copy(ref.border); c.font = copy.copy(ref.font)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    if r in K: c.value = K[r]
notes = ["【全区分に共通の確認事項】",
 "共通確認①：本調査結果は、どの時点の画面・PDFを書いたものか。案件台帳では、画面は案件196（詳細画面のレイアウト変更：終了日・患者共有済みの追加等、令和8年度3月希望）・案件213（傷病名の出力、12月実装予定）、PDFは案件143・144（レイアウト変更、令和8年度9月希望）・案件221（JLAC11優先、令和8年度3月希望。PDF/XMLのみで画面は対象外）が予定されている。リリースまでの間、画面とPDFで同じ情報が出ない項目があれば示してほしい。",
 "共通確認②：画面とPDFで、出てくるレコード（件数・対象）は同じか。各項目の出し方が同じでも、出てくる傷病名・検査・アレルギー自体が違えば同じ情報にならないため、"
 "画面とPDFで出すレコードの条件（例：否定された傷病名、検査中止（cancelled）の検査、active以外のアレルギーを出すか）を区分ごとに示してほしい。"]
for i, t in enumerate(notes):
    c = ws.cell(row=48+i, column=11, value=t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if i == 0: c.font = Font(bold=True)
wb.save(dst)
print('written', len(K))
