"""WBS確認会向け：詳細画面とPDFの名称の差（検査項目・薬剤名・アレルギー物質）の相談スライドを作る。
v5：XMLの名称欄＝code.text、コードセット＝displayと整理（正式回答・基金の確認結果・NTTデータの対応づけから）。機構への確認をXMLレイアウトの記載の不備に差し替え。
v4：③を「表示に差分が生じるか」に、④を厚労省の意向確認に組み替え。
v3：検査項目のPDF欄を正式回答（JLAC10のdisplay、なければcode.text）どおりに修正。
v2：①現状の差分→②出力元XMLレイアウトのルール→③判断→④改修要否 の2枚構成。
体裁は 06_会議/WBS確認会/20261009/02_output/20261008_血液型の登録表示抑制スライド_v1.pptx に合わせる。
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(sys.argv[1])
FONT = "Meiryo UI"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0xD9, 0xD9, 0xD9)
PINK = RGBColor(0xF2, 0xDC, 0xDB)
LIGHT = RGBColor(0xDE, 0xEB, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Emu(10689336)
prs.slide_height = Emu(7562088)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, paras, size=12, bold=False, color=BLACK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: str（改行区切り）または [(文字列, {bold,color,size})] のリスト"""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Cm(0.15))
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    items = paras.split("\n") if isinstance(paras, str) else paras
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = it if isinstance(it, list) else [(it, {})] if isinstance(it, str) else [it]
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
    return tb


def header(slide, title, sub):
    text(slide, 1.1, 0.8, 27.4, 1.4, title, size=22, bold=True)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(0.8), Cm(2.3), Cm(29.0), Cm(2.3))
    ln.line.color.rgb = NAVY
    ln.line.width = Pt(1.5)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(2.8), Cm(27.4), Cm(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    text(slide, 1.4, 2.8, 26.9, 1.1, sub, size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, col_w, rows, row_h=0.9, size=10.5, head_fill=NAVY, first_col_fill=GRAY, fills=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Cm(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Cm(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.15)
            c.margin_top = c.margin_bottom = Cm(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = it if isinstance(it, list) else [(it, {})]
                for s, opt in runs:
                    r = p.add_run()
                    r.text = s
                    r.font.name = FONT
                    r.font.size = Pt(opt.get("size", size))
                    r.font.bold = opt.get("bold", i == 0)
                    r.font.color.rgb = opt.get("color", WHITE if i == 0 else BLACK)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = head_fill
            elif fills and (i, j) in fills:
                c.fill.fore_color.rgb = fills[(i, j)]
            elif j == 0:
                c.fill.fore_color.rgb = first_col_fill
            else:
                c.fill.fore_color.rgb = WHITE
    return tbl


B = {"bold": True}
R = {"bold": True, "color": RED}
T = "X. 臨床情報詳細画面とPDF帳票の名称の差に係る相談"

# ---- スライド1：①〜④ ----
s = prs.slides.add_slide(BLANK)
header(s, T, "ご相談事項：検査項目・薬剤名・アレルギー物質の名称が、詳細画面とPDFで異なる件の改修要否")
text(s, 1.1, 4.1, 27.4, 1.3, [
    [("ToDo#1228を受けたDX機構の調査結果（10/6）と、出力元であるXMLレイアウト（CIP-IF-002 臨床情報閲覧結果）・FHIRの定義を突き合わせ、詳細画面とPDFで表示に差分が生じるかを整理した。", {})],
], size=11)
table(s, 1.1, 5.5, 27.4, [2.7, 3.4, 5.0, 8.6, 7.7], [
    ["項目", "① 現状：詳細画面", "① 現状：PDF帳票", "② XMLレイアウト・FHIRのルール", "③ 表示に差分が生じるか"],
    [["検査項目", "（感染症", "・検査）"], ["医療機関の登録名称", "（code.text）"],
     ["JLAC10のコードの名称", "（display）", "JLAC10のコードがない場合は", "code.text"],
     ["XMLは、名称欄（検査名）に登録名称（code.text）、コードセットにJLACのコードの名称（display）の両方を持つ", "両者を一致させる決まりはない"],
     [[("生じうる", R)], "登録名称がJLAC10の名称と違えば差分が出る", "（出ないのは、JLAC10のコードがない場合と、登録名称が同じ場合）"]],
    ["薬剤名", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYコードの名称", "（display）", "ダミーコードの場合はcode.text"],
     ["XMLは、名称欄（薬剤名）に登録名称（code.text）、コードセットにJFAGYのコードの名称（display）の両方を持つ", "両者を一致させる決まりはない"],
     [[("生じうる", R)], "登録名称がJFAGYの名称と違えば差分が出る", "（出ないのは、ダミーコードの場合と、登録名称が同じ場合）"]],
    ["アレルギー物質", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYコードの名称", "（display）", "ダミーコードの場合はcode.text"],
     ["薬剤名と同じ（名称欄＝アレルギー名）"],
     [[("生じうる", R)], "薬剤名と同じ"]],
    [["（参考）", "傷病名"], ["医療機関の登録名称", "（code.text）", "「の疑い」が付かない"],
     "前置修飾語＋傷病名＋後置修飾語＋「の疑い」",
     ["XMLの傷病名は「前置修飾語＋傷病名＋後置修飾語」に加工して出力。疑いは別の欄"],
     [[("生じる → 対応済み", B)], "12月リリース（案件213）で画面をPDFと同様に出力"]],
], row_h=1.6, size=9.5, fills={(4, j): PINK for j in range(5)})
text(s, 1.1, 14.0, 27.4, 1.1, [[("→ XMLは1件の中に登録名称とコードの名称の両方を持ち、画面は登録名称、PDFはコードの名称を表示している。両者を一致させる決まりはないため、3項目とも表示に差分は生じうる（医療機関の登録内容しだい）。", B)]], size=11)

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(15.3), Cm(27.4), Cm(0.9))
bar.fill.solid()
bar.fill.fore_color.rgb = NAVY
bar.line.fill.background()
text(s, 1.4, 15.3, 26.9, 0.9, "④ 厚労省にご意向を伺いたい事項", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.1, 16.3, 27.4, 3.3, [
    [("（1）", B), ("表示に差分が生じうるため、傷病名（案件213）と同様に、画面の名称をPDF（コードの名称）に揃える改修を機構に求める方向でよろしいでしょうか。なお、登録名称は必須項目のため、揃えても名称が空欄になることはありません。", {})],
    [("（2）", B), ("医療機関の登録名称が標準名称と同じで、実際には差分が出ていない場合もありますが、その場合でも念のため改修を求めますか。", {})],
    [("（3）", B), ("富士通からは「品質担保の観点から3月リリースでの修正は困難」との回答のため、改修の可否・時期・費用を機構に照会し、費用対効果を見て判断する進め方でよろしいでしょうか。あわせて、改修までの間、差分があることを医療機関に周知するか。", {})],
], size=11)
text(s, 1.1, 19.6, 27.4, 0.7, "出典：DX機構「PDFと詳細画面の差異に関する調査結果」（2026/10/6）、別紙3-2 XMLレイアウト Ver5.10（CIP-IF-002）、JP-CLINS v1.13.0、案件台帳（令和8年度）", size=9)

# ---- スライド2：参考（根拠） ----
s = prs.slides.add_slide(BLANK)
header(s, T, "（参考）判断の根拠：FHIRの定義とXMLの名称欄の中身")
text(s, 1.1, 4.1, 27.4, 0.8, [[("1. 詳細画面が表示しているcode.textの定義（JP-CLINS v1.13.0）", B)]], size=12)
table(s, 1.1, 4.9, 27.4, [6.6, 2.8, 18.0], [
    ["プロファイル", "code.text", "定義（抜粋）"],
    ["JP_Observation_LabResult_eCS", "1..1（必須）", "項目名。報告書などに記載する場合に使用する表示名。（注釈：多くの場合、coding.displayと同一になるが…）"],
    ["JP_AllergyIntolerance_eCS", "1..1（必須）", "電子カルテシステム等で登録され表示されている文字列をcode.textに必ず設定すること。"],
    [["（参考）", "JP_Condition_eCS"], "1..1（必須）", "前置修飾語、後置修飾語も含めて電子カルテ上での表示（入力）文字列全体を設定する。"],
], row_h=0.9, size=10)
text(s, 1.1, 8.9, 27.4, 0.7, "→ 必須のため、画面の名称が空欄になることはない。登録名称はコードの名称と一致させる決まりはない。", size=11)

text(s, 1.1, 9.8, 27.4, 0.8, [[("2. XMLレイアウト（CIP-IF-002）の名称欄に入っている値", B)]], size=12)
table(s, 1.1, 10.6, 27.4, [6.6, 4.6, 16.2], [
    ["XMLの欄", "入っている値", "根拠"],
    [["名称欄", "（検査名・感染症名・薬剤名・アレルギー名）"], "登録名称（code.text）",
     ["DX機構の調査結果：参照元を「検査/感染症名（code.text）」「アレルギー/薬剤名（code.text）」と記載", "基金の確認結果（2026/4）：「system値なし：検査名ObservationText」", "患者サマリーの対応づけ（NTTデータ）：「ConditionText：code.textから設定」"]],
    ["コードセットのdisplay値", ["コードの名称", "（coding.display）"],
     ["XMLレイアウト：「system値に相対する文字列を設定」（出現回数1..1）", "DX機構の調査結果：PDFは「コードセットのディスプレイ値」を表示"]],
], row_h=1.0, size=9.5)

text(s, 1.1, 15.0, 27.4, 0.8, [[("3. 機構への確認事項（ToDo#1228の照会に含める）", B)]], size=12)
text(s, 1.1, 15.8, 27.4, 3.7, [
    "・XMLレイアウトの検査名（KnName）・感染症名（KsName）の項目説明「FHIR識別名が設定される」は、実際の設定値（code.text）と異なるのではないか。",
    "・その他アレルギーのdisplay値（ArDisplay・項番75）の項目説明「JFAGY医薬品アレルギーコードに相対する文字列」は、食品・非食品非医薬品コードの誤記ではないか。",
    "・PDFの検査項目名は、調査結果では「JLAC10のディスプレイ値」だが、技術解説書の「FHIR項目名称に自動変換して出力」と同じものか。案件221（JLAC11優先）の後はどうなるか。",
    "・詳細画面の名称をPDFに揃える改修の可否、可能な場合の時期と費用。",
], size=10.5)
text(s, 1.1, 19.6, 27.4, 0.7, "出典：JP-CLINS v1.13.0（StructureDefinition）、別紙3-2 XMLレイアウト Ver5.10、DX機構 調査結果（2026/10/6）、技術解説書2.1.0版", size=9)

cp = prs.core_properties
for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
    setattr(cp, k, "")
prs.save(OUT)
print("saved")
