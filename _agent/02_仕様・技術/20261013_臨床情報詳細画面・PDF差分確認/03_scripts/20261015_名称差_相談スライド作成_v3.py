"""WBS確認会向け：詳細画面とPDFの名称の差（検査項目・薬剤名・アレルギー物質）の相談スライドを作る。
v3：検査項目のPDF欄を正式回答（JLAC10のdisplay、なければcode.text）どおりに修正。
v2：①現状の差分→②出力元XMLレイアウトのルール→③判断→④改修要否 の2枚構成。
体裁は 06_会議/WBS確認会/20261009/02_output/20261008_血液型の登録表示抑制スライド_v1.pptx に合わせる。
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(sys.argv[1])
FONT = "Meiryo UI"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0xD9, 0xD9, 0xD9)
PINK = RGBColor(0xF2, 0xDC, 0xDB)
LIGHT = RGBColor(0xDE, 0xEB, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Emu(10689336)
prs.slide_height = Emu(7562088)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, paras, size=12, bold=False, color=BLACK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: str（改行区切り）または [(文字列, {bold,color,size})] のリスト"""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Cm(0.15))
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    items = paras.split("\n") if isinstance(paras, str) else paras
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = it if isinstance(it, list) else [(it, {})] if isinstance(it, str) else [it]
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
    return tb


def header(slide, title, sub):
    text(slide, 1.1, 0.8, 27.4, 1.4, title, size=22, bold=True)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(0.8), Cm(2.3), Cm(29.0), Cm(2.3))
    ln.line.color.rgb = NAVY
    ln.line.width = Pt(1.5)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(2.8), Cm(27.4), Cm(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    text(slide, 1.4, 2.8, 26.9, 1.1, sub, size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, col_w, rows, row_h=0.9, size=10.5, head_fill=NAVY, first_col_fill=GRAY, fills=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Cm(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Cm(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.15)
            c.margin_top = c.margin_bottom = Cm(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = it if isinstance(it, list) else [(it, {})]
                for s, opt in runs:
                    r = p.add_run()
                    r.text = s
                    r.font.name = FONT
                    r.font.size = Pt(opt.get("size", size))
                    r.font.bold = opt.get("bold", i == 0)
                    r.font.color.rgb = opt.get("color", WHITE if i == 0 else BLACK)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = head_fill
            elif fills and (i, j) in fills:
                c.fill.fore_color.rgb = fills[(i, j)]
            elif j == 0:
                c.fill.fore_color.rgb = first_col_fill
            else:
                c.fill.fore_color.rgb = WHITE
    return tbl


B = {"bold": True}
R = {"bold": True, "color": RED}
T = "X. 臨床情報詳細画面とPDF帳票の名称の差に係る相談"

# ---- スライド1：①〜④ ----
s = prs.slides.add_slide(BLANK)
header(s, T, "ご相談事項：検査項目・薬剤名・アレルギー物質の名称が、詳細画面とPDFで異なる件の改修要否")
text(s, 1.1, 4.1, 27.4, 1.3, [
    [("ToDo#1228を受けたDX機構の調査結果（10/6）により、下表の差分が判明。出力元であるXMLレイアウト（CIP-IF-002 臨床情報閲覧結果）のルールに照らして、差分を許容できるかを整理した。", {})],
], size=11)
table(s, 1.1, 5.5, 27.4, [2.9, 4.8, 5.8, 7.2, 6.7], [
    ["項目", "① 現状：詳細画面", "① 現状：PDF帳票", "② 出力元XMLレイアウトのルール", "③ ②を踏まえた判断"],
    [["検査項目", "（感染症・検査）"], ["医療機関の登録名称", "（code.text）"],
     ["JLAC10のコードの名称（display）", "JLAC10のコードがない場合はcode.text"],
     ["検査名は「FHIR識別名」を設定（＝JLACの標準名称）", "JLACのコードセット（display）も必須"],
     [[("差分は許容できない", R)], "XMLでは標準名称（FHIR識別名）と定められており、画面だけが医療機関の名称になっている"]],
    ["薬剤名", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYコードの名称（display）", "ダミーコードの場合はcode.text"],
     ["薬剤名の設定内容は記載なし", "JFAGYのコードセット（display）は必須"],
     [[("差分は許容できない", R)], "XMLでは標準名称（display）が必ず出力され、PDF（技術解説書 表18）もdisplayを表示する"]],
    ["アレルギー物質", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYコードの名称（display）", "ダミーコードの場合はcode.text"],
     ["「アレルギー名を設定」", "JFAGYのコードセット（display）は必須"],
     [[("差分は許容できない", R)], "薬剤名と同じ"]],
    [["（参考）", "傷病名"], ["医療機関の登録名称", "（code.text）", "「の疑い」が付かない"],
     "前置修飾語＋傷病名＋後置修飾語＋「の疑い」",
     ["「前置修飾語＋傷病名＋後置修飾語」に加工して出力", "疑いは別の欄（病名疑い）"],
     [[("対応済み", B)], "12月リリース（案件213）で画面をPDFと同様に出力"]],
], row_h=1.75, size=9.5, fills={(4, j): PINK for j in range(5)})

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.1), Cm(14.6), Cm(27.4), Cm(0.9))
bar.fill.solid()
bar.fill.fore_color.rgb = NAVY
bar.line.fill.background()
text(s, 1.4, 14.6, 26.9, 0.9, "④ 改修が必要か", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.1, 15.7, 27.4, 3.6, [
    [("改修が必要と考える。", R), ("詳細画面は自施設の登録内容を確認する画面だが、画面にはcode.textしか出ないため、誤ったコードで登録しても、他の医療機関に別の名称が届いていることに気づけない。", {})],
    [("・進め方：", B), ("傷病名（案件213）と同様に、画面の名称をXMLレイアウト・PDFのルールに揃える改修を機構に求める。", {})],
    [("・時期・費用：", B), ("富士通は「品質担保の観点から3月リリースでの修正は困難」との回答。改修の可否・時期・費用を機構に照会し、費用対効果を見て判断する。", {})],
    [("・当面の扱い：", B), ("改修までの間、画面とPDFで名称が異なることを医療機関に周知するか。", {})],
], size=11)
text(s, 1.1, 19.6, 27.4, 0.7, "出典：DX機構「PDFと詳細画面の差異に関する調査結果」（2026/10/6）、別紙3-2 XMLレイアウト Ver5.10（CIP-IF-002）、技術解説書2.1.0版 表18、案件台帳（令和8年度）", size=9)

# ---- スライド2：参考（根拠） ----
s = prs.slides.add_slide(BLANK)
header(s, T, "（参考）判断の根拠：FHIRの定義と技術解説書の記載")
text(s, 1.1, 4.1, 27.4, 0.8, [[("1. 詳細画面が表示しているcode.textの定義（JP-CLINS v1.13.0）", B)]], size=12)
table(s, 1.1, 4.9, 27.4, [6.6, 2.8, 18.0], [
    ["プロファイル", "code.text", "定義（抜粋）"],
    ["JP_Observation_LabResult_eCS", "1..1（必須）", "項目名。報告書などに記載する場合に使用する表示名。"],
    ["JP_AllergyIntolerance_eCS", "1..1（必須）", "電子カルテシステム等で登録され表示されている文字列をcode.textに必ず設定すること。"],
    [["（参考）", "JP_Condition_eCS"], "1..1（必須）", "前置修飾語、後置修飾語も含めて電子カルテ上での表示（入力）文字列全体を設定する。"],
], row_h=1.0, size=10)
text(s, 1.1, 9.1, 27.4, 0.8, "→ 必須のため画面が空欄になることはない。code.textは医療機関の電子カルテ上の名称で、標準名称（コードの名称）とは一致しないことがある。", size=11)
text(s, 1.1, 10.3, 27.4, 0.8, [[("2. 閲覧時に表示する名称（技術解説書2.1.0版）", B)]], size=12)
text(s, 1.1, 11.1, 27.4, 2.6, [
    "・表24：FHIR項目名称は「臨床情報一覧やマイナポータル等の閲覧画面で利用者が直感的に理解しやすい形で表示するための検査項目名称」。",
    "・本文：臨床情報一覧のXMLレイアウトやPDF帳票上の検査項目名は、FHIR識別文字列を本サービス側でFHIR項目名称に自動変換して出力する。",
    "・表18：薬剤名・アレルギー物質は、JFAGYコードの表示名称（display値）を表示し、ダミーコードの場合のみtext値を表示する。",
], size=11)
text(s, 1.1, 14.0, 27.4, 0.8, [[("3. 機構への確認事項（ToDo#1228の照会に含める）", B)]], size=12)
text(s, 1.1, 14.8, 27.4, 3.4, [
    "・XMLレイアウトの薬剤名（YkName）・アレルギー名（ArName）に設定しているのは、code.textかdisplayか（項目説明に記載がないため）。",
    "・詳細画面の名称をXMLレイアウト・PDFに揃える改修の可否、可能な場合の時期と費用。",
    "・PDFの検査項目名は、調査結果では「JLAC10のディスプレイ値」だが、技術解説書の「FHIR項目名称に自動変換して出力」と同じものか。",
    "・3月リリース（案件221：JLAC11を優先して出力）の後、PDFの検査項目名はどうなるか（JLAC11だけの登録はcode.textになるのか）。詳細画面は変わらないという理解でよいか。",
], size=11)
text(s, 1.1, 19.6, 27.4, 0.7, "出典：JP-CLINS v1.13.0（StructureDefinition）、技術解説書2.1.0版（表18・表24・検体検査のJLACコードの記述）、別紙3-2 XMLレイアウト Ver5.10", size=9)

cp = prs.core_properties
for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
    setattr(cp, k, "")
prs.save(OUT)
print("saved")
