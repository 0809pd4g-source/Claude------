"""v6：厚労省推奨の構成（背景・状況→ご相談事項→参考）と体裁に合わせて作り直す。
体裁の手本：20260902_血液型の共有について.pptx（テンプレートの「7_タイトルとコンテンツ」、紺のセクションラベル、•箇条書き12pt、薄ピンク見出しの表）。
手本ファイルを土台にし、元のスライドを除いて同じレイアウトでスライドを作る。
使い方：python 本ファイル <手本pptx> <出力pptx>
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.oxml.ns import qn
from lxml import etree

REF, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(REF)
base = prs.slides[3]
LAYOUT = base.slide_layout
LABEL = BODY = None
for sh in base.shapes:
    if sh.name == "四角形: 角を丸くする 3":
        LABEL = copy.deepcopy(sh._element)
    if sh.name == "コンテンツ プレースホルダー 10" and sh.is_placeholder:
        BODY = copy.deepcopy(sh._element)
assert LABEL is not None and BODY is not None

# 手本のスライドを除く
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
RED = "C00000"
_id = [100]


def _next_id():
    _id[0] += 1
    return _id[0]


def new_slide(title):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1 and ph.placeholder_format.idx == 1:  # 本文プレースホルダーは自前で置く
            ph._element.getparent().remove(ph._element)
    s.shapes.title.text = title
    return s


def _set_xfrm(el, x, y, w, h):
    off = el.find(".//" + qn("a:off"))
    ext = el.find(".//" + qn("a:ext"))
    off.set("x", str(int(Cm(x)))); off.set("y", str(int(Cm(y))))
    ext.set("cx", str(int(Cm(w)))); ext.set("cy", str(int(Cm(h))))


def label(s, text, x, y, w=4.57):
    el = copy.deepcopy(LABEL)
    el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id()))
    el.find(".//" + qn("p:cNvPr")).set("name", f"ラベル {text}")
    _set_xfrm(el, x, y, w, 0.69)
    el.find(".//" + qn("a:t")).text = text
    s.shapes._spTree.append(el)


def _run(p, text, bold=False, color=None, size=1200):
    r = etree.SubElement(p, qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"), lang="ja-JP", altLang="en-US", sz=str(size))
    if bold:
        rpr.set("b", "1")
    if color:
        sf = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(r, qn("a:t")).text = text


def body(s, paras, x, y, w, h, size=1200):
    """paras: [str | [(text, {bold, color})]]。•付きの箇条書き（手本と同じ字下げ）"""
    el = copy.deepcopy(BODY)
    el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id()))
    el.find(".//" + qn("p:cNvPr")).set("name", "本文")
    _set_xfrm(el, x, y, w, h)
    tx = el.find(qn("p:txBody"))
    for p in tx.findall(qn("a:p")):
        tx.remove(p)
    for para in paras:
        p = etree.SubElement(tx, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"), marL="285750", indent="-285750")
        sa = etree.SubElement(ppr, qn("a:spcAft"))
        etree.SubElement(sa, qn("a:spcPts"), val="300")
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial", panose="020B0604020202020204", pitchFamily="34", charset="0")
        etree.SubElement(ppr, qn("a:buChar"), char="•")
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            _run(p, t, o.get("bold", False), o.get("color"), size)
    s.shapes._spTree.append(el)


def _border(tcPr):
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w="6350", cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val="000000")
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def table(s, x, y, col_w, rows, row_h=0.6, size=900, head_fill="F2CEEF", fills=None):
    shp = s.shapes.add_table(len(rows), len(col_w), Cm(x), Cm(y), Cm(sum(col_w)), Cm(row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Cm(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Cm(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            tf = c.text_frame
            tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = it if isinstance(it, list) else [(it, {})]
                for t, o in runs:
                    r = p.add_run()
                    r.text = t
                    r.font.size = Pt(size / 100)
                    r.font.name = "Meiryo UI"
                    rpr = r._r.get_or_add_rPr()
                    ea = etree.SubElement(rpr, qn("a:ea"), typeface="Meiryo UI")
                    r.font.bold = o.get("bold", i == 0)
                    from pptx.dml.color import RGBColor
                    r.font.color.rgb = RGBColor.from_string(o.get("color", "000000"))
            c.margin_left = c.margin_right = Cm(0.1)
            c.margin_top = Cm(0.05)
            c.margin_bottom = Cm(0.03)
            from pptx.enum.text import MSO_ANCHOR
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tcPr = c._tc.get_or_add_tcPr()
            for child in list(tcPr):
                tcPr.remove(child)
            _border(tcPr)
            sf = etree.SubElement(tcPr, qn("a:solidFill"))
            fill = head_fill if i == 0 else (fills or {}).get((i, j), "FFFFFF")
            etree.SubElement(sf, qn("a:srgbClr"), val=fill)
    return tbl


B = {"bold": True}
R = {"bold": True, "color": RED}
TITLE = "X.臨床情報詳細画面とPDF帳票の名称の差について"
X0, W0 = 1.37, 24.8

# ---- スライド1：背景・状況 → ご相談事項 ----
s = new_slide(TITLE)
label(s, "背景・状況", X0, 2.3)
body(s, [
    "ToDo#1228（10/5開発課題検討会）を受け、DX機構から「PDFと詳細画面の差異に関する調査結果」（10/6）が共有された。",
    [("検査項目・薬剤名・アレルギー物質は、", {}), ("詳細画面が医療機関の登録名称（code.text）、PDFがコードの名称（display）", B), ("を表示しており、同じ登録データでも名称が異なる場合がある。", {})],
    [("出力元のXMLレイアウト（CIP-IF-002）は、1件の中に登録名称とコードの名称の両方を持ち、両者を一致させる決まりはない。そのため、", {}), ("医療機関の登録内容しだいで表示に差分が生じうる", B), ("（下表）。", {})],
    "富士通からは「品質担保の観点から3月リリースでの修正は困難」との回答を受けている（10/5 DX機構）。",
], X0, 3.0, W0, 3.9)
table(s, 1.9, 7.0, [3.0, 4.6, 6.0, 5.4, 4.8], [
    ["項目", "詳細画面", "PDF帳票", "出力元XMLレイアウト", "表示の差分"],
    [["検査項目", "（感染症・検査）"], ["医療機関の登録名称", "（code.text）"],
     ["JLAC10のコードの名称（display）", "JLAC10のコードがない場合はcode.text"],
     ["名称欄（検査名）＝登録名称", "コードセット＝JLACのコードの名称"],
     [[("生じうる", R)], "（JLAC10のコードがない場合は同じ）"]],
    ["薬剤名", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYのコードの名称（display）", "ダミーコードの場合はcode.text"],
     ["名称欄（薬剤名）＝登録名称", "コードセット＝JFAGYのコードの名称"],
     [[("生じうる", R)], "（ダミーコードの場合は同じ）"]],
    ["アレルギー物質", ["医療機関の登録名称", "（code.text）"],
     ["JFAGYのコードの名称（display）", "ダミーコードの場合はcode.text"],
     ["名称欄（アレルギー名）＝登録名称", "コードセット＝JFAGYのコードの名称"],
     [[("生じうる", R)], "（ダミーコードの場合は同じ）"]],
    ["（参考）傷病名", ["医療機関の登録名称", "（code.text）", "「の疑い」が付かない"],
     "前置修飾語＋傷病名＋後置修飾語＋「の疑い」",
     ["傷病名＝修飾語を付けて加工", "疑いは別の欄"],
     [[("生じる", B)], "→ 12月リリース（案件213）で画面をPDFと同様に出力"]],
], row_h=0.9)

label(s, "ご相談事項", X0, 12.7)
body(s, [
    [("表示に差分が生じうるため、", {}), ("傷病名（案件213）と同様に、詳細画面の名称をPDF（コードの名称）に揃える改修を機構に求める方向", B), ("でよろしいでしょうか。なお、登録名称は必須項目のため、揃えても名称が空欄になることはありません。", {})],
    [("医療機関の登録名称がコードの名称と同じで、", {}), ("実際には差分が出ていない場合もありますが、その場合でも念のため改修を求めますか。", B)],
    "改修の可否・時期・費用を機構に照会し、費用対効果を見て判断する進め方でよろしいでしょうか。あわせて、改修までの間、差分があることを医療機関に周知するかについてご意向を伺いたい。",
], X0, 13.4, W0, 4.5)

# ---- スライド2：参考 → 機構への確認事項 ----
s = new_slide(TITLE)
label(s, "参考", X0, 2.3)
body(s, [[("XMLレイアウト（CIP-IF-002）の名称欄に入っている値", B), ("（DX機構の調査結果等から整理）", {})]], X0, 3.0, W0, 0.8)
table(s, 1.9, 3.8, [6.4, 4.4, 13.0], [
    ["XMLの欄", "入っている値", "根拠"],
    [["名称欄", "（検査名・感染症名・薬剤名・アレルギー名）"], "登録名称（code.text）",
     ["DX機構の調査結果：参照元を「検査/感染症名（code.text）」「アレルギー/薬剤名（code.text）」と記載",
      "基金の確認結果（2026/4）：「system値なし：検査名ObservationText」",
      "患者サマリーの対応づけ（NTTデータ）：「ConditionText：code.textから設定」"]],
    ["コードセットのdisplay値", ["コードの名称", "（coding.display）"],
     ["XMLレイアウト：「system値に相対する文字列を設定」（出現回数1..1）", "DX機構の調査結果：PDFは「コードセットのディスプレイ値」を表示"]],
], row_h=0.9)
body(s, [[("詳細画面が表示しているcode.textの定義", B), ("（JP-CLINS v1.13.0）", {})]], X0, 7.5, W0, 0.8)
table(s, 1.9, 8.3, [6.4, 2.4, 15.0], [
    ["プロファイル", "code.text", "定義（抜粋）"],
    ["JP_Observation_LabResult_eCS", "1..1（必須）", "項目名。報告書などに記載する場合に使用する表示名。（注釈：多くの場合、coding.displayと同一になるが…）"],
    ["JP_AllergyIntolerance_eCS", "1..1（必須）", "電子カルテシステム等で登録され表示されている文字列をcode.textに必ず設定すること。"],
    ["（参考）JP_Condition_eCS", "1..1（必須）", "前置修飾語、後置修飾語も含めて電子カルテ上での表示（入力）文字列全体を設定する。"],
], row_h=0.6)
body(s, ["登録名称は必須のため画面の名称が空欄になることはないが、コードの名称と一致させる決まりはない。"], X0, 11.0, W0, 0.8)

label(s, "機構への確認事項", X0, 12.2)
body(s, [
    "XMLレイアウトの検査名（KnName）・感染症名（KsName）の項目説明「FHIR識別名が設定される」は、実際の設定値（code.text）と異なるのではないか。",
    "その他アレルギーのdisplay値（ArDisplay・項番75）の項目説明「JFAGY医薬品アレルギーコードに相対する文字列」は、食品・非食品非医薬品コードの誤記ではないか。",
    "PDFの検査項目名は、調査結果では「JLAC10のディスプレイ値」だが、技術解説書の「FHIR項目名称に自動変換して出力」と同じものか。案件221（JLAC11優先）の後はどうなるか。",
    "詳細画面の名称をPDFに揃える改修の可否、可能な場合の時期と費用。",
], X0, 12.9, W0, 4.8, size=1100)

cp = prs.core_properties
for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
    setattr(cp, k, "")
pkg = prs.part.package
for rId, rel in list(pkg._rels.items()):
    if rel.reltype.endswith("/custom-properties"):
        pkg._rels.pop(rId)
prs.save(OUT)
print("saved")
