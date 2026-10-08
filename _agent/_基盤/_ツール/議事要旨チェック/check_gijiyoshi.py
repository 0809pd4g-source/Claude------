"""開発課題検討会・開発進捗連絡会議の議事要旨(案)pptxをチェックし、指摘候補をMarkdownで出力する。

使い方:
  python check_gijiyoshi.py --minutes 議事要旨(案).pptx --date 2026-10-05 \
      --materials 会議資料.pptx 会議資料_厚労省側.pptx --ledger 課題・TODO管理台帳.xlsx \
      --prev 前回議事要旨.pptx --notes 逐語録ベース議事録.md --transcript 逐語録.md --out 結果.md
台帳・会議資料・前回分・逐語録は省略可（省略した照合はスキップされる）。台帳は読み取り専用で開き、保存しない。
"""
import argparse
import collections
import datetime as dt
import difflib
import re
import unicodedata
from pathlib import Path

from pptx import Presentation

A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
ORG_ALIASES = {'基金': 'DX機構', '支払基金': 'DX機構'}

findings = []


def add(level, kind, where, msg, basis='', comment='', anchor=None, target=None):
    # anchor: 議事要旨の中で指摘箇所を探す手がかりの文字列 / target: 黄色ハイライトする文字列（anchorの位置以降で最初の出現）
    if any(f['where'] == where and f['msg'] == msg and f['comment'] == comment for f in findings):
        return
    findings.append(dict(level=level, kind=kind, where=where, msg=msg, basis=basis, comment=comment, anchor=anchor, target=target))


def clean(s):
    return (s or '').replace('​', '').replace('\x0b', '\n').replace('_x000B_', '\n').replace('\r', '')


def norm(s):
    s = unicodedata.normalize('NFKC', clean(s))
    for k, v in ORG_ALIASES.items():
        s = s.replace(k, v)
    s = s.replace('DX機構機構', 'DX機構')
    return re.sub(r'\s+', '', s)


def ratio(a, b):
    return difflib.SequenceMatcher(None, norm(a), norm(b)).ratio()


def bigram_sim(a, b):
    a, b = norm(a), norm(b)
    A = {a[i:i + 2] for i in range(len(a) - 1)}
    B = {b[i:i + 2] for i in range(len(b) - 1)}
    return len(A & B) / max(1, min(len(A), len(B)))


def diff_snippets(a, b, limit=4):
    a, b = norm(a), norm(b)
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
        if op != 'equal':
            out.append(f'「{a[max(0, i1 - 8):i2 + 8]}」⇔「{b[max(0, j1 - 8):j2 + 8]}」')
    return ' / '.join(out[:limit])


def changed_chars(mine, ref):
    """mine と ref の食い違い文字数。ref側にしかない※注記・補足・末尾の追記は「省略」とみなして数えない。"""
    a, b = norm(mine), norm(ref)
    n = 0
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
        if op == 'replace':
            n += max(i2 - i1, j2 - j1)
        elif op == 'delete':
            n += i2 - i1
        elif op == 'insert':
            if i1 >= len(a) or re.match(r'(※|■|①|・|なお|\(※\)|以下|下記)', b[j1:j2]):
                continue
            n += min(j2 - j1, 10)
    return n


def md(s):
    return clean(s).replace('|', '｜').replace('\n', '<br>')


def parse_date(s, year):
    s = clean(str(s))
    try:
        m = re.search(r'(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})', s)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})(?!\d)', s) or re.search(r'(\d{1,2})月(\d{1,2})日', s)
        if m:
            return dt.date(year, int(m[1]), int(m[2]))
    except ValueError:
        return None
    return None


def balanced_paren(text, start):
    """text[start] が開き括弧のとき、対応する閉じ括弧までの中身を返す。"""
    depth = 0
    for i in range(start, len(text)):
        c = text[i]
        if c in '(（':
            depth += 1
        elif c in ')）':
            depth -= 1
            if depth == 0:
                return text[start + 1:i]
    return text[start + 1:]


# ---------------------------------------------------------------- 議事要旨の読み込み
def load_minutes(path):
    prs = Presentation(path)
    slides = []
    for idx, sl in enumerate(prs.slides, 1):
        title, table, others = '', None, []
        for s in sl.shapes:
            if s.has_table and 'アジェンダ' in clean(s.table.cell(0, 1).text):
                table = s
            elif s.is_placeholder and s.has_text_frame and '議事要旨' in s.text_frame.text:
                title = clean(s.text_frame.text)
            else:
                others.append(s)
        slides.append(dict(idx=idx, slide=sl, title=title, table=table, others=others))
    rows = []
    for sd in slides:
        if not sd['table']:
            continue
        for ri, r in enumerate(sd['table'].table.rows):
            if ri == 0:
                continue
            cells = [clean(c.text) for c in r.cells]
            rows.append(dict(slide=sd['idx'], no=cells[0].strip(), agenda=cells[1].strip(),
                             consensus=cells[2], todo=cells[3] if len(cells) > 3 else '', row=r))
    return prs, slides, rows


def merge_agendas(rows):
    merged = collections.OrderedDict()
    for r in rows:
        key = (r['no'], r['agenda'])
        if key not in merged:
            merged[key] = dict(no=r['no'], agenda=r['agenda'], consensus='', todo='', slides=[])
        merged[key]['consensus'] += '\n' + r['consensus']
        merged[key]['todo'] += '\n' + r['todo']
        merged[key]['slides'].append(r['slide'])
    return list(merged.values())


SEC_RE = re.compile(r'^\s*(\d)[．.]\s*(.+?)\s*$')
ITEM_RE = re.compile(r'^\s*(#|ID)(\d+)\s*※?\s*[(（]')


def parse_ledger_sections(consensus, slides):
    """台帳確認の合意形成事項を「1．案件の新規起票」等の節と項目に分解する。"""
    sections, cur, item = [], None, None
    for line in consensus.split('\n'):
        m = SEC_RE.match(line)
        if m and len(m[2]) < 30 and not ITEM_RE.match(line):
            title = m[2]
            prev = next((s for s in sections if s['title'] == title), None)
            if prev:
                cur = prev
            else:
                cur = dict(title=title, declared=None, items=[], heading_line='', groups=[])
                sections.append(cur)
            item = None
            continue
        if cur is None:
            continue
        mc = re.search(r'[(（](\d+)件[)）]', line)
        if mc and not ITEM_RE.match(line):
            # 1つの節に「(N件)」の導入文が複数ある場合は、導入文ごとに件数を数える
            if cur['declared'] is None:
                cur['declared'] = int(mc[1])
                cur['heading_line'] = line.strip()
            cur['groups'].append(dict(declared=int(mc[1]), heading_line=line.strip(), n=0))
            continue
        mi = ITEM_RE.match(line)
        if mi:
            p = line.find('(') if '(' in line[:12] else line.find('（')
            item = dict(kind=mi[1], id=int(mi[2]), raw=line.strip(), body=balanced_paren(line, p), extra='')
            cur['items'].append(item)
            if cur['groups']:
                cur['groups'][-1]['n'] += 1
            continue
        if item is not None and line.strip() and line.strip() != '(続)':
            item['extra'] += line.strip() + '\n'
    for sec in sections:
        for it in sec['items']:
            ex = it['extra']
            m = re.search(r'(\d{4}/\d{1,2}/\d{1,2})\s*→\s*(\d{4}/\d{1,2}/\d{1,2})', ex)
            it['old'], it['new'] = (m[1], m[2]) if m else (None, None)
            mb = re.search(r'■[^\n]*背景\n?([^■]*)', ex)
            it['background'] = mb[1].strip() if mb else None
    return sections


def sec_type(title):
    if '新規起票' in title:
        return 'new'
    if '取り下げ' in title:
        return 'withdraw'
    if 'クローズ' in title:
        return 'close'
    if '延伸' in title:
        return 'extend'
    if '変更' in title or '更新' in title:
        return 'change'
    return 'other'


TODO_RE = re.compile(r'【\s*(厚労省|DX機構|基金|支払基金|ACN|開発事業者|富士通|NTTD?|[^】\s]{1,8})\s*(\d{1,2}/\d{1,2})\s*】')


def split_todos(cell):
    """ToDo列を【担当 M/D】で区切ってToDo単位に分ける（直後の※注記は同じToDoに含める）。"""
    text = clean(cell)
    todos, last = [], 0
    for m in TODO_RE.finditer(text):
        body = text[last:m.start()].strip()
        tail = text[m.end():]
        note = re.match(r'\s*(※[^\n]*(?:\n(?![\s]*\n)[^\n]*)*)', tail)
        end = m.end() + (note.end() if note else 0)
        todos.append(dict(body=body, org=m[1], due=m[2], note=note[1].strip() if note else '', raw=text[last:end].strip()))
        last = end
    rest = text[last:].strip()
    return todos, rest


# ---------------------------------------------------------------- 台帳
def load_ledger(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    todo = {}
    ws = wb['ToDo']
    for row in ws.iter_rows(min_row=3, max_col=12, values_only=True):
        try:
            n = int(row[1])
        except (TypeError, ValueError):
            continue
        todo[n] = dict(id=n, text=clean(str(row[2] or '')), src=clean(str(row[3] or '')), org=clean(str(row[4] or '')),
                       status=clean(str(row[6] or '')), due=row[7], progress=clean(str(row[8] or '')), dest=clean(str(row[0] or '')))
    issues = {}
    if '開発課題' in wb.sheetnames:
        ws = wb['開発課題']
        hdr = None
        for row in ws.iter_rows(min_row=1, max_row=4, max_col=20, values_only=True):
            vals = [clean(str(v or '')) for v in row]
            if '#' in vals and any('ステータス' in v for v in vals):
                hdr = vals
                break
        if hdr:
            col = lambda key: next(i for i, v in enumerate(hdr) if key in v)
            ci, ct, cd, cs, cp = col('#'), col('課題'), col('対応期限'), col('ステータス'), col('進捗')
            for row in ws.iter_rows(min_row=3, max_col=20, values_only=True):
                try:
                    n = int(row[ci])
                except (TypeError, ValueError):
                    continue
                issues[n] = dict(id=n, text=clean(str(row[ct] or '')), due=row[cd], status=clean(str(row[cs] or '')),
                                 progress=clean(str(row[cp] or '')))
    wb.close()
    return todo, issues


def to_date(v, year):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return parse_date(v, year) if v else None


def fmt(d):
    return f'{d.year}/{d.month}/{d.day}' if d else '（なし）'


# ---------------------------------------------------------------- 会議資料
def load_material(path):
    prs = Presentation(path)
    pages = []
    for sl in prs.slides:
        texts, tables = [], []
        for s in sl.shapes:
            stack = [s]
            while stack:
                x = stack.pop()
                if x.shape_type == 6:
                    stack.extend(x.shapes)
                elif x.has_table:
                    tables.append([[clean(c.text) for c in r.cells] for r in x.table.rows])
                elif x.has_text_frame and x.text_frame.text.strip():
                    texts.append(clean(x.text_frame.text))
        pages.append(dict(texts=texts, tables=tables))
    return dict(name=Path(path).name, stem=Path(path).stem, pages=pages)


def page_title(page):
    for t in page['texts']:
        if re.match(r'^\s*[0-9０-９]+[．.]|案件ID|^\s*参考', t) or len(t) < 80:
            return t.split('\n')[0][:60]
    return (page['texts'][0].split('\n')[0][:60]) if page['texts'] else ''


# ---------------------------------------------------------------- チェック本体
def check_structure(slides, rows):
    content = [s for s in slides if s['table']]
    n = len(content)
    for k, sd in enumerate(content, 1):
        m = re.search(r'[(（](\d+)/(\d+)[)）]', sd['title'])
        if not m:
            add('参考', '体裁', f"p{sd['idx']}", 'タイトルに(k/n)のページ番号がない', sd['title'])
        elif (int(m[1]), int(m[2])) != (k, n):
            add('要修正', '体裁', f"p{sd['idx']}", f'ページ番号が({m[1]}/{m[2]})だが、実際は({k}/{n})', sd['title'],
                f'細かい点で恐れ入りますが、ページ番号を({k}/{n})に修正いただけますでしょうか。')
    for sd in slides:
        for s in sd['others']:
            if s.left is not None and (s.left >= 9906000 * 0.97 or s.left + s.width <= 0):
                txt = clean(s.text_frame.text)[:40] if s.has_text_frame else ''
                add('要確認', '体裁', f"p{sd['idx']}", 'スライド外に図形が残っている（コメント枠の消し忘れ等）', txt)
    # (続) のつながり
    by_key = collections.defaultdict(list)
    for r in rows:
        by_key[(r['no'], r['agenda'])].append(r)
    for (no, ag), rs in by_key.items():
        for a, b in zip(rs, rs[1:]):
            if not a['consensus'].rstrip().endswith('(続)'):
                add('参考', '体裁', f"p{a['slide']}→p{b['slide']} 議題{no}", '次ページへ続く行の末尾に「(続)」がない')
            if not b['consensus'].lstrip().startswith('(続)'):
                add('参考', '体裁', f"p{b['slide']} 議題{no}", '続きの行の先頭に「(続)」がない')


def check_text_style(rows):
    blanks = collections.Counter()
    for r in rows:
        for col, txt in (('合意形成事項', r['consensus']), ('ToDo', r['todo'])):
            t = txt.strip()
            if re.fullmatch(r'[・•✓\s]*(特になし|なし|ー|―|-)[。．]?', t.split('\n')[0].strip()) and len(t) < 15:
                blanks[t.split('\n')[0].strip()] += 1
                if '\n' in t:
                    add('参考', '表記', f"p{r['slide']} 議題{r['no']} {col}", '「特になし」の後ろに空行・改行が残っている', repr(t))
    if len(blanks) > 1:
        add('要修正', '表記', '全体', '「特になし」の表記がゆれている: ' + '、'.join(f'「{k}」×{v}' for k, v in blanks.items()), '',
            '細かい点で恐れ入りますが、該当事項がない欄の表記を「特になし」に統一いただけますでしょうか。')
    alltext = '\n'.join(r['consensus'] + '\n' + r['todo'] for r in rows)
    patterns = [
        (r'。。|、、|。、|、。', '要修正', '句読点の重複'),
        (r'\s+[)）]', '参考', '閉じ括弧の前に空白'),
        (r'[(（]\s+', '参考', '開き括弧の後に空白'),
        (r'令和\d+年\d+末', '要確認', '「令和N年M末」（「月」の抜けの可能性）'),
        (r'(\S{4,})の\1', '要確認', '同じ語句の重複（脱字・コピペの可能性）'),
    ]
    for r in rows:
        for col, txt in (('合意形成事項', r['consensus']), ('ToDo', r['todo'])):
            for pat, lv, label in patterns:
                for m in re.finditer(pat, txt):
                    ctx = txt[max(0, m.start() - 25):m.end() + 15].replace('\n', ' ')
                    add(lv, '表記', f"p{r['slide']} 議題{r['no']} {col}", label, f'…{ctx}…',
                        f'細かい点で恐れ入りますが、「{m[0].strip() or "空白"}」の箇所をご確認いただけますでしょうか。' if lv != '参考' else '',
                        anchor=txt[max(0, m.start() - 25):m.start()].split('\n')[-1] + m[0], target=m[0])
            for line in txt.split('\n'):
                if re.match(r'^[ 　]+(#|ID)\d+', line):
                    add('参考', '体裁', f"p{r['slide']} 議題{r['no']} {col}", '項目の先頭に空白がある（字下げのずれ）', line[:40])
    # 括弧の全角半角ゆれ（ID直後）
    full = [m for m in re.finditer(r'(?:案件)?ID\d+（', alltext)]
    half = [m for m in re.finditer(r'(?:案件)?ID\d+\(', alltext)]
    if full and half:
        minor = full if len(full) <= len(half) else half
        add('参考', '表記', '全体', f'ID直後の括弧が全角{len(full)}件／半角{len(half)}件で混在',
            '少数側: ' + '、'.join(alltext[m.start():m.end() + 8] for m in minor[:5]),
            '細かい点で恐れ入りますが、括弧の全角・半角を統一いただけますでしょうか。')
    # 日付のゼロ埋めゆれ
    padded = re.findall(r'\d{4}/0\d/\d{1,2}|\d{4}/\d{1,2}/0\d', alltext)
    unpadded = re.findall(r'\d{4}/[1-9]/\d{1,2}|\d{4}/1[0-2]/[1-9](?!\d)', alltext)
    if padded and unpadded:
        add('参考', '表記', '全体', f'日付のゼロ埋めが混在（ゼロ埋め{len(padded)}件／なし{len(unpadded)}件）', '例: ' + '、'.join((padded + unpadded)[:4]))
    # 節見出しの句点ゆれ
    heads = re.findall(r'^.*下記.*[(（]\d+件[)）].*$', alltext, flags=re.M)
    with_dot = [h for h in heads if re.search(r'。\s*[(（]\d+件', h)]
    if heads and 0 < len(with_dot) < len(heads):
        add('参考', '表記', '議題2', '節の導入文で「。(N件)」と「(N件)」が混在',
            ' / '.join(h.strip()[:30] for h in heads))
    # 長文
    for r in rows:
        for col, txt in (('合意形成事項', r['consensus']), ('ToDo', r['todo'])):
            for sent in (x for line in txt.split('\n') for x in re.split(r'(?<=。)', line)):
                runs = re.split(r'[、,，]', sent)
                if len(sent) > 130 and max(len(x) for x in runs) > 70:
                    add('参考', '表記', f"p{r['slide']} 議題{r['no']} {col}", f'一文が長く読点が少ない（{len(sent)}字、読点なし最長{max(len(x) for x in runs)}字）',
                        sent[:60] + '…', '細かい点で恐れ入りますが、一文が長くなっているため、読点の追加または文の分割をご検討いただけますでしょうか。')
    # 参照先の書き方
    styles = collections.Counter()
    for m in re.finditer(r'※[^\n]{0,40}', alltext):
        s = m[0]
        styles['※詳細は、…' if s.startswith('※詳細は') else '※「…」' if s.startswith('※「') else 'その他※'] += 1
    inline = len(re.findall(r'[^※\n]{0,3}資料「[^」]+\.pptx」', alltext))
    if len([k for k in styles if styles[k]]) > 1 or (inline and styles):
        add('参考', '表記', '全体', '参照先の書き方が混在: ' + '、'.join(f'{k}×{v}' for k, v in styles.items()) + (f'、本文中に資料名×{inline}' if inline else ''),
            '', '細かい点で恐れ入りますが、資料の参照先の記載方法を統一いただけますでしょうか（前回は「※」で後置する形に統一）。')


def theme_fonts(prs):
    rt = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme'
    try:
        blob = prs.slide_master.part.part_related_by(rt).blob.decode('utf-8')
    except (KeyError, ValueError):
        return set()
    return set(re.findall(r'<a:(?:latin|ea) typeface="([^"]+)"', blob))


def theme_font_map(prs):
    rt = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme'
    try:
        blob = prs.slide_master.part.part_related_by(rt).blob.decode('utf-8')
    except (KeyError, ValueError):
        return {}
    out = {}
    for kind, short in (('major', 'mj'), ('minor', 'mn')):
        m = re.search(rf'<a:{kind}Font>.*?<a:latin typeface="([^"]*)".*?<a:ea typeface="([^"]*)"', blob, re.S)
        if m:
            out[f'+{short}-lt'], out[f'+{short}-ea'] = m[1], m[2]
    return out


def check_latin_font(prs, slides):
    """英数字に実際に使われるフォント（latin指定をテーマで解決したもの）が、本文の多数派と違う箇所を出す。
    例：大半が latin=+mn-ea（Meiryo UI）なのに一部だけ +mn-lt（Arial）だと、英数字の見た目が変わる。"""
    tmap = theme_font_map(prs)
    runs = []
    for sd in slides:
        t = sd['table']
        if not t:
            continue
        for ri, r in enumerate(t.table.rows):
            if ri == 0:
                continue
            for ci, c in enumerate(r.cells):
                for p in c.text_frame.paragraphs:
                    for run in p.runs:
                        if not re.search(r'[0-9A-Za-z#【】()]', run.text):
                            continue
                        rpr = run._r.find(A_NS + 'rPr')
                        lat = rpr.find(A_NS + 'latin') if rpr is not None else None
                        face = lat.get('typeface') if lat is not None else None
                        if face is None:
                            continue
                        runs.append((tmap.get(face, face), sd['idx'], ri, ci, run.text))
    cnt = collections.Counter(r[0] for r in runs)
    if len(cnt) < 2:
        return
    major = cnt.most_common(1)[0][0]
    seen = set()
    for face, si, ri, ci, text in runs:
        if face == major or (si, ri, ci, face) in seen:
            continue
        seen.add((si, ri, ci, face))
        sample = ''.join(x[4] for x in runs if x[1:4] == (si, ri, ci) and x[0] == face)[:40]
        col = {2: '合意形成事項', 3: 'ToDo'}.get(ci, f'列{ci}')
        add('要確認', '体裁', f'p{si} {col}', f'英数字のフォントが他と違う（{face}。本文の大半は{major}）', sample,
            '細かい点で恐れ入ります。他とフォントが異なるためご確認をお願いします。', anchor=text.strip()[:20], target=text.strip()[:20])


def check_format(prs, slides):
    W, H = prs.slide_width, prs.slide_height
    body_fonts, sizes, geo = collections.Counter(), collections.Counter(), []
    per_slide_font = collections.defaultdict(collections.Counter)
    for sd in slides:
        t = sd['table']
        if not t:
            continue
        geo.append((sd['idx'], t.left, t.width, tuple(c.width for c in t.table.columns), t.top + t.height, t.table.rows[0].height))
        if t.top + t.height > H:
            add('要修正', '体裁', f"p{sd['idx']}", '表がスライド下端からはみ出している')
        if t.left < 0 or t.left + t.width > W:
            add('要修正', '体裁', f"p{sd['idx']}", '表がスライド左右からはみ出している')
        for ri, r in enumerate(t.table.rows):
            for ci, c in enumerate(r.cells):
                for p in c.text_frame.paragraphs:
                    for run in p.runs:
                        if not run.text.strip():
                            continue
                        name = run.font.name or '(既定)'
                        size = run.font.size.pt if run.font.size else None
                        if ri > 0:
                            body_fonts[name] += len(run.text)
                            per_slide_font[sd['idx']][name] += len(run.text)
                        if size:
                            sizes[size] += len(run.text)
                        if run._r.find('.//' + A_NS + 'highlight') is not None:
                            add('要確認', '体裁', f"p{sd['idx']} 行{ri} 列{ci}", '蛍光ペン（ハイライト）が残っている', run.text[:40])
                        try:
                            ct = run.font.color.type if run.font.color else None
                            rgb = str(run.font.color.rgb) if ct == 1 else None
                        except Exception:
                            rgb = None
                        if ri > 0 and rgb and rgb not in ('000000',):
                            add('要確認', '体裁', f"p{sd['idx']} 行{ri} 列{ci}", f'本文に黒以外の文字色（#{rgb}）', run.text[:40])
    if len(sizes) > 1:
        major = sizes.most_common(1)[0][0]
        add('要確認', '体裁', '全体', '文字サイズが混在: ' + '、'.join(f'{k}pt×{v}字' for k, v in sizes.items()) + f'（多数派 {major}pt）')
    check_latin_font(prs, slides)
    if geo:
        lefts = collections.Counter(g[1] for g in geo)
        if len(lefts) > 1:
            add('参考', '体裁', '全体', '表の位置（左端）がページごとに異なる: ' + '、'.join(f"p{g[0]}={g[1] / 360000:.2f}cm" for g in geo),
                '', '細かい点で恐れ入りますが、ページをめくった際に表の位置がずれて見えるため、表の位置・幅をそろえていただけますでしょうか。')
        if len({g[5] for g in geo}) > 1:
            add('参考', '体裁', '全体', '見出し行（No./アジェンダ…）の高さがページごとに異なる: ' + '、'.join(f"p{g[0]}={g[5] / 360000:.2f}cm" for g in geo))
        todo_w = [(g[0], g[3][-1]) for g in geo if len(g[3]) >= 4]
        if todo_w and max(w for _, w in todo_w) > 2 * min(w for _, w in todo_w):
            add('参考', '体裁', '全体', '列幅（ToDo列）がページごとに大きく異なる: ' + '、'.join(f"p{s}={w / 360000:.1f}cm" for s, w in todo_w))
    # 箇条書き記号のゆれ（台帳確認の項目行）
    bullets = collections.Counter()
    examples = collections.defaultdict(list)
    for sd in slides:
        t = sd['table']
        if not t:
            continue
        for ri, r in enumerate(t.table.rows):
            if ri == 0:
                continue
            c = r.cells[2]
            for p in c.text_frame.paragraphs:
                txt = clean(''.join(run.text for run in p.runs)).strip()
                if not ITEM_RE.match(txt):
                    continue
                pPr = p._p.find(A_NS + 'pPr')
                bu = None
                if pPr is not None:
                    ch = pPr.find(A_NS + 'buChar')
                    bu = ch.get('char') if ch is not None else ('なし' if pPr.find(A_NS + 'buNone') is not None else None)
                key = (bu or '継承', p.level)
                bullets[key] += 1
                examples[key].append(f"p{sd['idx']}:{txt[:12]}")
    if len(bullets) > 1:
        major = bullets.most_common(1)[0][0]
        for k, v in bullets.items():
            if k != major:
                add('要確認', '体裁', '議題2', f'項目行の箇条書き記号・レベルが他と違う（記号{k[0]}・レベル{k[1]}、多数派は記号{major[0]}・レベル{major[1]}）',
                    '、'.join(examples[k][:6]), '細かい点で恐れ入りますが、箇条書きの記号・字下げを他の項目とそろえていただけますでしょうか。')


def check_sections(sections, meeting, ledger_todo, ledger_issue, prev_sections, materials, transcript):
    year = meeting.year
    for sec in sections:
        st = sec_type(sec['title'])
        n_items = len(sec['items'])
        for g in sec.get('groups') or []:
            if g['declared'] != g['n']:
                add('要修正', '件数', f"議題2 {sec['title']}", f"「({g['declared']}件)」と書かれているが、列挙は{g['n']}件", g['heading_line'],
                    f"細かい点で恐れ入りますが、{sec['title']}の件数が列挙と一致していないため（記載{g['declared']}件／列挙{g['n']}件）、ご確認いただけますでしょうか。",
                    anchor=g['heading_line'], target=f"{g['declared']}件")
        for it in sec['items']:
            label = f"議題2 {sec['title']} {it['kind']}{it['id']}"
            if st in ('extend', 'withdraw') and not it['background']:
                add('要修正', '記載漏れ', label, '延伸・取り下げの背景（■…背景）が書かれていない', it['raw'][:50],
                    '細かい点で恐れ入りますが、どのような合意により延伸・取り下げとなったかが読み取れないため、背景を記載いただけますでしょうか。',
                    anchor=f"{it['kind']}{it['id']}(")
            if st == 'extend' and not it['old']:
                add('要修正', '記載漏れ', label, '延伸期限（旧→新）が書かれていない', it['raw'][:50])
            if st == 'extend' and it['old'] and parse_date(it['new'], year) <= parse_date(it['old'], year):
                add('要修正', '期限', label, f"延伸後の期限が延伸前より後になっていない（{it['old']}→{it['new']}）")
            if it['kind'] != '#' or not ledger_todo:
                continue
            led = ledger_todo.get(it['id'])
            if not led:
                led_i = ledger_issue.get(it['id'])
                if not led_i:
                    add('要確認', '台帳', label, '台帳（ToDo・開発課題）に該当番号が見つからない')
                continue
            check_against_ledger(it, st, label, led, meeting, prev_sections, transcript)
            # ToDo本文 vs 台帳
            cc = changed_chars(it['body'], led['text'])
            if cc > 0:
                lvl = '要確認' if cc > 15 else '参考'
                add(lvl, '台帳', label, f'ToDo本文が台帳の記載と異なる（食い違い約{cc}字。台帳側の※注記・補足の省略は除外済み）', diff_snippets(it['body'], led['text']),
                    f"細かい点で恐れ入りますが、#{it['id']}の本文が台帳の記載と異なっているため、台帳の記載に合わせていただけますでしょうか。" if lvl == '要確認' else '',
                    anchor=f"#{it['id']}(")
        # 逐語録抜粋（背景の妥当性は人が見る）
        if transcript and st in ('extend', 'withdraw'):
            for it in sec['items']:
                ex = transcript_excerpt(transcript, it['id'])
                if ex:
                    it['transcript'] = ex


POST_ROWS = []  # 会議後の台帳更新の確認表（レポート用）
MEETING_KIND = ''  # 開発課題検討会 / 進捗連絡会議（議事要旨の表紙から判定）
ENTRY_RE = re.compile(r'(?m)^[ \t　]*(\d{4}/\d{1,2}/\d{1,2})[ \t　]*ACN[ \t　]*([^\s]*)')
PAIR_RE = re.compile(r'(\d{4}/\d{1,2}/\d{1,2})\s*→\s*(\d{4}/\d{1,2}/\d{1,2})')
ACTION_LABEL = {'extend': '延伸', 'close': '完了', 'withdraw': '取り下げ'}


def progress_entries(progress, year):
    """台帳の進捗欄を「YYYY/M/D ACN 名前」で始まる記載ごとに分ける（並び順は日付順とは限らない）。"""
    ms = list(ENTRY_RE.finditer(progress))
    out = []
    for k, m in enumerate(ms):
        text = progress[m.end():ms[k + 1].start() if k + 1 < len(ms) else len(progress)]
        pair = PAIR_RE.search(text)
        head = next((x for x in text.strip().split('\n') if x.strip()), '')  # 結果は1文目に書かれる
        act = None
        if '完了確認' in head or 'クローズとする' in head:
            act = '完了'
        elif '取り下げ' in head:
            act = '取り下げ'
        elif '延伸' in head or '期限を変更' in head:
            act = '延伸'
        out.append(dict(date=parse_date(m[1], year), who=m[2], text=text.strip(), act=act,
                        old=parse_date(pair[1], year) if pair else None, new=parse_date(pair[2], year) if pair else None,
                        result='予定' not in text.split('\n')[0]))
    return out


def transcript_dates(transcript, num, year):
    """逐語録で番号が最初に出た直後の「M月D日からM月D日」を、会議で話された延伸前→後として取る。"""
    if not transcript:
        return None
    m = re.search(rf'(?<!\d){num}(?!\d)', transcript)
    if not m:
        return None
    w = transcript[m.start():m.start() + 400]
    p = re.search(r'(\d{1,2})月(\d{1,2})日から(\d{1,2})月(\d{1,2})日', w)
    if not p:
        return None
    return dt.date(year, int(p[1]), int(p[2])), dt.date(year, int(p[3]), int(p[4]))


def check_against_ledger(it, st, label, led, meeting, prev_sections, transcript):
    """台帳は「会議前の記載＝正」。会議後の追記（会議を受けた更新）は、逐語録・議事要旨と食い違いがないかを見る。"""
    year = meeting.year
    md_ = f'{meeting.month}/{meeting.day}'
    entries = progress_entries(led['progress'], year)
    kind = MEETING_KIND or '(?:開発課題検討会|進捗連絡会議)'
    about = re.compile(rf'(?<!\d){re.escape(md_)}\s*(?:\(.\))?\s*(?:開発)?{kind}')
    # 会議前の記載：この会議に向けた「相談予定」等。会議後の追記：「M/D 会議名にて」で始まる、この会議の結果の記載
    pre = [e for e in entries if e['date'] and e['date'] < meeting and about.search(e['text'])]
    result_re = re.compile(about.pattern + r'\s*にて')
    post = [e for e in entries if e['date'] and e['date'] >= meeting and result_re.match(e['text']) and e['result']]
    latest = max((e['date'] for e in entries if e['date']), default=None)
    current = bool(post) and latest == max(e['date'] for e in post)  # 台帳の現在値がこの会議の結果を表しているか
    spoken = transcript_dates(transcript, it['id'], year) if st == 'extend' else None
    want = ACTION_LABEL.get(st)
    ldue = to_date(led['due'], year)
    row = dict(id=it['id'], minutes=want + (f"（{it['old']}→{it['new']}）" if st == 'extend' and it['old'] else ''),
               post='／'.join(f"{e['date'].month}/{e['date'].day} {e['who']}：{e['act'] or '?'}" + (f"（{fmt(e['old'])}→{fmt(e['new'])}）" if e['new'] else '') for e in post) or '（会議後の追記なし）',
               ledger=f"{led['status']}・期限{fmt(ldue)}", spoken=f"{fmt(spoken[0])}→{fmt(spoken[1])}" if spoken else '—', verdict='一致')
    # 延伸前の期限：逐語録 → 会議前の台帳記載 → 会議後の追記 → 前回議事要旨 の順に根拠とする
    if st == 'extend' and it['old']:
        old = parse_date(it['old'], year)
        pre_pair = next((e for e in sorted(pre, key=lambda e: e['date'], reverse=True) if e['old']), None)
        post_pair = next((e for e in post if e['old']), None)
        prev_new = next((parse_date(pit['new'], year) for ps in prev_sections for pit in ps['items']
                         if pit['kind'] == '#' and pit['id'] == it['id'] and pit['new']), None)
        for basis, src in ((spoken and spoken[0], '逐語録（会議での発言）'), (pre_pair and pre_pair['old'], f"台帳の会議前の記載（{pre_pair and fmt(pre_pair['date'])}）"),
                           (post_pair and post_pair['old'], '台帳の会議後の追記'), (prev_new, '前回議事要旨の延伸後期限')):
            if basis:
                if old != basis:
                    add('要確認' if src == '台帳の会議後の追記' else '要修正', '期限', label, f"延伸前の期限が {it['old']} だが、{src}では {fmt(basis)}", '',
                        f"細かい点で恐れ入りますが、#{it['id']}の延伸前の期限は{fmt(basis)}と認識しております。延伸期限を「{fmt(basis)}→{it['new']}」に修正いただけますでしょうか。",
                        anchor=f"#{it['id']}(", target=f"{it['old']}→")
                    row['verdict'] = '議事要旨の延伸前期限を修正'
                break
        # 延伸後の期限：会議での発言を正とし、議事要旨と会議後の台帳追記の両方を見る
        new = parse_date(it['new'], year)
        if spoken and new != spoken[1]:
            add('要修正', '期限', label, f"延伸後の期限 {it['new']} が、会議での発言（{fmt(spoken[1])}）と異なる", '',
                f"細かい点で恐れ入りますが、会議内では#{it['id']}の期限を{fmt(spoken[1])}とする発言があったと認識しております。ご確認いただけますでしょうか。",
                anchor=f"#{it['id']}(", target=it['new'])
            row['verdict'] = '議事要旨の延伸後期限を確認'
        truth = spoken[1] if spoken else new
        if current and ldue and truth and ldue != truth:
            add('要確認', '台帳', f"台帳#{it['id']}", f"会議後に更新された台帳の対応期限 {fmt(ldue)} が、{'会議での発言' if spoken else '議事要旨'}の延伸後期限 {fmt(truth)} と異なる", '',
                f"#{it['id']}：会議後に更新した台帳の対応期限（{fmt(ldue)}）が会議の結果（{fmt(truth)}）と異なるため、台帳側の修正要否をご確認ください。")
            row['verdict'] = '台帳（会議後の更新）を確認'
    # 会議後の追記の内容（延伸／完了／取り下げ）が会議の結果と合っているか
    for e in post:
        if e['act'] and want and e['act'] != want:
            add('要確認', '台帳', f"台帳#{it['id']}", f"会議後の台帳追記（{fmt(e['date'])} {e['who']}）は「{e['act']}」だが、議事要旨は「{want}」", e['text'][:80],
                f"#{it['id']}：{fmt(e['date'])}の追記では「{e['act']}」となっていますが、会議では「{want}」と認識しております。台帳側の追記の修正要否をご確認ください。")
            row['verdict'] = '台帳（会議後の追記）を確認'
    ok = {'close': ('完了',), 'withdraw': ('取り下げ',), 'extend': ('未着手', '対応中', '調整中')}.get(st)
    if current and ok and led['status'] not in ok:
        add('要確認', '台帳', f"台帳#{it['id']}", f"会議後の台帳のステータスが「{led['status']}」で、議事要旨の「{want}」と合わない", '',
            f"#{it['id']}：会議の結果（{want}）に対して台帳のステータスが「{led['status']}」のため、台帳側の修正要否をご確認ください。")
        row['verdict'] = '台帳（会議後の更新）を確認'
    if not post:
        row['verdict'] = '会議後の追記なし（未更新の可能性）'
    POST_ROWS.append(row)


def transcript_excerpt(transcript, num, width=170):
    pat = re.compile(rf'(?:No\.?|ナンバー|番号|#|ToDo|TODO|課題)?\s*{num}(?!\d)')
    out = []
    for m in pat.finditer(transcript):
        s = transcript[max(0, m.start() - 20):m.end() + width].replace('\n', ' ')
        out.append(s)
        if len(out) >= 2:
            break
    return out


def check_new_todos(agendas, meeting, ledger_todo):
    year = meeting.year
    md_ = f'{meeting.month}/{meeting.day}'
    todos = []
    for ag in agendas:
        ts, rest = split_todos(ag['todo'])
        for t in ts:
            t['agenda'] = f"議題{ag['no']}"
            t['slides'] = ag['slides']
        todos.extend(ts)
        if rest and not re.fullmatch(r'(?:[・\s]*(?:特になし|なし)[。]?\s*)+', rest):
            if ts:
                add('参考', '体裁', f"議題{ag['no']} ToDo", '【担当 期限】の後ろに補足文が続いている（補足なら可。独立したToDoなら担当・期限が必要）', rest[:60])
            else:
                add('要確認', '記載漏れ', f"議題{ag['no']} ToDo", 'ToDo列に【担当 期限】が付いていない記載がある', rest[:60],
                    '細かい点で恐れ入りますが、ToDoの担当・期限（【担当 M/D】）を記載いただけますでしょうか。')
    for t in todos:
        d = parse_date(t['due'], year)
        if d and d < meeting:
            add('要修正', '期限', t['agenda'], f"新規ToDoの期限 {t['due']} が会議日より前", t['body'][:50])
    if not ledger_todo:
        return todos
    born = [v for v in ledger_todo.values() if re.search(rf'(?<!\d){md_}(?!\d)', v['src'])
            and (MEETING_KIND in v['src'] or not MEETING_KIND)]
    used = set()
    for t in todos:
        sim = lambda v: max(ratio(t['body'], v['text']), bigram_sim(t['body'], v['text']))
        best = max(born, key=sim, default=None)
        if not best or sim(best) < 0.6:
            add('要確認', '台帳', t['agenda'], '議事要旨の新規ToDoに対応する台帳の行が見つからない（台帳未起票の可能性）', t['body'][:60])
            continue
        used.add(best['id'])
        t['ledger_id'] = best['id']
        ldue = to_date(best['due'], year)
        t['ledger_due'], t['ledger_org'] = ldue, best['org']
        # 会議で発生したToDoの台帳行は会議後の起票なので、食い違いはどちらが正しいか逐語録で確認する
        later = any(e['date'] and e['date'] > meeting and not re.match(rf'{re.escape(md_)}', e['text'])
                    for e in progress_entries(best['progress'], year))
        if later:
            ldue = None  # 後の会議で期限が変わっているため、この会議時点の期限としては比べない
        t['ledger_due'] = ldue
        if ldue and parse_date(t['due'], year) != ldue:
            add('要確認', '期限', f"{t['agenda']} → 台帳#{best['id']}", f"ToDo期限 {t['due']} が、会議後に起票された台帳の対応期限 {fmt(ldue)} と異なる（逐語録で確認）", t['body'][:50],
                f"細かい点で恐れ入りますが、本ToDoの期限が台帳（{fmt(ldue)}）と異なっているため、会議での合意内容をご確認いただけますでしょうか。",
                anchor=t['body'][:30], target=f"{t['due']}】")
        if norm(t['org']) not in norm(best['org']) and norm(best['org']) not in norm(t['org']):
            add('要確認', '担当', f"{t['agenda']} → 台帳#{best['id']}", f"担当 {t['org']} が、会議後に起票された台帳の対応者 {best['org']} と異なる（逐語録で確認）", t['body'][:50])
        POST_ROWS.append(dict(id=best['id'], minutes=f"新規ToDo【{t['org']} {t['due']}】", post=f"会議後の起票（{best['src']}）",
                              ledger=f"{best['org']}・期限{fmt(ldue)}", spoken='—', verdict='一致' if ldue == parse_date(t['due'], year) else '期限を確認'))
        cc = changed_chars(t['body'], best['text'])
        if cc > 0:
            add('要確認' if cc > 15 else '参考', '台帳', f"{t['agenda']} → 台帳#{best['id']}", f'ToDo本文が台帳と異なる（食い違い約{cc}字）', diff_snippets(t['body'], best['text']))
    for v in born:
        if v['id'] not in used:
            add('要確認', '台帳', f"台帳#{v['id']}", f'台帳では{md_}発生のToDoだが、議事要旨のToDo列に見当たらない', v['text'][:60])
    return todos


def check_ledger_coverage(sections, agendas, meeting, ledger_todo):
    """台帳上この会議で扱った（または期限が到来した）ToDoが議事要旨に出てくるか。"""
    if not ledger_todo:
        return
    year = meeting.year
    md_ = f'{meeting.month}/{meeting.day}'
    mentioned = {it['id'] for s in sections for it in s['items'] if it['kind'] == '#'}
    alltext = '\n'.join(a['consensus'] + a['todo'] + a['agenda'] for a in agendas)
    mentioned |= {int(x) for x in re.findall(r'(?:#|ToDo#?|TODO)\s*(\d{3,4})', alltext)}
    for v in ledger_todo.values():
        # 台帳の中だけの食い違い（進捗欄どうし・進捗欄とステータス）は指摘しない。議事要旨との突き合わせにのみ台帳を使う。
        handled_here = re.search(rf'(?<!\d){md_}\s*(開発課題検討会|進捗連絡会議)にて', v['progress'][:200])
        due = to_date(v['due'], year)
        overdue_open = (due and due <= meeting and v['status'] in ('未着手', '対応中', '調整中')
                        and ('開発課題' in v['dest'] or '進捗連絡' in v['dest']))
        if v['id'] in mentioned:
            continue
        if handled_here:
            add('要修正', '記載漏れ', f"台帳#{v['id']}", f"台帳の進捗欄では{md_}の会議で扱った（{v['status']}）とあるが、議事要旨に出てこない", v['text'][:60],
                f"細かい点で恐れ入りますが、#{v['id']}は台帳上{md_}に{v['status']}となっていますが、議事要旨に記載がないため、追記をご検討いただけますでしょうか。")
        elif overdue_open:
            add('要確認', '記載漏れ', f"台帳#{v['id']}", f"期限 {fmt(due)} が到来済み・未完了だが、議事要旨で扱われていない（クローズ・延伸漏れの可能性）", v['text'][:60])


def check_prev_todos(prev_agendas, meeting, ledger_todo, mentioned_ids):
    """前回議事要旨のToDoのうち期限が今回会議日までのものが、今回扱われているか。"""
    year = meeting.year
    for ag in prev_agendas:
        ts, _ = split_todos(ag['todo'])
        for t in ts:
            d = parse_date(t['due'], year)
            if not d or d > meeting:
                continue
            lid = None
            if ledger_todo:
                sim = lambda v: max(ratio(t['body'], v['text']), bigram_sim(t['body'], v['text']))
                best = max(ledger_todo.values(), key=sim)
                if sim(best) >= 0.6:
                    lid = best['id']
            if lid and lid not in mentioned_ids:
                st = ledger_todo[lid]['status']
                add('要確認' if st != '完了' or d == meeting else '参考', '前回ToDo', f"前回 議題{ag['no']} → 台帳#{lid}",
                    f"前回ToDo（期限{t['due']}・台帳ステータス「{st}」）が今回の議事要旨に出てこない", t['body'][:60])
            elif not lid:
                add('参考', '前回ToDo', f"前回 議題{ag['no']}", f"前回ToDo（期限{t['due']}）を台帳と対応付けられなかった", t['body'][:60])


def check_references(agendas, materials):
    names = {m['name']: m for m in materials}
    names.update({m['stem']: m for m in materials})
    alltext = [(a['no'], a['consensus'] + '\n' + a['todo']) for a in agendas]
    pat = re.compile(r'「?([^「」\s※]+?)(?:\.pptx)?」?\s*の?\s*(?:P|p|スライド)\s*(\d+)(?:\s*[～~\-－]\s*(\d+))?')
    for no, txt in alltext:
        for m in pat.finditer(txt):
            fname = re.sub(r'^(詳細は|詳細については)[、,]?', '', m[1].strip('「」 、,'))
            mat = names.get(fname) or names.get(fname + '.pptx')
            if not mat:
                cand = [x for x in materials if norm(fname) in norm(x['stem']) or norm(x['stem']) in norm(fname)]
                if not cand:
                    d = re.findall(r'\d{8}', fname)
                    cand = [x for x in materials if d and d[0] in x['stem'] and ('厚労省' in fname) == ('厚労省' in x['stem'])]
                    if len(cand) == 1:
                        add('参考', '参照', f'議題{no}', f"参照先の資料名「{fname}」が実際のファイル名「{cand[0]['name']}」と異なる", m[0],
                            '細かい点で恐れ入りますが、参照先の資料名を実際のファイル名に合わせていただけますでしょうか。')
                mat = cand[0] if len(cand) == 1 else None
            if not mat:
                if re.search(r'\d{8}', fname) or 'pptx' in m[0]:
                    add('参考', '参照', f'議題{no}', f'参照先資料「{fname}」が今回渡された資料にないため照合できない', m[0])
                continue
            a_, b_ = int(m[2]), int(m[3] or m[2])
            n = len(mat['pages'])
            if b_ > n or a_ < 1:
                add('要修正', '参照', f'議題{no}', f"参照ページ P{a_}-{b_} が資料「{mat['name']}」のページ数（{n}）を超える", m[0])
                continue
            titles = ' / '.join(f"P{i}「{page_title(mat['pages'][i - 1])}」" for i in sorted({a_, b_}))
            add('参考', '参照', f'議題{no}', f"参照先ページの見出し（目視確認用）: {mat['name']} {titles}", m[0])


def check_agendas(agendas, materials):
    if not materials:
        return
    planned = []
    for mat in materials:
        for page in mat['pages'][:3]:
            for tb in page['tables']:
                if tb and len(tb[0]) > 1 and 'アジェンダ' in tb[0][1]:
                    planned += [(r[0].strip(), r[1].strip()) for r in tb[1:] if r[1].strip()]
        if '厚労省' in mat['name'] and mat['pages']:
            for t in mat['pages'][0]['texts']:
                for line in t.split('\n'):
                    m = re.match(r'^\s*[0-9０-９]+[．.]\s*(.+)$', line)
                    if m:
                        planned.append(('厚労省', m[1].strip()))
                    else:
                        for part in re.split(r'(?=[0-9０-９][．.]\s)', line):
                            m2 = re.match(r'^\s*[0-9０-９]+[．.]\s*(.+)$', part)
                            if m2:
                                planned.append(('厚労省', m2[1].strip()))
    planned = [p for p in planned if p[1] != '厚労省アジェンダ']
    used = set()
    for no, title in planned:
        best = max(agendas, key=lambda a: ratio(title, re.sub(r'^厚労省アジェンダ_', '', a['agenda'])), default=None)
        if not best:
            continue
        r_ = ratio(title, re.sub(r'^厚労省アジェンダ_', '', best['agenda']))
        if r_ < 0.5:
            add('要確認', 'アジェンダ', f'資料のアジェンダ「{title}」', '議事要旨に対応する議題が見つからない')
            continue
        used.add(best['agenda'])
        if r_ < 1.0:
            add('参考', 'アジェンダ', f"議題{best['no']}", f'議題名が会議資料と異なる（一致率{r_:.0%}）', diff_snippets(title, re.sub(r'^厚労省アジェンダ_', '', best['agenda'])))
    for a in agendas:
        if a['agenda'] not in used and a['agenda'] not in ('ラップアップ',):
            add('参考', 'アジェンダ', f"議題{a['no']}", '会議資料のアジェンダ一覧に対応する項目が見つからない', a['agenda'][:40])


def page_kind(page):
    lead = ' '.join(t for t in page['texts'] if len(t) < 120)
    if '取り下げ' in lead:
        return 'withdraw'
    if re.search(r'案件[^。]*を?更新|に更新', lead) and '更新確認' not in lead:
        return 'update'
    if '新規起票' in lead or re.search(r'区分.*新規', lead):
        return 'new'
    return None


def collect_material_projects(materials):
    """会議資料から 案件ID→案件名 を集める。表（ID・案件名列）とスライド見出し（案件ID999_名称）を別々に持つ。"""
    table_ids, title_ids, listed = collections.defaultdict(list), {}, collections.defaultdict(set)
    for mat in materials:
        for pno, page in enumerate(mat['pages'], 1):
            for t in page['texts']:
                for m in re.finditer(r'案件ID\s*[＃#]?\s*(\d{3})\s*[_＿]\s*([^\n]+)', t):
                    title_ids.setdefault(int(m[1]), (m[2].strip(), f"{mat['name']} P{pno}"))
            kind = page_kind(page)
            for tb in page['tables']:
                hdr = [norm(h) for h in tb[0]] if tb else []
                ci = next((i for i, h in enumerate(hdr) if h in ('ID', '案件ID')), None)
                cn = next((i for i, h in enumerate(hdr) if '案件名' in h), None)
                ck = next((i for i, h in enumerate(hdr) if h == '区分'), None)
                if ci is None or cn is None:
                    continue
                for r in tb[1:]:
                    try:
                        pid = int(re.sub(r'\D', '', r[ci]))
                    except (ValueError, IndexError):
                        continue
                    table_ids[pid].append((r[cn].strip(), f"{mat['name']} P{pno}"))
                    k = kind or ('new' if ck is not None and '新規' in r[ck] else None)
                    if k:
                        listed[k].add(pid)
    return table_ids, title_ids, listed


def check_new_projects(sections, materials):
    """議題2の案件（新規起票・取り下げ）を会議資料と突き合わせる。"""
    table_ids, title_ids, listed = collect_material_projects(materials)
    if not table_ids and not title_ids:
        return
    by_type = collections.defaultdict(dict)
    for s in sections:
        for it in s['items']:
            if it['kind'] == 'ID':
                by_type[sec_type(s['title'])][it['id']] = it['body']
    minutes_names = {pid: body for d in by_type.values() for pid, body in d.items()}
    # 同じIDに別の案件名が付いていないか（資料どうし・資料と議事要旨）
    for pid, entries in table_ids.items():
        if pid not in minutes_names:
            continue  # 議事要旨に出てこない案件の、資料内だけの食い違いは指摘しない
        names = entries + ([(title_ids[pid][0], title_ids[pid][1])] if pid in title_ids else [])
        if pid in minutes_names:
            names.append((minutes_names[pid], '議事要旨'))
        for (n1, w1), (n2, w2) in ((a, b) for i, a in enumerate(names) for b in names[i + 1:]):
            if ratio(n1, n2) < 0.4 and bigram_sim(n1, n2) < 0.4:
                other = [tid for tid, (tn, _) in title_ids.items() if tid != pid and ratio(n1 if w2 == '議事要旨' else n2, tn) >= 0.7]
                if other:
                    add('要確認', '案件ID', f'ID{pid}', f'同じID{pid}が別の案件に使われている（資料の見出しでは「{n2 if w2 != "議事要旨" else n1}」に近い案件がID{other[0]}）',
                        f'{w1}「{n1}」／{w2}「{n2}」',
                        f'細かい点で恐れ入りますが、ID{pid}が「{n1[:25]}」と「{n2[:25]}」の両方に使われているように見受けられます。案件IDの採番をご確認いただけますでしょうか。',
                        anchor=f'ID{pid}(', target=f'ID{pid}')
                else:
                    add('要確認', '案件名', f'ID{pid}', f'ID{pid}の案件名が資料・議事要旨の間で一致していない（正式名称の確認が必要）', f'{w1}「{n1}」／{w2}「{n2}」',
                        f'細かい点で恐れ入りますが、ID{pid}の案件名が資料内で「{n1[:25]}」「{n2[:25]}」と異なっているため、案件台帳上の正式名称をご確認いただけますでしょうか。',
                        anchor=f'ID{pid}(')
                break
    for pid, body in by_type['new'].items():
        cands = list(table_ids.get(pid, [])) + ([title_ids[pid]] if pid in title_ids else [])
        ref = max(cands, key=lambda c: ratio(body, c[0])) if cands else None
        if not ref:
            add('参考', '案件名', f'議題2 新規起票 ID{pid}', '会議資料に同じIDの案件が見つからない（案件台帳で要確認）', body[:40])
            continue
        r_ = ratio(body, ref[0])
        if r_ < 0.9:
            add('要確認', '案件名', f'議題2 新規起票 ID{pid}', f'案件名が会議資料と異なる（一致率{r_:.0%}）',
                f'議事要旨「{body}」／資料（{ref[1]}）「{ref[0]}」',
                f'細かい点で恐れ入りますが、ID{pid}の案件名が会議資料の名称（{ref[0]}）と異なるため、案件台帳の正式名称に合わせていただけますでしょうか。',
                anchor=f'ID{pid}(', target=body)
        elif r_ < 1.0:
            add('参考', '案件名', f'議題2 新規起票 ID{pid}', '案件名が会議資料とわずかに異なる', diff_snippets(body, ref[0]))
    labels = {'new': '新規起票', 'withdraw': '取り下げ', 'update': '案件の更新（記載変更）'}
    by_type['update'] = by_type['change']
    for k, ids in listed.items():
        missing = sorted(ids - set(by_type[k]))
        if missing:
            names = '、'.join(f"ID{p}（{((table_ids.get(p) or [title_ids.get(p, ('',))])[0])[0][:20]}）" for p in missing)
            add('要確認', '記載漏れ', f'会議資料の{labels[k]}一覧', f'会議資料では{labels[k]}とされているが、議事要旨の{labels[k]}に含まれていない（{len(missing)}件）', names,
                f'細かい点で恐れ入りますが、会議資料で{labels[k]}となっている{len(missing)}件（{names[:80]}…）が議事要旨の{labels[k]}に含まれていないため、記載要否をご確認いただけますでしょうか。',
                anchor=next((s['heading_line'] for s in sections if sec_type(s['title']) == ('change' if k == 'update' else k) and s['heading_line']), None))


def flag_ledger_due(best, due, meeting):
    """逐語録の期限と、会議後に起票された台帳の期限が食い違う場合は台帳側も確認対象にする。"""
    d = parse_date(due, meeting.year) if re.search(r'\d', due) else None
    if best and d and best.get('ledger_id') and best.get('ledger_due') and best['ledger_due'] != d:
        add('要確認', '台帳', f"台帳#{best['ledger_id']}", f"会議後に起票された台帳の対応期限 {fmt(best['ledger_due'])} が、逐語録上の期限 {fmt(d)} と異なる", best['body'][:50],
            f"#{best['ledger_id']}：会議では期限{fmt(d)}の発言があったため、台帳の対応期限（{fmt(best['ledger_due'])}）の修正要否をご確認ください。")
        for r in POST_ROWS:
            if r['id'] == best['ledger_id']:
                r['spoken'], r['verdict'] = fmt(d), '議事要旨・台帳とも期限を確認'


def d_notes_ok(due, best, meeting):
    """逐語録側に具体的な期限があり、最も近いToDoの期限と食い違うか。"""
    d = parse_date(due, meeting.year) if re.search(r'\d', due) else None
    return bool(d and d != parse_date(best['due'], meeting.year))


def check_against_notes(todos, notes_path, meeting):
    """逐語録から作った議事録（★宿題）と議事要旨のToDoを突き合わせる。"""
    text = Path(notes_path).read_text(encoding='utf-8')
    hw = re.findall(r'【★宿題】([^【\n]+)【担当：([^、】]+)、期限：([^】]+)】', text)
    for body, who, due in hw:
        best = max(todos, key=lambda t: bigram_sim(body, t['body'] + t['note']), default=None)
        sim = bigram_sim(body, best['body'] + best['note']) if best else 0
        if sim < 0.35:
            near = f"／最も近いToDo（{best['agenda']}・【{best['org']} {best['due']}】）: {best['body'][:50]}" if best and sim >= 0.1 else ''
            if near and d_notes_ok(due, best, meeting):
                cmt = (f"細かい点で恐れ入りますが、会議内では期限を{due.strip()[:12]}とする発言があったと認識しております。"
                       f"{best['agenda']}のToDo（【{best['org']} {best['due']}】）の内容・期限が会議での合意と一致しているか、ご確認いただけますでしょうか。")
            else:
                cmt = '細かい点で恐れ入りますが、会議内で「' + body.strip()[:40] + '…」の依頼があったと認識しておりますが、ToDoへの記載要否をご確認いただけますでしょうか。'
            flag_ledger_due(best, due, meeting)
            add('要確認', '逐語録', '議事要旨ToDo列', f'逐語録では依頼・宿題の発言があるが、議事要旨のToDoに見当たらない（担当：{who}、期限：{due}）', body.strip()[:80] + near, cmt,
                anchor=best['body'][:30] if near else None, target=f"{best['due']}】" if near and d_notes_ok(due, best, meeting) else None)
            continue
        d_notes = parse_date(due, meeting.year) if re.search(r'\d', due) else None
        d_min = parse_date(best['due'], meeting.year)
        flag_ledger_due(best, due, meeting)
        if d_notes and d_min and d_notes != d_min:
            add('要確認', '逐語録', best['agenda'], f"ToDo期限 {best['due']} が逐語録上の期限 {due.strip()} と異なる", best['body'][:60],
                f"細かい点で恐れ入りますが、会議内では期限を{d_notes.month}/{d_notes.day}とする発言があったと認識しておりますが、{best['due']}で問題ないかご確認いただけますでしょうか。")


# ---------------------------------------------------------------- pptxへのコメント枠（ワッペン）挿入
RPR_AFTER_HIGHLIGHT = ('uLnTx', 'uLn', 'uFillTx', 'uFill', 'latin', 'ea', 'cs', 'sym', 'hlinkClick', 'hlinkMouseOver', 'rtl', 'extLst')


def cell_chars(cell):
    """セル内の文字を (run要素, run内位置) と対応づけて並べる。段落区切り・改行は空白1文字（runなし）として扱う。"""
    chars = []
    for pi, p in enumerate(cell.text_frame.paragraphs):
        if pi:
            chars.append((' ', None, None))
        for el in p._p:
            tag = el.tag.split('}')[1]
            if tag == 'r':
                t = el.find(A_NS + 't')
                for i, ch in enumerate(t.text or ''):
                    if ch != '​':
                        chars.append((ch, el, i))
            elif tag == 'br':
                chars.append((' ', None, None))
    return chars


def flat(s):
    return re.sub(r'\s', ' ', clean(s))


def locate(prs, anchor, target):
    """anchor を含むセルを探し、(スライド番号, 表図形, 行, 列, セル, ハイライト範囲) を返す。"""
    a = flat(anchor or '')
    if not a.strip():
        return None
    for si, sl in enumerate(prs.slides, 1):
        for shp in sl.shapes:
            if not shp.has_table:
                continue
            for ri, row in enumerate(shp.table.rows):
                for ci, cell in enumerate(row.cells):
                    chars = cell_chars(cell)
                    text = ''.join(c[0] for c in chars)
                    pos = text.find(a)
                    if pos < 0:
                        continue
                    rng = None
                    if target:
                        t = flat(target)
                        tp = text.find(t, pos)
                        if tp >= 0:
                            rng = (tp, tp + len(t))
                    return dict(slide=si, shape=shp, ri=ri, ci=ci, cell=cell, chars=chars, pos=pos, rng=rng, n=len(text))
    return None


def highlight(chars, rng):
    """chars[rng] に当たる run を「前・該当・後」に分割し、該当部分に黄色の蛍光ペンを付ける（書式は複製して保つ）。"""
    import copy
    from lxml import etree
    spans = collections.OrderedDict()
    for ch, run, i in chars[rng[0]:rng[1]]:
        if run is not None:
            spans.setdefault(run, []).append(i)
    for run, idx in spans.items():
        t = run.find(A_NS + 't').text
        a, b = min(idx), max(idx) + 1
        pre, mid, post = t[:a], t[a:b], t[b:]
        template = copy.deepcopy(run)
        if pre:
            r = copy.deepcopy(template)
            r.find(A_NS + 't').text = pre
            run.addprevious(r)
        if post:
            r = copy.deepcopy(template)
            r.find(A_NS + 't').text = post
            run.addnext(r)
        run.find(A_NS + 't').text = mid
        rpr = run.find(A_NS + 'rPr')
        if rpr is None:
            rpr = etree.Element(A_NS + 'rPr')
            run.insert(0, rpr)
        old = rpr.find(A_NS + 'highlight')
        if old is not None:
            rpr.remove(old)
        h = etree.Element(A_NS + 'highlight')
        etree.SubElement(h, A_NS + 'srgbClr', val='FFFF00')
        nxt = next((c for c in rpr if c.tag.split('}')[1] in RPR_AFTER_HIGHLIGHT), None)
        if nxt is not None:
            nxt.addprevious(h)
        else:
            rpr.append(h)


LINE_H = 143000      # 9pt本文1行の高さ（実測：PowerPoint描画で約11.3pt）
CHAR_W_FULL = 114300  # 9pt全角1文字の幅


def text_lines(text, width, indent=228600):
    """セル幅 width（EMU）に text を流し込んだときの行数の見積もり。indent は箇条書きの字下げ分。"""
    usable = max(CHAR_W_FULL, width - 2 * 91440 - indent)
    n = 0
    for seg in text.split('\n'):
        w = sum(CHAR_W_FULL * (0.55 if ord(c) < 0x2000 else 1.0) for c in seg)
        n += max(1, -(-int(w) // usable))
    return n


def cell_text_lines(cell, width):
    segs = []
    for p in cell.text_frame.paragraphs:
        cur = ''
        for el in p._p:
            tag = el.tag.split('}')[1]
            if tag == 'r':
                cur += (el.find(A_NS + 't').text or '').replace('\u200b', '')
            elif tag == 'br':
                segs.append(cur)
                cur = ''
        segs.append(cur)
    return segs


def y_of(shp, ri, ci, char_pos):
    """表の ri 行 ci 列で、セル内 char_pos 文字目がある縦位置（EMU）を見積もる。"""
    cols = [c.width for c in shp.table.columns]
    y = shp.top
    for k, row in enumerate(shp.table.rows):
        heights = [text_lines('\n'.join(cell_text_lines(c, cols[j])), cols[j], 228600 if j >= 2 else 0) * LINE_H + 91440
                   for j, c in enumerate(row.cells)]
        h = max([row.height] + heights)
        if k == ri:
            segs = cell_text_lines(row.cells[ci], cols[ci])
            before, used = [], 0
            for seg in segs:
                if used + len(seg) + 1 > char_pos:
                    before.append(seg[:max(0, char_pos - used)])
                    break
                before.append(seg)
                used += len(seg) + 1
            lines = text_lines('\n'.join(before), cols[ci]) if before else 1
            return int(y + 45720 + (lines - 0.5) * LINE_H), h
        y += h
    return y, 0


def annotate(minutes_path, out_path, items):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt
    prs = Presentation(minutes_path)
    W, H = prs.slide_width, prs.slide_height
    box_w, gap = 2700000, 60000
    per_slide = collections.defaultdict(list)
    unplaced = []
    for no, f in items:
        loc = locate(prs, f['anchor'], f['target'])
        if not loc:
            unplaced.append((no, f))
            continue
        if loc['rng']:
            highlight(loc['chars'], loc['rng'])
        shp = loc['shape']
        y_t, _ = y_of(shp, loc['ri'], loc['ci'], loc['rng'][0] if loc['rng'] else loc['pos'])
        cols = list(shp.table.columns)
        x_t = shp.left + sum(c.width for c in cols[:loc['ci']]) + min(cols[loc['ci']].width - 91440, 400000)
        if loc['rng']:
            pre = ''.join(c[0] for c in loc['chars'][:loc['rng'][0]]).split(' ')[-1]
            x_t = shp.left + sum(c.width for c in cols[:loc['ci']]) + 91440 + 228600 + int(len(pre) * CHAR_W_FULL * 0.9) % max(1, cols[loc['ci']].width - 400000)
        lines = sum(max(1, -(-len(x) // 22)) for x in (f'【内部向けコメント】（No.{no}）', f['comment']))
        h = int(lines * LINE_H + 2 * 54000 + 30000)
        per_slide[loc['slide']].append(dict(no=no, f=f, y_t=y_t, x_t=x_t, h=h))
    for si, boxes in per_slide.items():
        boxes.sort(key=lambda b: b['y_t'])
        y = 120000
        for b in boxes:
            b['y'] = max(y, b['y_t'] - b['h'] // 2)
            y = b['y'] + b['h'] + gap
        bottom = H - 60000
        for b in reversed(boxes):
            if b['y'] + b['h'] > bottom:
                b['y'] = bottom - b['h']
            bottom = b['y'] - gap
        sl = prs.slides[si - 1]
        for b in boxes:
            box = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, W + 150000, b['y'], box_w, b['h'])
            box.name = f"ACNコメント_No{b['no']}"
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(0xE4, 0xDC, 0xF0)
            box.line.color.rgb = RGBColor(0x7F, 0x6F, 0x9F)
            box.line.width = Pt(0.75)
            box.shadow.inherit = False
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
                setattr(tf, side, Emu(54000))
            for k, txt in enumerate((f"【内部向けコメント】（No.{b['no']}）", b['f']['comment'])):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.alignment = PP_ALIGN.LEFT
                r = para.add_run()
                r.text = txt
                r.font.size = Pt(9)
                r.font.bold = k == 0
                r.font.color.rgb = RGBColor(0, 0, 0)
            ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, W + 150000, b['y'] + min(b['h'] // 2, 200000), b['x_t'], b['y_t'])
            ln.name = f"ACNコメント線_No{b['no']}"
            ln.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
            ln.line.width = Pt(0.75)
    prs.save(out_path)
    return unplaced


# ---------------------------------------------------------------- 出力
def write_report(out, args, sections, todos, transcript_used):
    fs = findings
    cnt = collections.Counter(f['level'] for f in fs)
    L = [f"# 議事要旨チェック結果｜{Path(args.minutes).name}", '',
         f"> 実行日: {dt.date.today().isoformat()}　会議日: {args.date}",
         f"> 照合した資料: 会議資料 {len(args.materials or [])}件 / 台帳 {'あり' if args.ledger else 'なし'} / 前回議事要旨 {'あり' if args.prev else 'なし'} / 逐語録 {'あり' if transcript_used else 'なし'}",
         '> 生成: `_基盤/_ツール/議事要旨チェック/check_gijiyoshi.py`（機械チェック。指摘は候補であり、採否は人が判断する）', '',
         '## サマリー', '', f"- 要修正 {cnt['要修正']}件 / 要確認 {cnt['要確認']}件 / 参考 {cnt['参考']}件", '']
    L += ['### 台帳確認（議題2）の構成', '', '| 節 | 記載件数 | 列挙件数 | 項目 |', '|---|---|---|---|']
    for s in sections:
        L.append(f"| {s['title']} | {s['declared'] if s['declared'] is not None else '-'} | {len(s['items'])} | {'、'.join(i['kind'] + str(i['id']) for i in s['items'])} |")
    L += ['', f'### 新規ToDo（ToDo列）{len(todos)}件', '', '| 議題 | 担当 | 期限 | 台帳# | 内容 |', '|---|---|---|---|---|']
    for t in todos:
        L.append(f"| {t['agenda']} | {t['org']} | {t['due']} | {t.get('ledger_id', '—')} | {md(t['body'][:70])} |")
    if POST_ROWS:
        L += ['', '### 会議後の台帳更新の確認', '',
              '台帳は会議前の記載を正とし、会議後の追記・起票（会議を受けた更新）だけを議事要旨・逐語録と突き合わせた。', '',
              '| # | 議事要旨 | 台帳の会議後の追記 | 台帳の現在値 | 逐語録（期限） | 判定 |', '|---|---|---|---|---|---|']
        for r in sorted(POST_ROWS, key=lambda r: (r['verdict'] == '一致', r['id'])):
            L.append(f"| {r['id']} | {md(r['minutes'])} | {md(r['post'])} | {md(r['ledger'])} | {r['spoken']} | {'✅ 一致' if r['verdict'] == '一致' else '⚠️ ' + r['verdict']} |")
    L += ['', '## 指摘一覧', '', '| No | 重要度 | 種別 | 場所 | 指摘 | 根拠・該当箇所 |', '|---|---|---|---|---|---|']
    for i, f in enumerate(fs, 1):
        L.append(f"| {i} | {f['level']} | {f['kind']} | {md(f['where'])} | {md(f['msg'])} | {md(f['basis'])} |")
    L += ['', '## コメント案（要修正・要確認のうち文案があるもの）', '',
          '議事要旨(案)の右余白に貼る「【内部向けコメント】」の下書き。採用するものだけ使う。', '']
    for i, f in enumerate(fs, 1):
        if f['comment'] and f['level'] != '参考':
            internal = f['where'].startswith('台帳#')
            head = '> 【ACN内・台帳の修正確認】（議事要旨には貼らない）  ' if internal else '> 【内部向けコメント】  '
            mark = '　※pptxに自動で貼れなかったため手で貼る' if f.get('unplaced') else ''
            L += [f"**No.{i}（{md(f['where'])}）**{mark}", '', head, f"> {md(f['comment'])}", '']
    if transcript_used:
        L += ['## 延伸・取り下げの背景と逐語録の該当箇所（人が確認）', '',
              '背景が会議での実際のやり取りと合っているかは自動判定しない。逐語録で該当番号が出てくる箇所を並べる。', '']
        for s in sections:
            if sec_type(s['title']) not in ('extend', 'withdraw'):
                continue
            for it in s['items']:
                L.append(f"- **#{it['id']}**（{s['title']}）背景：{md(it['background'] or '（なし）')}")
                for ex in it.get('transcript', []) or ['（逐語録に番号の言及が見つからない）']:
                    L.append(f'  - 逐語録：…{md(ex)}…')
        L.append('')
    if any(t.get('excerpt') for t in todos):
        L += ['## 新規ToDoと逐語録の近い箇所（会議で依頼があったかを人が確認）', '',
              '議事要旨のToDoが会議での依頼・合意に基づくかは自動判定しない（語句の一致では区別できないため）。逐語録で最も近い箇所を並べる。', '']
        for t in todos:
            L += [f"- **{t['agenda']}【{t['org']} {t['due']}】** {md(t['body'][:60])}…", f"  - 逐語録：…{md(t.get('excerpt', ''))}…"]
        L.append('')
    L += ['## このチェックでやっていないこと', '',
          '- 文意の分かりやすさ・主語の明確さの判定（長文・重複語句などの機械的な候補のみ）',
          '- 延伸・取り下げ背景が会議の実態と合っているかの判定（逐語録の該当箇所を並べるまで）',
          '- 案件台帳（案件ID・案件名）の原本との照合（会議資料に載っている名称との照合のみ）',
          '- 実際の表示上のはみ出し（行の高さは保存値で判定。PowerPointで開いての目視は別途）', '']
    Path(out).write_text('\n'.join(L), encoding='utf-8')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--minutes', required=True)
    ap.add_argument('--date', required=True)
    ap.add_argument('--materials', nargs='*', default=[])
    ap.add_argument('--ledger')
    ap.add_argument('--prev')
    ap.add_argument('--notes')
    ap.add_argument('--transcript')
    ap.add_argument('--out', required=True)
    ap.add_argument('--annotate', help='コメント枠と黄色ハイライトを入れたpptxの出力先（元ファイルは変更しない）')
    args = ap.parse_args()
    meeting = dt.date.fromisoformat(args.date)

    prs, slides, rows = load_minutes(args.minutes)
    global MEETING_KIND
    cover = ' '.join(clean(s.text_frame.text) for s in prs.slides[0].shapes if s.has_text_frame)
    MEETING_KIND = '進捗連絡会議' if '進捗連絡' in cover else '開発課題検討会' if '開発課題検討会' in cover else ''
    agendas = merge_agendas(rows)
    ledger_ag = next((a for a in agendas if '台帳' in a['agenda']), None)
    sections = parse_ledger_sections(ledger_ag['consensus'], slides) if ledger_ag else []

    prev_agendas, prev_sections, prows = [], [], []
    if args.prev:
        _, _, prows = load_minutes(args.prev)
        prev_agendas = merge_agendas(prows)
        pl = next((a for a in prev_agendas if '台帳' in a['agenda']), None)
        prev_sections = parse_ledger_sections(pl['consensus'], []) if pl else []
    ledger_todo, ledger_issue = load_ledger(args.ledger) if args.ledger else ({}, {})
    materials = [load_material(m) for m in args.materials]
    transcript = Path(args.transcript).read_text(encoding='utf-8') if args.transcript else ''

    check_structure(slides, rows)
    check_text_style(rows)
    check_format(prs, slides)
    check_sections(sections, meeting, ledger_todo, ledger_issue, prev_sections, materials, transcript)
    todos = check_new_todos(agendas, meeting, ledger_todo)
    check_ledger_coverage(sections, agendas, meeting, ledger_todo)
    mentioned = {it['id'] for s in sections for it in s['items'] if it['kind'] == '#'} | {t.get('ledger_id') for t in todos}
    if prev_agendas:
        check_prev_todos(prev_agendas, meeting, ledger_todo, mentioned)
        blank = lambda rs: collections.Counter(t for r in rs for t in (r['consensus'].strip(), r['todo'].strip()) if re.fullmatch(r'特になし[。]?', t))
        cur_b, prev_b = blank(rows), blank(prows)
        if cur_b and prev_b and cur_b.most_common(1)[0][0] != prev_b.most_common(1)[0][0]:
            add('参考', '表記', '全体', f"「特になし」の書き方が前回と異なる（前回「{prev_b.most_common(1)[0][0]}」→今回「{cur_b.most_common(1)[0][0]}」）")
    check_references(agendas, materials)
    check_agendas(agendas, materials)
    check_new_projects(sections, materials)
    if args.notes:
        check_against_notes(todos, args.notes, meeting)
    order = {'要修正': 0, '要確認': 1, '参考': 2}
    findings.sort(key=lambda f: (order[f['level']], f['kind']))
    unplaced = []
    if args.annotate:
        items = [(i, f) for i, f in enumerate(findings, 1)
                 if f['comment'] and f['level'] != '参考' and not f['where'].startswith('台帳#')]
        unplaced = annotate(args.minutes, args.annotate, items)
        for no, f in unplaced:
            f['unplaced'] = True
        print(f"{args.annotate}: コメント枠 {len(items) - len(unplaced)}件（貼れなかったもの {len(unplaced)}件: " + '、'.join(f'No.{n}' for n, _ in unplaced) + '）')
    if transcript:
        for t in todos:
            best = max(range(0, max(1, len(transcript) - 300), 60), key=lambda i: bigram_sim(t['body'], transcript[i:i + 300]))
            t['excerpt'] = re.sub(r'\s+', ' ', transcript[best:best + 220])
    write_report(args.out, args, sections, todos, bool(transcript))
    c = collections.Counter(f['level'] for f in findings)
    print(f"{args.out}: 要修正 {c['要修正']} / 要確認 {c['要確認']} / 参考 {c['参考']}")


if __name__ == '__main__':
    main()
