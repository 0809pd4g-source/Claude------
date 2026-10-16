"""
artifacts.py — 「実物」と「流れ」を見せる部品と、人が手で足す注釈表現。

参考資料（人が作ったアクセンチュアの資料 31 本）の図形データで多かったのに、compose に無かった表現を足す。
  実物:   M-CODE（コード窓）/ M-SCREEN-MOCK（画面＋番号注釈）/ M-PROMPTBOX（入力→出力）/ M-CHAT-MOCK / M-FILETREE
  流れ:   M-TIMELINE（1 本の時間軸に発火点）/ M-PIPELINE（工程のレールに成果物とゲート）/ M-ACTORS（人物と関係線）
  列:     M-NUMCOL（大きな番号＋縦罫線＋見出し＋本文）
  注釈:   V-ANNOT（引き出し線付きの短い言い切り）/ V-HILITE（領域に重ねる強調枠）/ V-LEGEND（凡例）
寸法の既定値は実測（grammar/motifs.json）に合わせる: 線 0.75〜1pt、要素の間隔 0.1〜0.2in、番号は楕円。
data の形は references/compose-api.md を参照。
"""

import json
import os

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as S
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import helpers as _h

MONO = "Consolas"
_MOTIFS = None


def motif(name, key, default):
    """grammar/motifs.json の実測寸法（中央値）を返す。無ければ default。"""
    global _MOTIFS
    if _MOTIFS is None:
        p = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "grammar", "motifs.json")
        try:
            _MOTIFS = json.load(open(p, encoding="utf-8")).get("params", {})
        except Exception:
            _MOTIFS = {}
    return _MOTIFS.get(name, {}).get(key, default)


def _c():
    import compose
    return compose


def _leftover(ctx, name, used, h):
    """部品が使った高さと枠の高さを比べ、下に大きな余りがあれば知らせる（中身を足すか枠を詰める）。"""
    if h > 0 and (h - used) / h > 0.35:
        ctx.warn(f"{name}: 枠の下 {int((h - used) / h * 100)}% が空く（締め・注釈・根拠を足すか、別の部品と組む）")


def _badge(slide, cx, cy, n, d=None, fill=None, color=None, size=12):
    c = _c()
    d = d or motif("number_badge", "d", 0.32)
    b = c._rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill or _h.DARK_PURPLE, shape=S.OVAL)
    c._shape_text(b, str(n), size, color or _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    return b


def _mono_box(slide, x, y, w, h, lines, size=12, color=None, numbers=False, hl=(), hl_fill=None):
    """等幅の行を 1 行ずつ描く（行番号・強調行つき）。行の高さを返す。"""
    c = _c()
    lh = size * 1.45 / 72
    n = len(lines)
    if n * lh > h:
        lh = h / max(n, 1)
    num_w = 0.38 if numbers else 0.0
    for i, ln in enumerate(lines):
        ly = y + i * lh
        if (i + 1) in hl:
            c._rect(slide, x, ly, w, lh, fill=hl_fill or _h.LIGHTEST_PURPLE)
            c._rect(slide, x, ly, 0.05, lh, fill=_h.CORE_PURPLE)
        if numbers:
            t = c._textbox(slide, x + 0.02, ly, num_w, lh, [str(i + 1)], size, _h.MID_GRAY, anchor="m",
                           align=PP_ALIGN.RIGHT, inset=(0, 0, 0.06, 0), space_after=0)
            _fontify(t, MONO)
        t = c._textbox(slide, x + num_w + 0.06, ly, w - num_w - 0.08, lh, [ln.replace(" ", " ")], size,
                       color or _h.TEXT_BODY, anchor="m", inset=(0.02, 0, 0.02, 0), space_after=0)
        t.text_frame.word_wrap = False
        _fontify(t, MONO)
    return lh


def _fontify(shape, font):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = font


def _window(ctx, r, title, kind="code"):
    """窓（タイトルバー＋本体）。本体の矩形を返す。角丸は使わない。"""
    c = _c()
    x, y, w, h = r
    bar_h = 0.34
    c._rect(ctx.slide, x, y, w, h, fill=_h.WHITE, line=_h.LIGHT_GRAY, line_w=1.0)
    bar = c._rect(ctx.slide, x, y, w, bar_h, fill=_h.OFF_WHITE if kind != "terminal" else _h.DARKEST_PURPLE)
    tcol = _h.TEXT_SUB if kind != "terminal" else _h.WHITE
    for k in range(3):
        c._rect(ctx.slide, x + 0.12 + k * 0.17, y + bar_h / 2 - 0.05, 0.10, 0.10,
                fill=_h.LIGHT_GRAY if kind != "terminal" else _h.LIGHT_PURPLE, shape=S.OVAL)
    if title:
        t = c._textbox(ctx.slide, x + 0.7, y, w - 0.8, bar_h, [title], 12, tcol, anchor="m", space_after=0)
        _fontify(t, MONO)
    c._hline(ctx.slide, x, y + bar_h, w)
    return (x + 0.08, y + bar_h + 0.08, w - 0.16, h - bar_h - 0.16)


def _side_notes(ctx, notes, anchors, nx, nw, top, bottom, size=14):
    """番号付きの注釈を右側の列に並べ、各番号から対象の行へ細い線を引く。"""
    c = _c()
    n = len(notes)
    if not n:
        return
    gap = 0.12
    hh = (bottom - top - gap * (n - 1)) / n
    for i, (note, (ax, ay)) in enumerate(zip(notes, anchors)):
        ny = top + i * (hh + gap)
        text = note if isinstance(note, str) else note.get("text", "")
        head = None if isinstance(note, str) else note.get("head")
        paras = ([(head, {"bold": True, "color": _h.DARKEST_PURPLE, "space_after": 2})] if head else []) + [text]
        sz, ok = c.fit_size(paras, nw - 0.5, hh, (size, 12), 4)
        if not ok:
            ctx.warn("注釈の文字が多い")
        nb = _badge(ctx.slide, nx + 0.17, ny + 0.17, i + 1)
        c._textbox(ctx.slide, nx + 0.42, ny, nw - 0.42, hh, paras, sz, space_after=2, emphasis=ctx.emphasis)
        ab = _badge(ctx.slide, ax, ay, i + 1, d=0.26, size=12)
        import connect          # 対象の番号 → 注釈の番号を接続付きの線で（横に並べば直線、高さが違えばカギ線）
        connect.link(ctx.slide, ab, nb, "r", "l", _h.MID_GRAY, 0.75)


# ═══════════════════════════════════════════════════════════════
#  実物
# ═══════════════════════════════════════════════════════════════
def m_code(ctx, r, d):
    """{"title": "mods/audit.ts", "lines": [...], "highlight": [3,4], "notes": [{"line": 3, "text": ".."}], "width": 0.6}"""
    x, y, w, h = r
    notes = d.get("notes") or []
    ww = w * d.get("width", 0.6 if notes else 1.0)
    lines = d.get("lines") or []
    size = 12 if len(lines) > 12 or max([len(l) for l in lines] + [0]) * 14 * 0.55 / 72 > ww - 0.7 else 14
    lh0 = min(size * 2.0 / 72, (h - 0.5) / max(len(lines), 1))
    wh = min(h, len(lines) * lh0 + 0.5)
    inner = _window(ctx, (x, y, ww, wh), d.get("title"), d.get("kind", "code"))
    lh = _mono_box(ctx.slide, *inner, lines, size, numbers=d.get("numbers", True), hl=set(d.get("highlight") or []))
    used = wh
    if notes:
        anchors = [(x + ww - 0.2, inner[1] + (n.get("line", 1) - 0.5) * lh) for n in notes]
        nh = min(h, max(wh, len(notes) * 0.95))
        _side_notes(ctx, notes, anchors, x + ww + 0.35, w - ww - 0.35, y + 0.1, y + nh - 0.1)
        used = max(used, nh)
    _leftover(ctx, "M-CODE", used, h)


def m_screen_mock(ctx, r, d):
    """{"title": "Claude Code", "kind": "terminal|app", "rows": [{"text": "> /plugin", "style": "prompt|out|tool|muted|accent"}],
        "notes": [{"row": 1, "head": "..", "text": ".."}], "width": 0.58}"""
    c = _c()
    x, y, w, h = r
    notes = d.get("notes") or []
    ww = w * d.get("width", 0.58 if notes else 1.0)
    kind = d.get("kind", "terminal")
    rows = [it if isinstance(it, dict) else {"text": str(it)} for it in d.get("rows") or []]
    n = max(len(rows), 1)
    rh = min(0.5, (h - 0.5) / n)
    wh = min(h, n * rh + 0.5)
    inner = _window(ctx, (x, y, ww, wh), d.get("title"), "terminal" if kind == "terminal" else "app")
    style_col = {"prompt": _h.DARKEST_PURPLE, "out": _h.TEXT_BODY, "tool": _h.DARK_PURPLE, "muted": _h.MID_GRAY,
                 "accent": _h.CORE_PURPLE}
    for i, it in enumerate(rows):
        ry = inner[1] + i * rh
        st = it.get("style", "out")
        if st == "tool":
            c._rect(ctx.slide, inner[0], ry + 0.03, inner[2], rh - 0.06, fill=_h.OFF_WHITE)
        t = c._textbox(ctx.slide, inner[0] + 0.05, ry, inner[2] - 0.1, rh, [it.get("text", "")], 12,
                       style_col.get(st, _h.TEXT_BODY), anchor="m", bold=st in ("prompt", "accent"), space_after=0)
        t.text_frame.word_wrap = False
        _fontify(t, MONO)
    used = wh
    if notes:
        anchors = [(x + ww - 0.2, inner[1] + (nt.get("row", 1) - 0.5) * rh) for nt in notes]
        nh = min(h, max(wh, len(notes) * 0.95))
        _side_notes(ctx, notes, anchors, x + ww + 0.35, w - ww - 0.35, y + 0.1, y + nh - 0.1)
        used = max(used, nh)
    _leftover(ctx, "M-SCREEN-MOCK", used, h)


def m_promptbox(ctx, r, d):
    """{"label_in": "入力（プロンプト）", "prompt": "...", "label_out": "出力", "output": "...", "mid": "Claude"}"""
    c = _c()
    x, y, w, h = r
    gap = 0.75
    pw = (w - gap) / 2
    need = max(c.est_height(v if isinstance(v, list) else str(v).splitlines(), pw - 0.3, 14, 4)
               for v in (d.get("prompt", ""), d.get("output", "")))
    full_h = h
    h = min(h, need + 0.75)
    for k, (lab, body, mono) in enumerate(((d.get("label_in", "入力"), d.get("prompt", ""), True),
                                          (d.get("label_out", "出力"), d.get("output", ""), False))):
        px = x + k * (pw + gap)
        tag = c._rect(ctx.slide, px, y, 1.6, 0.34, fill=_h.DARKEST_PURPLE if k == 0 else _h.DARK_PURPLE)
        c._shape_text(tag, lab, 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        c._rect(ctx.slide, px, y + 0.34, pw, h - 0.34, fill=_h.OFF_WHITE if k == 0 else _h.WHITE,
                line=_h.LIGHT_GRAY, line_w=1.0)
        paras = body if isinstance(body, list) else str(body).split("\n")
        sz, ok = c.fit_size(paras, pw - 0.3, h - 0.6, (14, 12), 4)
        if not ok:
            ctx.warn("M-PROMPTBOX: 文字が多い")
        t = c._textbox(ctx.slide, px + 0.12, y + 0.46, pw - 0.24, h - 0.56, paras, sz, space_after=4,
                       emphasis=ctx.emphasis)
        if mono:
            _fontify(t, MONO)
    c._arrow_h(ctx.slide, x + pw + 0.12, y + h / 2, gap - 0.24)
    if d.get("mid"):
        c._textbox(ctx.slide, x + pw, y + h / 2 - 0.55, gap, 0.3, [d["mid"]], 12, _h.MID_GRAY,
                   align=PP_ALIGN.CENTER, space_after=0)
    _leftover(ctx, "M-PROMPTBOX", h, full_h)


def m_chat_mock(ctx, r, d):
    """{"messages": [{"role": "user|ai", "text": ".."}], "title": "..."}"""
    c = _c()
    x, y, w, h = r
    msgs = d.get("messages") or []
    inner = _window(ctx, r, d.get("title"), "app") if d.get("title") else (x, y, w, h)
    ix, iy, iw, ih = inner
    gap = 0.14
    needs = [c.est_height([m.get("text", "")], iw * 0.7 - 0.2, 14, 2) + 0.12 for m in msgs]
    total = sum(needs) + gap * (len(msgs) - 1)
    k = min(1.0, ih / total) if total else 1
    cy = iy
    for m, nh in zip(msgs, needs):
        bh = nh * k
        user = m.get("role", "user") == "user"
        bw = iw * 0.7
        bx = ix + iw - bw if user else ix
        c._rect(ctx.slide, bx, cy, bw, bh, fill=_h.OFF_WHITE if user else _h.LIGHTEST_PURPLE)
        lab = m.get("name") or ("ユーザー" if user else "AI")
        c._textbox(ctx.slide, bx + 0.08, cy, bw - 0.16, bh, [(lab, {"bold": True, "size": 12, "space_after": 0,
                                                                     "color": _h.MID_GRAY if user else _h.DARK_PURPLE}),
                                                             m.get("text", "")], 14 if k == 1 else 12,
                   space_after=2, emphasis=ctx.emphasis)
        cy += bh + gap


def m_filetree(ctx, r, d):
    """{"items": [{"path": "skills/acnpptx/SKILL.md", "note": ".."}, ...]} … path の "/" の深さで字下げする。"""
    c = _c()
    x, y, w, h = r
    raw = [it if isinstance(it, dict) else {"path": str(it)} for it in d.get("items") or []]
    # 抜けている親フォルダを補う（skills/review/SKILL.md だけ書いても skills/review/ を出す）。
    # 「~/.claude/」のような別の根は、そのまま 1 つの根として扱う
    items, seen = [], set()
    for it in raw:
        p_ = it["path"]
        root_sep = p_.startswith("~/")
        parts = p_.rstrip("/").split("/")
        if root_sep:
            parts = ["~/" + parts[1]] + parts[2:] if len(parts) > 1 else parts
        for k in range(1, len(parts)):
            anc = "/".join(parts[:k]) + "/"
            if anc not in seen:
                seen.add(anc)
                items.append({"path": anc, "_parts": parts[:k]})
        key = "/".join(parts) + ("/" if p_.endswith("/") else "")
        if key in seen:
            for x_ in items:          # 補った親に、後から書かれた注記を付ける
                if x_["path"] == key:
                    x_.update({k_: v for k_, v in it.items() if k_ != "path"})
            continue
        seen.add(key)
        items.append({**it, "path": key, "_parts": parts})
    n = max(len(items), 1)
    rh = min(0.42, h / n)
    name_w = w * d.get("name_width", 0.45)
    prev_depth_y = {}
    for i, it in enumerate(items):
        parts = it.get("_parts") or it["path"].rstrip("/").split("/")
        depth = len(parts) - 1
        label = parts[-1] + ("/" if it["path"].endswith("/") else "")
        ry = y + i * rh
        ix = x + 0.1 + depth * 0.32
        if depth:
            px = ix - 0.2
            top = prev_depth_y.get(depth - 1, ry)
            c._connector(ctx.slide, px, top + rh * 0.5, px, ry + rh / 2, color=_h.LIGHT_GRAY, width=0.75, glyph=True)
            c._connector(ctx.slide, px, ry + rh / 2, ix - 0.04, ry + rh / 2, color=_h.LIGHT_GRAY, width=0.75, glyph=True)
        prev_depth_y[depth] = ry
        is_dir = it["path"].endswith("/")
        t = c._textbox(ctx.slide, ix, ry, name_w - (ix - x), rh, [label], 14 if rh >= 0.38 else 12,
                       _h.DARKEST_PURPLE if is_dir or it.get("emph") else _h.TEXT_BODY, anchor="m",
                       bold=is_dir or bool(it.get("emph")), space_after=0)
        t.text_frame.word_wrap = False
        _fontify(t, MONO)
        if it.get("note"):
            c._textbox(ctx.slide, x + name_w + 0.1, ry, w - name_w - 0.1, rh, [it["note"]], 12 if rh < 0.38 else 14,
                       _h.TEXT_SUB, anchor="m", space_after=0, emphasis=ctx.emphasis)
            c._connector(ctx.slide, x + name_w - 0.05, ry + rh / 2, x + name_w + 0.05, ry + rh / 2,
                         color=_h.LIGHT_GRAY, width=0.75, dash=True)


# ═══════════════════════════════════════════════════════════════
#  流れ
# ═══════════════════════════════════════════════════════════════
def m_timeline(ctx, r, d):
    """{"events": [{"label": "PreToolUse", "text": "..", "at": 0.2, "emph": true}], "phases": [{"label": "..", "from": 0, "to": 0.5}],
        "axis_label": ["開始", "終了"]}  … 1 本の太い軸の上に発火点を置き、ラベルを上下交互に置く。"""
    c = _c()
    x, y, w, h = r
    evs = d.get("events") or []
    n = max(len(evs), 1)
    phases = d.get("phases") or []
    ph_h = 0.36 if phases else 0
    full_h = h
    slot0 = (w - 0.6) / max(len(evs), 1)
    bw0 = min(slot0 * 1.9, 2.6)
    lab_need = max([c.est_height([e.get("label", "")] + ([e["text"]] if e.get("text") else []), bw0 - 0.1, 14, 2)
                    for e in evs] + [0.5])
    lab_max = min(1.35, lab_need + 0.05)
    co_h = 0.55 if any(e.get("callout") for e in evs) else 0
    need = ph_h + 0.12 + 2 * (lab_max + 0.22) + co_h
    if h > need:
        y, h = y + (h - need) / 3, need
    ay = y + ph_h + 0.12 + lab_max + 0.22
    lx0, lx1 = x + 0.3, x + w - 0.3
    for p in phases:
        px0 = lx0 + (lx1 - lx0) * p.get("from", 0)
        px1 = lx0 + (lx1 - lx0) * p.get("to", 1)
        band = c._rect(ctx.slide, px0, y, px1 - px0 - 0.04, ph_h - 0.06, fill=_h.LIGHTEST_PURPLE)
        c._connector(ctx.slide, px0, y + ph_h - 0.06, px0, ay, color=_h.LIGHT_GRAY, width=0.75, dash=True)
        c._shape_text(band, p.get("label", ""), 12, _h.DARKEST_PURPLE, True, inset=(0.05, 0, 0.05, 0))
    c._connector(ctx.slide, lx0, ay, lx1, ay, color=_h.DARKEST_PURPLE, width=motif("timeline_axis", "w", 2.25), arrow=True)
    al = d.get("axis_label")
    if al:
        c._textbox(ctx.slide, x, ay + 0.06, 1.6, 0.3, [al[0]], 12, _h.MID_GRAY, space_after=0)
        c._textbox(ctx.slide, x + w - 1.6, ay + 0.06, 1.6, 0.3, [al[-1]], 12, _h.MID_GRAY, align=PP_ALIGN.RIGHT,
                   space_after=0)
    slot = (lx1 - lx0) / n
    lab_h = lab_max
    dot = motif("timeline_dot", "d", 0.2)
    for i, e in enumerate(evs):
        ex = lx0 + (lx1 - lx0) * e["at"] if "at" in e else lx0 + slot * (i + 0.5)
        emph = e.get("emph")
        c._rect(ctx.slide, ex - dot / 2, ay - dot / 2, dot, dot, fill=_h.CORE_PURPLE if emph else _h.WHITE,
                line=_h.CORE_PURPLE if emph else _h.DARKEST_PURPLE, line_w=1.5, shape=S.OVAL)
        up = i % 2 == 0
        ty0 = ay - 0.22 - lab_h if up else ay + 0.22
        c._connector(ctx.slide, ex, ay - dot / 2 if up else ay + dot / 2, ex, ty0 + lab_h if up else ty0,
                     color=_h.LIGHT_GRAY, width=0.75)
        bw = min(slot * 1.9, 2.6)
        paras = [(e.get("label", ""), {"bold": True, "color": _h.CORE_PURPLE if emph else _h.DARKEST_PURPLE,
                                       "size": 14, "space_after": 2})]
        if e.get("text"):
            paras.append(e["text"])
        sz, ok = c.fit_size(paras, bw - 0.1, lab_h, (14, 12), 2)
        if not ok:
            ctx.warn(f"M-TIMELINE: 事象{i + 1}の説明が長い")
        # 端の事象は枠の中に収める（本文の左右の余白からはみ出さない）。寄せた側に文字もそろえる
        bx = min(max(x, ex - bw / 2), x + w - bw)
        al = PP_ALIGN.LEFT if bx > ex - bw / 2 + 0.05 else PP_ALIGN.RIGHT if bx < ex - bw / 2 - 0.05 else PP_ALIGN.CENTER
        c._textbox(ctx.slide, bx, ty0, bw, lab_h, paras, sz, anchor="b" if up else "t",
                   align=al, space_after=2, emphasis=ctx.emphasis)
        if e.get("callout"):    # その事象への短い言い切り（引き出し線付き）
            cy0 = ay + 0.22 + lab_h + 0.12
            cw_ = 2.8
            cx0 = min(max(x, ex - cw_ / 2), x + w - cw_)
            c._textbox(ctx.slide, cx0, cy0, cw_, 0.4, [e["callout"]], 16, _h.DARK_PURPLE, bold=True, align=PP_ALIGN.CENTER,
                       space_after=0)
            if up:
                c._connector(ctx.slide, ex, ay + dot / 2, ex, cy0, color=_h.DARK_PURPLE, width=1.0)


def m_pipeline(ctx, r, d):
    """{"stages": [{"label": "Plan", "artifact": "intent.md", "text": "..", "gate": "レビュー"}], "loop_back": "本番→次の intent"}
    … 工程のレール（ホームベース）に、成果物（文書形）と工程間のゲート（菱形）を載せる。"""
    c = _c()
    x, y, w, h = r
    st = d.get("stages") or []
    n = max(len(st), 1)
    has_gate = any(s_.get("gate") for s_ in st)
    gap = 0.3 if has_gate else 0.1
    sw = (w - gap * (n - 1)) / n
    band_h = 0.46
    din = 0.22 if has_gate else 0.12     # 文書形を工程の幅より内側に置き、ゲートのラベルの場所を空ける
    loop_h = 0.5 if d.get("loop_back") else 0
    doc_h = min(0.95, (h - band_h - loop_h) * 0.32)
    tneed = max([c.est_height(([s_["text"]] if s_.get("text") else []) + list(s_.get("items", [])), sw - 0.15, 14, 3)
                 for s_ in st] + [0.3])
    full_h = h
    h = min(h, band_h + 0.2 + doc_h + 0.12 + tneed + 0.1 + loop_h)
    last = []           # 工程ごとの一番下の図形（戻り線の接続先）
    for i, s in enumerate(st):
        sx = x + i * (sw + gap)
        shp = c.homeplate(ctx.slide, sx, y, sw, band_h, fill=_h.DARKEST_PURPLE if s.get("emph") else _h.DARK_PURPLE)
        c._shape_text(shp, s.get("label", ""), 14, _h.WHITE, True, inset=(0.06, 0, 0.04, 0))
        last.append(shp)
        dy = y + band_h + 0.2
        if s.get("artifact"):
            doc = c._rect(ctx.slide, sx + din, dy, sw - 2 * din, doc_h, fill=_h.WHITE, line=_h.DARK_PURPLE, line_w=1.0,
                          shape=S.FLOWCHART_DOCUMENT)
            c._shape_text(doc, s["artifact"], 12, _h.DARKEST_PURPLE, True, anchor="t", inset=(0.04, 0.08, 0.04, 0))
            _fontify(doc, MONO)
            last[-1] = doc
        ty = dy + doc_h + 0.12
        th = y + h - loop_h - ty
        if s.get("text") or s.get("items"):
            paras = [s["text"]] if s.get("text") else []
            paras += [(t, {"bullet": "•"}) for t in s.get("items", [])]
            sz, ok = c.fit_size(paras, sw - 0.15, th, (14, 12), 3)
            if not ok:
                ctx.warn(f"M-PIPELINE: 工程{i + 1}の説明が長い")
            last[-1] = c._textbox(ctx.slide, sx, ty, sw, th, paras, sz, space_after=3, emphasis=ctx.emphasis)
        if s.get("gate") and i < n - 1:
            gx = sx + sw + gap / 2
            gd = 0.34
            g = c._rect(ctx.slide, gx - gd / 2, dy + 0.02, gd, gd, fill=_h.CORE_PURPLE, shape=S.DIAMOND)
            lw_ = max(gap + 2 * din - 0.04, sum(c._char_w(ch, 12) for ch in s["gate"]) * 1.1 + 0.08)
            c._textbox(ctx.slide, gx - lw_ / 2, dy + gd + 0.04, lw_, 0.3, [s["gate"]], 12, _h.DARK_PURPLE,
                       align=PP_ALIGN.CENTER, bold=True, space_after=0, inset=(0, 0, 0, 0))
    for i, s_ in enumerate(st):     # 工程に付けた強調枠（座標を指定しなくても、その工程の列を囲む）
        if s_.get("hilite"):
            sx = x + i * (sw + gap)
            fy1 = y + h - loop_h - 0.05
            c._rect(ctx.slide, sx - 0.06, y - 0.06, sw + 0.12, fy1 - y + 0.12, fill=None, line=_h.CORE_PURPLE, line_w=2.25)
            lw_ = min(sw + 0.12, 0.3 + len(s_["hilite"]) * 0.19)
            tag = c._rect(ctx.slide, sx - 0.06, fy1 + 0.06, lw_, 0.3, fill=_h.CORE_PURPLE)
            c._shape_text(tag, s_["hilite"], 12, _h.WHITE, True, inset=(0.04, 0, 0.04, 0), wrap=False)
    if d.get("loop_back"):
        ly = y + h - loop_h / 2
        import connect      # 最後の工程の下辺 → 最初の工程の下辺の U 字（接続付き）
        connect.link(ctx.slide, last[-1], last[0], "b", "b", _h.DARK_PURPLE, 1.5, arrow=True, mid=ly)
        c._textbox(ctx.slide, x + w / 2 - 2.5, ly - 0.32, 5.0, 0.3, [d["loop_back"]], 12, _h.DARK_PURPLE,
                   align=PP_ALIGN.CENTER, bold=True, space_after=0)
    _leftover(ctx, "M-PIPELINE", h, full_h)


def _person(slide, cx, top, size, fill):
    """人物の図: 頭（楕円）＋胴（台形）。参考資料で多い「楕円＋胴体」の縦積み。"""
    c = _c()
    hd = size * 0.42
    head = c._rect(slide, cx - hd / 2, top, hd, hd, fill=fill, shape=S.OVAL)
    body = c._rect(slide, cx - size * 0.38, top + hd + size * 0.04, size * 0.76, size * 0.5, fill=fill, shape=S.TRAPEZOID)
    _person.last = body          # 線の接続先（胴）
    return top + hd + size * 0.54


def m_actors(ctx, r, d):
    """{"actors": [{"name": "人", "role": "判断点を握る", "emph": true, "icon": {"asset": ".."}}],
        "links": [{"from": 0, "to": 1, "label": "依頼"}], "layout": "row|hub"}"""
    c = _c()
    x, y, w, h = r
    acts = d.get("actors") or []
    n = max(len(acts), 1)
    layout = d.get("layout", "hub" if n >= 5 else "row")
    size = min(0.9, h * 0.22)
    pos = []
    if layout == "hub" and n >= 3:
        import math
        cx, cy = x + w / 2, y + h / 2 - size / 2
        pos.append((cx, cy - size / 2))
        rx, ry = w / 2 - 1.3, h / 2 - size - 0.35
        for i in range(1, n):
            a = -math.pi / 2 + 2 * math.pi * (i - 1) / (n - 1)
            pos.append((cx + rx * math.cos(a), cy + ry * math.sin(a) - size / 2))
    else:
        for i in range(n):
            pos.append((x + w * (i + 0.5) / n, y + h * 0.18))
    tops, figs = [], []
    for i, (a, (px, py)) in enumerate(zip(acts, pos)):
        fill = _h.DARK_PURPLE if a.get("emph") else _h.MID_GRAY
        if a.get("icon"):
            fig = ctx.slide.shapes[-1] if c._place_visual(ctx, a["icon"], (px - size / 2, py, size, size)) is not None else None
            bottom = py + size
        else:
            bottom = _person(ctx.slide, px, py, size, fill)
            fig = _person.last
        tops.append((px, py, bottom))
        figs.append(fig)
        paras = [(a.get("name", ""), {"bold": True, "color": _h.DARKEST_PURPLE if a.get("emph") else _h.TEXT_BODY,
                                      "size": 14, "space_after": 0})]
        if a.get("role"):
            paras.append((a["role"], {"size": 12, "color": _h.TEXT_SUB}))
        c._textbox(ctx.slide, px - 1.2, bottom + 0.02, 2.4, 0.62, paras, 14, align=PP_ALIGN.CENTER, space_after=0)
    for lk in d.get("links") or []:
        a, b = tops[lk["from"]], tops[lk["to"]]
        ya, yb = (a[1] + a[2]) / 2, (b[1] + b[2]) / 2
        dx = 1 if b[0] > a[0] else -1
        x1, x2 = a[0] + dx * size * 0.55, b[0] - dx * size * 0.55
        fa_, fb_ = figs[lk["from"]], figs[lk["to"]]
        col_ = _h.DARK_PURPLE if lk.get("emph") else _h.MID_GRAY
        if fa_ is not None and fb_ is not None:     # 人どうしを接続付きの線で（横に並べば直線、高さが違えばカギ線）
            import connect
            connect.link(ctx.slide, fa_, fb_, None, None, col_, 1.0, arrow=lk.get("arrow", True), dash=lk.get("dash", False))
        else:
            c._connector(ctx.slide, x1, ya, x2, yb, color=col_, width=1.0, arrow=lk.get("arrow", True), dash=lk.get("dash", False))
        if lk.get("label"):
            c._textbox(ctx.slide, (x1 + x2) / 2 - 1.0, (ya + yb) / 2 - 0.32, 2.0, 0.28, [lk["label"]], 12,
                       _h.DARK_PURPLE, align=PP_ALIGN.CENTER, bold=True, space_after=0)


def m_numcol(ctx, r, d):
    """{"items": [{"head": "..", "text": "..", "items": [..], "icon": {..}}], "focus": 1}
    … 大きな番号『01』＋左の縦罫線＋見出し＋本文（参考資料の M-NUMCOL）。"""
    c = _c()
    x, y, w, h = r
    its = d.get("items") or []
    n = max(len(its), 1)
    gap = 0.3
    cw = (w - gap * (n - 1)) / n
    focus = d.get("focus")
    def _p(it):
        return ([(it.get("head", ""), {})] + ([it["text"]] if it.get("text") else []) + list(it.get("items", [])))
    need = max([c.est_height(_p(it), cw - 0.3, 14, 4) + 0.15 for it in its] + [0.5]) + 0.75 +         (1.0 if any(it.get("icon") for it in its) else 0)
    full_h = h
    h = min(h, max(need, 1.6))
    _leftover(ctx, "M-NUMCOL", h, full_h)
    for i, it in enumerate(its):
        cx = x + i * (cw + gap)
        on = focus is None or focus == i
        col = _h.CORE_PURPLE if on else _h.LIGHT_PURPLE
        c._rect(ctx.slide, cx, y, 0.03, h, fill=col)
        c._textbox(ctx.slide, cx + 0.15, y - 0.05, cw - 0.15, 0.62, [f"{i + 1:02d}"], 36, col, bold=True, space_after=0)
        paras = [(it.get("head", ""), {"bold": True, "size": 16, "color": _h.DARKEST_PURPLE, "space_after": 6})]
        if it.get("text"):
            paras.append(it["text"])
        paras += [(t, {"bullet": "•"}) for t in it.get("items", [])]
        ih = 1.0 if it.get("icon") else 0
        sz, ok = c.fit_size(paras, cw - 0.3, h - 0.75 - ih, (14, 12), 4)
        if not ok:
            ctx.warn(f"M-NUMCOL: 列{i + 1}の文字が多い")
        c._textbox(ctx.slide, cx + 0.15, y + 0.68, cw - 0.15, h - 0.75 - ih, paras, sz, space_after=4,
                   emphasis=ctx.emphasis)
        if it.get("icon"):
            c._place_visual(ctx, it["icon"], (cx + 0.15, y + h - 0.9, 0.85, 0.85))


# ═══════════════════════════════════════════════════════════════
#  注釈（本文の上に重ねる）
# ═══════════════════════════════════════════════════════════════
def _rel(box, v):
    x, y, w, h = box
    return x + w * v[0], y + h * v[1]


def v_annot(ctx, dv, box):
    """{"text": "現有人材の可視化が一丁目一番地", "at": [0.62, 0.78], "w": 0.3, "to": [0.4, 0.6]}
    … 短い言い切り（紫の太字）と、対象への引き出し線。座標は本文領域に対する比率。"""
    c = _c()
    tx, ty = _rel(box, dv.get("at", [0.65, 0.8]))
    tw = box[2] * dv.get("w", 0.3)
    size = dv.get("size", 16)
    th = c.est_height([dv["text"]], tw, size, 0) + 0.04
    c._textbox(ctx.slide, tx, ty, tw, th, [dv["text"]], size, _h.DARK_PURPLE, bold=True, space_after=0,
               inset=(0.02, 0.0, 0.02, 0.0))
    if dv.get("to"):
        px, py = _rel(box, dv["to"])
        sx = tx if px < tx else tx + tw
        c._connector(ctx.slide, sx, ty + th / 2, px, py, color=_h.DARK_PURPLE, width=1.0)
        d = 0.09
        c._rect(ctx.slide, px - d / 2, py - d / 2, d, d, fill=_h.DARK_PURPLE, shape=S.OVAL)


def v_hilite(ctx, dv, box):
    """{"box": [0.0, 0.55, 0.45, 0.4], "label": "今回の対象", "dash": false} … 対象の領域に強調枠を重ねる。"""
    c = _c()
    rx, ry, rw, rh = dv["box"]
    x, y = _rel(box, (rx, ry))
    w, h = box[2] * rw, box[3] * rh
    c._rect(ctx.slide, x, y, w, h, fill=None, line=_h.CORE_PURPLE, line_w=2.25, dash=dv.get("dash", False))
    if dv.get("label"):
        lw = min(w, 0.25 + len(dv["label"]) * 0.2)
        ty = y - 0.3 if y - 0.3 >= box[1] - 0.02 else y      # 上に余地が無ければ枠の内側に置く
        tag = c._rect(ctx.slide, x, ty, lw, 0.3, fill=_h.CORE_PURPLE)
        c._shape_text(tag, dv["label"], 12, _h.WHITE, True, inset=(0.04, 0, 0.04, 0), wrap=False)


def v_legend(ctx, dv, box):
    """{"items": [{"fill": "purple", "label": "効率化"}, {"line": "purple", "label": "現状"}]} … 本文の右上に凡例。"""
    c = _c()
    import skeleton
    items = dv.get("items") or []
    x, y, w, h = box
    lw = sum(0.45 + len(it.get("label", "")) * 0.17 for it in items) + 0.6
    lx = x + w - lw
    ly = y - 0.36
    t = c._textbox(ctx.slide, lx, ly, 0.55, 0.3, ["凡例"], 12, _h.TEXT_BODY, bold=True, anchor="m", space_after=0)
    cx = lx + 0.6
    for it in items:
        if it.get("line"):
            c._rect(ctx.slide, cx, ly + 0.08, 0.3, 0.14, fill=None, line=skeleton.color(it["line"], "line"), line_w=1.5)
        else:
            c._rect(ctx.slide, cx, ly + 0.08, 0.3, 0.14, fill=skeleton.color(it.get("fill", "purple"), "fill"))
        c._textbox(ctx.slide, cx + 0.34, ly, len(it.get("label", "")) * 0.17 + 0.1, 0.3, [it.get("label", "")], 12,
                   _h.TEXT_BODY, anchor="m", space_after=0)
        cx += 0.45 + len(it.get("label", "")) * 0.17


def register(modules, devices):
    modules.update({
        "M-CODE": m_code, "M-SCREEN-MOCK": m_screen_mock, "M-PROMPTBOX": m_promptbox, "M-CHAT-MOCK": m_chat_mock,
        "M-FILETREE": m_filetree, "M-TIMELINE": m_timeline, "M-PIPELINE": m_pipeline, "M-ACTORS": m_actors,
        "M-ORG-ACTORS": m_actors, "M-NUMCOL": m_numcol,
    })
    devices.update({"V-ANNOT": v_annot, "V-HILITE": v_hilite, "V-LEGEND": v_legend})
