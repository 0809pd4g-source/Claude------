"""
craft_metrics.py — スライドを「図形のデータ」として読み、人が作った資料と同じ物差しで測る。

craft_check.py（AI 感リンター）と、参考資料の分析（パッケージ外）の両方が使う共通の計測ロジック。
  mine_pptx(path)        : 1 図形 1 レコード（座標は 13.333x7.5 に正規化、テーマ色は RGB と役割に解決）
  split_regions(shapes)  : タイトル / メッセージ / パンくず / フッター / 本文 に振り分ける
  slide_metrics(shapes)  : 本文の指標（図形・線・長方形以外・画像・文字・充填率・空き箱・均等な並び・文字サイズの種類 など）
"""
import collections
import colorsys
import math

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as T
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

E = 914400
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}
SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


# ── 色 ─────────────────────────────────────────────────────────
def _hex_to_rgb(h):
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def _rgb_to_hex(r, g, b):
    return "%02X%02X%02X" % tuple(max(0, min(255, round(v * 255))) for v in (r, g, b))


def apply_mods(hexc, el):
    """lumMod / lumOff / tint / shade を適用する。"""
    r, g, b = _hex_to_rgb(hexc)
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        tag = m.tag.split("}")[1]
        v = int(m.get("val", "100000")) / 100000
        if tag == "lumMod":
            l *= v
        elif tag == "lumOff":
            l += v
        elif tag == "tint":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (1 - (1 - c) * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif tag == "shade":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = max(0, min(1, l))
    return _rgb_to_hex(*colorsys.hls_to_rgb(h, l, s))


def theme_colors(master):
    try:
        th = master.part.part_related_by(RT.THEME)
        from lxml import etree
        import safexml
        root = safexml.fromstring(th.blob)
    except Exception:
        return {}
    cs = root.find(".//a:clrScheme", NS)
    out = {}
    if cs is None:
        return out
    for c in cs:
        name = c.tag.split("}")[1]
        v = c.find("a:srgbClr", NS)
        if v is not None:
            out[name] = v.get("val").upper()
        else:
            v = c.find("a:sysClr", NS)
            if v is not None:
                out[name] = (v.get("lastClr") or "000000").upper()
    return out


def color_of(holder, theme):
    """solidFill などの子要素（srgbClr / schemeClr / sysClr / prstClr）を RGB にする。"""
    if holder is None:
        return None
    for c in holder:
        tag = c.tag.split("}")[1]
        if tag == "srgbClr":
            base = c.get("val").upper()
        elif tag == "schemeClr":
            nm = SCHEME_ALIAS.get(c.get("val"), c.get("val"))
            base = theme.get(nm)
            if base is None:
                return "scheme:" + c.get("val")
        elif tag == "sysClr":
            base = (c.get("lastClr") or ("FFFFFF" if c.get("val") == "window" else "000000")).upper()
        elif tag == "prstClr":
            base = {"black": "000000", "white": "FFFFFF", "red": "FF0000"}.get(c.get("val"), "808080")
        else:
            continue
        return apply_mods(base, c)
    return None


def role(hexc):
    """RGB を色の役割に分類する（紫系 / 灰系 / 白 / 黒 / その他の色）。"""
    if not hexc or len(hexc) != 6:
        return hexc or "none"
    r, g, b = _hex_to_rgb(hexc)
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if s < 0.18 or l > 0.97 or l < 0.06:
        if l > 0.97:
            return "white"
        if l > 0.85:
            return "gray_light"
        if l > 0.55:
            return "gray_mid"
        if l > 0.2:
            return "gray_dark"
        return "black"
    hue = h * 360
    if 255 <= hue <= 305:
        return "purple_light" if l > 0.8 else "purple_dark" if l < 0.3 else "purple"
    for lo, hi, nm in ((0, 20, "red"), (20, 50, "orange"), (50, 70, "yellow"), (70, 170, "green"),
                       (170, 255, "blue"), (305, 345, "magenta"), (345, 361, "red")):
        if lo <= hue < hi:
            return nm + ("_light" if l > 0.8 else "")
    return "other"


# ── 図形 ───────────────────────────────────────────────────────
def geom(el):
    pg = el.find(".//a:prstGeom", NS)
    if pg is not None:
        return pg.get("prst")
    return "cust" if el.find(".//a:custGeom", NS) is not None else None


def style_color(el, ref, theme):
    st = el.find("p:style", NS)
    if st is None:
        return None
    r = st.find("a:" + ref, NS)
    if r is None or r.get("idx") in (None, "0"):
        return None
    return color_of(r, theme)


def fill_info(el, theme):
    sp = el.find("p:spPr", NS)
    if sp is not None:
        if sp.find("a:noFill", NS) is not None:
            return "none"
        sf = sp.find("a:solidFill", NS)
        if sf is not None:
            return color_of(sf, theme)
        if sp.find("a:gradFill", NS) is not None:
            return "grad"
        if sp.find("a:blipFill", NS) is not None:
            return "blip"
        if sp.find("a:pattFill", NS) is not None:
            return "patt"
        if sp.find("a:grpFill", NS) is not None:
            return "none"     # グループの塗りを継承（グループは普通は塗らない）
    c = style_color(el, "fillRef", theme)
    return c or "inh"


def line_info(el, theme):
    ln = el.find("p:spPr/a:ln", NS)
    color = None
    w = dash = None
    arrows = [None, None]
    if ln is not None:
        if ln.find("a:noFill", NS) is not None:
            return None
        if ln.get("w"):
            w = round(int(ln.get("w")) / 12700, 2)
        sf = ln.find("a:solidFill", NS)
        color = color_of(sf, theme) if sf is not None else None
        d = ln.find("a:prstDash", NS)
        dash = d.get("val") if d is not None else None
        for i, t in enumerate(("headEnd", "tailEnd")):
            e = ln.find("a:" + t, NS)
            if e is not None and e.get("type", "none") != "none":
                arrows[i] = e.get("type")
    if color is None:
        color = style_color(el, "lnRef", theme)
        if color is None and ln is None:
            return None
    return {"w": w or 0.75, "c": color, "role": role(color), "dash": dash or "solid",
            "head": arrows[0], "tail": arrows[1]}


def text_info(shape, theme):
    tf = shape.text_frame
    el = shape._element
    bp = el.find(".//a:bodyPr", NS)
    sizes, bold, chars, colors, aligns = {}, 0, 0, {}, {}
    bullets = 0
    paras = 0
    for p in tf.paragraphs:
        t = "".join(r.text for r in p.runs).strip()
        if not t:
            continue
        paras += 1
        ppr = p._p.find("a:pPr", NS)
        if ppr is not None and (ppr.find("a:buChar", NS) is not None or ppr.find("a:buAutoNum", NS) is not None):
            bullets += 1
        al = (ppr.get("algn") if ppr is not None else None) or "inh"
        aligns[al] = aligns.get(al, 0) + len(t)
        for r in p.runs:
            n = len(r.text.strip())
            if not n:
                continue
            chars += n
            sz = r.font.size.pt if r.font.size else None
            sizes[str(round(sz)) if sz else "inh"] = sizes.get(str(round(sz)) if sz else "inh", 0) + n
            if r.font.bold:
                bold += n
            rpr = r._r.find("a:rPr", NS)
            sf = rpr.find("a:solidFill", NS) if rpr is not None else None
            c = role(color_of(sf, theme)) if sf is not None else "inh"
            colors[c] = colors.get(c, 0) + n
    if not chars:
        return None
    return {"chars": chars, "paras": paras, "bullets": bullets, "sizes": sizes, "bold": bold,
            "colors": colors, "align": max(aligns, key=aligns.get) if aligns else "inh",
            "anchor": (bp.get("anchor") if bp is not None else None) or "inh",
            "vert": (bp.get("vert") if bp is not None else None), "head": tf.text.strip()[:40]}


def kind_of(s, el, g, fill, line):
    st = s.shape_type
    if st == T.GROUP:
        return "group"
    if st == T.PICTURE or el.tag.endswith("}pic"):
        return "pic"
    if getattr(s, "has_table", False) and s.has_table:
        return "table"
    if getattr(s, "has_chart", False) and s.has_chart:
        return "chart"
    if el.tag.endswith("graphicFrame"):
        return "frame"
    if s.is_placeholder:
        return "ph"
    if el.tag.endswith("cxnSp") or g in ("line", "straightConnector1") or (g and "Connector" in g):
        return "line"
    has_text = getattr(s, "has_text_frame", False) and s.has_text_frame and s.text_frame.text.strip()
    if st == T.TEXT_BOX or (g == "rect" and fill in ("none", "inh") and line is None and has_text):
        return "tb"
    return "shape"


def mine_pptx(path, deck="deck", theme_cache=None):
    theme_cache = {} if theme_cache is None else theme_cache
    prs = Presentation(path)
    sx, sy = 13.333 / (prs.slide_width / E), 7.5 / (prs.slide_height / E)
    shapes_out, slides_out = [], []
    n = len(prs.slides)
    for si, sl in enumerate(prs.slides, 1):
        master = sl.slide_layout.slide_master
        key = id(master)
        if key not in theme_cache:
            theme_cache[key] = theme_colors(master)
        theme = theme_cache[key]
        z = [0]

        def walk(shapes, parent, depth, off):
            for s in shapes:
                el = s._element
                try:
                    x, y, w, h = ((s.left or 0) / E, (s.top or 0) / E, (s.width or 0) / E, (s.height or 0) / E)
                except Exception:
                    continue
                # グループ内の座標は子座標系（chOff/chExt）を親の枠に写像する
                if off is not None:
                    gx, gy, gw, gh, cx, cy, cw, ch = off
                    fx, fy = (gw / cw if cw else 1), (gh / ch if ch else 1)
                    x, y, w, h = gx + (x - cx) * fx, gy + (y - cy) * fy, w * fx, h * fy
                g = geom(el)
                fill = fill_info(el, theme) if not el.tag.endswith(("graphicFrame", "grpSp")) else None
                line = line_info(el, theme) if not el.tag.endswith(("graphicFrame", "grpSp", "pic")) else None
                k = kind_of(s, el, g, fill, line)
                z[0] += 1
                rec = {"deck": deck, "slide": si, "id": s.shape_id, "z": z[0], "parent": parent, "depth": depth,
                       "k": k, "g": g, "x": round(x * sx, 3), "y": round(y * sy, 3), "w": round(w * sx, 3),
                       "h": round(h * sy, 3), "rot": round(getattr(s, "rotation", 0) or 0, 1)}
                if fill is not None:
                    rec["fill"] = fill
                    rec["fill_role"] = role(fill) if fill not in ("none", "inh", "grad", "blip", "patt") else fill
                if line:
                    rec["line"] = line
                if s.is_placeholder:
                    try:
                        rec["ph"] = str(s.placeholder_format.type).split(".")[-1].split(" ")[0]
                        rec["ph_idx"] = s.placeholder_format.idx
                    except Exception:
                        pass
                if k == "line":
                    nv = el.find(".//p:nvCxnSpPr/p:cNvCxnSpPr", NS)
                    if nv is not None:
                        st_, en_ = nv.find("a:stCxn", NS), nv.find("a:endCxn", NS)
                        rec["cxn"] = [int(st_.get("id")) if st_ is not None else None,
                                      int(en_.get("id")) if en_ is not None else None]
                    xf = el.find(".//a:xfrm", NS)
                    if xf is not None:
                        rec["flip"] = [xf.get("flipH") == "1", xf.get("flipV") == "1"]
                if getattr(s, "has_text_frame", False) and s.has_text_frame:
                    ti = text_info(s, theme)
                    if ti:
                        rec["text"] = ti
                if k == "table":
                    tbl = s.table
                    rec["table"] = {"rows": len(tbl.rows), "cols": len(tbl.columns)}
                    chars = sum(len(c.text) for r in tbl.rows for c in r.cells)
                    rec["table"]["chars"] = chars
                if k == "chart":
                    try:
                        rec["chart"] = str(s.chart.chart_type).split(".")[-1].split(" ")[0]
                    except Exception:
                        rec["chart"] = "?"
                if k == "pic":
                    rec["pic_px"] = None
                    try:
                        rec["pic_px"] = list(s.image.size)
                    except Exception:
                        pass
                shapes_out.append(rec)
                if k == "group":
                    xf = el.find("p:grpSpPr/a:xfrm", NS)
                    sub = off
                    if xf is not None and xf.find("a:chOff", NS) is not None:
                        o, e_, co, ce = (xf.find("a:off", NS), xf.find("a:ext", NS),
                                         xf.find("a:chOff", NS), xf.find("a:chExt", NS))
                        vals = [int(v) / E for v in (o.get("x"), o.get("y"), e_.get("cx"), e_.get("cy"),
                                                     co.get("x"), co.get("y"), ce.get("cx"), ce.get("cy"))]
                        if off is not None:   # 入れ子のグループ: 親の写像をさらに掛ける
                            gx, gy, gw, gh, cx, cy, cw, ch = off
                            fx, fy = (gw / cw if cw else 1), (gh / ch if ch else 1)
                            vals[0], vals[1] = gx + (vals[0] - cx) * fx, gy + (vals[1] - cy) * fy
                            vals[2], vals[3] = vals[2] * fx, vals[3] * fy
                        sub = tuple(vals)
                    walk(s.shapes, s.shape_id, depth + 1, sub)

        walk(sl.shapes, None, 0, None)
        hidden = sl._element.get("show") == "0"
        slides_out.append({"deck": deck, "slide": si, "n": n, "pos": round(si / n, 3), "hidden": hidden,
                           "layout": sl.slide_layout.name})
    return shapes_out, slides_out


def ink(png_path):
    """描画画像のインク量（白でない画素の割合）: (本文全体, 本文の下 1/3, インクのある行の割合)。
    図形の枠が下まで伸びていても、中身（文字・図・塗り）が少なければ小さくなる。"""
    from PIL import Image
    im = np.asarray(Image.open(png_path).convert("L").resize((480, 270)), dtype=np.float32)
    y0, y1 = int(270 * 1.4 / 7.5), int(270 * 6.9 / 7.5)
    b = im[y0:y1, int(480 * 0.03):int(480 * 0.97)] < 235
    third = b.shape[0] // 3
    return float(b.mean()), float(b[2 * third:].mean()), float((b.mean(axis=1) > 0.02).mean())


# ── 領域と指標 ────────────────────────────────────────────────
BODY_Y0, BODY_Y1, BODY_X0, BODY_X1 = 1.25, 7.05, 0.25, 13.08
FOOTER_PH = {"FOOTER", "SLIDE_NUMBER", "DATE"}
TITLE_PH = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}


def txt(r):
    return r.get("text") or {}


def chars(r):
    return txt(r).get("chars", 0) + (r.get("table") or {}).get("chars", 0)


def main_size(r):
    s = txt(r).get("sizes") or {}
    s = {k: v for k, v in s.items() if k != "inh"}
    return int(max(s, key=s.get)) if s else None


# ── 領域の振り分け ────────────────────────────────────────────
def split_regions(rs):
    """title / lead（メッセージ）/ crumb（パンくず）/ footer / body に分ける。"""
    out = collections.defaultdict(list)
    leaves = [r for r in rs if r["k"] != "group"]
    for r in leaves:
        cy = r["y"] + r["h"] / 2
        ph = r.get("ph")
        if ph in FOOTER_PH or r["y"] > 7.0 or (r["k"] == "pic" and r["y"] > 6.7 and r["w"] < 1.5):
            out["footer"].append(r)
        elif ph in TITLE_PH:
            out["title"].append(r)
        elif r["y"] < 0.55 and r["h"] < 0.6 and chars(r) and r["w"] > 1:
            out["crumb"].append(r)
        elif 0.55 <= r["y"] < 1.75 and r["w"] > 6 and chars(r) >= 15 and r["k"] in ("ph", "tb", "shape") \
                and not out["lead"]:
            out["lead"].append(r)
        elif BODY_Y0 - 0.15 <= cy <= BODY_Y1 and r["w"] > 0.01 or r["k"] == "line":
            if r["k"] == "pic" and r["w"] < 1.3 and r["h"] < 0.6 and (r["y"] < 0.5 or r["y"] > 6.6):
                out["footer"].append(r)   # ロゴ
            else:
                out["body"].append(r)
        else:
            out["other"].append(r)
    return out


# ── 指標 ──────────────────────────────────────────────────────
def coverage(body):
    gx, gy = 64, 32
    m = np.zeros((gy, gx), bool)
    for r in body:
        if r["k"] == "line":
            continue
        i0 = int(max(0, (r["x"] - BODY_X0) / (BODY_X1 - BODY_X0) * gx))
        i1 = int(min(gx, math.ceil((r["x"] + r["w"] - BODY_X0) / (BODY_X1 - BODY_X0) * gx)))
        j0 = int(max(0, (r["y"] - BODY_Y0) / (BODY_Y1 - BODY_Y0) * gy))
        j1 = int(min(gy, math.ceil((r["y"] + r["h"] - BODY_Y0) / (BODY_Y1 - BODY_Y0) * gy)))
        if i1 > i0 and j1 > j0:
            m[j0:j1, i0:i1] = True
    return float(m.mean())


def text_area(r):
    """文字が占める面積の概算（in²）。全角 1 字 ≒ (pt/72)² × 1.15（行間込み）。"""
    t = txt(r)
    if not t:
        return 0.0
    area = 0.0
    for k, n in (t.get("sizes") or {}).items():
        pt = 14 if k == "inh" else int(k)
        area += n * (pt / 72) ** 2 * 1.15
    return area


def is_container(r):
    return (r["k"] == "shape" and r.get("fill_role") not in (None, "none", "inh", "white") and
            r["w"] * r["h"] >= 0.6) or (r["k"] == "shape" and r.get("line") and r["w"] * r["h"] >= 1.0)


def find_arrays(body):
    """同じ大きさの図形が 3 つ以上、行または列に揃って並ぶものを返す。"""
    cand = [r for r in body if r["k"] in ("shape", "pic") and r["w"] > 0.25 and r["h"] > 0.18]
    arrays = []
    for axis in ("row", "col"):
        groups = collections.defaultdict(list)
        for r in cand:
            key = (r["k"], r.get("g"), round(r["w"] / 0.08), round(r["h"] / 0.08),
                   round(r["y"] / 0.08) if axis == "row" else round(r["x"] / 0.08))
            groups[key].append(r)
        for v in groups.values():
            if len(v) < 3:
                continue
            v = sorted(v, key=lambda r: r["x"] if axis == "row" else r["y"])
            gaps = [(b["x"] - a["x"] - a["w"]) if axis == "row" else (b["y"] - a["y"] - a["h"]) for a, b in zip(v, v[1:])]
            if min(gaps) < -0.02 or (max(gaps) - min(gaps)) > 0.15:
                continue
            fills = collections.Counter(r.get("fill") for r in v)
            focus = len(fills) == 2 and min(fills.values()) == 1
            arrays.append({"axis": axis, "n": len(v), "w": v[0]["w"], "h": v[0]["h"], "gap": float(np.median(gaps)),
                           "g": v[0].get("g"), "k": v[0]["k"], "focus": focus, "fill_role": v[0].get("fill_role"),
                           "with_text": sum(1 for r in v if chars(r)) / len(v), "members": [r["id"] for r in v]})
    return arrays


def geom_role(r, body):
    """長方形以外の図形が何の役をしているか（形と大きさと文字から推定）。"""
    g, w, h, t = r.get("g"), r["w"], r["h"], txt(r)
    head = (t.get("head") or "").strip()
    digits = head.replace(".", "").replace("0", "0").isdigit() or (len(head) <= 3 and any(c in head for c in "①②③④⑤⑥⑦⑧⑨"))
    if g == "ellipse":
        if w < 0.2 and h < 0.2:
            return "dot"
        if digits and w < 0.8:
            return "number_badge"
        below = [b for b in body if b is not r and abs((b["x"] + b["w"] / 2) - (r["x"] + r["w"] / 2)) < r["w"] * 0.4
                 and 0 <= b["y"] - (r["y"] + r["h"]) < 0.12 and b.get("g") in ("cust", "triangle", "flowChartDelay",
                                                                               "round2SameRect", "roundRect", "trapezoid", "rect")
                 and b["w"] < r["w"] * 3]
        if below and w < 1.2 and not head:
            return "person_head"
        if head and w >= 0.8:
            return "circle_label"
        if not head and w >= 1.5:
            return "venn_or_ring"
        return "ellipse_other"
    if g in ("homePlate", "chevron"):
        return "step" if head else "step_blank"
    if g in ("rightArrow", "downArrow", "leftArrow", "upArrow", "notchedRightArrow", "upDownArrow",
             "leftRightArrow", "blockArc", "curvedRightArrow", "curvedDownArrow", "circularArrow"):
        return "block_arrow"
    if g in ("triangle", "rtTriangle"):
        return "triangle_marker" if w < 0.5 else "triangle_big"
    if g in ("leftBracket", "rightBracket", "leftBrace", "rightBrace", "bracketPair", "bracePair"):
        return "bracket"
    if g and ("Callout" in g or g in ("wedgeRectCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout")):
        return "callout"
    if g in ("flowChartDocument", "foldedCorner", "flowChartMultidocument"):
        return "document"
    if g in ("can", "flowChartMagneticDisk"):
        return "database"
    if g == "cust":
        return "freeform"
    if g in ("roundRect", "round2SameRect", "round1Rect", "flowChartAlternateProcess"):
        return "rounded_box"
    if g in ("diamond", "flowChartDecision"):
        return "decision"
    if g in ("arc", "donut", "pie", "chord"):
        return "arc_ring"
    if g in ("mathPlus", "mathMultiply", "mathEqual", "mathMinus"):
        return "math_symbol"
    if g in ("star5", "star4", "irregularSeal1", "irregularSeal2"):
        return "star_burst"
    return "other:" + str(g)


def line_kind(r):
    w, h = r["w"], r["h"]
    o = "h" if h < 0.03 else "v" if w < 0.03 else "diag"
    ln = r.get("line") or {}
    return o, bool(ln.get("head") or ln.get("tail")), ln.get("dash", "solid") != "solid", ln.get("w", 0.75), \
        ln.get("role"), bool(r.get("cxn") and any(r["cxn"]))


def slide_metrics(rs):
    reg = split_regions(rs)
    body = reg["body"]
    lead = reg["lead"][0] if reg["lead"] else None
    m = {}
    m["objects"] = len(body)
    m["lines"] = sum(r["k"] == "line" for r in body)
    shapes = [r for r in body if r["k"] == "shape"]
    m["autoshapes"] = len(shapes)
    m["nonrect"] = sum(r.get("g") not in ("rect", None) for r in shapes)
    m["ellipse"] = sum(r.get("g") == "ellipse" for r in shapes)
    pics = [r for r in body if r["k"] == "pic"]
    m["pics"] = len(pics)
    m["icons"] = sum(r["w"] < 1.3 and r["h"] < 1.3 for r in pics)
    m["shots"] = sum(r["w"] * r["h"] > 4 for r in pics)
    m["tables"] = sum(r["k"] == "table" for r in body)
    m["charts"] = sum(r["k"] == "chart" for r in body)
    m["groups"] = sum(r["k"] == "group" and r["depth"] == 0 for r in rs)
    m["chars"] = sum(chars(r) for r in body)
    m["lead_chars"] = chars(lead) if lead else 0
    m["lead_size"] = main_size(lead) if lead else None
    m["lead_bold"] = (txt(lead).get("bold", 0) / max(1, txt(lead).get("chars", 1))) if lead else None
    m["lead_color"] = (max(txt(lead).get("colors", {"inh": 1}), key=txt(lead).get("colors", {"inh": 1}).get)
                       if lead else None)
    m["title_chars"] = sum(chars(r) for r in reg["title"])
    m["crumb"] = bool(reg["crumb"])
    m["coverage"] = coverage(body)
    conts = [r for r in body if is_container(r)]
    fills = [min(1.5, text_area(r) / (r["w"] * r["h"])) for r in conts if r["w"] * r["h"] > 0 and chars(r) and r["h"] >= 0.45]
    m["containers"] = len(conts)
    m["container_fill_med"] = float(np.median(fills)) if fills else None
    # 空き箱: 文字の入った箱なのに、中がほとんど空いている（文字の無い背景・枠・ガントのマスは数えない）
    m["empty_boxes"] = sum(1 for r in conts if chars(r) and r["h"] >= 0.5 and r["w"] * r["h"] >= 1.5 and text_area(r) / (r["w"] * r["h"]) < 0.12
                          and not any(o is not r and o["x"] >= r["x"] - 0.02 and o["y"] >= r["y"] - 0.02 and
                                      o["x"] + o["w"] <= r["x"] + r["w"] + 0.02 and o["y"] + o["h"] <= r["y"] + r["h"] + 0.02
                                      for o in body))
    sizes = collections.Counter()
    bold = tot = 0
    for r in body:
        t = txt(r)
        for k, n in (t.get("sizes") or {}).items():
            if k != "inh":
                sizes[int(k)] += n
        bold += t.get("bold", 0)
        tot += t.get("chars", 0)
    m["size_kinds"] = sum(1 for k, v in sizes.items() if v >= 0.05 * max(1, sum(sizes.values())))
    m["sizes"] = dict(sizes)
    m["bold_ratio"] = bold / tot if tot else None
    area = collections.Counter()
    for r in shapes:
        fr = r.get("fill_role")
        if fr and fr not in ("none", "inh"):
            fam = fr.split("_")[0] if fr not in ("white", "black") else fr
            area[fam] += r["w"] * r["h"]
    body_area = (BODY_X1 - BODY_X0) * (BODY_Y1 - BODY_Y0)
    m["area"] = {k: round(v / body_area, 3) for k, v in area.items()}
    arrays = find_arrays(body)
    m["arrays"] = [{k: a[k] for k in ("axis", "n", "w", "h", "gap", "g", "k", "focus", "fill_role", "with_text")}
                   for a in arrays]
    # 均等カード: 文字の入った長方形が横に並ぶもの（工程の矢羽根・積層・表の行は数えない）
    m["array_equal_cards"] = max([a["n"] for a in arrays if a["k"] == "shape" and a["g"] in ("rect", None)
                                  and a["axis"] == "row" and a["with_text"] >= 0.6 and a["w"] * a["h"] >= 1.0] or [0])
    m["geom_roles"] = dict(collections.Counter(geom_role(r, body) for r in shapes if r.get("g") not in ("rect", None)))
    lk = [line_kind(r) for r in body if r["k"] == "line"]
    m["line_kinds"] = {"h": sum(x[0] == "h" for x in lk), "v": sum(x[0] == "v" for x in lk),
                       "diag": sum(x[0] == "diag" for x in lk), "arrow": sum(x[1] for x in lk),
                       "dash": sum(x[2] for x in lk), "connected": sum(x[5] for x in lk)}
    lefts = sorted({round(r["x"] / 0.05) for r in body if r["k"] != "line"})
    m["distinct_lefts"] = len(lefts)
    if body:
        xs0 = min(r["x"] for r in body)
        xs1 = max(r["x"] + r["w"] for r in body)
        ys1 = max(r["y"] + r["h"] for r in body)
        m["margin_l"], m["margin_r"], m["body_bottom"] = round(xs0, 2), round(13.333 - xs1, 2), round(ys1, 2)
        largest = max((r["w"] * r["h"] for r in body if r["k"] != "line"), default=0)
        m["largest_share"] = round(largest / body_area, 3)
    # 帯（見出し帯・結論帯）
    m["header_bands"] = sum(1 for r in shapes if chars(r) and 0.22 <= r["h"] <= 0.6 and r["w"] >= 1.2 and
                            r.get("fill_role") in ("purple", "purple_dark", "black", "gray_dark", "blue"))
    # 下端の締めの帯: 本文幅いっぱいの「塗った」帯に文字があるもの（出典の文字だけの枠は数えない）
    m["bottom_band"] = any(r["w"] >= 8 and r["y"] >= 5.4 and chars(r) and r["h"] <= 1.0 and r["k"] == "shape"
                           and r.get("fill_role") not in (None, "none", "inh", "white") for r in body)
    m["stamp"] = any(r["k"] in ("shape", "tb") and r["y"] < 1.9 and r["x"] > 9.5 and r["w"] < 3 and
                     any(s in (txt(r).get("head") or "") for s in ("イメージ", "Image", "案", "暫定", "参考", "例", "Draft", "仮"))
                     for r in rs)
    return m


