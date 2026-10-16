"""
diagrams.py — 旧パターン N / Z / AA / AC / AE / AG / AK / AL / AM の部品を compose の部品として描く。

  M-PROFILE  人物カード（顔写真か人物の図＋名前＋役割＋経歴）。quote を付けるとインタビューカード（旧 AM）
  M-ROADMAP  成熟度の段階（階段状のステージ＋現在地・目標の印）（旧 Z）
  M-2AXIS    2 軸ポジショニング（2×2 の象限＋対象の点）（旧 AA）
  M-PYRAMID  積み上げピラミッド（台形の段＋右の説明）（旧 AC）
  M-VENN     ベン図（2〜4 円の重なり＋中央の統合概念）（旧 AE）
  M-ARCH     システム構成図（層ごとの要素＋要素間の線）（旧 AG）
  M-TILEMAP  タイルマップ（列ごとのタイル・該当を塗る）。layout="canvas" でビジネスモデルキャンバス（旧 AL）、
             layout="calendar" でカレンダー（旧 AK）
人の資料の作法に合わせる: 面は塗りすぎない／細い線（0.75〜1pt）／強調は 1 か所／箱は中身の大きさ。
data の形は references/compose-api.md を参照。
"""

import math

from lxml import etree
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as S
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

import helpers as _h


def _c():
    import compose
    return compose


def _alpha(shape, pct):
    """塗りの透明度（0〜100%）。重なりを見せるベン図用。"""
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf[0]
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    al = etree.SubElement(clr, qn("a:alpha"))
    al.set("val", str(int((100 - pct) * 1000)))


def _person(slide, cx, top, size, fill):
    import artifacts
    return artifacts._person(slide, cx, top, size, fill)


# ═══════════════════════════════════════════════════════════════
#  M-PROFILE（旧 N・AM）
# ═══════════════════════════════════════════════════════════════
def m_profile(ctx, r, d):
    """{"people": [{"name": "..", "role": "..", "org": "..", "items": [".."], "photo": {"image": path} | {"asset": id},
                    "emph": true}], "cols": 3, "quote": "..（1 人のときのインタビューカード）"}"""
    c = _c()
    x, y, w, h = r
    people = d.get("people") or ([d] if d.get("name") else [])
    n = max(len(people), 1)
    quote = d.get("quote") or (people[0].get("quote") if n == 1 and people else None)
    if n == 1 and quote:
        return _interview(ctx, r, people[0] if people else d, quote)
    cols = d.get("cols") or (n if n <= 4 else 3 if n in (5, 6, 9) else 4)
    rows = math.ceil(n / cols)
    gap_x, gap_y = 0.25, 0.22
    cw = (w - gap_x * (cols - 1)) / cols
    ch = (h - gap_y * (rows - 1)) / rows
    pic = min(1.1, ch * 0.42, cw * 0.36)
    for i, p in enumerate(people):
        cx = x + (i % cols) * (cw + gap_x)
        cy = y + (i // cols) * (ch + gap_y)
        on = p.get("emph")
        c._rect(ctx.slide, cx, cy, cw, 0.04, fill=_h.CORE_PURPLE if on else _h.DARK_PURPLE)
        if p.get("photo"):
            c._place_visual(ctx, p["photo"], (cx, cy + 0.14, pic, pic))
        else:
            _person(ctx.slide, cx + pic / 2, cy + 0.16, pic * 0.9, _h.DARK_PURPLE if on else _h.MID_GRAY)
        tx = cx + pic + 0.15
        head = [(p.get("name", ""), {"bold": True, "size": 16, "color": _h.DARKEST_PURPLE, "space_after": 2})]
        if p.get("role"):
            head.append((p["role"], {"size": 12, "color": _h.DARK_PURPLE, "bold": True, "space_after": 2}))
        if p.get("org"):
            head.append((p["org"], {"size": 12, "color": _h.TEXT_SUB, "space_after": 2}))
        c._textbox(ctx.slide, tx, cy + 0.12, cw - pic - 0.15, pic + 0.1, head, 14, anchor="m", space_after=2)
        items = p.get("items") or ([p["text"]] if p.get("text") else [])
        if items:
            paras = [(t, {"bullet": "•"}) for t in items]
            by = cy + 0.14 + pic + 0.12
            bh = ch - (by - cy)
            sz, ok = c.fit_size(paras, cw - 0.3, bh, (14, 12), 3)
            if not ok:
                ctx.warn(f"M-PROFILE: {p.get('name', '')} の経歴が長い")
            c._textbox(ctx.slide, cx, by, cw, bh, paras, sz, space_after=3, emphasis=ctx.emphasis)


def _interview(ctx, r, p, quote):
    """インタビューカード（旧 AM）: 左に人物、右に大きな引用と要点。"""
    c = _c()
    x, y, w, h = r
    lw = min(3.2, w * 0.3)
    pic = min(1.6, lw * 0.6)
    if p.get("photo"):
        c._place_visual(ctx, p["photo"], (x + (lw - pic) / 2, y + 0.1, pic, pic))
    else:
        _person(ctx.slide, x + lw / 2, y + 0.12, pic * 0.9, _h.DARK_PURPLE)
    info = [(p.get("name", ""), {"bold": True, "size": 16, "color": _h.DARKEST_PURPLE, "space_after": 2, "align": PP_ALIGN.CENTER})]
    for k in ("role", "org"):
        if p.get(k):
            info.append((p[k], {"size": 12, "color": _h.TEXT_SUB, "align": PP_ALIGN.CENTER, "space_after": 2}))
    c._textbox(ctx.slide, x, y + pic + 0.25, lw, 1.0, info, 14, space_after=2)
    qx = x + lw + 0.35
    qw = w - lw - 0.35
    c._rect(ctx.slide, qx, y + 0.1, 0.06, min(h - 0.2, 1.8), fill=_h.CORE_PURPLE)
    qs, _ = c.fit_size([quote], qw - 0.4, 1.8, (28, 18, 16), 4)
    c._textbox(ctx.slide, qx + 0.25, y + 0.05, qw - 0.25, 1.9, [f"「{quote}」"], qs, _h.DARKEST_PURPLE, bold=True, anchor="m",
               space_after=0)
    items = p.get("items") or []
    if items:
        paras = [(t, {"bullet": "•"}) for t in items]
        sz, ok = c.fit_size(paras, qw - 0.3, h - 2.2, (14, 12), 4)
        c._textbox(ctx.slide, qx + 0.25, y + 2.15, qw - 0.25, h - 2.2, paras, sz, space_after=4, emphasis=ctx.emphasis)


# ═══════════════════════════════════════════════════════════════
#  M-ROADMAP（旧 Z: 成熟度モデル）
# ═══════════════════════════════════════════════════════════════
def m_roadmap(ctx, r, d):
    """{"stages": [{"label": "Lv1 個別対応", "text": "..", "items": [..]}], "current": 1, "target": 3,
        "axis": ["成熟度", "時間"]} … 右上がりの階段。現在地と目標に印を付ける。"""
    c = _c()
    x, y, w, h = r
    st = d.get("stages") or []
    n = max(len(st), 1)
    ax = d.get("axis") or []
    ox = x + (0.35 if ax else 0)
    gw = w - (ox - x)
    sw = gw / n
    base_text = 1.15          # いちばん低い段にも説明が入る高さを残す
    step_h = (h - 0.85 - base_text) / max(n - 1, 1)
    cur, tgt = d.get("current"), d.get("target")
    for i, s in enumerate(st):
        sx = ox + i * sw
        top = y + h - 0.35 - base_text - 0.42 - step_h * i
        on = i == tgt
        fill = _h.DARKEST_PURPLE if on else (_h.DARK_PURPLE if cur is not None and i <= cur else _h.LIGHTEST_PURPLE)
        tcol = _h.WHITE if (on or (cur is not None and i <= cur)) else _h.DARKEST_PURPLE
        band = c._rect(ctx.slide, sx, top, sw - 0.06, 0.42, fill=fill)
        c._shape_text(band, s.get("label", ""), 14, tcol, True, inset=(0.06, 0, 0.06, 0))
        c._hline(ctx.slide, sx, top + 0.42, sw - 0.06, _h.LIGHT_GRAY)
        c._connector(ctx.slide, sx, top + 0.42, sx, y + h - 0.35, color=_h.LIGHT_GRAY, width=0.75, dash=True)
        paras = ([s["text"]] if s.get("text") else []) + [(t, {"bullet": "•"}) for t in s.get("items", [])]
        if paras:
            bh = y + h - 0.35 - (top + 0.5)
            sz, ok = c.fit_size(paras, sw - 0.3, bh, (14, 12), 3)
            if not ok:
                ctx.warn(f"M-ROADMAP: 段階{i + 1}の説明が長い")
            c._textbox(ctx.slide, sx, top + 0.5, sw - 0.1, bh, paras, sz, space_after=3, emphasis=ctx.emphasis)
        for mark, idx, col in (("現在", cur, _h.MID_GRAY), ("目標", tgt, _h.CORE_PURPLE)):
            if idx == i:
                tag = c._rect(ctx.slide, sx + sw - 0.86, top - 0.34, 0.8, 0.28, fill=col)
                c._shape_text(tag, mark, 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    c._connector(ctx.slide, ox, y + h - 0.3, x + w, y + h - 0.3, color=_h.DARKEST_PURPLE, width=1.5, arrow=True)
    if ax:
        c._connector(ctx.slide, x + 0.15, y + h - 0.3, x + 0.15, y, color=_h.DARKEST_PURPLE, width=1.5, arrow=True)
        t = c._textbox(ctx.slide, x - 0.15, y + h / 2 - 0.6, 0.3, 1.2, [ax[0]], 12, _h.MID_GRAY, space_after=0)
        t.text_frame._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
        if len(ax) > 1:
            c._textbox(ctx.slide, x + w - 1.6, y + h - 0.28, 1.6, 0.28, [ax[1]], 12, _h.MID_GRAY, align=PP_ALIGN.RIGHT,
                       space_after=0)


# ═══════════════════════════════════════════════════════════════
#  M-2AXIS（旧 AA）
# ═══════════════════════════════════════════════════════════════
def m_2axis(ctx, r, d):
    """{"x_axis": ["低", "実現性", "高"], "y_axis": ["低", "効果", "高"],
        "quadrants": ["様子見", "優先", "後回し", "早く着手"]（左上・右上・左下・右下）, "highlight": 1,
        "items": [{"label": "..", "x": 0.8, "y": 0.7, "size": 1.0, "emph": true}]}"""
    c = _c()
    x, y, w, h = r
    side = d.get("side")      # 右に凡例・説明を置く場合の幅（比率）
    gw = w * (1 - side) if side else w
    lx, ly, lw, lh = x + 0.45, y + 0.05, gw - 0.55, h - 0.5
    q = d.get("quadrants") or []
    focus = d.get("highlight")
    for k in range(4):
        qx = lx + (k % 2) * lw / 2
        qy = ly + (k // 2) * lh / 2
        on = focus == k
        box = c._rect(ctx.slide, qx, qy, lw / 2, lh / 2, fill=_h.LIGHTEST_PURPLE if on else None,
                      line=_h.LIGHT_GRAY, line_w=0.75)
        if k < len(q) and isinstance(q[k], dict):     # 象限ごとに見出し＋説明（2×2 の整理表）
            qd = q[k]
            paras = [(qd.get("title", ""), {"bold": True, "size": 16, "color": _h.DARKEST_PURPLE if on else _h.TEXT_BODY,
                                            "space_after": 4})]
            if qd.get("sub"):
                paras.append((qd["sub"], {"size": 12, "color": _h.DARK_PURPLE, "bold": True, "space_after": 4}))
            body = [(t, {"bullet": "•"}) for t in qd.get("items") or []]
            sz, ok = c.fit_size(body or [" "], lw / 2 - 0.45, lh / 2 - 0.75, None, 3)
            if not ok:
                ctx.warn(f"M-2AXIS: 象限「{qd.get('title', '')}」の文字が多い")
            c._textbox(ctx.slide, qx + 0.12, qy + 0.1, lw / 2 - 0.24, lh / 2 - 0.2, paras + [(t, {**o, "size": sz}) for t, o in body],
                       sz, space_after=3, emphasis=ctx.emphasis)
        elif k < len(q) and q[k]:
            c._textbox(ctx.slide, qx + 0.08, qy + 0.06, lw / 2 - 0.16, 0.34, [q[k]], 14,
                       _h.DARKEST_PURPLE if on else _h.MID_GRAY, bold=True, space_after=0,
                       align=PP_ALIGN.RIGHT if k % 2 else PP_ALIGN.LEFT)
    c._connector(ctx.slide, lx, ly + lh, lx + lw + 0.1, ly + lh, color=_h.DARKEST_PURPLE, width=1.5, arrow=True, glyph=True)
    c._connector(ctx.slide, lx, ly + lh, lx, ly - 0.05, color=_h.DARKEST_PURPLE, width=1.5, arrow=True, glyph=True)
    xa, ya = d.get("x_axis") or [], d.get("y_axis") or []
    if xa:
        c._textbox(ctx.slide, lx, ly + lh + 0.04, lw, 0.3, [xa[1] if len(xa) > 2 else xa[0]], 12, _h.TEXT_BODY,
                   align=PP_ALIGN.CENTER, bold=True, space_after=0)
        if len(xa) > 2:
            c._textbox(ctx.slide, lx, ly + lh + 0.04, 1.0, 0.3, [xa[0]], 12, _h.MID_GRAY, space_after=0)
            c._textbox(ctx.slide, lx + lw - 1.0, ly + lh + 0.04, 1.0, 0.3, [xa[2]], 12, _h.MID_GRAY, align=PP_ALIGN.RIGHT,
                       space_after=0)
    if ya:
        mid = ya[1] if len(ya) > 2 else ya[0]
        mh = len(mid) * 0.17 + 0.1
        t = c._textbox(ctx.slide, x, ly + lh / 2 - mh / 2, 0.35, mh, [mid], 12, _h.TEXT_BODY, bold=True, space_after=0,
                       inset=(0, 0, 0, 0))
        t.text_frame._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
        if len(ya) > 2:       # 縦軸の両端（上＝高・下＝低）。横軸と同じく端のラベルを置く
            for txt, top_end in ((ya[2], True), (ya[0], False)):
                eh = len(txt) * 0.17 + 0.08
                e = c._textbox(ctx.slide, x, ly - 0.02 if top_end else ly + lh - eh, 0.35, eh, [txt], 12, _h.MID_GRAY,
                               space_after=0, inset=(0, 0, 0, 0))
                e.text_frame._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
    for it in d.get("items") or []:
        px = lx + lw * max(0.03, min(0.97, it.get("x", 0.5)))
        py = ly + lh * (1 - max(0.03, min(0.97, it.get("y", 0.5))))
        dd = 0.22 * it.get("size", 1.0) ** 0.5
        on = it.get("emph")
        c._rect(ctx.slide, px - dd / 2, py - dd / 2, dd, dd, fill=_h.CORE_PURPLE if on else _h.DARK_PURPLE, shape=S.OVAL)
        c._textbox(ctx.slide, px + dd / 2 + 0.03, py - 0.16, 2.2, 0.32, [it.get("label", "")], 12,
                   _h.DARKEST_PURPLE if on else _h.TEXT_BODY, bold=bool(on), space_after=0)
    if side and d.get("notes"):
        paras = [(t, {"bullet": "•"}) for t in d["notes"]]
        sz, _ = c.fit_size(paras, w * side - 0.4, h, (14, 12), 4)
        c._textbox(ctx.slide, x + gw + 0.2, y, w * side - 0.2, h, paras, sz, space_after=4, emphasis=ctx.emphasis)


# ═══════════════════════════════════════════════════════════════
#  M-PYRAMID（旧 AC）
# ═══════════════════════════════════════════════════════════════
def m_pyramid(ctx, r, d):
    """{"levels": [{"label": "経営", "text": ".."}, ...]（上から）, "highlight": 0, "share": ["5%", "25%", "70%"]}"""
    c = _c()
    x, y, w, h = r
    lv = d.get("levels") or []
    n = max(len(lv), 1)
    pw = min(w * 0.5, h * 1.25)
    gap = 0.06
    lh = (h - gap * (n - 1)) / n
    focus = d.get("highlight")
    shades = [_h.DARKEST_PURPLE, _h.DARK_PURPLE, _h.CORE_PURPLE, _h.LIGHT_PURPLE, _h.LIGHTEST_PURPLE]
    cx = x + pw / 2
    for i, L in enumerate(lv):
        ty = y + i * (lh + gap)
        top_w = pw * i / n
        bot_w = pw * (i + 1) / n
        ff = ctx.slide.shapes.build_freeform(Inches(cx - top_w / 2), Inches(ty))
        ff.add_line_segments([(Inches(cx + top_w / 2), Inches(ty)), (Inches(cx + bot_w / 2), Inches(ty + lh)),
                              (Inches(cx - bot_w / 2), Inches(ty + lh))])
        shp = ff.convert_to_shape()
        col = shades[min(i, len(shades) - 1)] if focus is None or focus == i else _h.LIGHT_GRAY
        shp.fill.solid()
        shp.fill.fore_color.rgb = col
        shp.line.fill.background()
        c._strip_style(shp)
        light = col in (_h.LIGHT_PURPLE, _h.LIGHTEST_PURPLE, _h.LIGHT_GRAY)
        lab_w = max(0.8, (top_w + bot_w) / 2 - 0.2)
        c._textbox(ctx.slide, cx - lab_w / 2, ty, lab_w, lh, [L.get("label", "")], 14,
                   _h.DARKEST_PURPLE if light else _h.WHITE, bold=True, anchor="m", align=PP_ALIGN.CENTER, space_after=0)
        share = (d.get("share") or [None] * n)[i] if i < len(d.get("share") or []) else None
        tx = x + pw + 0.3
        c._connector(ctx.slide, cx + (top_w + bot_w) / 4 + 0.05, ty + lh / 2, tx - 0.05, ty + lh / 2,
                     color=_h.LIGHT_GRAY, width=0.75, dash=True)
        paras = []
        if share:
            paras.append((share, {"bold": True, "size": 18, "color": _h.CORE_PURPLE, "space_after": 0}))
        if L.get("text"):
            paras.append(L["text"])
        paras += [(t, {"bullet": "•"}) for t in L.get("items", [])]
        if paras:
            sz, ok = c.fit_size(paras, w - pw - 0.5, lh, (14, 12), 2)
            if not ok:
                ctx.warn(f"M-PYRAMID: 段{i + 1}の説明が長い")
            c._textbox(ctx.slide, tx, ty, w - pw - 0.3, lh, paras, sz, anchor="m", space_after=2, emphasis=ctx.emphasis)


# ═══════════════════════════════════════════════════════════════
#  M-VENN（旧 AE）
# ═══════════════════════════════════════════════════════════════
def m_venn(ctx, r, d):
    """{"sets": [{"label": "技術", "items": [..]}, ...]（2〜4）, "center": "統合概念", "notes": [..]（右の説明）}"""
    c = _c()
    x, y, w, h = r
    sets = d.get("sets") or []
    n = max(2, min(4, len(sets)))
    notes = d.get("notes")
    vw = w * (0.6 if notes else 1.0)
    cx0, cy0 = x + vw / 2, y + h / 2
    rad = min(vw, h) * (0.30 if n == 2 else 0.27)
    off = rad * (0.62 if n == 2 else 0.68)
    angles = {2: [180, 0], 3: [210, 330, 90], 4: [225, 315, 45, 135]}[n]
    fills = [_h.CORE_PURPLE, _h.DARK_PURPLE, _h.LIGHT_PURPLE, _h.DARKEST_PURPLE]
    for i, s in enumerate(sets[:n]):
        a = math.radians(angles[i])
        ccx, ccy = cx0 + off * math.cos(a) * (1.25 if n == 2 else 1.0), cy0 - off * math.sin(a)
        circ = c._rect(ctx.slide, ccx - rad, ccy - rad, rad * 2, rad * 2, fill=fills[i], line=_h.WHITE, line_w=1.0,
                       shape=S.OVAL)
        _alpha(circ, 72)
        lx = ccx + off * 0.7 * math.cos(a) * (1.25 if n == 2 else 1.0)
        ly = ccy - off * 0.7 * math.sin(a)
        paras = [(s.get("label", ""), {"bold": True, "size": 16, "color": _h.DARKEST_PURPLE, "space_after": 2,
                                       "align": PP_ALIGN.CENTER})]
        paras += [(t, {"size": 12, "align": PP_ALIGN.CENTER, "space_after": 0}) for t in s.get("items", [])]
        c._textbox(ctx.slide, lx - rad * 0.75, ly - rad * 0.5, rad * 1.5, rad, paras, 14, anchor="m", space_after=0)
    if d.get("center"):
        cw_ = rad * 0.9
        tag = c._rect(ctx.slide, cx0 - cw_ / 2, cy0 - 0.25, cw_, 0.5, fill=_h.DARKEST_PURPLE)
        c._shape_text(tag, d["center"], 14, _h.WHITE, True, inset=(0.04, 0, 0.04, 0))
    if notes:
        paras = [(t, {"bullet": "•"}) for t in notes]
        sz, _ = c.fit_size(paras, w - vw - 0.4, h, (14, 12), 4)
        c._textbox(ctx.slide, x + vw + 0.2, y, w - vw - 0.2, h, paras, sz, anchor="m", space_after=6, emphasis=ctx.emphasis)


# ═══════════════════════════════════════════════════════════════
#  M-ARCH（旧 AG）
# ═══════════════════════════════════════════════════════════════
def m_arch(ctx, r, d):
    """{"zones": [{"label": "利用者", "items": ["社員", {"label": "台帳", "kind": "db", "col": 2, "icon": "q:database"}]}, ...]（上から）,
        "links": [["社員", "ポータル", "SSO"], {"from": "A", "to": "B", "label": "記録", "dash": true}],
        "emph": ["基盤 API"], "scope": {"items": ["保全計画", "設備データ基盤"], "label": "ご提案範囲"},
        "direction": "rows|cols"}
    … 層（帯）の中に要素の箱を並べ、名前で指定した要素どうしを線でつなぐ。
    v6.4: 要素は層をまたいで同じ列にそろえ（col で列を指定できる）、線は直角に通す（斜めの線は引かない）。
    線は 2pt の濃い紫で矢じり付き、ラベルは線の横に白地で置く。"""
    c = _c()
    s = ctx.slide
    x, y, w, h = r
    zones = d.get("zones") or []
    nz = max(len(zones), 1)
    rows = d.get("direction", "rows") == "rows"
    gap = 0.3 if rows else 0.18
    lab = 1.3
    pos, zone_of, zbox = {}, {}, []
    emph = set(d.get("emph") or [])

    shp_of = {}       # 要素名 → 箱の図形（コネクタの接続先）
    def draw_item(it, bx, by, iw, ih):
        name = it if isinstance(it, str) else it.get("label", "")
        on = name in emph
        db = not isinstance(it, str) and it.get("kind") == "db"
        box = c._rect(s, bx, by, iw, ih, fill=_h.LIGHTEST_PURPLE if on else _h.WHITE,
                      line=_h.DARK_PURPLE if on else _h.MID_GRAY, line_w=1.5 if on else 1.0,
                      shape=S.CAN if db else S.RECTANGLE)
        tx, tw = bx, iw
        ic = None if isinstance(it, str) else it.get("icon")
        if ic and ih >= 0.5 and not db:
            isz = ih - 0.18
            if c._place_visual(ctx, ic, (bx + 0.1, by + 0.09, isz, isz)) is not None:
                tx, tw = bx + isz + 0.14, iw - isz - 0.18
        size = 12 if sum(c._char_w(ch, 14) for ch in name) * 1.08 > tw - 0.12 else 14
        if tw < iw:
            c._textbox(s, tx, by, tw, ih, [name], size, _h.DARKEST_PURPLE, anchor="m", bold=True, space_after=0,
                       inset=(0.02, 0, 0.04, 0))
        else:
            c._shape_text(box, name, size, _h.DARKEST_PURPLE, True, inset=(0.04, 0.06 if db else 0, 0.04, 0))
        pos[name] = (bx, by, iw, ih)
        shp_of[name] = box
        return name

    if rows:
        zh = (h - gap * (nz - 1)) / nz
        ncol = max([len(z.get("items") or []) for z in zones] + [1])
        ncol = max([ncol] + [it.get("col", 0) + 1 for z in zones for it in z.get("items") or [] if isinstance(it, dict)])
        cg = 0.25
        iw = min(2.6, (w - lab - 0.4 - cg * (ncol - 1)) / ncol)
        span = ncol * iw + (ncol - 1) * cg
        sx = x + lab + (w - lab - span) / 2
        for zi, z in enumerate(zones):
            zy = y + zi * (zh + gap)
            band = c._rect(s, x, zy, lab, zh, fill=_h.DARKEST_PURPLE)    # 層の見出しは同じ濃い紫（順番で色を変えない）
            c._shape_text(band, z.get("label", ""), 14, _h.WHITE, True, inset=(0.05, 0, 0.05, 0))
            c._panel(s, x + lab, zy, w - lab, zh)
            zbox.append((x, zy, w, zh))
            items = z.get("items") or []
            off = (ncol - len(items)) // 2
            ih = min(0.66, zh - 0.36)
            for k, it in enumerate(items):
                col = it.get("col", off + k) if isinstance(it, dict) else off + k
                name = draw_item(it, sx + col * (iw + cg), zy + (zh - ih) / 2, iw, ih)
                zone_of[name] = zi
    else:
        zw = (w - gap * (nz - 1)) / nz
        for zi, z in enumerate(zones):
            zx = x + zi * (zw + gap)
            band = c._rect(s, zx, y, zw, 0.42, fill=_h.DARKEST_PURPLE)
            c._shape_text(band, z.get("label", ""), 14, _h.WHITE, True)
            c._panel(s, zx, y + 0.42, zw, h - 0.42)
            zbox.append((zx, y, zw, h))
            items = z.get("items") or []
            m = max(len(items), 1)
            ih = min(0.66, (h - 0.7 - 0.25 * (m - 1)) / m)
            for k, it in enumerate(items):
                name = draw_item(it, zx + 0.2, y + 0.6 + k * (ih + 0.25), zw - 0.4, ih)
                zone_of[name] = zi

    import connect          # 線は接続付きのコネクタ（端点が要素につながる。直線を継いで曲げない）
    segs, pending, labels = [], [], []       # 引いた線の線分・あとで置くラベル（文字, 通り道）・置いたラベルの矩形
    site_use = {}

    def _free(rr):
        """ラベルの矩形 rr が、要素の箱・線・ほかのラベルに重ならないか（線で文字がつぶれない）。"""
        if rr[0] < x + (lab if rows else 0) + 0.02 or rr[0] + rr[2] > x + w - 0.02:
            return False
        inflate = (rr[0] - 0.03, rr[1] - 0.03, rr[2] + 0.06, rr[3] + 0.06)
        if any(_ovl(inflate, b_) for b_ in pos.values()) or any(_ovl(inflate, l_) for l_ in labels):
            return False
        return not any(connect.seg_hits(p_, q_, rr, m=-0.03) for p_, q_ in segs)

    def _ovl(r1, r2):
        return r1[0] < r2[0] + r2[2] and r2[0] < r1[0] + r1[2] and r1[1] < r2[1] + r2[3] and r2[1] < r1[1] + r1[3]

    def _place_labels():
        """ラベルは自分の線の脇（縦の線なら右か左、横の線なら上か下）に置き、要素の箱・ほかの線・ほかのラベルに
        重ならない位置を選ぶ（線が文字の上を通らない）。長い区間の中ほどから順に試す。"""
        for text, pts in pending:
            lw = 0.16 + sum(c._char_w(ch, 12) for ch in text) * 1.08
            lh = 0.28
            cands = []
            sg = sorted(zip(pts, pts[1:]), key=lambda pq: -abs(pq[0][0] - pq[1][0]) - abs(pq[0][1] - pq[1][1]))
            for t in (0.5, 0.35, 0.65, 0.2, 0.8):
                for p_, q_ in sg:
                    mx, my = p_[0] + (q_[0] - p_[0]) * t, p_[1] + (q_[1] - p_[1]) * t
                    if abs(p_[0] - q_[0]) < 1e-3:      # 縦の線: 右・左
                        cands += [(mx + 0.07, my - lh / 2, lw, lh), (mx - 0.07 - lw, my - lh / 2, lw, lh)]
                    else:                               # 横の線: 上・下
                        cands += [(mx - lw / 2, my - 0.05 - lh, lw, lh), (mx - lw / 2, my + 0.05, lw, lh)]
            rr = next((c_ for c_ in cands if _free(c_)), None)
            if rr is None:
                ctx.warn(f"M-ARCH: 線のラベル「{text}」を線・箱に重ならずに置けない（要素の間を広げるか、ラベルを短く）")
                rr = cands[0]
            labels.append(rr)
            c._textbox(s, rr[0], rr[1], rr[2], rr[3], [text], 12, _h.DARK_PURPLE, bold=True, anchor="m",
                       align=PP_ALIGN.CENTER, space_after=0, inset=(0.02, 0, 0.02, 0))

    def _score(pts, fa, fb, sa_, sb_):
        """通り道の悪さ: ほかの箱を横切る（100）・引いた線と重なる（60）・交わる（6）・同じ接続点を使う（10）・曲がり・長さ。"""
        sc_ = 0.0
        for p_, q_ in zip(pts, pts[1:]):
            sc_ += 100 * sum(1 for nm, r_ in pos.items() if nm not in (fa, fb) and connect.seg_hits(p_, q_, r_))
            for u_, v_ in segs:
                sc_ += 60 if connect.seg_overlap(p_, q_, u_, v_) else 6 if connect.seg_cross(p_, q_, u_, v_) else 0
            sc_ += 0.15 * (abs(p_[0] - q_[0]) + abs(p_[1] - q_[1]))
        sc_ += 10 * (site_use.get((fa, sa_), 0) + site_use.get((fb, sb_), 0)) + 1.5 * (len(pts) - 2)
        return sc_

    for lk in d.get("links") or []:
        if isinstance(lk, dict):
            fa, fb, text, dash = lk.get("from"), lk.get("to"), lk.get("label"), lk.get("dash", False)
        else:
            fa, fb, text, dash = lk[0], lk[1], (lk[2] if len(lk) > 2 else None), (len(lk) > 3 and lk[3] == "dash")
        a, b = pos.get(fa), pos.get(fb)
        if not a or not b:
            ctx.warn(f"M-ARCH: 線の端の要素が見つからない: {[fa, fb]}")
            continue
        sc_names = set((d.get("scope") or {}).get("items", []) if isinstance(d.get("scope"), dict) else d.get("scope") or [])
        # 範囲を示す図では、範囲の中どうしの線だけ濃い紫、範囲外の線は灰色（強調を範囲に集める）
        col = _h.DARK_PURPLE if not sc_names or (fa in sc_names and fb in sc_names) else _h.MID_GRAY
        lw_ = 2.0 if col == _h.DARK_PURPLE else 1.25
        A_, B_ = shp_of.get(fa), shp_of.get(fb)
        # 通り道の候補（向かい合う辺の直線・Z 字・L 字・層の間の溝を通る U 字）を調べ、箱を横切らず、
        # ほかの線と重ならず、接続点を共有しないものを選ぶ（1 本の DB への出入りが同じ点で重なると、線が分かれて見える）
        right = b[0] + b[2] / 2 > a[0] + a[2] / 2
        hs = ("r", "l") if right else ("l", "r")
        apart = b[0] >= a[0] + a[2] or b[0] + b[2] <= a[0]
        xg = ((a[0] + a[2] + b[0]) / 2 if right else (b[0] + b[2] + a[0]) / 2) if apart else None
        cands = []
        if abs((a[1] + a[3] / 2) - (b[1] + b[3] / 2)) <= 0.2:            # 同じ層
            cands.append((hs[0], hs[1], None))
            cands.append(("b", "b", max(a[1] + a[3], b[1] + b[3]) + gap / 2))
            cands.append(("t", "t", min(a[1], b[1]) - gap / 2))
        else:
            down = a[1] < b[1]
            vs = ("b", "t") if down else ("t", "b")
            za = zone_of.get(fa, 0)
            y1, y2 = (a[1] + a[3], b[1]) if down else (a[1], b[1] + b[3])
            ymid = (y1 + y2) / 2
            if rows and zbox:                    # 横の区間は層と層の間の溝を通す
                zi = max(0, min(za if down else za - 1, len(zbox) - 1))
                yg = zbox[zi][1] + zbox[zi][3] + gap / 2
                if min(y1, y2) < yg < max(y1, y2):
                    ymid = yg
            cands.append((vs[0], vs[1], ymid))
            if xg is not None:
                cands.append((hs[0], hs[1], xg))
            cands += [(hs[0], vs[1], None), (vs[0], hs[1], None)]
        best = None
        for sa_, sb_, mid_ in cands:
            pts = connect.route(A_, B_, sa_, sb_, mid_)
            if pts is None:
                continue
            sc_ = _score(pts, fa, fb, sa_, sb_)
            if best is None or sc_ < best[0]:
                best = (sc_, sa_, sb_, mid_, pts)
        if best is None:
            continue
        _, sa_, sb_, mid_, pts = best
        if best[0] >= 100:
            ctx.warn(f"M-ARCH: 「{fa}」→「{fb}」の線がほかの要素を横切る（列をそろえるか、要素の並びを変える）")
        connect.link(s, A_, B_, sa_, sb_, col, lw_, arrow=True, dash=dash, mid=mid_)
        segs.extend(zip(pts, pts[1:]))
        site_use[(fa, sa_)] = site_use.get((fa, sa_), 0) + 1
        site_use[(fb, sb_)] = site_use.get((fb, sb_), 0) + 1
        if text:
            pending.append((text, pts))

    _place_labels()
    sc = d.get("scope")
    if sc:
        names = sc.get("items") if isinstance(sc, dict) else sc
        inside = [pos[n] for n in names if n in pos]
        groups = {}
        if inside:
            bx0, by0 = min(b_[0] for b_ in inside), min(b_[1] for b_ in inside)
            bx1, by1 = max(b_[0] + b_[2] for b_ in inside), max(b_[1] + b_[3] for b_ in inside)
            others = [v for k, v in pos.items() if k not in names and v[0] < bx1 and v[0] + v[2] > bx0
                      and v[1] < by1 and v[1] + v[3] > by0]
            if not others:               # まとめて囲んでも範囲外の要素が入らなければ、枠は 1 つ
                groups = {0: inside}
        for n in names if not groups else []:   # 入るなら層ごとに枠を分ける
            if n in pos:
                groups.setdefault(zone_of.get(n, 0), []).append(pos[n])
        for gi, (zk, boxes) in enumerate(sorted(groups.items())):
            fx = min(b_[0] for b_ in boxes) - 0.1
            fy = min(b_[1] for b_ in boxes) - 0.1
            fw = max(b_[0] + b_[2] for b_ in boxes) + 0.1 - fx
            fh = max(b_[1] + b_[3] for b_ in boxes) + 0.1 - fy
            c._rect(s, fx, fy, fw, fh, fill=None, line=_h.CORE_PURPLE, line_w=2.0, dash=True)
            if gi < len(groups) - 1:
                continue
            labt = sc.get("label", "ご提案範囲") if isinstance(sc, dict) else "ご提案範囲"
            tw = 0.25 + len(labt) * 0.17
            tag = c._rect(s, fx + fw - tw, fy + fh, tw, 0.28, fill=_h.CORE_PURPLE)
            c._shape_text(tag, labt, 12, _h.WHITE, True, inset=(0.03, 0, 0.03, 0), wrap=False)


# ═══════════════════════════════════════════════════════════════
#  M-TILEMAP（旧 AK・AL）
# ═══════════════════════════════════════════════════════════════
_BMC = [("key_partners", "パートナー", 0, 0, 2, 4), ("key_activities", "主要活動", 2, 0, 2, 2),
        ("key_resources", "リソース", 2, 2, 2, 2), ("value", "価値提案", 4, 0, 2, 4),
        ("relationships", "顧客との関係", 6, 0, 2, 2), ("channels", "チャネル", 6, 2, 2, 2),
        ("segments", "顧客セグメント", 8, 0, 2, 4), ("cost", "コスト構造", 0, 4, 5, 1.6), ("revenue", "収益の流れ", 5, 4, 5, 1.6)]


def m_tilemap(ctx, r, d):
    """layout="columns"（既定）: {"columns": [{"head": "..", "tiles": [{"text": "..", "on": true}, ".."]}]}
       layout="canvas"（ビジネスモデルキャンバス）: {"canvas": {"value": [..], "segments": [..], ...}, "highlight": "value"}
       layout="calendar": {"days": ["月", ...], "weeks": ["第1週", ...], "events": [{"week": 0, "day": 2, "text": "..", "on": true}]}"""
    lay = d.get("layout") or ("canvas" if d.get("canvas") else "calendar" if d.get("events") is not None else "columns")
    if lay == "canvas":
        return _canvas(ctx, r, d)
    if lay == "calendar":
        return _calendar(ctx, r, d)
    c = _c()
    x, y, w, h = r
    cols = d.get("columns") or []
    nc = max(len(cols), 1)
    gap = 0.12
    cw = (w - gap * (nc - 1)) / nc
    mx = max([len(cl.get("tiles") or []) for cl in cols] + [1])
    th = min(0.55, (h - 0.55 - 0.08 * (mx - 1)) / mx)
    for i, cl in enumerate(cols):
        cx = x + i * (cw + gap)
        c.pane_header(ctx, (cx, y, cw, h), cl.get("head", ""), style="rule")
        for k, t in enumerate(cl.get("tiles") or []):
            t = t if isinstance(t, dict) else {"text": str(t)}
            on = t.get("on", d.get("default_on", False))
            ty = y + 0.55 + k * (th + 0.08)
            box = c._rect(ctx.slide, cx, ty, cw, th, fill=_h.DARK_PURPLE if on else _h.WHITE,
                          line=None if on else _h.LIGHT_GRAY, line_w=0.75)
            c._shape_text(box, t.get("text", ""), 12, _h.WHITE if on else _h.MID_GRAY, bool(on), inset=(0.05, 0, 0.05, 0))
    if d.get("legend", True) and any((t.get("on") if isinstance(t, dict) else False) for cl in cols for t in cl.get("tiles") or []):
        import artifacts
        artifacts.v_legend(ctx, {"items": [{"fill": "purple", "label": d.get("on_label", "対象")}]}, r)


def _canvas(ctx, r, d):
    c = _c()
    x, y, w, h = r
    cv = d.get("canvas") or {}
    focus = d.get("highlight")
    ux, uy = w / 10, h / 5.6
    for key, label, gx, gy, gw, gh in _BMC:
        bx, by, bw, bh = x + gx * ux, y + gy * uy, gw * ux - 0.06, gh * uy - 0.06
        on = focus == key
        c._rect(ctx.slide, bx, by, bw, bh, fill=_h.LIGHTEST_PURPLE if on else None,
                line=_h.CORE_PURPLE if on else _h.LIGHT_GRAY, line_w=1.5 if on else 0.75)
        items = cv.get(key) or []
        paras = [(label, {"bold": True, "size": 14, "color": _h.DARKEST_PURPLE, "space_after": 4})]
        paras += [(t, {"bullet": "•"}) for t in items]
        sz, ok = c.fit_size(paras[1:] or [" "], bw - 0.35, bh - 0.45, (14, 12), 2)
        if not ok:
            ctx.warn(f"M-TILEMAP(canvas): {label} の文字が多い")
        paras = [paras[0]] + [(t, {**o, "size": sz}) for t, o in paras[1:]]
        c._textbox(ctx.slide, bx + 0.06, by + 0.04, bw - 0.12, bh - 0.08, paras, sz, space_after=2, emphasis=ctx.emphasis)


def _calendar(ctx, r, d):
    c = _c()
    x, y, w, h = r
    days = d.get("days") or ["月", "火", "水", "木", "金"]
    weeks = d.get("weeks") or ["第1週", "第2週", "第3週", "第4週"]
    lab = 1.0
    cw = (w - lab) / len(days)
    hh = 0.36
    rh = (h - hh) / len(weeks)
    for j, dn in enumerate(days):
        hb = c._rect(ctx.slide, x + lab + j * cw, y, cw - 0.04, hh, fill=_h.DARKEST_PURPLE)
        c._shape_text(hb, dn, 14, _h.WHITE, True, inset=(0, 0, 0, 0))
    for i, wk in enumerate(weeks):
        wy = y + hh + i * rh
        c._textbox(ctx.slide, x, wy, lab - 0.05, rh, [wk], 14, _h.DARKEST_PURPLE, bold=True, anchor="m", space_after=0)
        for j in range(len(days)):
            c._rect(ctx.slide, x + lab + j * cw, wy + 0.04, cw - 0.04, rh - 0.04, fill=None, line=_h.LIGHT_GRAY, line_w=0.75)
    slots = {}
    for ev in d.get("events") or []:
        i, j = ev.get("week", 0), ev.get("day", 0)
        k = slots.get((i, j), 0)
        slots[(i, j)] = k + 1
        ex, ey = x + lab + j * cw + 0.06, y + hh + i * rh + 0.1 + k * 0.36
        on = ev.get("on")
        tag = c._rect(ctx.slide, ex, ey, cw - 0.16, 0.32, fill=_h.CORE_PURPLE if on else _h.LIGHTEST_PURPLE)
        c._shape_text(tag, ev.get("text", ""), 12, _h.WHITE if on else _h.DARKEST_PURPLE, bool(on), inset=(0.04, 0, 0.04, 0))
        if ey + 0.32 > y + hh + (i + 1) * rh:
            ctx.warn("M-TILEMAP(calendar): 1 日の予定が多く、枠からはみ出す")


# ═══════════════════════════════════════════════════════════════
#  M-SCORE（公式 Starter Pack の Scoring – Harvey Balls）
# ═══════════════════════════════════════════════════════════════
def _harvey(slide, cx, cy, d, v):
    """ハーベイボール（0〜4 の 5 段階）。白い円に、上から時計回りに v/4 を塗る。"""
    c = _c()
    c._rect(slide, cx - d / 2, cy - d / 2, d, d, fill=_h.WHITE, line=_h.DARK_PURPLE, line_w=1.0, shape=S.OVAL)
    v = max(0, min(4, int(round(v))))
    if v == 4:
        c._rect(slide, cx - d / 2, cy - d / 2, d, d, fill=_h.DARK_PURPLE, line=_h.DARK_PURPLE, line_w=1.0, shape=S.OVAL)
    elif v:
        p = c._rect(slide, cx - d / 2, cy - d / 2, d, d, fill=_h.DARK_PURPLE, shape=S.PIE)
        p.adjustments[0] = 270.0 * 0.6                 # python-pptx は角度を 1/100000 で書く（OOXML は 1/60000 度）
        p.adjustments[1] = ((270.0 + 90.0 * v) % 360) * 0.6


def m_score(ctx, r, d):
    """{"criteria": ["効果", "実現性", "コスト"], "options": [{"name": "案 A", "scores": [4, 3, 1], "notes": ["..", "..", ".."],
        "emph": true, "total": "推奨"}], "scale": ["なし", "低", "中", "高", "最大"], "total_head": "総合評価"}
    … 選択肢（行）×評価基準（列）の表。セルにハーベイボールと短いコメント。推す行だけ薄く塗る。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    crit = d.get("criteria") or []
    opts = d.get("options") or []
    has_total = any(o.get("total") for o in opts)
    nc = len(crit) + (1 if has_total else 0)
    lab_w = min(2.2, w * 0.2)
    cw = (w - lab_w) / max(nc, 1)
    hdr = 0.45
    rh = (h - hdr - 0.05) / max(len(opts), 1)
    ball = min(0.34, rh * 0.4)
    for j, cr in enumerate(crit + ([d.get("total_head", "総合評価")] if has_total else [])):
        c._textbox(s, x + lab_w + j * cw, y, cw, hdr - 0.04, [cr], 14, _h.DARKEST_PURPLE, bold=True, anchor="b",
                   align=PP_ALIGN.CENTER, space_after=0)
    c._hline(s, x, y + hdr, w, _h.DARK_PURPLE, 1.5)
    for i, o in enumerate(opts):
        oy = y + hdr + 0.05 + i * rh
        if o.get("emph"):
            c._rect(s, x, oy, w, rh - 0.04, fill=_h.LIGHTEST_PURPLE)
        c._textbox(s, x, oy, lab_w - 0.1, rh - 0.04, [o.get("name", "")], 16, _h.DARKEST_PURPLE, bold=True,
                   anchor="m", space_after=0)
        notes = o.get("notes") or []
        for j, sc in enumerate(o.get("scores") or []):
            cx = x + lab_w + j * cw
            note = notes[j] if j < len(notes) else ""
            if note:
                _harvey(s, cx + 0.12 + ball / 2, oy + rh / 2 - 0.02, ball, sc)
                sz, ok = c.fit_size([note], cw - ball - 0.3, rh - 0.12, (14, 12), 0)
                if not ok:
                    ctx.warn(f"M-SCORE: {o.get('name', '')} の「{crit[j] if j < len(crit) else ''}」のコメントが長い")
                c._textbox(s, cx + ball + 0.2, oy, cw - ball - 0.25, rh - 0.04, [note], sz, anchor="m", space_after=0)
            else:
                _harvey(s, cx + cw / 2, oy + rh / 2 - 0.02, ball, sc)
        if has_total:
            tx = x + lab_w + len(crit) * cw
            t = o.get("total", "")
            c._textbox(s, tx, oy, cw, rh - 0.04, [t], 16 if len(t) <= 6 else 14,
                       _h.CORE_PURPLE if o.get("emph") else _h.TEXT_BODY, bold=bool(o.get("emph")), anchor="m",
                       align=PP_ALIGN.CENTER, space_after=0)
        if i < len(opts) - 1:
            c._hline(s, x, oy + rh - 0.02, w, _h.LIGHT_GRAY, 0.75, dash=True)
    if d.get("legend", True):           # 凡例（右上）: ○◔◑◕● と段階の名前
        scale = d.get("scale") or ["なし", "低", "中", "高", "最大"]
        lx = x + w
        for v in range(4, -1, -1):
            lab = scale[v] if v < len(scale) else ""
            tw = 0.12 + len(lab) * 0.17
            lx -= tw
            c._textbox(s, lx, y - 0.36, tw, 0.3, [lab], 12, _h.TEXT_SUB, anchor="m", space_after=0, inset=(0.02, 0, 0, 0))
            lx -= 0.24
            _harvey(s, lx + 0.1, y - 0.21, 0.18, v)
            lx -= 0.06


# ═══════════════════════════════════════════════════════════════
#  M-PHASEGRID（公式 Starter Pack の Phase Matrix）
# ═══════════════════════════════════════════════════════════════
PG_GAP = 0.08


def m_phasegrid(ctx, r, d):
    """{"phases": ["構想", "設計", "構築", "展開"], "rows": [{"label": "業務", "cells": ["..", ["..", ".."], ..]}], "current": 1}
    … 工程の五方向の矢印を列見出しにし、行（領域）×工程の格子に作業を書く。現在の工程の列だけ薄く塗る。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    ph = d.get("phases") or []
    rows = d.get("rows") or []
    cur = d.get("current")
    lab_w = 1.25
    n = max(len(ph), 1)
    cw = (w - lab_w) / n
    hh = 0.5
    for j, p in enumerate(ph):
        px = x + lab_w + j * cw
        on = cur == j
        shp = c.homeplate(s, px, y, cw - (c.PHASE_GAP if j < n - 1 else 0), hh,      # 五方向の矢印を隙間を空けて並べる
                          fill=_h.CORE_PURPLE if on else _h.DARK_PURPLE if cur is None else _h.LIGHT_GRAY)
        c._shape_text(shp, p, 14, _h.WHITE if (on or cur is None) else _h.TEXT_BODY, True, inset=(0.08, 0, 0.04, 0))
    # 行の高さは中身の量に比例させる（短い行の箱が高すぎて空かないように）。極端に低くはしない
    need = [max([c.est_height([(t_, {}) for t_ in (cell if isinstance(cell, list) else [cell] if cell else [])] or [" "],
                              cw - 0.2, 14, 2) for cell in (rw.get("cells") or [])] or [0.4]) + 0.2 for rw in rows] or [1]
    avail = h - hh - 0.12
    tot = sum(need)
    rhs = [max(avail * k / tot, avail / max(len(rows), 1) * 0.6) for k in need]
    rhs = [r_ * avail / sum(rhs) for r_ in rhs]
    rh = min(rhs) if rhs else avail
    # 格子の文字は全セルで同じ大きさにする（セルごとに縮めると字の大きさが不揃いになり、手作り感が崩れる）
    cells = [[(t, {"bullet": "•"}) for t in (cell if isinstance(cell, list) else [cell] if cell else [])]
             for rw in rows for cell in (rw.get("cells") or [])]
    sz = min([c.fit_size(pp, cw - 0.2, rh - 0.12, None, 2)[0] for pp in cells if pp] or [14])
    ry_next = y + hh + 0.12
    for i, rw in enumerate(rows):
        ry, rh = ry_next, rhs[i]
        ry_next += rh
        band = c._rect(s, x, ry, lab_w - 0.1, rh - 0.08, fill=_h.DARKEST_PURPLE)
        c._shape_text(band, rw.get("label", ""), 14, _h.WHITE, True, inset=(0.05, 0, 0.05, 0))
        for j, cell in enumerate(rw.get("cells") or []):
            items = cell if isinstance(cell, list) else [cell] if cell else []
            paras = [(t, {"bullet": "•"}) for t in items]
            if not paras:
                continue
            if not c.fit_size(paras, cw - 0.2, rh - 0.12, (sz,), 2)[1]:
                ctx.warn(f"M-PHASEGRID: {rw.get('label', '')}×{ph[j] if j < len(ph) else j} の文字が多い")
            if cur == j:            # 現在の工程の列は、セルごとの淡い紫の箱（文字は箱に直接。列をまたぐ帯にしない）
                c._rect(s, x + lab_w + j * cw + 0.02, ry, cw - PG_GAP, rh - 0.08, fill=_h.LIGHTEST_PURPLE)
            c._textbox(s, x + lab_w + j * cw + 0.06, ry, cw - 0.12, rh - 0.08, paras, sz, anchor="m", space_after=2,
                       emphasis=ctx.emphasis)
        if i < len(rows) - 1:
            c._hline(s, x + lab_w, ry + rh - 0.04, w - lab_w, _h.LIGHT_GRAY, 0.75, dash=True)


# ═══════════════════════════════════════════════════════════════
#  M-MATURITY（公式 Starter Pack の Maturity Scale）
# ═══════════════════════════════════════════════════════════════
def m_maturity(ctx, r, d):
    """{"levels": ["Lv1 個別対応", "Lv2 標準化", ..], "rows": [{"label": "業務プロセス", "current": 1, "target": 3, "note": ".."}],
        "note_head": "打ち手"} … 領域ごとに現状（●）と目標（◇）を段階の上に置き、差を矢印で示す。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    lv = d.get("levels") or []
    rows = d.get("rows") or []
    has_note = any(rw.get("note") for rw in rows)
    lab_w = min(2.2, w * 0.2)
    note_w = w * 0.26 if has_note else 0
    gw = w - lab_w - note_w - (0.2 if has_note else 0)
    n = max(len(lv), 1)
    cw = gw / n
    hdr = 0.5
    for j, L in enumerate(lv):
        c._textbox(s, x + lab_w + j * cw, y, cw, hdr - 0.04, [L], 14, _h.DARKEST_PURPLE, bold=True, anchor="b",
                   align=PP_ALIGN.CENTER, space_after=0)
    if has_note:
        c._textbox(s, x + w - note_w, y, note_w, hdr - 0.04, [d.get("note_head", "打ち手")], 14, _h.DARKEST_PURPLE,
                   bold=True, anchor="b", space_after=0)
    c._hline(s, x, y + hdr, w, _h.DARK_PURPLE, 1.5)
    rh = (h - hdr - 0.05) / max(len(rows), 1)
    mk = min(0.3, rh * 0.4)
    for i, rw in enumerate(rows):
        ry = y + hdr + 0.05 + i * rh
        cy = ry + rh / 2
        c._textbox(s, x, ry, lab_w - 0.1, rh, [rw.get("label", "")], 14, _h.TEXT_BODY, bold=True, anchor="m",
                   space_after=0)
        c._hline(s, x + lab_w + cw / 2, cy, cw * (n - 1), _h.LIGHT_GRAY, 0.75)      # 段階の目盛り
        for j in range(n):
            tx = x + lab_w + j * cw + cw / 2
            c._rect(s, tx - 0.04, cy - 0.04, 0.08, 0.08, fill=_h.LIGHT_GRAY, shape=S.OVAL)
        cur, tgt = rw.get("current"), rw.get("target")
        if cur is not None and tgt is not None and tgt != cur:
            x1 = x + lab_w + cur * cw + cw / 2 + (mk / 2 if tgt > cur else -mk / 2)
            x2 = x + lab_w + tgt * cw + cw / 2 - (mk / 2 + 0.02 if tgt > cur else -mk / 2 - 0.02)
            c._connector(s, x1, cy, x2, cy, _h.CORE_PURPLE, 1.5, arrow=True)
        if tgt is not None:
            tx = x + lab_w + tgt * cw + cw / 2
            c._rect(s, tx - mk / 2, cy - mk / 2, mk, mk, fill=_h.WHITE, line=_h.CORE_PURPLE, line_w=1.5, shape=S.DIAMOND)
        if cur is not None:
            tx = x + lab_w + cur * cw + cw / 2
            c._rect(s, tx - mk / 2, cy - mk / 2, mk, mk, fill=_h.DARK_PURPLE, shape=S.OVAL)
        if rw.get("note"):
            sz, _ = c.fit_size([rw["note"]], note_w, rh - 0.08, None, 0)
            c._textbox(s, x + w - note_w, ry, note_w, rh, [rw["note"]], sz, anchor="m", space_after=0, emphasis=ctx.emphasis)
        if i < len(rows) - 1:
            c._hline(s, x, ry + rh, w, _h.LIGHT_GRAY, 0.75, dash=True)
    if d.get("legend", True):
        c._rect(s, x + w - 2.15, y - 0.29, 0.18, 0.18, fill=_h.DARK_PURPLE, shape=S.OVAL)
        c._textbox(s, x + w - 1.93, y - 0.36, 0.7, 0.3, ["現状"], 12, _h.TEXT_SUB, anchor="m", space_after=0)
        c._rect(s, x + w - 1.1, y - 0.29, 0.18, 0.18, fill=_h.WHITE, line=_h.CORE_PURPLE, line_w=1.25, shape=S.DIAMOND)
        c._textbox(s, x + w - 0.88, y - 0.36, 0.85, 0.3, [d.get("target_label", "目標")], 12, _h.TEXT_SUB, anchor="m",
                   space_after=0)


# ═══════════════════════════════════════════════════════════════
#  M-HOUSE（公式 Starter Pack の Pillars House）
# ═══════════════════════════════════════════════════════════════
def m_house(ctx, r, d):
    """{"roof": "目指す姿", "roof_text": "..", "pillars": [{"label": "..", "items": [..]}], "base": ["土台 1", "土台 2"]}
    … 屋根（目指す姿）・柱（施策）・基盤（土台）。柱は文字＋下線の見出しで、面は塗りすぎない。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    pil = d.get("pillars") or []
    base = d.get("base") or []
    base = base if isinstance(base, list) else [base]
    roof_h = min(1.15, h * 0.22)
    base_h = 0.42
    bh_total = len(base) * (base_h + 0.06)
    # 屋根（三角）
    ff = s.shapes.build_freeform(Inches(x), Inches(y + roof_h))
    ff.add_line_segments([(Inches(x + w / 2), Inches(y)), (Inches(x + w), Inches(y + roof_h))])
    roof = ff.convert_to_shape()
    roof.fill.solid()
    roof.fill.fore_color.rgb = _h.DARKEST_PURPLE
    roof.line.fill.background()
    c._strip_style(roof)
    paras = [(d.get("roof", ""), {"bold": True, "size": 16, "color": _h.WHITE, "space_after": 0})]
    if d.get("roof_text"):
        paras.append((d["roof_text"], {"size": 12, "color": _h.WHITE}))
    c._textbox(s, x + w * 0.3, y + roof_h * 0.38, w * 0.4, roof_h * 0.6, paras, 16, align=PP_ALIGN.CENTER, anchor="m",
               space_after=0)
    # 柱
    py = y + roof_h + 0.12
    ph = h - roof_h - 0.12 - bh_total - 0.06
    n = max(len(pil), 1)
    gap = 0.18
    pw = (w - 0.3 - gap * (n - 1)) / n
    for i, p in enumerate(pil):
        px = x + 0.15 + i * (pw + gap)
        c._panel(s, px, py, pw, ph)
        inner = c.pane_header(ctx, (px + 0.1, py + 0.06, pw - 0.2, ph - 0.1), p.get("label", ""), style="rule")
        items = p.get("items") or []
        paras = [(t, {"bullet": "•"}) for t in items]
        if paras:
            sz, ok = c.fit_size(paras, inner[2], inner[3], None, 3)
            if not ok:
                ctx.warn(f"M-HOUSE: 柱「{p.get('label', '')}」の文字が多い")
            c._textbox(s, inner[0], inner[1], inner[2], inner[3], paras, sz, space_after=3, emphasis=ctx.emphasis)
    # 基盤
    by = y + h - bh_total
    for k, b in enumerate(base):
        band = c._rect(s, x, by + k * (base_h + 0.06), w, base_h, fill=_h.LIGHTEST_PURPLE if k else _h.DARK_PURPLE)
        c._shape_text(band, b, 14, _h.DARKEST_PURPLE if k else _h.WHITE, True)


# ═══════════════════════════════════════════════════════════════
#  M-CASE（事例ページ。公式 Credential Template と提案書の事例ページ）
# ═══════════════════════════════════════════════════════════════
def m_case(ctx, r, d):
    """{"client": "大手製造業 A 社", "attrs": [["業界", "製造"], ["規模", "従業員 1 万人"], ["期間", "6 か月"]],
        "headline": "保全計画を自動化し、計画工数を 30% 削減",
        "sections": [{"head": "背景・課題", "items": [..]}, {"head": "支援内容", "items": [..]},
                     {"head": "成果", "kpis": [{"value": "30%", "label": "計画工数の削減"}], "items": [..]}],
        "image": {"asset": ".."}（任意。属性の上に置く）}
    … 左に事例の属性、右に見出し（縦罫線）と「課題 → 支援 → 成果」。成果は数字を大きく。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    side = min(2.6, w * 0.22)
    ay = y
    if d.get("image"):
        ih = min(1.5, h * 0.3)
        c._place_visual(ctx, d["image"], (x, ay, side, ih))
        ay += ih + 0.15
    if d.get("client"):
        c._textbox(s, x, ay, side, 0.5, [d["client"]], 16, _h.DARKEST_PURPLE, bold=True, anchor="b", space_after=0)
        ay += 0.5
        c._hline(s, x, ay + 0.02, side, _h.DARK_PURPLE, 1.5)
        ay += 0.12
    attrs = d.get("attrs") or []
    hs = [c.est_height([v], side, 14, 0) + 0.26 for _, v in attrs]
    # 属性の行は枠の下まで均等に広げる（左の列だけ上に詰まって下が空かないように。広げすぎない）
    pad = min(0.45, max(0.12, (y + h - ay - sum(hs)) / max(len(attrs), 1) - 0.04)) if attrs else 0.12
    for (k, v), hh in zip(attrs, hs):
        paras = [(k, {"size": 12, "bold": True, "color": _h.DARK_PURPLE, "space_after": 0}), (v, {"size": 14})]
        c._textbox(s, x, ay, side, hh, paras, 14, space_after=0)
        ay += hh + pad - 0.08
        c._hline(s, x, ay, side, _h.LIGHT_GRAY, 0.75)
        ay += 0.08
    # 右: 見出し（左に縦罫線）＋区画
    mx = x + side + 0.4
    mw = w - side - 0.4
    hl = d.get("headline", "")
    hsz, _ = c.fit_size([hl], mw - 0.3, 0.95, (18, 16), 0)
    hh = c.est_height([hl], mw - 0.3, hsz, 0)
    c._rect(s, mx, y + 0.02, 0.06, hh, fill=_h.CORE_PURPLE)
    c._textbox(s, mx + 0.18, y, mw - 0.18, hh + 0.04, [hl], hsz, _h.DARKEST_PURPLE, bold=True, anchor="m", space_after=0)
    secs = d.get("sections") or []
    res = next((sec for sec in secs if sec.get("kpis")), None)
    band_h = 0
    if res is not None and len(secs) >= 3:       # 成果（数字）は下の帯に横並び。紙面の下まで使い、結論を最後に読ませる
        secs = [sec for sec in secs if sec is not res]
        band_h = 1.55
        by = y + h - band_h
        inner = c.pane_header(ctx, (mx, by, mw, band_h), res.get("head", "成果"), style="rule")
        kx = mx
        for kp in res.get("kpis") or []:
            kw = min(4.0, max(2.2, sum(c._char_w(ch, 36) for ch in kp.get("value", "")) + 0.3,
                              sum(c._char_w(ch, 12) for ch in kp.get("label", "")) + 0.3))
            c._textbox(s, kx, inner[1] - 0.05, kw, 0.7, [kp.get("value", "")], 36, _h.CORE_PURPLE, bold=True,
                       anchor="b", space_after=0)
            c._textbox(s, kx, inner[1] + 0.67, kw, 0.4, [kp.get("label", "")], 12, _h.TEXT_SUB, space_after=0)
            kx += kw + 0.3
        rest = [(t, {"bullet": "•"}) for t in res.get("items") or []]
        if rest and kx < mx + mw - 2:
            c._textbox(s, kx, inner[1], mx + mw - kx, inner[3], rest, 14, anchor="m", space_after=4, emphasis=ctx.emphasis)
        h_sec = h - band_h - 0.25
    else:
        h_sec = h
    sy = y + hh + 0.3
    n = max(len(secs), 1)
    gap = 0.3
    sw = (mw - gap * (n - 1)) / n
    for i, sec in enumerate(secs):
        sx = mx + i * (sw + gap)
        inner = c.pane_header(ctx, (sx, sy, sw, y + h_sec - sy), sec.get("head", ""), style="rule")
        iy = inner[1]
        for kp in sec.get("kpis") or []:
            c._textbox(s, sx, iy, sw, 0.62, [kp.get("value", "")], 32, _h.CORE_PURPLE, bold=True, anchor="b", space_after=0)
            c._textbox(s, sx, iy + 0.62, sw, 0.32, [kp.get("label", "")], 12, _h.TEXT_SUB, space_after=0)
            iy += 1.0
        paras = [(t, {"bullet": "•"}) for t in sec.get("items") or []]
        if paras:
            hh2 = y + h_sec - iy
            sz, ok = c.fit_size(paras, sw, hh2, None, 4)
            if not ok:
                ctx.warn(f"M-CASE: 「{sec.get('head', '')}」の文字が多い")
            need = min(hh2, c.est_height(paras, sw, sz, 4) + 0.1)     # 中身の高さの枠（下が空けば大きく描き直される）
            c._textbox(s, sx, iy, sw, need, paras, sz, space_after=4, emphasis=ctx.emphasis)
            if need < hh2 * 0.6:
                ctx.warn("M-CASE: 枠の下が空く（各区画に根拠・数字を足す）")
        if i < n - 1:     # 区画の間の縦線（線は使った高さの計算に入らない）
            c._connector(s, sx + sw + gap / 2, sy + 0.1, sx + sw + gap / 2, y + h_sec - 0.1, _h.LIGHT_GRAY, 0.75)


def register(modules):
    modules.update({"M-PROFILE": m_profile, "M-ROADMAP": m_roadmap, "M-2AXIS": m_2axis, "M-PYRAMID": m_pyramid,
                    "M-VENN": m_venn, "M-ARCH": m_arch, "M-TILEMAP": m_tilemap,
                    "M-SCORE": m_score, "M-PHASEGRID": m_phasegrid, "M-MATURITY": m_maturity, "M-HOUSE": m_house,
                    "M-CASE": m_case})
