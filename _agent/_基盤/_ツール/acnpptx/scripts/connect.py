"""
connect.py — 図形どうしを「接続付きのコネクタ」でつなぐ（v6.6）。

人の資料では、カギ線（bentConnector3）の 80% は両端が図形の接続点につながっている（PowerPoint で選ぶと端点が緑。
箱を動かすと線が付いてくる）。直線を 3 本継いで曲げた線は、動かすと崩れるので使わない。

  link(slide, a, b, "r", "l")            a の右辺の中点 → b の左辺の中点（直線かカギ線を自動で選ぶ）
  link(..., mid=4.2)                     曲がる位置（横に出る線なら x、縦に出る線なら y。束ね線の位置）
  link(..., arrow=True / both=True)      終点に矢印 / 両端に矢印

接続点の番号（人の資料の実測）: 矩形系 上0 左1 下2 右3 ／ 楕円 上0 左2 下4 右6 ／ 円柱 上1 左2 下3 右4。
コネクタの形: 一直線なら straightConnector1、直交する辺どうしなら bentConnector2、
平行な辺どうし（Z 字・U 字）なら bentConnector3。局所座標で始点 (0,0)・終点 (w,h) とし、反転してから中心まわりに回転する。

つなぐ／つながないの線引き（v6.7）:
  つなぐ  … 図の要素（文字の入った箱・円・円柱）どうしの関係の線。構成図・組織図・ツリー・ハブ・フローの戻り。曲げる線はすべて。
  つながない … アイコン・画像・人の図・番号札の並びの順番を示す短い直線の矢印、区切り線・下線・軸（compose._connector）。
  一直線にそろう線はカギ線にしない（直線）。線は文字（ラベル・箱の文字）の上を通さない（route で通り道を調べて選ぶ）。
"""

import math

from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU = 914400
DIRS = {"t": (0, -1), "b": (0, 1), "l": (-1, 0), "r": (1, 0)}
SITES = {"ellipse": {"t": 0, "l": 2, "b": 4, "r": 6},
         "can": {"t": 1, "l": 2, "b": 3, "r": 4},
         "flowChartMagneticDisk": {"t": 1, "l": 2, "b": 3, "r": 4}}
RECT_SITES = {"t": 0, "l": 1, "b": 2, "r": 3}


def _prst(shape):
    g = shape._element.find(".//" + qn("a:prstGeom"))
    return g.get("prst") if g is not None else "rect"


def site(shape, side):
    """辺（t/l/b/r）の接続点の番号と座標（インチ）。"""
    x, y = shape.left / EMU, shape.top / EMU
    w, h = shape.width / EMU, shape.height / EMU
    idx = SITES.get(_prst(shape), RECT_SITES)[side]
    px, py = {"t": (x + w / 2, y), "b": (x + w / 2, y + h), "l": (x, y + h / 2), "r": (x + w, y + h / 2)}[side]
    return idx, px, py


def auto_sides(a, b):
    """a と b の位置関係から、向かい合う辺を選ぶ。"""
    ax, ay, aw, ah = a.left, a.top, a.width, a.height
    bx, by, bw, bh = b.left, b.top, b.width, b.height
    if bx >= ax + aw:
        return "r", "l"
    if bx + bw <= ax:
        return "l", "r"
    return ("b", "t") if by >= ay + ah else ("t", "b")


def _map(rot, fh, v):
    x, y = v
    if fh:
        x = -x
    a = math.radians(rot)
    return (round(x * math.cos(a) - y * math.sin(a)), round(x * math.sin(a) + y * math.cos(a)))


def _frame(ux, uy):
    """局所の +x を ux、+y を uy に写す (rot, flipH)。"""
    for rot in (0, 90, 180, 270):
        for fh in (False, True):
            if _map(rot, fh, (1, 0)) == ux and _map(rot, fh, (0, 1)) == uy:
                return rot, fh
    raise ValueError((ux, uy))


def _dot(p, q):
    return p[0] * q[0] + p[1] * q[1]


def _geometry(s, ds, e, de, mid):
    """(prst, rot, flipH, w, h, adj1, swapped)。s/e は端点、ds/de は辺の外向き、mid は曲がる位置の座標（None で既定）。"""
    v = (e[0] - s[0], e[1] - s[1])

    def dist(p):          # p から曲がる位置までの、ds 方向の距離
        return (mid - p[0]) * ds[0] + (mid - p[1]) * ds[1]
    if _dot(ds, de) == 0:                                  # 直交する辺: 曲がり 1 回
        ux, uy = ds, (-de[0], -de[1])
        w, h = _dot(v, ux), _dot(v, uy)
        if w > 0.005 and h > 0.005:
            rot, fh = _frame(ux, uy)
            return "bentConnector2", rot, fh, w, h, None, False
        return None
    perp = (-ds[1], ds[0])
    along, side = _dot(v, ds), _dot(v, perp)
    if abs(side) < 0.04 and _dot(ds, de) < 0 and along > 0:  # 向かい合って一直線（0.04in 未満のずれはカギにしない）
        return "straightConnector1", None, None, None, None, None, False
    uy = perp if side >= 0 else (-perp[0], -perp[1])
    h = abs(side)
    if _dot(ds, de) < 0:                                   # Z 字（向かい合う辺）
        if along <= 0.005:
            return None
        rot, fh = _frame(ds, uy)
        m = along / 2 if mid is None else dist(s)
        return "bentConnector3", rot, fh, along, h, m / along, False
    # U 字（同じ向きの辺）: 遠い方の端から出る向きで描く（局所の幅を正にするため、必要なら端を入れ替える）
    swapped = along < 0
    w = max(abs(along), 0.01)
    if swapped:
        uy = (-uy[0], -uy[1])
    rot, fh = _frame(ds, uy)
    out = abs(along) + 0.3 if mid is None else dist(e if swapped else s)
    return "bentConnector3", rot, fh, w, h, out / w, swapped


def route(a, b, sa, sb, mid=None):
    """link が描く線の通り道（インチの点の列）。曲げてつなげない配置なら None（link は斜めの直線にする）。"""
    _, sx, sy = site(a, sa)
    _, ex, ey = site(b, sb)
    s, e = (sx, sy), (ex, ey)
    g = _geometry(s, DIRS[sa], e, DIRS[sb], mid)
    if g is None:
        return None
    prst, rot, fh, w, h, adj, swapped = g
    if prst == "straightConnector1":
        return [s, e]
    if swapped:
        s = e
    ux, uy = _map(rot, fh, (1, 0)), _map(rot, fh, (0, 1))

    def P(lx, ly):
        return (s[0] + ux[0] * lx + uy[0] * ly, s[1] + ux[1] * lx + uy[1] * ly)
    if prst == "bentConnector2":
        return [s, P(w, 0), P(w, h)]
    m = adj * w
    return [s, P(m, 0), P(m, h), P(w, h)]


def seg_hits(p, q, r, m=0.02):
    """線分 p→q が矩形 r=(x, y, w, h) の内側（縁から m 内側）を通るか。"""
    x0, y0, x1, y1 = r[0] + m, r[1] + m, r[0] + r[2] - m, r[1] + r[3] - m
    if x1 <= x0 or y1 <= y0:
        return False
    for k in range(41):
        t = k / 40
        px, py = p[0] + (q[0] - p[0]) * t, p[1] + (q[1] - p[1]) * t
        if x0 < px < x1 and y0 < py < y1:
            return True
    return False


def seg_overlap(p, q, u, v, tol=0.04):
    """2 本の線分が同じ線の上で重なるか（縦どうし・横どうしで、0.05in 以上重なる）。"""
    if abs(p[0] - q[0]) < 1e-3 and abs(u[0] - v[0]) < 1e-3 and abs(p[0] - u[0]) < tol:
        return min(max(p[1], q[1]), max(u[1], v[1])) - max(min(p[1], q[1]), min(u[1], v[1])) > 0.05
    if abs(p[1] - q[1]) < 1e-3 and abs(u[1] - v[1]) < 1e-3 and abs(p[1] - u[1]) < tol:
        return min(max(p[0], q[0]), max(u[0], v[0])) - max(min(p[0], q[0]), min(u[0], v[0])) > 0.05
    return False


def seg_cross(p, q, u, v):
    """縦の線分と横の線分が、端以外で交わるか。"""
    for (a, b), (c, d) in (((p, q), (u, v)), ((u, v), (p, q))):
        if abs(a[0] - b[0]) < 1e-3 and abs(c[1] - d[1]) < 1e-3:
            return (min(c[0], d[0]) + 0.02 < a[0] < max(c[0], d[0]) - 0.02
                    and min(a[1], b[1]) + 0.02 < c[1] < max(a[1], b[1]) - 0.02)
    return False


def link(slide, a, b, sa=None, sb=None, color=None, width=1.0, arrow=False, both=False, dash=False, mid=None):
    """図形 a の辺 sa と図形 b の辺 sb を、接続付きのコネクタでつなぐ。mid は曲がる位置の座標
    （横に出る線なら x、縦に出る線なら y）。曲げてつなげない配置（逆向き）は、接続付きの直線にする。"""
    import compose as c
    import helpers as _h
    if sa is None or sb is None:
        sa, sb = auto_sides(a, b)
    ia, sx, sy = site(a, sa)
    ib, ex, ey = site(b, sb)
    ds, de = DIRS[sa], DIRS[sb]
    g = _geometry((sx, sy), ds, (ex, ey), de, mid) or ("straightConnector1",) + (None,) * 5 + (False,)
    prst, rot, fh, w, h, adj, swapped = g
    if prst == "straightConnector1":
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(sx * EMU)), Emu(int(sy * EMU)),
                                          Emu(int(ex * EMU)), Emu(int(ey * EMU)))
    else:
        if swapped:
            (ia, sx, sy, a), (ib, ex, ey, b) = (ib, ex, ey, b), (ia, sx, sy, a)
        cx, cy = (sx + ex) / 2, (sy + ey) / 2
        conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, Emu(0), Emu(0), Emu(10), Emu(10))
        sp = conn._element.spPr
        xfrm = sp.find(qn("a:xfrm"))
        for k in ("flipH", "flipV", "rot"):
            if k in xfrm.attrib:
                del xfrm.attrib[k]
        if rot:
            xfrm.set("rot", str(int(rot * 60000)))
        if fh:
            xfrm.set("flipH", "1")
        xfrm.find(qn("a:off")).set("x", str(int((cx - w / 2) * EMU)))
        xfrm.find(qn("a:off")).set("y", str(int((cy - h / 2) * EMU)))
        xfrm.find(qn("a:ext")).set("cx", str(max(1, int(w * EMU))))
        xfrm.find(qn("a:ext")).set("cy", str(max(1, int(h * EMU))))
        geom = sp.find(qn("a:prstGeom"))
        geom.set("prst", prst)
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for gd in list(av):
            av.remove(gd)
        if adj is not None:
            gd = etree.SubElement(av, qn("a:gd"))
            gd.set("name", "adj1")
            gd.set("fmla", f"val {int(round(adj * 100000))}")
    # 接続（端点が緑になる）
    cnv = conn._element.find(".//" + qn("p:cNvCxnSpPr"))
    for tag in ("a:stCxn", "a:endCxn"):
        old = cnv.find(qn(tag))
        if old is not None:
            cnv.remove(old)
    # グループ（SVG のアイコンなど）には接続点が無いので、その端だけ接続しない
    if not a._element.tag.endswith("}grpSp"):
        st = etree.SubElement(cnv, qn("a:stCxn"))
        st.set("id", str(a.shape_id)); st.set("idx", str(ia))
    if not b._element.tag.endswith("}grpSp"):
        en = etree.SubElement(cnv, qn("a:endCxn"))
        en.set("id", str(b.shape_id)); en.set("idx", str(ib))
    conn.line.color.rgb = color or _h.MID_GRAY
    conn.line.width = Pt(width)
    c._strip_style(conn)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow or both:
        ln = conn.line._get_or_add_ln()
        ends = ["headEnd", "tailEnd"] if both else (["headEnd"] if swapped else ["tailEnd"])   # 順序は head → tail（スキーマの順）
        for t in ends:
            e = etree.SubElement(ln, qn("a:" + t))
            e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    return conn


def anchor(slide, x, y, w=0.02, h=0.02):
    """線の端を図形でないところ（束ね線の途中など）に置くための、見えない小さな点。使わずに済むなら使わない。"""
    import compose as c
    return c._rect(slide, x - w / 2, y - h / 2, w, h, fill=None)
