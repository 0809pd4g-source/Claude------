"""
compose.py — デザイン文法（配置 × 部品 × 関係記号 × 締め × ナビ表現）でスライドを描画する。

スライドを「1 パターン = 1 関数」ではなく、grammar/ の語彙の組み合わせとして描く。
  grid（配置）       : G-1 / G-11 / G-23 / G-32 / G-N3 / G-N4 / G-N5 / G-3W / G-TB / G-RH / G-L / G-2x2 / G-IO / G-SIDE / G-GRID / G-MATRIX / G-TL / G-FREE / G-HUB
  panes[].module（部品）: MODULES 参照（M-BULLETS ほか）
  connectors（関係記号）: R-CAUSE / R-SOLVE / R-ADD / R-MUL / R-EQ / R-VS / R-MUTUAL / R-CONV / R-BRACKET / R-ZOOM / R-PAR
  closing（締め）    : C-CONCLUSION / C-POINTS / C-NEXTACTION / C-NOTE / C-GOAL / C-TARGET / C-NONE
  devices（ナビ表現） : V-NAVTAB / V-INDEX / V-SCOPE / V-STAMP / V-STATUS

使い方（デッキ全体）:
    import helpers as _h; _h.set_lang("ja"); _h.load_theme("acn-technology-graphik")
    from compose import build_deck
    report = build_deck(outline, "out.pptx")          # outline は outline.json v2 形式（references/outline-schema.md）
使い方（1 枚）:
    slide = new_content_slide(prs, title=..., message=..., breadcrumb=...)
    render_body(slide, {"grid": "G-11", "panes": [...], "connectors": ["R-CAUSE"], "closing": {...}})
CLI:
    python compose.py outline.json out.pptx [--theme NAME]

部品ごとの data の形は references/compose-api.md を参照。
座標はインチの float、色は helpers の現在値（load_theme 後）を使う。
"""

import math
import os
import re
import sys

_SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)

from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from lxml import etree  # noqa: E402

import helpers as _h  # noqa: E402  属性は呼び出し時に参照する（load_theme 後の値を使うため）
import paths as _paths  # noqa: E402

GAP = 0.18          # ペイン間の標準間隔（人の資料の並びの間隔は中央値 0.13in、p75 0.32in）
GAP_CONN = 0.60     # 関係記号を置くときの間隔
HDR_H = 0.42        # ペイン見出しの高さ
CLOSE_H = 0.48      # 締めバーの高さ（人の資料の総括帯は本文より背が低い）
NOTE_H = 0.30       # 出典・注記の高さ
SIZES = (14, 12)    # 本文の文字サイズ候補（13pt は使わない）。見出し帯・強調文は 16pt まで使える


# ═══════════════════════════════════════════════════════════════
#  基本図形・テキスト
# ═══════════════════════════════════════════════════════════════
def _font():
    return _h.FONT


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, dash=False, shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(max(h, 0.01)))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s.shadow.inherit = False
    _strip_style(s)
    return s


PHASE_GAP = 0.08     # 工程見出しの五方向の矢印どうしの隙間（人の資料: 81% が五方向だけの並び、隙間の中央値 0.03・p75 0.10）


def homeplate(slide, x, y, w, h, fill=None, line=None, line_w=0.75, tip=None):
    """矢印：五方向（homePlate）。工程の見出しはすべてこの形を隙間 PHASE_GAP で並べる（シェブロンを重ねない）。
    矢じりの深さは高さの 0.35 倍で、並びの全段で同じ形になる。戻り値の .tip_in に矢じりの深さ（インチ）を入れる。"""
    s = _rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w, shape=MSO_AUTO_SHAPE_TYPE.PENTAGON)
    t = tip if tip is not None else h * 0.35
    s.adjustments[0] = min(t / max(min(w, h), 0.01), 0.9 * w / max(min(w, h), 0.01))
    s.tip_in = t
    return s


def _hline(slide, x, y, w, color=None, width=0.75, dash=False):
    """横の罫線を線（コネクタ）で引く。細い長方形で代用しない（人の資料は線で引き、後から編集しやすい）。"""
    return _connector(slide, x, y, x + w, y, color=color or _h.LIGHT_GRAY, width=width, dash=dash)


def _panel(slide, x, y, w, h, line=None):
    """塗らずに細い線で囲む箱。人の資料は本文の面を塗らない（紫の面積は中央値 0.8%、灰色は 0%）。"""
    return _rect(slide, x, y, w, h, fill=None, line=line or _h.LIGHT_GRAY, line_w=0.75)


def _strip_style(shape):
    """python-pptx が付ける p:style（テーマ accent1 参照）を外し、明示した色だけにする。"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _bodypr(tf, anchor="t", inset=(0.05, 0.03, 0.05, 0.03), wrap=True):
    tf.word_wrap = wrap
    bp = tf._txBody.find(qn("a:bodyPr"))
    bp.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for k, v in zip(("lIns", "tIns", "rIns", "bIns"), inset):
        bp.set(k, str(int(v * 914400)))


_MARK = re.compile(r"(\*\*|==)(.+?)\1")      # **太字** ／ ==太字＋濃い紫==


def _plain(text):
    """強調の記法を外した文（幅の見積もり用）。"""
    return _MARK.sub(r"\2", text)


def _runs(p, text, size, color, bold=False, emphasis=None, emph_color=None, mapped=True):
    """**太字** ／ ==太字＋濃い紫== の記法と emphasis 語リストを run に分けて書く。和欧間の空白は取る（tidy_ja）。
    人の資料の部分強調は 75% が色なしの太字。色を付けるのは 1 ブロック 1 句（==語==）だけ。"""
    emph_color = emph_color or _h.DARK_PURPLE
    text = _h.tidy_ja(text)
    parts = []
    pos = 0
    for m in _MARK.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(2), "color" if m.group(1) == "==" else True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    if emphasis:
        out = []
        pat = re.compile("(" + "|".join(re.escape(w) for w in sorted(emphasis, key=len, reverse=True)) + ")")
        for seg, em in parts:
            if em:
                out.append((seg, True))
                continue
            for piece in pat.split(seg):
                if piece:
                    out.append((piece, piece in emphasis))
        parts = out
    for seg, em in parts:
        for k, piece in enumerate(seg.split("\n")):     # 文中の改行（_balance が入れる）は行内改行にする
            if k:
                p.add_line_break()
            if not piece:
                continue
            r = p.add_run()
            r.text = piece
            r.font.size = Pt(_eff(size) if mapped else size)
            r.font.name = _font()
            r.font.bold = bool(bold or em)
            r.font.color.rgb = emph_color if em == "color" else color


def _marked_breaks(text, w_in, size, mapped):
    """**強調** を含む文: 記法を外した文で改行位置を決め、記法付きの文に同じ位置で改行を入れる
    （強調の途中で改行するときは、いったん閉じて次の行で開き直す）。"""
    plain = _plain(text)
    bold = sum(len(m.group(2)) for m in _MARK.finditer(text))
    out = _balance_at(plain, w_in / (1 + 0.08 * bold / max(len(plain), 1)), size, mapped)   # 太字は 8% 広い
    if out == plain or "**" in plain or "==" in plain:
        return text
    breaks, i, j = [], 0, 0
    while j < len(out):
        if out[j] == "\n":
            breaks.append(i)
            j += 1
            while i < len(plain) and plain[i] == " " and (j >= len(out) or out[j] != " "):
                i += 1
            continue
        if i < len(plain) and plain[i] == out[j]:
            i, j = i + 1, j + 1
        elif i < len(plain) and plain[i] == " ":
            i += 1
        else:
            return text
    bset, res, pi, k, inside = set(breaks), [], 0, 0, None
    while k < len(text):
        mk = text[k:k + 2]
        if mk in ("**", "==") and (inside is None or inside == mk):
            if inside is None and pi in bset:
                res.append("\n")
                bset.discard(pi)
            inside = None if inside else mk
            res.append(mk)
            k += 2
            continue
        if pi in bset:
            res.append(f"{inside}\n{inside}" if inside else "\n")
            bset.discard(pi)
            if text[k] == " ":
                k, pi = k + 1, pi + 1
                continue
        res.append(text[k])
        k, pi = k + 1, pi + 1
    return "".join(res).replace("****", "").replace("====", "")


def _balance_at(text, w_in, size, mapped):
    """_balance を、実際に描く文字サイズ（写像するか否か）で見積もる。泣き別れが無くても、語の途中で
    折り返す行（「申し／込み」「120／分」「サブエージェン／ト」）は _wrap_safe で語の切れ目に改行を移す。"""
    if ("**" in text or "==" in text) and "\n" not in text:
        return _marked_breaks(text, w_in, size, mapped)
    saved = _DECK["body"]
    _DECK["body"] = saved and mapped
    try:
        out = _balance(text, w_in, size)
        if out == text:
            return _wrap_safe(text, w_in, size)
        # 泣き別れを直した改行の前の部分も、語の途中で折り返さないように語の切れ目で改行する
        head, tail = out.split("\n", 1)
        return _wrap_safe(head, w_in, size) + "\n" + tail
    finally:
        _DECK["body"] = saved


_NO_LINE_START = set("、。，．）」』】ー・ぁぃぅぇぉっゃゅょァィゥェォッャュョ%％")


def _jtype(c):
    o = ord(c)
    if 0x3040 <= o <= 0x309F:
        return "hira"
    if 0x30A0 <= o <= 0x30FF:
        return "kata"
    if 0x4E00 <= o <= 0x9FFF or c == "々":
        return "kanji"
    if c.isascii() and (c.isalnum() or c in "$.,+#_-"):
        return "alnum"
    return "other"


def _good_break(text, i):
    """text[i-1] と text[i] の間で改行してよいか（語の途中・送り仮名・数と単位の間・行頭禁則を避ける）。"""
    a, b = text[i - 1], text[i]
    if b in _NO_LINE_START:
        return False
    if b in "をがはにでへもの" and _jtype(a) in ("hira", "kanji", "kata"):   # 助詞は前の語に付ける（「まで／を」「ごと／の」）
        return False
    if a in "全各新旧再非未同本当毎前後約計総両諸" and (i < 2 or _jtype(text[i - 2]) != "kanji") \
            and _jtype(b) in ("kata", "kanji", "alnum"):        # 接頭の 1 字は次の語に付ける（「全／チーム」「各／工場」）
        return False
    if a in " /" or b == " ":
        return True
    if a in _PARTICLE_END or b in "（「『【":
        return True
    ta, tb = _jtype(a), _jtype(b)
    if ta == tb and ta in ("alnum", "kata", "kanji"):
        return False
    if ta == "alnum" and tb in ("kanji", "hira"):      # 「120分」「Codeの」
        return False
    if ta == "hira" and tb == "kanji":                 # 「申し込み」の送り仮名は切らない。かなが 2 字以上続いた後なら切れ目
        run, j = 0, i - 1
        while j >= 0 and _jtype(text[j]) == "hira":
            run, j = run + 1, j - 1
        return run >= 2
    if ta in ("hira", "kanji", "kata") and tb in ("kata", "alnum") and ta != tb:
        return True
    if ta == "kata" and tb == "kanji":
        return True
    return False


def _wrap_safe(text, w_in, size):
    """折り返し位置が語の途中に来る行だけ、手前の語の切れ目で改行する（行数は増やさない）。"""
    if "\n" in text or "**" in text or len(text) > 200 or len(text) < 6:
        return text
    n0 = est_lines(text, w_in, size)
    for k in (0.97, 1.0):         # 安全側の幅で行が増えるなら、見積もりどおりの幅でもう一度
        out = _wrap_lines(text, max(w_in - 0.12, 0.2) * k, size)
        if out is None:
            return text
        if len(out) <= n0:
            return "\n".join(out)
    return text


def _wrap_lines(text, usable, size):
    """幅 usable で語の切れ目に改行した行のリスト（折り返さないなら None）。"""
    widths = [_char_w(c, size) for c in text]
    if sum(widths) <= usable:
        return None
    out, s = [], 0
    while s < len(text):
        acc, e = 0.0, s
        while e < len(text) and acc + widths[e] <= usable:
            acc += widths[e]
            e += 1
        if e >= len(text):
            out.append(text[s:])
            break
        e = max(e, s + 1)
        b = e if _good_break(text, e) else None
        if b is None:
            for k in range(e - 1, s, -1):
                if sum(widths[s:k]) < usable * 0.5:
                    break
                if _good_break(text, k):
                    b = k
                    break
        if b is None:
            # 窓（行の後ろ半分）に語の切れ目が無い（英単語・カタカナ語・漢字の熟語の途中）: 行が短くなっても、
            # 手前の語の切れ目で改行する（「Accen／ture」「各チーム／リーダー」を避ける）
            b = next((k for k in range(e - 1, s, -1) if _good_break(text, k)), None)
        b = b or e
        out.append(text[s:b].rstrip())
        s = b + (1 if b < len(text) and text[b] == " " else 0)
    # 見積もりは安全側（3% 狭い）なので、見積もりの折り返し位置が語の切れ目でも、実際の描画は 1 字先（語の途中）で
    # 折り返すことがある（「知見を共／有している」）。折り返す文は、語の切れ目の改行を常に明示する（行数は増やさない）
    return out


_PARTICLE_END = set("、。・のがをはにでとへもや）」』")


def _balance(text, w_in, size):
    """折り返した最終行が 1〜2 文字になる文を、文節の切れ目（助詞・読点の後、かなと漢字の境目）で改行する。
    人が手で直す「大き／い」「影響が大／きい」のような泣き別れを防ぐ。強調記法（**）を含む文は触らない。
    2 行の文は前後が釣り合う位置で、3 行以上の文は最終行が 3 文字以上になる、なるべく後ろの切れ目で改行する。"""
    if "\n" in text or "**" in text or len(text) > 160:
        return text
    n = est_lines(text, w_in, size)
    if n < 2 or not has_orphan([text], w_in, size):
        return text
    usable = max(w_in - 0.12, 0.2)
    cum = [0.0]
    for ch in text:
        cum.append(cum[-1] + _char_w(ch, size))
    best = None
    for i in range(2, len(text) - 2):
        a, b = text[i - 1], text[i]
        kana_to_kanji = 0x3040 <= ord(a) <= 0x30FF and ord(b) >= 0x4E00
        if not (a in _PARTICLE_END or kana_to_kanji):
            continue
        head, tail = cum[i], cum[-1] - cum[i]
        if n == 2:
            if head <= usable and tail <= usable:
                d = abs(head - tail)
                if best is None or d < best[0]:
                    best = (d, i)
        else:   # 前を n-1 行に収め、後ろ（最終行）を 3 文字分以上にする
            if head <= (n - 1) * usable and 3 * size / 72.0 <= tail <= usable:
                d = -head
                if best is None or d < best[0]:
                    best = (d, i)
    return text if best is None else text[:best[1]] + "\n" + text[best[1]:]


def _bullet(p, char="•", indent_in=0.18, level=0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int((indent_in * (level + 1)) * 914400)))
    pPr.set("indent", str(int(-indent_in * 914400)))
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", char)


def _textbox(slide, x, y, w, h, paras, size=14, color=None, anchor="t", align=None, bullet=None,
             space_after=6, bold=False, emphasis=None, inset=(0.05, 0.03, 0.05, 0.03), line_spacing=None):
    """paras: str のリスト（または (text, opts) のタプル）。"""
    color = color or _h.TEXT_BODY
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.1)), Inches(max(h, 0.1)))
    tf = tb.text_frame
    _bodypr(tf, anchor, inset)
    first = True
    for item in paras:
        text, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align or o.get("align"):
            p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        b = o.get("bullet", bullet)
        if b:
            _bullet(p, b, level=o.get("level", 0))
        sz, bd = o.get("size", size), o.get("bold", bold)
        tw = (w - inset[0] - inset[2] + 0.12 - (0.25 if b else 0)) / (1.08 if bd else 1.0)
        # 1 行分の高さしか無い小さな文字枠（軸ラベル・凡例・タグ）は、投影用でも文字を大きくしない（枠の幅が固定のため）
        _runs(p, _balance_at(_h.tidy_ja(text), tw, sz, h >= 0.45), sz, o.get("color", color), bd, emphasis, mapped=h >= 0.45)
    return tb


def _shape_text(shape, text, size=14, color=None, bold=False, align=PP_ALIGN.CENTER, anchor="m",
                inset=(0.05, 0.03, 0.05, 0.03), wrap=True, emphasis=None):
    tf = shape.text_frame
    _bodypr(tf, anchor, inset, wrap)
    lines = text.split("\n") if isinstance(text, str) else list(text)
    tw = (shape.width / 914400 - inset[0] - inset[2] + 0.12) / (1.08 if bold else 1.0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        line = _h.tidy_ja(line)
        mp = shape.height / 914400 >= 0.45
        _runs(p, _balance_at(line, tw, size, mp) if wrap else line, size, color or _h.TEXT_BODY, bold, emphasis, mapped=mp)


def _char_w(ch, size):
    """1 文字の幅（インチ）。字幅は helpers.char_em（Meiryo UI のフォントファイルから実測）。"""
    return _eff(size) / 72.0 * _h.char_em(ch)


def est_lines(text, w_in, size):
    """行数の見積もり。折り返す文は語の切れ目で改行したときの行数で数える（_wrap_safe と同じ数え方にして、
    語の切れ目の改行で 1 行増える文も枠の大きさの見積もりに入れる）。"""
    usable = max(w_in - 0.12, 0.2)
    total = 0
    for para in str(text).split("\n"):
        pl = _plain(para)
        width = sum(_char_w(c, size) for c in pl)
        n = max(1, math.ceil(width / usable))
        if n > 1 and 6 <= len(pl) <= 200:
            n = max(n, len(_wrap_lines(pl, usable * 0.97, size) or [pl]))
        total += n
    return total


def has_orphan(paras, w_in, size):
    """折り返した最終行が 1〜2 文字だけになる段落があるか（「…ディレクタ／ー」のような泣き別れ）。"""
    usable = max(w_in - 0.12, 0.2)
    for p in paras:
        for para in str(p if isinstance(p, str) else p[0]).split("\n"):
            width = sum(_char_w(c, size) for c in _plain(para))
            n = math.ceil(width / usable)
            if n >= 2 and width - (n - 1) * usable < 2.2 * size / 72.0:
                return True
    return False


def est_height(paras, w_in, size, space_after=6):
    lines = sum(est_lines(p if isinstance(p, str) else p[0], w_in, size) for p in paras)
    return lines * _eff(size) * 1.25 / 72.0 + len(paras) * space_after / 72.0 + 0.08


_CUR = {"sizes": SIZES}   # 描画中のスライドの本文サイズ候補（密度 high のスライドは 12pt から）
_DECK = {"mode": "read", "body": False, "boost": 0}
# 投影用デッキ（mode=present。LT・勉強会・講演）は、本文の部品の文字を 1 段ずつ大きく描く（12→14・14→16・16→18）。
# 部品のコードはそのまま（14 / 12pt で書いたまま）で、描画と寸法の見積もりの両方にこの写像を掛けるので、レイアウトは崩れない。
# boost は「中身が少なく枠の下が空く部品」を描き直すときの追加の段（_draw_filled）
_LADDER = (12, 14, 16, 18)


def _eff(size):
    if not _DECK["body"]:
        return size
    steps = (1 if _DECK["mode"] == "present" else 0) + _DECK["boost"]
    if not steps or size not in _LADDER:
        return size
    return _LADDER[min(_LADDER.index(size) + steps, len(_LADDER) - 1)]


_OVERFLOW = ("長い", "多い", "収まらない", "はみ出", "狭い")


def _xfrm(el):
    """図形の位置と大きさの要素（a:xfrm、表・グラフは p:xfrm）。"""
    x = el.find(qn("p:xfrm")) if el.tag.endswith("}graphicFrame") else el.find(".//" + qn("a:xfrm"))
    if x is None:
        return None, None
    off, ext = x.find(qn("a:off")), x.find(qn("a:ext"))
    if off is None or ext is None or ext.get("cy") is None:
        return None, None
    return off, ext


def _shape_bottom(el):
    off, ext = _xfrm(el)
    if off is None:
        return None
    return (int(off.get("y")) + int(ext.get("cy"))) / 914400


def _draw_filled(ctx, fn, r, pane):
    """部品を描き、中身が少なく枠の下が大きく空くなら、文字を 1 段（投影用は 2 段まで）大きくして描き直す。
    人の資料は捨てた分だけ主図を大きくし、紙面を下まで使う（描画のインク量は本文の下 1/3 でも中央値 0.32）。
    大きくして文字があふれる（「長い」「多い」の警告が出る）なら、1 つ前の大きさに戻す。"""
    tree = ctx.slide.shapes._spTree
    n0, w0, i0 = len(tree), len(ctx.warnings), len(ctx.icons)

    def draw(boost):
        for el in list(tree)[n0:]:
            tree.remove(el)
        del ctx.warnings[w0:]
        del ctx.icons[i0:]
        _DECK["boost"] = boost
        try:
            fn(ctx, r, pane)
        finally:
            _DECK["boost"] = 0

    def state():
        new = ctx.warnings[w0:]
        bots = [b for b in (_shape_bottom(el) for el in list(tree)[n0:]) if b is not None]
        free = (r[1] + r[3] - max(bots)) / r[3] if bots and r[3] > 0 else 0
        return any(k in m for m in new for k in _OVERFLOW), any("空く" in m for m in new) or free >= 0.28

    draw(0)
    over, sparse = state()
    best = drawn = 0
    for b in range(1, (2 if _DECK["mode"] == "present" else 1) + 1):
        if over or not sparse:
            break
        draw(b)
        drawn = b
        over, sparse = state()
        if over:
            break
        best = b
    if drawn != best:          # 最後に描いた大きさであふれたら、あふれない大きさで描き直す
        draw(best)
    # それでも下が大きく空くなら、要素の間を縦に広げて枠を使う（小さな図形は大きさを保って動かすだけ）
    els = list(tree)[n0:]
    # 縦の線（ガントの現在線など）は枠の下まで伸びるので、使った高さの計算から外す
    bots = [b for b in (_shape_bottom(el) for el in els if not el.tag.endswith("}cxnSp")) if b is not None]
    if bots and pane.get("module") in STRETCH_OK and not pane.get("no_stretch"):
        used = max(bots) - r[1]
        if used > 0.5 and (r[3] - used) / r[3] >= 0.22:
            _stretch(els, r[1], min(1.45, r[3] / used))


# 縦に広げてよい部品（行の格子でできていて、行の間を広げても空き箱にならないもの）。
# 箱の中に文字を入れる部品（行見出し・カード）は広げると箱の中が空くので対象外
STRETCH_OK = {"M-TABLE", "M-MATRIX-SYM", "M-GANTT", "M-SCORE", "M-PHASEGRID", "M-MATURITY", "M-TIMELINE"}


def _stretch(els, top, k):
    """縦方向に k 倍に広げる。高さ 0.6in 以上の図形・文字枠・表は高さも k 倍、それ未満（丸・点・棒・帯）は中心を動かすだけ。"""
    E = 914400
    for el in els:
        tag = el.tag.split("}")[1]
        xfrm = el.find(qn("p:grpSpPr") + "/" + qn("a:xfrm")) if tag == "grpSp" else \
            el.find(qn("p:xfrm")) if tag == "graphicFrame" else el.find(".//" + qn("a:xfrm"))
        if xfrm is None:
            continue
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if off is None or ext is None:
            continue
        y, h = int(off.get("y")) / E, int(ext.get("cy")) / E
        if tag in ("graphicFrame",) and el.find(".//" + qn("a:tbl")) is not None:
            for tr in el.iter(qn("a:tr")):
                tr.set("h", str(int(int(tr.get("h")) * k)))
            off.set("y", str(int((top + (y - top) * k) * E)))
            ext.set("cy", str(int(h * k * E)))
        elif tag == "cxnSp" or (h >= 0.6 and tag not in ("grpSp", "pic")):
            off.set("y", str(int((top + (y - top) * k) * E)))
            ext.set("cy", str(int(h * k * E)))
        else:
            cy = top + (y + h / 2 - top) * k
            off.set("y", str(int((cy - h / 2) * E)))


def fit_size(paras, w_in, h_in, sizes=None, space_after=6):
    """収まる最大の文字サイズを返す。収まらなければ (最小サイズ, False)。"""
    sizes = sizes or _CUR["sizes"]
    fits = [s for s in sizes if est_height(paras, w_in, s, space_after) <= h_in]
    for s in fits:                      # 収まるサイズのうち、行末に 1〜2 文字だけ残らない最大のもの
        if not has_orphan(paras, w_in, s):
            return s, True
    if fits:
        return fits[0], True
    return sizes[-1], False


class Ctx:
    """描画中の状態（警告・番号付けなど）。"""
    def __init__(self, slide, prs=None, spec=None):
        self.slide, self.prs, self.spec = slide, prs, spec or {}
        self.warnings = []
        self.emphasis = self.spec.get("emphasis")
        self.index = any(_dev_type(d) == "V-INDEX" for d in self.spec.get("devices", []))
        self.icons = []          # [(検索語, 素材 ID)] このスライドで使ったアイコン（同じ絵を 2 度使わない・一覧に残す）

    def warn(self, msg):
        self.warnings.append(msg)


def _dev_type(d):
    return d if isinstance(d, str) else d.get("type")


# ═══════════════════════════════════════════════════════════════
#  配置（grid）
# ═══════════════════════════════════════════════════════════════
def _split_h(box, ratios, gaps):
    x, y, w, h = box
    total = w - sum(gaps)
    out, cx = [], x
    s = float(sum(ratios))
    for i, r in enumerate(ratios):
        ww = total * r / s
        out.append((cx, y, ww, h))
        cx += ww + (gaps[i] if i < len(gaps) else 0)
    return out


def _split_v(box, ratios, gaps):
    x, y, w, h = box
    total = h - sum(gaps)
    out, cy = [], y
    s = float(sum(ratios))
    for i, r in enumerate(ratios):
        hh = total * r / s
        out.append((x, cy, w, hh))
        cy += hh + (gaps[i] if i < len(gaps) else 0)
    return out


HORIZONTAL = {"G-11", "G-23", "G-32", "G-N3", "G-N4", "G-N5", "G-3W", "G-IO", "G-SIDE"}
VERTICAL = {"G-TB"}


def grid_rects(grid, n, box, ratios=None, gaps=None):
    """配置 ID とペイン数から、各ペインの矩形 (x, y, w, h) を返す。"""
    n = max(n, 1)
    if grid in ("G-1", "G-RH", "G-MATRIX", "G-TL", "G-FREE", "G-HUB", "G-BLEED", "G-ZOOM") and n == 1:
        return [box]
    if grid in HORIZONTAL or (grid in ("G-1", "G-RH", "G-MATRIX", "G-TL", "G-FREE", "G-HUB") and n > 1):
        default = {"G-23": [2, 3], "G-32": [3, 2], "G-SIDE": [7, 3], "G-3W": [1, 1.4, 1], "G-IO": None}.get(grid)
        if grid == "G-IO":
            default = [1] * (n - 1) + [1.3]
        r = ratios or (default if default and len(default) == n else [1] * n)
        g = gaps or [GAP] * (n - 1)
        return _split_h(box, r, g)
    if grid in VERTICAL:
        r = ratios or [1] * n
        g = gaps or [GAP] * (n - 1)
        return _split_v(box, r, g)
    if grid == "G-L":
        x, y, w, h = box
        left_w = (w - GAP) * 0.36
        lefts = _split_v((x, y, left_w, h), [1] * (n - 1), [0.18] * (n - 2)) if n > 1 else []
        return lefts + [(x + left_w + (gaps[0] if gaps else GAP), y, w - left_w - (gaps[0] if gaps else GAP), h)]
    if grid in ("G-2x2", "G-GRID", "G-2x3"):
        cols = 2 if grid == "G-2x2" else 3 if grid == "G-2x3" else (n if n <= 4 else (3 if n in (5, 6, 9) else 4))
        rows = math.ceil(n / cols)
        gg = 0.14 if _CUR.get("max") else 0.22        # 格子の隙間（max は 3 段階の下限）
        rws = _split_v(box, [1] * rows, [gg] * (rows - 1))
        out = []
        for r in rws:
            out += _split_h(r, [1] * cols, [gg] * (cols - 1))
        return out[:n]
    return _split_h(box, [1] * n, [GAP] * (n - 1))


# ═══════════════════════════════════════════════════════════════
#  ペイン見出し
# ═══════════════════════════════════════════════════════════════
def pane_header(ctx, r, text, number=None, status=None, dark=False, style=None):
    """ペイン上部の見出し。残りの矩形を返す。
    style="rule"（既定）: 太字＋細い下線。人の資料は列見出しを塗らず「文字＋下線」にすることが多い
    style="band": 塗りの見出し帯（上位の見出し・強調したい見出しだけに使う）。dark=True で濃色"""
    x, y, w, h = r
    s = ctx.slide
    style = style or ("band" if dark or number is not None or status else "rule")
    if style == "rule":
        size, ok = fit_size([text], w - 0.1, HDR_H, (16, 14, 12))
        if not ok:
            ctx.warn(f"見出しが長い: {text[:20]}")
        _textbox(s, x, y, w, HDR_H - 0.04, [text], size, _h.DARKEST_PURPLE, bold=True, anchor="b",
                 inset=(0.02, 0, 0.02, 0.02), space_after=0)
        _hline(s, x, y + HDR_H - 0.01, w, _h.DARK_PURPLE, 1.5)
        return (x, y + HDR_H + 0.1, w, h - HDR_H - 0.1)
    fill = _h.DARKEST_PURPLE if dark else _h.LIGHTEST_PURPLE
    tcol = _h.WHITE if dark else _h.DARKEST_PURPLE
    band = _rect(s, x, y, w, HDR_H, fill=fill)
    tx = x
    if number is not None:
        d = 0.30
        b = _rect(s, x + 0.08, y + (HDR_H - d) / 2, d, d, fill=_h.DARK_PURPLE if not dark else _h.WHITE,
                  shape=MSO_AUTO_SHAPE_TYPE.OVAL)
        _shape_text(b, str(number), 12, _h.WHITE if not dark else _h.DARKEST_PURPLE, True,
                    inset=(0, 0, 0, 0), wrap=False)
        tx = x + 0.08 + d
    st_w = 0
    if status:
        st_w = 0.9
        col = {"完了": _h.DARK_PURPLE, "進行中": _h.CORE_PURPLE, "未着手": _h.MID_GRAY}.get(status, _h.MID_GRAY)
        tag = _rect(s, x + w - st_w - 0.06, y + 0.07, st_w, HDR_H - 0.14, fill=col)
        _shape_text(tag, status, 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    size, ok = fit_size([text], w - (tx - x) - st_w - 0.1, HDR_H, (14, 12))
    if not ok:
        ctx.warn(f"見出しが長い: {text[:20]}")
    _shape_text(band, text, size, tcol, True, PP_ALIGN.LEFT,
                inset=(tx - x + 0.08, 0.02, st_w + 0.1, 0.02))
    return (x, y + HDR_H + 0.08, w, h - HDR_H - 0.08)


# ═══════════════════════════════════════════════════════════════
#  部品（modules）
#  各関数: (ctx, r=(x,y,w,h), d=data) → None。d["header"] はペイン見出し（render_body 側で処理）
# ═══════════════════════════════════════════════════════════════
def _items(d, key="items"):
    v = d.get(key) or []
    return [i if isinstance(i, dict) else {"text": str(i)} for i in v]


def m_bullets(ctx, r, d):
    x, y, w, h = r
    paras = []
    for it in _items(d):
        paras.append((it.get("text", ""), {"bullet": "•"}))
        for sub in it.get("sub", []):
            paras.append((sub, {"bullet": "–", "level": 1, "size_delta": -2}))
    size, ok = fit_size(paras, w - 0.25, h, space_after=8)
    if not ok:
        ctx.warn("M-BULLETS: 文字量が枠を超える")
    paras = [(t, {**o, "size": size - 2 if o.get("size_delta") and size > 12 else size}) for t, o in paras]
    anchor = "m" if est_height(paras, w - 0.25, size, 8) < h * 0.6 else "t"
    _textbox(ctx.slide, x, y, w, h, paras, size, anchor=anchor, space_after=8, emphasis=ctx.emphasis)


def m_checklist(ctx, r, d):
    d2 = {"items": [{"text": "✓ " + it.get("text", "")} for it in _items(d)]}
    x, y, w, h = r
    paras = [it["text"] for it in d2["items"]]
    size, _ = fit_size(paras, w, h, space_after=10)
    _textbox(ctx.slide, x, y, w, h, paras, size, anchor="m", space_after=10, emphasis=ctx.emphasis)


def m_numlist(ctx, r, d):
    x, y, w, h = r
    items = _items(d)
    n = max(len(items), 1)
    gap = 0.14
    row_h = (h - gap * (n - 1)) / n
    d_badge = min(0.46, row_h * 0.8)
    for i, it in enumerate(items):
        ry = y + i * (row_h + gap)
        b = _rect(ctx.slide, x, ry + (row_h - d_badge) / 2, d_badge, d_badge, fill=_h.DARK_PURPLE,
                  shape=MSO_AUTO_SHAPE_TYPE.OVAL)
        _shape_text(b, str(it.get("number", i + 1)), 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        tx, tw = x + d_badge + 0.15, w - d_badge - 0.15
        paras = []
        if it.get("title"):
            paras.append((it["title"], {"bold": True, "color": _h.DARKEST_PURPLE, "space_after": 2}))
        if it.get("text"):
            paras.append(it["text"])
        size, ok = fit_size(paras, tw, row_h)
        if not ok:
            ctx.warn(f"M-NUMLIST: {i + 1} 行目が長い")
        _textbox(ctx.slide, tx, ry, tw, row_h, paras, size, anchor="m", space_after=2, emphasis=ctx.emphasis)


def m_rowhead(ctx, r, d):
    """行見出し（左に見出しセル・右に説明）。CB-17: 見出しは淡色、本文が主役。
    columns / lead / badge / detail があれば Why 型（dense.m_whyrows）で描く。"""
    import dense
    if dense.is_why(d):
        return dense.m_whyrows(ctx, r, d)
    x, y, w, h = r
    rows = d.get("rows") or []
    n = max(len(rows), 1)
    gap = 0.10
    # 見出し列の幅は一番長い見出しに合わせる（短い見出しに広い塗りの列を取ると、塗りの箱の中が空く）
    longest = max([sum(_char_w(c, 14) for c in str(rw.get("head", ""))) * 1.08 for rw in rows] + [0.6])
    head_w = d.get("head_w") or min(2.6, max(1.1, longest + 0.4), max(w * 0.3, longest + 0.4))
    bw0 = w - head_w - 0.2

    def _paras(row):
        body = row.get("items") or ([row["text"]] if row.get("text") else [])
        return [(t, {"bullet": "•"}) if row.get("items") else t for t in body]
    # 行の高さは中身の量に比例させる（全部同じ高さの箱にしない）。伸ばしすぎて箱の中が空かないよう上限を付ける
    need = [max(est_height(_paras(rw) or [" "], bw0 - 0.2, 14, 4), est_height([rw.get("head", "")], head_w - 0.2, 14, 0)) + 0.12
            for rw in rows] or [0.5]
    avail = h - gap * (n - 1)
    tot = sum(need)
    heights = [avail * k / tot for k in need] if tot else [avail / n] * n
    heights = [min(hh, nd * 1.5 + 0.15) for hh, nd in zip(heights, need)]
    if sum(heights) + gap * (n - 1) < h * 0.7:
        ctx.warn("M-ROWHEAD: 行の中身が少なく、下が空く（行を足すか、締め・注釈を足す）")
    # 見出しと本文の文字は全行で同じ大きさにする（行ごとに縮めると字の大きさが不揃いになる）
    hs = min([fit_size([rw.get("head", "")], head_w - 0.1, hh, (14, 12))[0] for rw, hh in zip(rows, heights)] or [14])
    size = min([fit_size(_paras(rw), bw0 - (0.2 if rw.get("items") else 0), hh, space_after=4)[0]
                for rw, hh in zip(rows, heights) if _paras(rw)] or [14])
    ry = y
    for i, row in enumerate(rows):
        row_h = heights[i]
        if i:
            ry += heights[i - 1] + gap
        em = row.get("emph")
        hc = _rect(ctx.slide, x, ry, head_w, row_h, fill=_h.CORE_PURPLE if em else _h.LIGHTEST_PURPLE)
        _shape_text(hc, row.get("head", ""), hs, _h.WHITE if em else _h.DARKEST_PURPLE, True, PP_ALIGN.LEFT,
                    inset=(0.12, 0.03, 0.08, 0.03))
        if i < len(rows) - 1:
            _hline(ctx.slide, x + head_w, ry + row_h + gap / 2, w - head_w, dash=True)
        paras = _paras(row)
        bw = bw0
        if not fit_size(paras, bw - (0.2 if row.get("items") else 0), row_h, (size,), space_after=4)[1]:
            ctx.warn(f"M-ROWHEAD: 行「{row.get('head', '')[:10]}」の説明が長い")
        _textbox(ctx.slide, x + head_w + 0.15, ry, bw, row_h, paras, size, anchor="m", space_after=4,
                 emphasis=ctx.emphasis)


def m_table(ctx, r, d, symbols=False):
    x, y, w, h = r
    cols = d.get("columns") or []
    rows = d.get("rows") or []
    nr, nc = len(rows) + 1, max(len(cols), 1)
    row_h = min(0.8, h / nr)
    gf = ctx.slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ("bandRow", "firstRow"):
        tblPr.set(k, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid
    ratios = d.get("col_widths") or ([1.6] + [1] * (nc - 1) if d.get("first_col_head", True) and nc > 2 else [1] * nc)
    s = sum(ratios)
    for j, rr in enumerate(ratios[:nc]):
        tbl.columns[j].width = Inches(w * rr / s)
    hl = d.get("highlight_col")
    hr = d.get("highlight_row")
    gray_rows = set(d.get("gray_rows") or [])
    tstyle = d.get("style", "rule")   # rule: 見出しは文字＋下線、行は点線で区切る（人の資料の型）/ grid: 塗りの見出し＋格子
    longest = max([len(str(c)) for row in rows for c in row] + [1])
    size = _CUR["sizes"][0] if longest <= 24 and nr <= 9 else _CUR["sizes"][-1]
    # 行の高さは中身の行数に比例させる（全行同じ高さにしない）。合計は枠に合わせて広げるが、広げすぎない
    colw = [w * rr / s for rr in ratios[:nc]]

    def _cell(i, j):
        return str(cols[j]) if i == 0 else (str(rows[i - 1][j]) if j < len(rows[i - 1]) else "")
    need = []
    for i in range(nr):
        ln = max(est_lines(_plain(_cell(i, j)) if _cell(i, j) != "^" else "", max(0.3, colw[j] - 0.16), size)
                 for j in range(nc))
        need.append(ln * _eff(size) * 1.3 / 72 + 0.1)
    tot = sum(need)
    k_ = min(h / tot, 1.7) if tot else 1
    heights = [max(0.3, n_ * max(1.0, k_)) for n_ in need]
    row_h = sum(heights) / nr
    # 列ごとの揃え: 数値だけの列は右、記号・短い値の列は中央、文章の列は左。見出しも同じ揃えにする
    num_re = re.compile(r"[-+−0-9.,%（）()人月万円億件秒倍]{1,12}")
    sym_re = re.compile(r"[◎○△×✓—－\s]{1,3}")
    col_align = []
    for j in range(nc):
        vals = [_plain(_cell(i, j)) for i in range(1, nr) if _cell(i, j) not in ("", "^")]
        if j > 0 and vals and all(num_re.fullmatch(v.replace(" ", "")) for v in vals):
            col_align.append(PP_ALIGN.RIGHT)
        elif j > 0 and vals and (symbols or all(sym_re.fullmatch(v) for v in vals)):
            col_align.append(PP_ALIGN.CENTER)
        else:
            col_align.append(PP_ALIGN.LEFT)
    for i in range(nr):
        tbl.rows[i].height = Inches(heights[i])
        for j in range(nc):
            cell = tbl.cell(i, j)
            text = _cell(i, j)
            if text == "^":     # 上のセルと結合する印（区分の縦帯など）
                text = ""
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = 1 + 2   # MSO_ANCHOR.MIDDLE
            if i == 0 and tstyle == "rule":
                cell.fill.background()
                col = _h.DARKEST_PURPLE
            elif i == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = _h.DARKEST_PURPLE
                col = _h.WHITE
            elif i - 1 in gray_rows:          # 今回の範囲外の行（役割分担表で、対象外の工程を灰色にする）
                cell.fill.solid(); cell.fill.fore_color.rgb = _h.OFF_WHITE
                col = _h.MID_GRAY
            elif (hl is not None and j == hl) or (hr is not None and i - 1 == hr):
                cell.fill.solid(); cell.fill.fore_color.rgb = _h.LIGHTEST_PURPLE
                col = _h.TEXT_BODY
            elif j == 0 and d.get("first_col_head", True) and nc > 2:
                cell.fill.solid(); cell.fill.fore_color.rgb = _h.OFF_WHITE
                col = _h.DARKEST_PURPLE
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = _h.WHITE
                col = _h.TEXT_BODY
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = col_align[j]
            sz = 18 if (symbols and i > 0 and j > 0 and len(text) <= 2) else size
            _runs(p, text, sz, col, bold=(i == 0 or (j == 0 and nc > 2)), emphasis=ctx.emphasis if i else None)
    _table_borders(tbl, nr, nc, tstyle)
    for j in range(nc):           # "^" の続く範囲を縦に結合する
        i = 1
        while i < nr:
            if _cell(i, j) == "^":
                st = i - 1
                while i < nr and _cell(i, j) == "^":
                    i += 1
                tbl.cell(st, j).merge(tbl.cell(i - 1, j))
            else:
                i += 1
    sc = sorted(d.get("scope_rows") or [])
    if sc:                        # 今回の範囲の行を紫の破線で囲み、右下にラベル
        ty = y + sum(heights[:sc[0] + 1])
        th = sum(heights[sc[0] + 1:sc[-1] + 2])
        _rect(ctx.slide, x - 0.04, ty - 0.02, w + 0.08, th + 0.04, fill=None, line=_h.CORE_PURPLE, line_w=2.0, dash=True)
        lab = d.get("scope_label", "今回のご提案範囲")       # 凡例として表の右下に置く（表の文字・メッセージに重ねない）
        lw = 0.3 + len(lab) * 0.17
        by = y + sum(heights) + 0.08
        _rect(ctx.slide, x + w - lw - 0.42, by + 0.07, 0.36, 0.17, fill=None, line=_h.CORE_PURPLE, line_w=1.5, dash=True)
        _textbox(ctx.slide, x + w - lw, by, lw, 0.3, [lab], 12, _h.TEXT_SUB, anchor="m", space_after=0)
    if row_h * nr < h * 0.6:
        ctx.warn("M-TABLE: 表が枠に対して小さい（行を増やすか枠を縮める）")


def _table_borders(tbl, nr, nc, style="grid"):
    hexc = str(_h.LIGHT_GRAY)

    def _ln(tcPr, side, w=None, color=None, dash=None):
        ln = etree.SubElement(tcPr, qn(side))
        if w is None:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
            return
        ln.set("w", str(int(w * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color or hexc)
        if dash:
            pd = etree.SubElement(ln, qn("a:prstDash")); pd.set("val", dash)
    for i in range(nr):
        for j in range(nc):
            tcPr = tbl.cell(i, j)._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                for e in tcPr.findall(qn(side)):
                    tcPr.remove(e)
                if style == "rule":
                    if side in ("a:lnL", "a:lnR") or (side == "a:lnT" and i == 0):
                        _ln(tcPr, side)
                    elif side == "a:lnT":
                        _ln(tcPr, side, 0.75, hexc, None if i == 1 else "sysDot")
                    elif i == 0:
                        _ln(tcPr, side, 1.5, str(_h.DARK_PURPLE))
                    else:
                        _ln(tcPr, side, 0.75, hexc, None if i == nr - 1 else "sysDot")
                    continue
                ln = etree.SubElement(tcPr, qn(side)); ln.set("w", "6350")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", hexc)
            # 線要素は fill 要素より前に置く必要がある
            fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
            for f in fills:
                tcPr.remove(f); tcPr.append(f)


_SYM_LEGEND = {"role": "◎ 主担当　○ 支援・確認　－ 対象外", "eval": "◎ 十分　○ 一部対応　△ 限定的　× 対象外"}


def m_matrix_sym(ctx, r, d):
    """M-TABLE と同じ＋記号のセル。凡例（legend）は必ず右上に置く。無ければ legend_kind（role|eval）の既定を使う。"""
    m_table(ctx, r, d, symbols=True)
    leg = d.get("legend")
    if not leg:
        used = {ch for row in d.get("rows") or [] for c in row for ch in str(c) if ch in "◎○△×－"}
        if used:
            leg = _SYM_LEGEND.get(d.get("legend_kind", "role" if "－" in used or "△" not in used else "eval"))
            ctx.warn(f"M-MATRIX-SYM: 凡例が無いので既定の凡例「{leg}」を付けた（意味が違うなら legend で指定する）")
    if leg:
        _textbox(ctx.slide, r[0], r[1] - 0.34, r[2], 0.3, [leg], 12, _h.TEXT_SUB, align=PP_ALIGN.RIGHT, anchor="m",
                 space_after=0)


def m_callout(ctx, r, d):
    x, y, w, h = r
    text = d.get("text", "")
    style = d.get("style", "box")
    if style == "key":
        size, _ = fit_size([text], w - 0.6, h * 0.7, (28, 24, 18))
        _rect(ctx.slide, x, y + h * 0.15, 0.10, h * 0.7, fill=_h.CORE_PURPLE)
        paras = [(text, {"bold": True, "color": _h.DARKEST_PURPLE, "size": size})]
        if d.get("sub"):
            paras.append((d["sub"], {"size": 14, "color": _h.TEXT_BODY}))
        _textbox(ctx.slide, x + 0.35, y, w - 0.35, h, paras, size, anchor="m", space_after=10, emphasis=ctx.emphasis)
        return
    paras = [(d["title"], {"bold": True, "color": _h.DARKEST_PURPLE})] if d.get("title") else []
    paras += [text] if text else []
    paras += [(t, {"bullet": "•"}) for t in d.get("items", [])]
    size, ok = fit_size(paras, w - 0.4, h - 0.2)
    # 枠は中身の高さに合わせ、枠の中を空けない（余りは上下に振り分ける）
    need = est_height(paras, w - 0.4, size) + 0.3
    if need < h * 0.8:
        y, h = y + (h - need) / 2, need
    box = _rect(ctx.slide, x, y, w, h, fill=_h.WHITE, line=_h.CORE_PURPLE, line_w=1.5)
    if not ok:
        ctx.warn("M-CALLOUT: 文字量が多い")
    _textbox(ctx.slide, x + 0.15, y + 0.08, w - 0.3, h - 0.16, paras, size, anchor="m", emphasis=ctx.emphasis)


def m_quote(ctx, r, d):
    x, y, w, h = r
    _textbox(ctx.slide, x, y, 0.8, 0.9, ["“"], 44, _h.CORE_PURPLE, bold=True)
    size, _ = fit_size([d.get("text", "")], w - 1.0, h - 0.7, (24, 18, 14))
    paras = [(d.get("text", ""), {"size": size, "color": _h.DARKEST_PURPLE, "bold": True})]
    if d.get("by"):
        paras.append((f"— {d['by']}", {"size": 14, "color": _h.MID_GRAY}))
    _textbox(ctx.slide, x + 0.9, y, w - 0.9, h, paras, size, anchor="m", space_after=12)


def _card(ctx, r, title, body_paras, icon=None, accent="top", fill=None, rich_body=None):
    if rich_body is not None:
        import dense
        x, y, w, h = r
        _card(ctx, r, None, [], icon, accent, fill)          # 枠とアイコンだけ描く（文字は下でリッチ本文に）
        top = y + 0.16 + ((min(0.6, h * 0.28) + 0.08) if icon else 0)
        ix = x + (0.22 if accent == "left" else 0.15)
        body = ([{"h": title}] if title else []) + list(rich_body if isinstance(rich_body, list) else [rich_body])
        dense.rich_box(ctx, ix - 0.04, top, w - (ix - x) - 0.1, y + h - top - 0.08, body,
                       warn_label=f"カード「{(title or '')[:12]}」")
        return
    x, y, w, h = r
    s = ctx.slide
    if fill:
        _rect(s, x, y, w, h, fill=fill)
    elif accent != "left":       # 左の縦線 1 本でカードにする（4 辺を囲まない）
        _panel(s, x, y, w, h)
    if accent == "top":
        _rect(s, x, y, w, 0.05, fill=_h.CORE_PURPLE)
    elif accent == "left":
        _rect(s, x, y, 0.07, h, fill=_h.CORE_PURPLE)
    ix = x + (0.22 if accent == "left" else 0.15)
    iw = w - (ix - x) - 0.15
    cy = y + 0.16
    if icon:
        isz = min(0.6, h * 0.28)
        if _place_visual(ctx, icon, (ix, cy, isz, isz)) is not None:
            cy += isz + 0.08
    paras = []
    if title:
        paras.append((title, {"bold": True, "color": _h.DARKEST_PURPLE, "space_after": 6}))
    paras += body_paras
    size, ok = fit_size(paras, iw, y + h - cy - 0.1)
    if not ok:
        ctx.warn(f"カード「{(title or '')[:12]}」の文字量が多い")
    _textbox(s, ix, cy, iw, y + h - cy - 0.1, paras, size, anchor="t", space_after=6, emphasis=ctx.emphasis)


def m_card(ctx, r, d):
    import dense
    if dense.is_strength(d):          # 番号タブ＋紫の見出し＋灰の面の要点（5 つの強み型）
        return dense.m_strength(ctx, r, d)
    if d.get("body") is not None:     # リッチ本文（小見出し・箇条書き 2 段・結論）
        _card(ctx, r, d.get("title"), [], d.get("icon"), rich_body=d["body"])
        return
    body = ([d["text"]] if d.get("text") else []) + [(t, {"bullet": "•"}) for t in d.get("items", [])]
    _card(ctx, r, d.get("title"), body, d.get("icon"))


def m_cardstack(ctx, r, d):
    x, y, w, h = r
    cards = d.get("cards") or []
    n = max(len(cards), 1)
    gap = 0.14
    ch = (h - gap * (n - 1)) / n
    for i, c in enumerate(cards):
        body = [(t, {"bullet": "•"}) for t in c.get("items", [])] + ([c["text"]] if c.get("text") else [])
        _card(ctx, (x, y + i * (ch + gap), w, ch), c.get("title"), body, c.get("icon"), accent="left")


def m_headcol(ctx, r, d):
    x, y, w, h = r
    if d.get("body") is not None:     # リッチ本文
        import dense
        _panel(ctx.slide, x, y, w, h)
        dense.rich_box(ctx, x + 0.08, y + 0.1, w - 0.16, h - 0.2, d["body"], warn_label="M-HEADCOL")
        return
    paras = [(t, {"bullet": "•"}) for t in d.get("items", [])] + ([d["text"]] if d.get("text") else [])
    _panel(ctx.slide, x, y, w, h)
    if d.get("icon") and _place_visual(ctx, d["icon"], (x + w / 2 - 0.3, y + 0.12, 0.6, 0.6)) is not None:
        y, h = y + 0.8, h - 0.8
    size, ok = fit_size(paras, w - 0.45, h - 0.2, space_after=8)
    if not ok:
        ctx.warn("M-HEADCOL: 列の文字量が多い")
    _textbox(ctx.slide, x + 0.12, y + 0.1, w - 0.24, h - 0.2, paras, size, anchor="t", space_after=8,
             emphasis=ctx.emphasis)


def m_badgecol(ctx, r, d):
    x, y, w, h = r
    s = ctx.slide
    paras = []
    if d.get("title"):
        paras.append((d["title"], {"bold": True, "color": _h.DARKEST_PURPLE, "align": PP_ALIGN.CENTER, "size": 18}))
    if d.get("text"):
        paras.append((d["text"], {"align": PP_ALIGN.CENTER}))
    paras += [(t, {"bullet": "•"}) for t in d.get("items", [])]
    head = min(0.8, w * 0.4, h * 0.3) if d.get("icon") else 0.62
    size, ok = fit_size(paras, w - 0.2, h - head - 0.12)
    if not ok:
        ctx.warn("M-BADGECOL: 文字量が多い")
    body_h = min(h - head - 0.12, est_height(paras, w - 0.2, size) + 0.1)
    top = y + max(0.0, (h - head - 0.12 - body_h) / 2)   # 縦中央に寄せる（並列の記号と高さを揃える）
    if d.get("icon") and _place_visual(ctx, d["icon"], (x + (w - head) / 2, top, head, head)) is not None:
        pass
    else:
        b = _rect(s, x + (w - head) / 2, top, head, head, fill=_h.DARK_PURPLE, shape=MSO_AUTO_SHAPE_TYPE.OVAL)
        _shape_text(b, str(d.get("badge", "")), 18, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    top += head + 0.12
    _textbox(s, x, top, w, body_h, paras, size, anchor="t", space_after=6, emphasis=ctx.emphasis)


def m_kpi(ctx, r, d, vertical=False):
    x, y, w, h = r
    tiles = d.get("tiles") or [d]
    n = max(len(tiles), 1)
    rects = _split_v(r, [1] * n, [0.15] * (n - 1)) if vertical else _split_h(r, [1] * n, [0.2] * (n - 1))
    # 数字の大きさは並ぶタイルで揃える（桁の多いタイルだけ小さくしない）。emph のタイルがあれば、それ以外は黒字にする
    def _vs(t, th):
        val = str(t.get("value", ""))
        return 44 if len(val) <= 5 and th > 1.6 else (36 if len(val) <= 7 else 28)
    vsize = min(_vs(t, th) for (_, _, _, th), t in zip(rects, tiles))
    has_emph = any(t.get("emph") for t in tiles)
    for (tx, ty, tw, th), t in zip(rects, tiles):
        val = str(t.get("value", ""))
        vcol = _h.CORE_PURPLE if (t.get("emph") or not has_emph) else _h.TEXT_BODY
        paras = [(val + (t.get("unit") or ""), {"size": vsize, "bold": True, "color": vcol, "space_after": 2})]
        if t.get("label"):
            paras.append((t["label"], {"bold": True, "color": _h.DARKEST_PURPLE, "size": 14}))
        if t.get("detail"):
            paras.append((t["detail"], {"size": 12, "color": _h.TEXT_BODY}))
        _textbox(ctx.slide, tx + 0.2, ty, tw - 0.3, th, paras, 14, anchor="m", space_after=4)
        ch = min(th, sum(est_height([p_[0]], tw - 0.3, p_[1].get("size", 14), 4) for p_ in paras))
        _rect(ctx.slide, tx, ty + (th - ch) / 2, 0.06, ch, fill=vcol if has_emph else _h.CORE_PURPLE)   # 縦線は数字と説明の高さだけ


def m_chevron(ctx, r, d):
    x, y, w, h = r
    s = ctx.slide
    steps = _items(d, "steps")
    n = max(len(steps), 1)
    has_text = any(st.get("text") or st.get("items") for st in steps)
    ch_h = min(0.75, h * 0.28) if has_text else min(0.9, h)
    gap = PHASE_GAP
    cw = (w - gap * (n - 1)) / n
    detail_h = 0.0
    if has_text:   # 説明ボックスは中身に合わせた高さにし、まとまり全体を縦中央に置く（下半分の空白を避ける）
        need = max(est_height([(t, {}) for t in st.get("items", [])]
                              + ([st["text"]] if st.get("text") and st.get("label") else []), cw - 0.3, 14)
                   for st in steps)
        detail_h = min(h - ch_h - 0.18, max(need + 0.5, (h - ch_h - 0.18) * 0.55))
    group = ch_h + (0.18 + detail_h if has_text else 0)
    cy = y + (h - group) / 2
    # 説明の文字は全工程で同じ大きさにする（工程ごとに縮めない）
    sizes_all = [fit_size(([st["text"]] if st.get("text") and st.get("label") else []) +
                          [(t, {"bullet": "•"}) for t in st.get("items", [])], cw - 0.3, detail_h - 0.1,
                          (12,) if n >= 6 else None)[0] for st in steps] if has_text else [14]
    for i, st in enumerate(steps):
        cx = x + i * (cw + gap)
        # 強調する工程があれば、その工程だけ紫にして他は灰色にする（全部同じ濃い紫だと強調が効かない）
        any_hi = any(x_.get("highlight") for x_ in steps)
        hi = st.get("highlight")
        col = _h.CORE_PURPLE if hi else (_h.LIGHT_GRAY if any_hi else _h.DARK_PURPLE)
        c = homeplate(s, cx, cy, cw, ch_h, fill=col)       # 全工程を五方向の矢印で、隙間を空けて並べる
        label = st.get("label") or st.get("title") or st.get("text", "")
        tw = sum(_char_w(ch, 14) for ch in label)   # 矢羽根の文字領域は見た目より狭いので、短いラベルは折り返さない
        size = 14 if tw <= (cw - c.tip_in) * 0.85 else 12
        _shape_text(c, label, size, _h.TEXT_BODY if (any_hi and not hi) else _h.WHITE, True,
                    inset=(0.08, 0.02, 0.04, 0.02), wrap=tw > (cw - c.tip_in) * 0.95)
        if has_text:
            paras = ([st["text"]] if st.get("text") and st.get("label") else []) + [(t, {"bullet": "•"}) for t in st.get("items", [])]
            ty = cy + ch_h + 0.18
            th = detail_h
            size, ok = fit_size(paras, cw - 0.3, th - 0.1, (min(sizes_all),))
            if not ok:
                ctx.warn(f"M-CHEVRON: 工程{i + 1}の説明が長い")
            _textbox(s, cx + 0.06, ty + 0.06, cw - 0.12, th - 0.12, paras, size, anchor="t", space_after=6,
                     emphasis=ctx.emphasis)


def m_phase(ctx, r, d):
    """フェーズ列: 上にフェーズ帯（矢羽根）、下に各フェーズの実施事項。"""
    phases = d.get("phases") or []
    steps = [{"label": p.get("label", ""), "items": p.get("items", []), "highlight": p.get("highlight")} for p in phases]
    m_chevron(ctx, r, {"steps": steps})


def m_flow(ctx, r, d):
    """処理フロー: ボックス＋矢印。5 個を超えたら 2 段に折り返す。style: "icon" はアイコン＋ラベルの概念図（figures.py）。"""
    if d.get("style") == "icon":
        import figures
        return figures.m_iconflow(ctx, r, d)
    x, y, w, h = r
    s = ctx.slide
    nodes = _items(d, "nodes")
    n = len(nodes)
    if n == 0:
        return
    per_row = n if n <= 5 else math.ceil(n / 2)
    rows = math.ceil(n / per_row)
    agap = 0.45
    bw = (w - agap * (per_row - 1)) / per_row
    rh = (h - 0.4 * (rows - 1)) / rows
    bh = min(rh, 1.6 if any(nd.get("text") for nd in nodes) else 0.9)
    flow_boxes, wraps = [], []
    for i, nd in enumerate(nodes):
        row, col = divmod(i, per_row)
        bx = x + col * (bw + agap)
        by = y + row * (rh + 0.4) + (rh - bh) / 2
        hl = nd.get("highlight")
        box = _rect(s, bx, by, bw, bh, fill=_h.DARKEST_PURPLE if hl else _h.LIGHTEST_PURPLE)
        flow_boxes.append(box)
        paras = [(nd.get("label", nd.get("text", "")), {"bold": True, "color": _h.WHITE if hl else _h.DARKEST_PURPLE})]
        if nd.get("label") and nd.get("text"):
            paras.append((nd["text"], {"color": _h.WHITE if hl else _h.TEXT_BODY, "size": 12}))
        size, ok = fit_size(paras, bw - 0.15, bh)
        if not ok:
            ctx.warn(f"M-DIAGRAM-FLOW: ノード{i + 1}が長い")
        tf_paras = [(t, {**o, "align": PP_ALIGN.CENTER, "size": min(o.get("size", size), size)}) for t, o in paras]
        _textbox(s, bx, by, bw, bh, tf_paras, size, anchor="m", space_after=2)
        if col < per_row - 1 and i < n - 1:
            _arrow_h(s, bx + bw + 0.08, by + bh / 2, agap - 0.16)
        elif i < n - 1:   # 折り返し（下段の先頭へ）。次の箱を描いてから接続付きのカギ線でつなぐ
            wraps.append(i)
    import connect
    for i in wraps:
        connect.link(s, flow_boxes[i], flow_boxes[i + 1], "b", "t", _h.MID_GRAY, 1.75, arrow=True)


def m_layer(ctx, r, d):
    x, y, w, h = r
    layers = d.get("layers") or []
    n = max(len(layers), 1)
    gap = 0.08
    lh = (h - gap * (n - 1)) / n
    longest = max([sum(_char_w(c, 14) for c in L.get("label", "")) for L in layers] + [0.8])
    lab_w = min(w * 0.45, max(w * 0.22, longest + 0.35))
    bodies = []
    # 箇条にするか「 / 」でつなぐかは全層でそろえる（層ごとに書き方を変えない）
    joined = all(len(" ".join(L.get("items") or [])) < 36 for L in layers)
    for L in layers:
        paras = [L["text"]] if L.get("text") else []
        if L.get("items"):
            paras = ["　/　".join(L["items"])] if joined else [(t, {"bullet": "•"}) for t in L["items"]]
        bodies.append(paras)
    body_size = min([fit_size(p, w - lab_w - 0.3, lh)[0] for p in bodies if p] + [_CUR["sizes"][0]])   # 全層で同じサイズ
    # 層の高さは中身に合わせる（伸ばしすぎると箱の中が空く）。余った分は下に残し、警告で知らせる
    need = max([est_height(p, w - lab_w - 0.3, body_size, 2) for p in bodies if p] + [0.4])
    grow = 1.8 if _DECK["mode"] == "present" else 1.3      # 投影用は主図を大きく（層を高く）する
    lh_fit = min(lh, max(0.55, need * grow + 0.15))
    if (lh - lh_fit) * n > h * 0.3:
        ctx.warn("M-LAYER: 層の中身が少なく、下が空く（各層に根拠や例を足すか、別の部品と組む）")
    lh = lh_fit
    for i, L in enumerate(layers):
        ly = y + i * (lh + gap)
        # 層の見出しは同じ濃い紫でそろえる（下の層ほど明るくすると、強調していない層が目立つ）。強調する層だけ明るい紫
        shade = _h.CORE_PURPLE if L.get("emph") else _h.DARKEST_PURPLE
        lab = _rect(ctx.slide, x, ly, lab_w, lh, fill=shade)
        lsize = 14 if sum(_char_w(c, 14) for c in L.get("label", "")) <= lab_w - 0.2 else 12
        _shape_text(lab, L.get("label", ""), lsize, _h.WHITE, True, wrap=False, inset=(0.05, 0, 0.05, 0))
        body = _panel(ctx.slide, x + lab_w, ly, w - lab_w, lh)
        _textbox(ctx.slide, x + lab_w + 0.15, ly, w - lab_w - 0.25, lh, bodies[i], body_size, anchor="m", space_after=2,
                 emphasis=ctx.emphasis)


def m_gantt(ctx, r, d):
    """ガント（CB-18: バー長は期間に比例）。tasks: [{label, start, end, group?, highlight?}]、start/end は periods の添字。"""
    if any(t.get("group") for t in d.get("tasks") or []) or any(d.get(k) for k in ("phases", "milestones", "legend", "scope")):
        import patterns          # 作業の流れ×月（行見出し・バーの中の作業名・凡例・マイルストーン・範囲の枠）
        return patterns.gantt_ws(ctx, r, d)
    x, y, w, h = r
    s = ctx.slide
    periods = d.get("periods") or []
    tasks = d.get("tasks") or []
    lab_w = min(3.0, w * 0.27)
    np_ = max(len(periods), 1)
    pw = (w - lab_w) / np_
    hdr = 0.40
    for j, p in enumerate(periods):
        c = _rect(s, x + lab_w + j * pw, y, pw, hdr, fill=_h.DARKEST_PURPLE, line=_h.WHITE, line_w=0.5)
        _shape_text(c, str(p), 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    nt = max(len(tasks), 1)
    rh = min(0.8, (h - hdr) / nt)
    lab_size = 14 if all(sum(_char_w(c, 14) for c in t.get("label", "")) <= lab_w - 0.15 for t in tasks) else 12
    for i, t in enumerate(tasks):
        ty = y + hdr + i * rh
        _hline(s, x, ty + rh - 0.005, w, dash=True)
        _textbox(s, x, ty, lab_w - 0.1, rh, [t.get("label", "")], lab_size, anchor="m", space_after=0)
        st, en = float(t.get("start", 0)), float(t.get("end", t.get("start", 0)))
        bx = x + lab_w + st * pw + 0.04
        bw = (en - st + 1) * pw - 0.08
        # 人の資料のガント: タスクはホームベース形。強調するものだけ濃く塗り、ほかは白地に細枠
        on = t.get("highlight")
        _rect(s, bx, ty + rh * 0.22, bw, rh * 0.56, fill=_h.CORE_PURPLE if on else _h.WHITE,
              line=None if on else _h.DARK_PURPLE, line_w=1.0,
              shape=MSO_AUTO_SHAPE_TYPE.PENTAGON if bw > 0.6 else MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        if t.get("note"):      # 注記はバーの右。右に入らなければ左、左にも入らなければバーの中
            need = sum(_char_w(c, 12) for c in t["note"]) + 0.2
            right = x + w - (bx + bw) - 0.05
            left = bx - (x + lab_w) - 0.05
            if right >= need or (left < need and bw - 0.25 < need):
                _textbox(s, bx + bw + 0.05, ty, max(0.5, right), rh, [t["note"]], 12, _h.MID_GRAY, anchor="m",
                         space_after=0)
            elif left >= need:
                _textbox(s, bx - 0.05 - need, ty, need, rh, [t["note"]], 12, _h.MID_GRAY, anchor="m",
                         align=PP_ALIGN.RIGHT, space_after=0)
            else:
                _textbox(s, bx + 0.05, ty, bw - 0.25, rh, [t["note"]], 12, _h.WHITE if on else _h.DARK_PURPLE,
                         anchor="m", space_after=0)
    if d.get("now") is not None:
        nx = x + lab_w + float(d["now"]) * pw
        _connector(s, nx, y + hdr, nx, y + hdr + nt * rh, _h.PINK, width=1.5, dash=True)
        _textbox(s, nx - 0.4, y + hdr + nt * rh, 0.8, 0.28, ["現在"], 12, _h.PINK, align=PP_ALIGN.CENTER)


def m_tree(ctx, r, d):
    """ロジックツリー（イシューツリー・ドライバーツリー）。段の数は自由。
    {"root": "売上を伸ばす", "children": [{"label": "客数", "children": ["新規", {"label": "既存", "note": "..", "emph": true}]}, ..],
     "direction": "right|down", "note_head": "示唆"}
    right（既定）: 左→右。葉に note があれば右端にコメントの列を作り、葉の行に揃える。
    down: 上→下。組織図と同じ結線（親の下辺 → 横の束ね線 → 子の上辺）。"""
    if d.get("style") == "kpi":
        import patterns          # KPI の木（値の列・具体例の箇条書き・行の間の点線）
        return patterns.kpi_tree(ctx, r, d)
    x, y, w, h = r
    s = ctx.slide

    def norm(n):
        n = n if isinstance(n, dict) else {"label": str(n)}
        return {**n, "children": [norm(k) for k in (n.get("children") or []) if k is not None]}

    root = norm({"label": d.get("root", ""), "children": d.get("children") or [], "emph": d.get("root_emph")})
    leaves, depth = [], [0]

    def walk(n, lv):
        n["_lv"] = lv
        depth[0] = max(depth[0], lv)
        if not n["children"]:
            leaves.append(n)
        for k in n["children"]:
            walk(k, lv + 1)

    walk(root, 0)
    D = depth[0] + 1
    nleaf = max(len(leaves), 1)

    def style(n):
        if n["_lv"] == 0:
            return _h.DARKEST_PURPLE, _h.WHITE, True
        if n.get("emph"):
            return _h.DARK_PURPLE, _h.WHITE, True
        if n["_lv"] == 1 and D > 2:
            return _h.LIGHTEST_PURPLE, _h.DARKEST_PURPLE, True
        return None, _h.TEXT_BODY, False

    def box(n, bx, by, bw, bh):
        fill, col, bold = style(n)
        b = _rect(s, bx, by, bw, bh, fill=fill, line=None if fill else _h.LIGHT_GRAY, line_w=0.75)
        sz, _ = fit_size([n["label"]], bw - 0.1, bh, (16, 14, 12) if n["_lv"] == 0 else (14, 12))
        _shape_text(b, n["label"], sz, col, bold, PP_ALIGN.CENTER if d.get("direction") == "down" or n["_lv"] == 0
                    else PP_ALIGN.LEFT, inset=(0.08, 0.02, 0.05, 0.02))
        n["_box"] = (bx, by, bw, bh)
        n["_shp"] = b

    import connect

    def links(n, sa, sb):
        """親 → 子を接続付きのカギ線で（子ごとに 1 本。重なった部分が束ね線に見える）。"""
        for k in n["children"]:
            mid = None
            if sa == "b":
                mid = n["_box"][1] + n["_box"][3] + (k["_box"][1] - n["_box"][1] - n["_box"][3]) / 2
            else:
                mid = n["_box"][0] + n["_box"][2] + (k["_box"][0] - n["_box"][0] - n["_box"][2]) / 2
            connect.link(s, n["_shp"], k["_shp"], sa, sb, _h.MID_GRAY, 1.0, mid=mid)
            links(k, sa, sb)

    if d.get("direction") == "down":
        lh = min(0.9, h / (D * 1.5))                                  # 段の高さと段の間を、枠の高さに合わせて広げる
        gap_y = min(1.1, (h - lh * D) / max(D - 1, 1))
        slot = w / nleaf
        bw = min(2.3, slot - 0.15)
        for i, lf in enumerate(leaves):
            lf["_cx"] = x + slot * (i + 0.5)

        def cx_of(n):
            if not n["children"]:
                return n["_cx"]
            xs = [cx_of(k) for k in n["children"]]
            n["_cx"] = (xs[0] + xs[-1]) / 2
            return n["_cx"]

        cx_of(root)

        def draw(n):
            by = y + n["_lv"] * (lh + gap_y)
            box(n, n["_cx"] - bw / 2, by, bw, lh)
            for k in n["children"]:
                draw(k)

        draw(root)
        links(root, "b", "t")
        return

    has_note = any(lf.get("note") for lf in leaves)
    note_w = w * 0.3 if has_note else 0
    tw = w - note_w - (0.3 if has_note else 0)
    gap_x = 0.45
    colw = (tw - gap_x * (D - 1)) / D
    hdr = 0.42 if has_note and d.get("note_head") else 0
    top = y + hdr
    lh = min(0.8, (h - hdr - 0.1 * (nleaf - 1)) / nleaf)
    gap = (h - hdr - lh * nleaf) / max(nleaf - 1, 1) if nleaf > 1 else 0
    for i, lf in enumerate(leaves):
        lf["_cy"] = top + i * (lh + gap) + lh / 2

    def cy_of(n):
        if not n["children"]:
            return n["_cy"]
        ys = [cy_of(k) for k in n["children"]]
        n["_cy"] = (ys[0] + ys[-1]) / 2
        return n["_cy"]

    cy_of(root)

    def draw_r(n):
        bx = x + n["_lv"] * (colw + gap_x)
        bh = min(h - hdr, max(lh, 1.0)) if n["_lv"] == 0 else lh
        box(n, bx, n["_cy"] - bh / 2, colw, bh)
        for k in n["children"]:
            draw_r(k)

    draw_r(root)
    links(root, "r", "l")
    if has_note:
        nx = x + w - note_w
        if hdr:
            _textbox(s, nx, y, note_w, hdr - 0.04, [d["note_head"]], 14, _h.DARKEST_PURPLE, bold=True, anchor="b",
                     space_after=0)
            _hline(s, nx, y + hdr - 0.01, note_w, _h.DARK_PURPLE, 1.5)
        for lf in leaves:
            if not lf.get("note"):
                continue
            bx, by, bw, bh = lf["_box"]
            _connector(s, bx + bw + 0.05, lf["_cy"], nx - 0.08, lf["_cy"], _h.LIGHT_GRAY, 0.75, dash=True)
            sz, ok = fit_size([lf["note"]], note_w, lh + gap * 0.8, None, 0)
            if not ok:
                ctx.warn(f"M-TREE: 「{lf['label']}」のコメントが長い")
            _textbox(s, nx, lf["_cy"] - (lh + gap * 0.8) / 2, note_w, lh + gap * 0.8, [lf["note"]], sz,
                     _h.DARK_PURPLE if lf.get("emph") else _h.TEXT_BODY, bold=bool(lf.get("emph")), anchor="m",
                     space_after=0, emphasis=ctx.emphasis)


def m_cycle(ctx, r, d):
    """循環図: 円弧の矢印 N 本の円（figures.m_cycle。style arrows / panels）。style: "free" で旧来の円周配置。"""
    if d.get("style") != "free":
        import figures
        return figures.m_cycle(ctx, r, d)
    x, y, w, h = r
    s = ctx.slide
    nodes = _items(d, "nodes")
    n = max(len(nodes), 3)
    cx, cy = x + w / 2, y + h / 2
    nd_w = min(2.2, w * (0.30 if n <= 4 else 0.24))
    nd_h = min(0.9, h * (0.18 if n <= 4 else 0.14))
    rx, ry = w / 2 - nd_w / 2 - 0.05, h / 2 - nd_h / 2 - 0.05
    centers = []
    for i in range(len(nodes)):
        a = -math.pi / 2 + 2 * math.pi * i / n
        centers.append((cx + rx * math.cos(a), cy + ry * math.sin(a)))
    def edge(cx_, cy_, dx, dy, pad=0.04):
        """中心 (cx_, cy_) から方向 (dx, dy) に進んで箱の縁を出る点。"""
        tx = (nd_w / 2 + pad) / abs(dx) if dx else 1e9
        ty = (nd_h / 2 + pad) / abs(dy) if dy else 1e9
        t = min(tx, ty)
        return cx_ + dx * t, cy_ + dy * t
    arrows = []
    for i, (px, py) in enumerate(centers):
        qx, qy = centers[(i + 1) % len(centers)]
        dx, dy = qx - px, qy - py
        sx, sy = edge(px, py, dx, dy)
        ex, ey = edge(qx, qy, -dx, -dy)
        if (ex - sx) * dx + (ey - sy) * dy > 0.05:   # 箱どうしが近すぎて縁が重なる場合は描かない
            arrows.append((sx, sy, ex, ey))
        else:
            ctx.warn("M-CYCLE: ノードが近すぎて矢印を描けない（ペインを広げる）")
    for nd, (px, py) in zip(nodes, centers):
        hl = nd.get("highlight")
        b = _rect(s, px - nd_w / 2, py - nd_h / 2, nd_w, nd_h, fill=_h.DARKEST_PURPLE if hl else _h.LIGHTEST_PURPLE)
        tx, tw = px - nd_w / 2, nd_w
        if nd.get("icon"):          # アイコンは箱の左に置き、文字を右に寄せる
            isz = min(0.5, nd_h - 0.2)
            if _place_visual(ctx, nd["icon"], (tx + 0.1, py - isz / 2, isz, isz), color=_h.WHITE if hl else None) is not None:
                tx, tw = tx + isz + 0.14, tw - isz - 0.18
        al = PP_ALIGN.LEFT if tw < nd_w else PP_ALIGN.CENTER
        paras = [(nd.get("label", nd.get("text", "")), {"bold": True, "color": _h.WHITE if hl else _h.DARKEST_PURPLE, "align": al})]
        if nd.get("label") and nd.get("text"):
            paras.append((nd["text"], {"size": 12, "align": al, "color": _h.WHITE if hl else _h.TEXT_BODY}))
        _textbox(s, tx, py - nd_h / 2, tw, nd_h, paras, 14, anchor="m", space_after=0)
    for sx, sy, ex, ey in arrows:
        _connector(s, sx, sy, ex, ey, _h.DARK_PURPLE, width=2.25, arrow=True)
    if d.get("center"):
        c = _rect(s, cx - 0.9, cy - 0.35, 1.8, 0.7, fill=None)
        _shape_text(c, d["center"], 18, _h.DARKEST_PURPLE, True)


def m_ngok(ctx, r, d, labels=("NG", "OK")):
    x, y, w, h = r
    left, right = d.get("ng") or d.get("before") or {}, d.get("ok") or d.get("after") or {}
    rects = _split_h(r, [1, 1], [0.6])
    for (rx, ry, rw, rh), side, lab, col in ((rects[0], left, labels[0], _h.MID_GRAY),
                                              (rects[1], right, labels[1], _h.CORE_PURPLE)):
        tag = _rect(ctx.slide, rx, ry, 0.9, 0.42, fill=col)
        _shape_text(tag, lab, 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        _textbox(ctx.slide, rx + 1.0, ry, rw - 1.0, 0.42, [side.get("title", "")], 14, _h.DARKEST_PURPLE,
                 bold=True, anchor="m")
        _rect(ctx.slide, rx, ry + 0.5, rw, rh - 0.5, fill=_h.OFF_WHITE, line=col, line_w=1.0)
        paras = [(t, {"bullet": "•"}) for t in side.get("items", [])] + ([side["text"]] if side.get("text") else [])
        size, ok = fit_size(paras, rw - 0.5, rh - 0.7)
        _textbox(ctx.slide, rx + 0.12, ry + 0.6, rw - 0.24, rh - 0.7, paras, size, space_after=8, emphasis=ctx.emphasis)
    _arrow_h(ctx.slide, rects[0][0] + rects[0][2] + 0.1, y + h / 2, 0.4)


def m_before_after(ctx, r, d):
    m_ngok(ctx, r, d, labels=(d.get("before_label", "Before"), d.get("after_label", "After")))


def m_defbox(ctx, r, d):
    x, y, w, h = r
    _rect(ctx.slide, x, y, w, h, fill=_h.WHITE, line=_h.DARK_PURPLE, line_w=1.25)
    paras = [(d.get("term", ""), {"size": 18, "bold": True, "color": _h.DARKEST_PURPLE})]
    if d.get("definition"):
        paras.append(d["definition"])
    paras += [(t, {"bullet": "•"}) for t in d.get("items", [])]
    size, _ = fit_size(paras, w - 0.4, h - 0.2)
    _textbox(ctx.slide, x + 0.2, y + 0.1, w - 0.4, h - 0.2, paras, size, anchor="m", space_after=8, emphasis=ctx.emphasis)


def m_chart(ctx, r, d):
    import charts
    x, y, w, h = r
    kind = d.get("chart_type", "column")
    title = d.get("title")
    if title and d.get("unit"):
        title = f"{title}（{d['unit']}）"
    args = dict(x=x, y=y, w=w, h=h - (0.3 if d.get("source") else 0))
    if d.get("highlight") is not None and d.get("series"):   # 強調する系列だけ主色、他はグレー（CB-07）
        hi = d["highlight"]
        args["colors"] = [_h.CORE_PURPLE if k == hi else _h.LIGHT_GRAY for k in range(len(d["series"]))]
    if kind == "waterfall":
        # {"chart_type": "waterfall", "steps": [{"label": "2024", "value": 100, "total": true}, {"label": "価格", "value": 20}, ...]}
        charts.add_waterfall_chart(ctx.slide, title, d.get("steps") or [], args["x"], args["y"], args["w"], args["h"],
                                   unit=d.get("value_unit", ""))
    elif kind == "doughnut":
        labels, vals = d.get("labels") or d.get("categories"), d.get("values") or []
        hl = d.get("highlight")
        if hl is not None and not d.get("series"):   # 強調する区分（添字か添字の配列）だけ紫、他は灰色の濃淡
            hl = set(hl if isinstance(hl, list) else [hl])
            purples, grays = [_h.CORE_PURPLE, _h.DARK_PURPLE, _h.DARKEST_PURPLE], [_h.MID_GRAY] * 3   # 白い数値ラベルが読める濃さ
            pi = gi = 0
            cols = []
            for k in range(len(vals)):
                if k in hl:
                    cols.append(purples[pi % 3]); pi += 1
                else:
                    cols.append(grays[gi % 3]); gi += 1
            args["colors"] = cols
        ch_w = min(args["w"], args["h"] * 1.6)
        charts.add_doughnut_chart(ctx.slide, title, labels, vals, args["x"], args["y"], ch_w, args["h"],
                                  colors=args.get("colors"))
        center = d.get("center") or (charts._fmt(sum(vals)) + d.get("value_unit", "") if vals else "")
        if center:      # 中央に合計（凡例が右に付くので、リングの中心はグラフ枠の左寄り）
            ring_d = min(args["h"] - (0.4 if title else 0), ch_w * 0.62)
            cx = args["x"] + ch_w * 0.36
            cy = args["y"] + (0.4 if title else 0) + (args["h"] - (0.4 if title else 0)) / 2
            _textbox(ctx.slide, cx - ring_d * 0.3, cy - 0.32, ring_d * 0.6, 0.64,
                     [(center, {"bold": True, "size": 28, "color": _h.DARKEST_PURPLE}),
                      (d.get("center_label", "合計"), {"size": 12, "color": _h.TEXT_SUB})],
                     28, align=PP_ALIGN.CENTER, anchor="m", space_after=0)
    elif kind == "combo":
        # {"chart_type": "combo", "categories": [..], "series": [棒..], "line_series": [折れ線（第 2 軸）..]}
        charts.add_combo_chart(ctx.slide, title, d.get("categories", []), d.get("series", []), d.get("line_series", []),
                               args["x"], args["y"], args["w"], args["h"])
    elif kind == "pie":
        charts.add_pie_chart(ctx.slide, title, d.get("labels") or d.get("categories"), d.get("values"), **args)
    else:
        fn = {"column": charts.add_column_chart, "bar": charts.add_bar_chart, "line": charts.add_line_chart,
              "stacked": charts.add_stacked_column_chart, "area": charts.add_area_chart}.get(kind, charts.add_column_chart)
        n_cat, n_ser = len(d.get("categories", [])), len(d.get("series", []))
        extra = {}
        if kind in ("column", "bar"):
            # 人の資料のグラフ: 値を棒に直接書き、目盛線は薄く（または無く）する
            extra["show_data_labels"] = d.get("data_labels", n_cat <= 8 and n_ser <= 2)
        ch = fn(ctx.slide, title, d.get("categories", []), d.get("series", []),
                show_legend=n_ser > 1, **args, **extra)
        try:
            va = ch.value_axis
            if extra.get("show_data_labels"):
                va.has_major_gridlines = False
            else:
                va.major_gridlines.format.line.color.rgb = _h.LIGHT_GRAY
                va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            ch.category_axis.format.line.color.rgb = _h.LIGHT_GRAY
            if kind == "bar":      # 横棒は上から順に読む（先頭のカテゴリを一番上に）
                ch.category_axis.reverse_order = True
            if kind in ("bar", "column"):
                ch.plots[0].gap_width = 60       # 棒を太くする（細い棒は紙面が薄く見える）
                if d.get("data_labels"):         # 値を棒の先に直接書く（目盛を読ませない）
                    from pptx.enum.chart import XL_LABEL_POSITION
                    fmt = "#,##0" if all(float(v).is_integer() for s_ in d.get("series", [])
                                         for v in s_.get("values", [])) else "General"
                    for ser_ in ch.plots[0].series:     # 系列ごとに付ける（系列の設定がプロットの設定より優先されるため）
                        dl = ser_.data_labels
                        dl.show_value = True
                        dl.number_format = fmt
                        dl.number_format_is_linked = False
                        dl.font.size = Pt(12)
                        dl.font.color.rgb = _h.TEXT_BODY
                        dl.position = XL_LABEL_POSITION.OUTSIDE_END
                    va.visible = False
                    va.has_major_gridlines = False
            hp = d.get("highlight_point")
            if hp is not None and n_ser == 1:   # 1 系列の中の 1 本だけ紫、他は灰（凡例は出さない）
                hp = set(hp if isinstance(hp, list) else [hp])
                ser = ch.plots[0].series[0]
                for k in range(n_cat):
                    pt = ser.points[k]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _h.CORE_PURPLE if k in hp else _h.LIGHT_GRAY
        except Exception:
            pass
    if d.get("source"):
        _textbox(ctx.slide, x, y + h - 0.28, w, 0.28, [d["source"]], 12, _h.MID_GRAY)


def m_image(ctx, r, d):
    """M-PHOTO / M-ILLUST / M-SCREEN / M-FIGREF: image（パス）/ asset（素材 ID）/ svg（SVG 文字列・パス）。"""
    x, y, w, h = r
    cap = d.get("caption")
    src = d.get("source")
    ih = h - (0.32 if cap else 0) - (0.28 if src else 0)
    _place_visual(ctx, d, (x, y, w, ih))
    yy = y + ih
    if cap:
        _textbox(ctx.slide, x, yy + 0.02, w, 0.3, [cap], 12, _h.TEXT_BODY, align=PP_ALIGN.CENTER)
        yy += 0.32
    if src:
        _textbox(ctx.slide, x, yy, w, 0.28, [src], 12, _h.MID_GRAY)


def m_icongrid(ctx, r, d):
    import dense
    if dense.is_stat(d):              # アイコン（小）＋数字＋説明の格子（実績型）
        return dense.m_statgrid(ctx, r, d)
    x, y, w, h = r
    items = _items(d)
    n = max(len(items), 1)
    cols = n if n <= 4 else (3 if n in (5, 6, 9) else 4)
    rows = math.ceil(n / cols)
    cell_h = min((h - 0.2 * (rows - 1)) / rows, 2.6)
    y0 = y + (h - (cell_h * rows + 0.2 * (rows - 1))) / 2
    cells = []
    for rr in _split_v((x, y0, w, cell_h * rows + 0.2 * (rows - 1)), [1] * rows, [0.2] * (rows - 1)):
        cells += _split_h(rr, [1] * cols, [0.25] * (cols - 1))
    for (cx, cy, cw, chh), it in zip(cells, items):
        isz = min(0.8, chh * 0.4, cw * 0.4)
        if it.get("icon"):
            _place_visual(ctx, it["icon"], (cx + (cw - isz) / 2, cy + 0.05, isz, isz))
        paras = [(it.get("label", ""), {"bold": True, "color": _h.DARKEST_PURPLE, "align": PP_ALIGN.CENTER,
                                        "space_after": 4, "size": 18})]
        if it.get("text"):
            paras.append((it["text"], {"align": PP_ALIGN.CENTER}))
        size, ok = fit_size(paras, cw, chh - isz - 0.15)
        _textbox(ctx.slide, cx, cy + isz + 0.15, cw, chh - isz - 0.15, paras, size, space_after=2)


MODULES = {
    "M-BULLETS": m_bullets, "M-CHECKLIST": m_checklist, "M-NUMLIST": m_numlist, "M-ROWHEAD": m_rowhead,
    "M-TABLE": m_table, "M-MATRIX-SYM": m_matrix_sym, "M-CALLOUT": m_callout, "M-QUOTE": m_quote,
    "M-CARD": m_card, "M-CARDSTACK": m_cardstack, "M-HEADCOL": m_headcol, "M-BADGECOL": m_badgecol,
    "M-KPI-TILE": m_kpi, "M-KPI-STACK": lambda c, r, d: m_kpi(c, r, d, vertical=True),
    "M-CHEVRON": m_chevron, "M-PHASE": m_phase, "M-DIAGRAM-FLOW": m_flow, "M-LAYER": m_layer,
    "M-GANTT": m_gantt, "M-TREE": m_tree, "M-CYCLE": m_cycle, "M-NGOK": m_ngok,
    "M-BEFORE-AFTER": m_before_after, "M-DEFBOX": m_defbox, "M-CHART": m_chart,
    "M-PHOTO": m_image, "M-ILLUST": m_image, "M-SCREEN": m_image, "M-FIGREF": m_image,
    "M-ICONGRID": m_icongrid,
}
SUPPORTED_MODULES = sorted(MODULES)


# ═══════════════════════════════════════════════════════════════
#  画像・作図の配置
# ═══════════════════════════════════════════════════════════════
ICON_MIN_SCORE = 6.0        # 検索語からアイコンを選ぶときの下限（これ未満は入れない。合わない絵は無いより悪い）
ICON_AUTO_MIN = 10.0        # ラベルから自動で選ぶときの下限（日本語の語句は当たりが弱いので厳しめ）


def _icon_query(ctx, q, auto=False):
    """検索語（英語の概念語が当たりやすい: "calendar schedule"）から素材 ID を選ぶ。同じスライドで使った絵は避ける。"""
    import asset_search
    used = {a for _, a in ctx.icons}
    floor = ICON_AUTO_MIN if auto else ICON_MIN_SCORE
    for sc, a in asset_search.search(q, top=12, kind="icon"):
        if sc < floor:
            break
        if a["id"] not in used:
            ctx.icons.append((q, a["id"]))
            return a["id"]
    if not auto:
        ctx.warn(f"アイコン: 「{q}」に合う素材が無い（英語の概念語で言い換える。例: calendar / checklist / handshake）")
    return None


def _recolor(path, rgb):
    """紫の線画アイコンを指定色に塗り替えた PNG（濃い地の上は白にする）。透過はそのまま。"""
    import paths as _paths
    from PIL import Image
    hexs = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    out_dir = _paths.ensure_data_dir(os.path.join("cache", "icons"))
    out = os.path.join(out_dir, f"{os.path.splitext(os.path.basename(path))[0]}_{hexs}.png")
    if not os.path.isfile(out):
        im = Image.open(path).convert("RGBA")
        solid = Image.new("RGBA", im.size, (rgb[0], rgb[1], rgb[2], 255))
        solid.putalpha(im.getchannel("A"))
        solid.save(out)
    return out


def icon_spec(item, auto=False, label_key="label"):
    """部品の項目からアイコンの指定を作る。item["icon"] を優先し、無ければ auto のとき item["icon_q"] かラベルで探す。"""
    if not isinstance(item, dict):
        return None
    ic = item.get("icon")
    if ic in (None, "", False, "none"):
        if ic is None and auto:
            q = item.get("icon_q") or item.get(label_key) or item.get("title") or item.get("name")
            return {"q": q, "auto": not item.get("icon_q")} if q else None
        return None
    return ic


def _place_visual(ctx, v, box, color=None):
    """v: {"asset": "0054"} / {"q": "検索語"} / "q:検索語" / {"image": path} / {"svg": "<svg..>" or path} / 素材 ID・パス・SVG。
    color: アイコンを塗り替える色（RGBColor か "white"）。濃い地の上では白にする。合う素材が無ければ何も置かず None。"""
    x, y, w, h = box
    if isinstance(v, str):
        if v.startswith("q:"):
            v = {"q": v[2:].strip()}
        else:
            v = {"svg": v} if v.lstrip().startswith("<svg") or v.endswith(".svg") else (
                {"asset": v} if re.fullmatch(r"[\w\-]+:?\d{3,5}|\d{3,5}", v) else {"image": v})
    if v.get("q"):
        aid = _icon_query(ctx, v["q"], v.get("auto", False))
        if not aid:
            return None
        v = {**v, "asset": aid}
    elif v.get("asset"):
        ctx.icons.append(("", str(v["asset"])))
    if v.get("svg"):
        from svg_to_shapes import add_svg
        g = add_svg(ctx.slide, v["svg"], x, y, w, h, name=v.get("name", "作図"))
        for wmsg in getattr(g, "svg_warnings", []):
            ctx.warn("SVG: " + wmsg)
        return g
    path = v.get("image")
    if v.get("asset"):
        import asset_search
        path = asset_search.asset_path(v["asset"])
    if not path or not os.path.isfile(path):
        ctx.warn(f"画像が見つからない: {v}")
        ph = _rect(ctx.slide, x, y, w, h, fill=_h.OFF_WHITE, line=_h.LIGHT_GRAY, dash=True)
        _shape_text(ph, "画像", 12, _h.MID_GRAY)
        return ph
    color = v.get("color", color)
    if color is not None and v.get("asset"):
        if isinstance(color, str):
            color = {"white": _h.WHITE, "dark": _h.DARKEST_PURPLE, "gray": _h.MID_GRAY}.get(color, _h.WHITE)
        try:
            path = _recolor(path, color)
        except Exception as e:
            ctx.warn(f"アイコンの色替えに失敗: {e!r}")
    return _h.add_image_fit(ctx.slide, path, x, y, w, h)


# ═══════════════════════════════════════════════════════════════
#  関係記号（connectors）
# ═══════════════════════════════════════════════════════════════
def _connector(slide, x1, y1, x2, y2, color=None, width=1.75, arrow=False, elbow=False, dash=False, glyph=False):
    """図形をつながない線（下線・区切り・軸・ファイルツリーの罫）。図形どうしをつなぐ線は connect.link を使う。
    glyph=True は、線を継いで形を描く記号（ツリーの罫・座標軸）。名前に Glyph を付け、LINE-02 の対象から外す。"""
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    c = slide.shapes.add_connector(kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if glyph:
        c.name = f"Glyph Line {c.shape_id}"
    c.line.color.rgb = color or _h.MID_GRAY
    c.line.width = Pt(width)
    _strip_style(c)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def _arrow_h(slide, x, ymid, w, color=None, big=False):
    hh = 0.5 if big else 0.34
    a = _rect(slide, x, ymid - hh / 2, w, hh, fill=color or (_h.DARKEST_PURPLE if big else _h.CORE_PURPLE),
              shape=MSO_AUTO_SHAPE_TYPE.CHEVRON if big else MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE)
    if not big:
        a.rotation = 90
        # 回転後に見た目の幅＝高さになるよう調整
        a.width, a.height = Inches(hh), Inches(min(w, 0.32))
        a.left = Inches(x + (w - hh) / 2)
        a.top = Inches(ymid - min(w, 0.32) / 2)
    return a


def _arrow_v(slide, xmid, y, h, color=None):
    w = 0.42
    a = _rect(slide, xmid - w / 2, y, w, min(h, 0.32), fill=color or _h.CORE_PURPLE,
              shape=MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE)
    a.rotation = 180
    return a


def _symbol(slide, x, y, w, h, text):
    d = min(w, h, 0.5)
    c = _rect(slide, x + (w - d) / 2, y + (h - d) / 2, d, d, fill=_h.DARK_PURPLE, shape=MSO_AUTO_SHAPE_TYPE.OVAL)
    _shape_text(c, text, 18 if len(text) == 1 else 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)


def draw_connector(ctx, ctype, a, b, vertical=False):
    """ペイン a と b（矩形）の間に関係記号を描く。"""
    s = ctx.slide
    if ctype in (None, "R-PAR"):
        return
    if vertical:
        gx, gy, gw, gh = a[0], a[1] + a[3], a[2], b[1] - (a[1] + a[3])
        if ctype in ("R-CAUSE", "R-SOLVE", "R-CONV"):
            _arrow_v(s, gx + gw / 2, gy + (gh - 0.32) / 2, gh)
        elif ctype in ("R-ADD", "R-MUL", "R-EQ", "R-VS"):
            _symbol(s, gx, gy, gw, gh, {"R-ADD": "+", "R-MUL": "×", "R-EQ": "=", "R-VS": "VS"}[ctype])
        return
    gx, gw = a[0] + a[2], b[0] - (a[0] + a[2])
    ymid = (max(a[1], b[1]) + min(a[1] + a[3], b[1] + b[3])) / 2
    if ctype == "R-CAUSE":
        _arrow_h(s, gx + 0.08, ymid, gw - 0.16)
    elif ctype == "R-SOLVE":
        _arrow_h(s, gx + 0.04, ymid, gw - 0.08, big=True)
    elif ctype in ("R-ADD", "R-MUL", "R-EQ", "R-VS"):
        _symbol(s, gx, ymid - 0.3, gw, 0.6, {"R-ADD": "+", "R-MUL": "×", "R-EQ": "=", "R-VS": "VS"}[ctype])
    elif ctype == "R-MUTUAL":
        c = _rect(s, gx + 0.05, ymid - 0.17, gw - 0.1, 0.34, fill=_h.CORE_PURPLE, shape=MSO_AUTO_SHAPE_TYPE.LEFT_RIGHT_ARROW)
    elif ctype == "R-BRACKET":
        br = _rect(s, gx + 0.02, a[1], 0.18, a[3], line=_h.DARK_PURPLE, line_w=1.5, shape=MSO_AUTO_SHAPE_TYPE.RIGHT_BRACE)
        _arrow_h(s, gx + 0.22, a[1] + a[3] / 2, gw - 0.3)
    elif ctype == "R-ZOOM":
        _connector(s, a[0] + a[2], a[1], b[0], b[1], _h.MID_GRAY, dash=True)
        _connector(s, a[0] + a[2], a[1] + a[3], b[0], b[1] + b[3], _h.MID_GRAY, dash=True)
    elif ctype == "R-CONV":
        _arrow_h(s, gx + 0.08, ymid, gw - 0.16)
    elif ctype in ("R-WEDGE", "R-WEDGE-EMPH", "R-TRI"):
        # 列と列の間の背の高い細い三角（人の資料: 目的→骨子、課題→要諦。高さは面の 60〜70%、幅は隙間の 4 割前後、灰）
        top, bot = max(a[1], b[1]), min(a[1] + a[3], b[1] + b[3])
        mid, half = (top + bot) / 2, (bot - top) * 0.33
        top, bot = mid - half, mid + half
        inset = max(0.06, (gw - max(0.3, gw * 0.5)) / 2)
        gx, gw = gx + inset - 0.06, gw - 2 * inset + 0.12
        ff = s.shapes.build_freeform(Inches(gx + 0.06), Inches(top))
        ff.add_line_segments([(Inches(gx + gw - 0.06), Inches((top + bot) / 2)), (Inches(gx + 0.06), Inches(bot))])
        sh = ff.convert_to_shape()
        sh.fill.solid()
        sh.fill.fore_color.rgb = _h.CORE_PURPLE if ctype == "R-WEDGE-EMPH" else _h.MID_GRAY
        sh.line.fill.background()
        _strip_style(sh)
    else:
        ctx.warn(f"関係記号 {ctype} は未対応（記号なしで描画）")


# ═══════════════════════════════════════════════════════════════
#  締め（closings）・ナビ表現（devices）
# ═══════════════════════════════════════════════════════════════
def draw_closing(ctx, c, box):
    """box の下端に締めを描き、使った高さを返す。"""
    if not c or c.get("type", "C-NONE") == "C-NONE":
        return 0.0
    s = ctx.slide
    x, y, w, h = box
    t = c["type"]
    text = c.get("text", "")
    if t == "C-WHY":
        items = c.get("items") or [{"h": c.get("head", ""), "p": text}]
        import dense
        paras = dense.rich([{"h": it.get("h", ""), "p": it.get("p", "")} if isinstance(it, dict) else it for it in items])
        lab_w = 1.5
        sz, sp, ok = dense.fit(paras, w - lab_w - 0.3, 1.3, (14, 12, 11))
        hh = min(1.4, dense.height(sp, w - lab_w - 0.3, sz) + 0.12)
        by = y + h - hh
        tag = _rect(s, x, by, lab_w, hh, fill=_h.DARK_PURPLE)
        lab = c.get("label", "Why Accenture")
        # 英単語の途中で切れない大きさにし、語ごとに改行する（「Why Accen／ture」を避ける）
        lsz = next((z for z in (16, 14, 12) if max(sum(_char_w(ch, z) for ch in wd) for wd in lab.split()) * 1.12
                    <= lab_w - 0.2), 12)
        _shape_text(tag, lab.replace(" ", "\n") if " " in lab and lab.isascii() else lab, lsz, _h.WHITE, True)
        _rect(s, x + lab_w + 0.08, by, w - lab_w - 0.08, hh, fill=_h.OFF_WHITE)
        if not ok:
            ctx.warn("C-WHY: 文字が多い（2 項目・各 1 行に）")
        dense.draw_text(ctx, x + lab_w + 0.2, by, w - lab_w - 0.3, hh, sp, sz, anchor="m")
        return hh
    if t == "C-NOTE":
        _textbox(s, x, y + h - NOTE_H, w, NOTE_H, [text], 11, _h.MID_GRAY, anchor="m")
        return NOTE_H
    size, ok = fit_size([text], w - 1.6, CLOSE_H - 0.06, (18, 16, 14))
    if not ok:
        ctx.warn("締めの文が長い（60字以内に）")
    by = y + h - CLOSE_H
    if t in ("C-CONCLUSION", "C-GOAL", "C-TARGET", "C-OUTCOME"):
        bar = _rect(s, x, by, w, CLOSE_H, fill=_h.DARKEST_PURPLE)
        label = {"C-GOAL": "ゴール", "C-TARGET": "目標", "C-OUTCOME": "効果"}.get(t, c.get("label"))
        if label:
            tag = _rect(s, x, by, 1.2, CLOSE_H, fill=_h.CORE_PURPLE)
            _shape_text(tag, label, 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        _shape_text(bar, text, size, _h.WHITE, True, PP_ALIGN.LEFT if label else PP_ALIGN.CENTER,
                    inset=((1.35 if label else 0.2), 0.02, 0.2, 0.02))
    elif t in ("C-POINTS", "C-NEXTACTION", "C-CAVEAT", "C-BRIDGE", "C-QUESTION"):
        label = {"C-POINTS": "Point", "C-NEXTACTION": "Next", "C-CAVEAT": "留意点", "C-BRIDGE": "次へ",
                 "C-QUESTION": "問い"}[t]
        _rect(s, x, by, w, CLOSE_H, fill=_h.LIGHTEST_PURPLE)
        tag = _rect(s, x, by, 1.2, CLOSE_H, fill=_h.DARKEST_PURPLE)
        _shape_text(tag, label, 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        items = c.get("items") or [text]
        line = "　／　".join(items) if len(items) > 1 else items[0]
        size, ok = fit_size([line], w - 1.6, CLOSE_H - 0.06, (18, 16, 14))
        _textbox(s, x + 1.35, by, w - 1.5, CLOSE_H, [line], size, _h.DARKEST_PURPLE, bold=True, anchor="m",
                 emphasis=ctx.emphasis)
    else:
        ctx.warn(f"締め {t} は未対応（C-CONCLUSION で代用）")
        return draw_closing(ctx, {**c, "type": "C-CONCLUSION"}, box)
    return CLOSE_H


def draw_navtab(ctx, dev, box):
    tabs = dev.get("tabs") or []
    cur = dev.get("current", 0)
    x, y, w, h = box
    th = 0.34
    for i, (tx, ty, tw, _) in enumerate(_split_h((x, y, w, th), [1] * len(tabs), [0.04] * (len(tabs) - 1))):
        on = i == cur
        t = _rect(ctx.slide, tx, ty, tw, th, fill=_h.DARKEST_PURPLE if on else _h.OFF_WHITE)
        _shape_text(t, tabs[i], 12, _h.WHITE if on else _h.MID_GRAY, on, inset=(0.02, 0, 0.02, 0))
    return th + 0.15


def draw_scope(ctx, dev, rects):
    i = dev.get("pane", 0)
    if i >= len(rects):
        return
    x, y, w, h = rects[i]
    _rect(ctx.slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, line=_h.CORE_PURPLE, line_w=2.0, dash=True)
    if dev.get("label"):
        tag = _rect(ctx.slide, x + w + 0.08 - max(1.6, len(dev["label"]) * 0.2 + 0.3), y - 0.08,
                    max(1.6, len(dev["label"]) * 0.2 + 0.3), 0.32, fill=_h.CORE_PURPLE)
        _shape_text(tag, dev["label"], 12, _h.WHITE, True, inset=(0.05, 0, 0.05, 0), wrap=False)


# 「イメージ」の札を付けてよいのは、これから作る物の見本（成果物・画面・構築イメージ・UI 案）だけ。
# 要旨・実績・現状・課題・効果・見積・体制などの中身のページに付けると、提案そのものが仮の話に見える。
# 試算値・仮置きの値は札ではなく、出典の行（source）に「試算の前提: …」と書く。
STAMP_MODULES = {"M-OUTPUTS", "M-SCREEN-MOCK", "M-CHAT-MOCK", "M-PROMPTBOX", "M-FILETREE", "M-CODE", "M-ILLUST"}
STAMP_STEPS = {"成果物", "作成物", "画面", "構築イメージ", "UI案", "画面案", "見本", "サンプル"}
STAMP_TITLE = re.compile(r"成果物|作成物|アウトプット|画面|UI|モック|構築イメージ|見本|サンプル|テンプレート|ダッシュボード")
STAMP_NG_TITLE = re.compile(r"要旨|サマリ|骨子|実績|事例|現状|課題|効果|見積|体制|会議体|スケジュール|前提|理由|Why|"
                            r"市場|動向|目的|範囲|リスク|ステップ")


def stamp_allowed(spec):
    """そのページに「イメージ」の札を付けてよいか（成果物・画面などの見本のページだけ True）。"""
    title = spec.get("title", "") or ""
    if STAMP_NG_TITLE.search(title) or spec.get("story_step") in ("要旨", "事例", "効果", "費用", "体制", "計画",
                                                                    "背景", "課題", "原因", "前提", "根拠"):
        return False
    mods = {p.get("module") for p in spec.get("panes") or []}
    return bool(mods & STAMP_MODULES) or spec.get("story_step") in STAMP_STEPS or bool(STAMP_TITLE.search(title))


def draw_stamp(ctx, dev, box):
    """右上の「イメージ」の札。タイトルの行の右端に置く（メッセージの幅を削らない）。"""
    x, y, w, h = box
    text = dev.get("text", "イメージ")
    sw, sh = 1.3, 0.32
    sx, sy = _h.ML + _h.CW - sw, 0.50
    title_w = 0.0
    for ph in ctx.slide.placeholders:
        if ph.placeholder_format.idx == 0 and ph.has_text_frame:
            title_w = sum(_h.char_em(c, True) for c in ph.text_frame.text) * 28 / 72
    if title_w > _h.CW - sw - 0.3:              # タイトルが長くて重なるときだけ、本文の右上に置く
        sx, sy = x + w - sw, y - 0.42
    E = 914400
    for el in ctx.slide.shapes._spTree:          # 右上の凡例・注記と重なるなら、その左へずらす
        b = _shape_bottom(el)
        off, ext = _xfrm(el)
        if b is None:
            continue
        ox, oy, ow = int(off.get("x")) / E, int(off.get("y")) / E, int(ext.get("cx")) / E
        if oy < sy + sh and b > sy and ox < sx + sw and ox + ow > sx and oy > 0.9 and ow < 6:
            sx = min(sx, ox - sw - 0.1)
    st = _rect(ctx.slide, sx, sy, sw, sh, fill=None, line=_h.MID_GRAY, line_w=1.0)
    _shape_text(st, text, 12, _h.MID_GRAY, True, inset=(0, 0, 0, 0), wrap=False)


# ═══════════════════════════════════════════════════════════════
#  本文の描画
# ═══════════════════════════════════════════════════════════════
def _conn_list(spec, n):
    """connectors を「ペイン間ごとの種類」のリストに正規化。"""
    cons = spec.get("connectors") or []
    gaps = [None] * max(n - 1, 0)
    for c in cons:
        if isinstance(c, str):
            if c == "R-CONV":
                continue
            for i in range(len(gaps)):
                if gaps[i] is None:
                    gaps[i] = c
        else:
            i = c.get("from", 0)
            if 0 <= i < len(gaps):
                gaps[i] = c.get("type")
    return gaps


_NOT_TEXT = {"module", "grid", "icon", "icon_q", "style", "type", "header_style", "col_style", "fill", "mark",
             "image", "asset", "svg", "chart_type", "layout", "legend_kind", "skeleton", "density", "intent"}


def _spec_chars(v, key=None):
    """スライドの本文の字数（ペイン・締めの文字列の合計。部品名やアイコン指定は数えない）。"""
    if isinstance(v, str):
        return 0 if key in _NOT_TEXT else len(v.replace("**", ""))
    if isinstance(v, dict):
        return sum(_spec_chars(x, k) for k, x in v.items() if k not in ("title", "message", "breadcrumb", "chapter"))
    if isinstance(v, list):
        return sum(_spec_chars(x, key) for x in v)
    return 0


def render_body(slide, spec, prs=None, box=None):
    """本文（CY〜BY）を spec に従って描く。警告のリストを返す。"""
    ctx = Ctx(slide, prs, spec)
    # 密度 high のスライドは本文 12pt を基本にする（人の資料は 12pt 以下が本文の約半分）
    dens = spec.get("density")
    high = dens in ("high", "max")
    _CUR["max"] = dens == "max" and _DECK["mode"] != "present"     # 最も密なページ（本文 600〜900 字）
    if _DECK["mode"] == "present":   # 投影用: 16→18pt まで使い、主図を大きく描く（_eff が 1 段上げる）
        _CUR["sizes"] = (12,) if high else (16, 14, 12)
    else:
        _CUR["sizes"] = (12,) if high else SIZES
    _CUR["allow11"] = _DECK["mode"] != "present" and high and _spec_chars(spec) >= (300 if _CUR["max"] else 400)
    if slide.slide_id in _MSG_WARN:
        ctx.warn(_MSG_WARN.pop(slide.slide_id))
    x, y, w, h = box or body_box(slide)
    devices = [d if isinstance(d, dict) else {"type": d} for d in spec.get("devices", [])]
    for dv in devices:
        if dv["type"] == "V-NAVTAB":
            used = draw_navtab(ctx, dv, (x, y, w, h))
            y, h = y + used, h - used
    # 締め・出典
    closing = spec.get("closing")
    if isinstance(closing, str):
        closing = {"type": closing, "text": spec.get("closing_text", "")}
    src = spec.get("source")
    if src:
        _textbox(slide, x, y + h - NOTE_H, w, NOTE_H, [src], 11, _h.MID_GRAY, anchor="m")
        h -= NOTE_H + 0.02
    used = draw_closing(ctx, closing, (x, y, w, h))
    if used:
        h -= used + (0.45 if "R-CONV" in (spec.get("connectors") or []) else 0.18)
    if spec.get("skeleton"):
        # 人が作ったスライドの骨格に中身を流し込む（grammar/skeletons.json）
        import skeleton as _sk
        try:
            _sk.render(ctx, _sk.get(spec["skeleton"]), spec.get("content") or {}, (x, y, w, h))
        except Exception as e:
            ctx.warn(f"骨格 {spec['skeleton']} の描画でエラー: {e!r}")
        _finish(ctx, spec, devices, [(x, y, w, h)], (x, y, w, h))
        _LAST_ICONS[:] = ctx.icons
        return ctx.warnings
    panes = spec.get("panes") or []
    grid = spec.get("grid", "G-1")
    gaps_types = _conn_list(spec, len(panes))
    gaps = [GAP_CONN if t else (0.12 if _CUR["max"] else GAP) for t in gaps_types]
    if grid in ("G-2x2", "G-GRID", "G-2x3"):
        gaps = None
    rects = grid_rects(grid, len(panes), (x, y, w, h), spec.get("ratios") or _content_ratios(spec, panes, grid), gaps)
    # 横に並ぶゾーンの本文の大きさをそろえる（見積もりの最小）。見出しのないゾーンは面の上端をそろえる
    _CUR["rich_base"] = None
    _CUR["strength"] = None
    import dense
    cards = [(p, r_) for p, r_ in zip(panes, rects) if p.get("module") == "M-CARD" and dense.is_strength(p)]
    if len(cards) >= 2:         # 同じ物を並べたら同じ大きさ（要点の面の高さ・文字の大きさを並びでそろえる）
        ms = [dense.strength_sizes(r_, p) for p, r_ in cards]
        _CUR["strength"] = {"lead": max(m["lead"] for m in ms), "head": min(m["head"] for m in ms),
                            "text": min(m["text"] for m in ms)}
    zones = [(p, r_) for p, r_ in zip(panes, rects) if p.get("module") == "M-ZONE"]
    if len(zones) == 1 and len(panes) >= 2 and not zones[0][0].get("sections"):
        # 混在の並びの単独ゾーン: 中身が面の 8 割未満なら、面を中身の高さに縮めて縦中央に置く
        p0, r0 = zones[0]
        need = dense.zone_need(r0, p0, dense.zone_base(r0, p0))
        if need < r0[3] * 0.8:
            k = rects.index(r0)
            rects[k] = (r0[0], r0[1] + (r0[3] - need) / 2, r0[2], need)
    if len(zones) == 1 and len(panes) == 1:
        panes = [{**panes[0], "_solo": True}]
    if len(zones) >= 2:
        import dense
        _CUR["rich_base"] = min(dense.zone_base(r_, p) for p, r_ in zones)
        if any(p.get("title") for p, _ in zones):
            panes = [{**p, "_hdr_pad": dense.Z()["hdr"]} if p.get("module") == "M-ZONE" and not p.get("title") and not p.get("sections")
                     else p for p in panes]
        # 縦に分けたゾーン（sections）は列いっぱいに描くので、並びの塊の高さの基準に入れない
        flat = [(p, r_) for p, r_ in zones if not p.get("sections")] or zones
        need = max(dense.zone_need(r_, p, _CUR["rich_base"]) for p, r_ in flat)
        # 並ぶゾーンの中身は、いちばん多いゾーンに合わせて同じだけ下げる（兄弟の本文の上端をそろえる）
        if all(abs(r_[3] - rects[0][3]) < 0.01 for _, r_ in zones):
            panes = [{**p, "_blk": need} if p.get("module") == "M-ZONE" else p for p in panes]
        if all(p.get("module") == "M-ZONE" for p in panes) and not any(p.get("sections") for p, _ in zones)                 and all(abs(r_[3] - rects[0][3]) < 0.01 for r_ in rects)                 and need < rects[0][3] * 0.88:
            # 面を中身の高さ（全ゾーンで最大）にそろえ、並び全体（格子なら全行）を本文の縦中央より少し上に置く
            ys = sorted({round(r_[1], 3) for r_ in rects})
            gy = (ys[1] - ys[0] - rects[0][3]) if len(ys) > 1 else 0
            block = need * len(ys) + gy * (len(ys) - 1)
            top = ys[0] + (ys[-1] + rects[0][3] - ys[0] - block) * 0.4
            rects = [(r_[0], top + ys.index(round(r_[1], 3)) * (need + gy), r_[2], need) for r_ in rects]
    for i, (pane, r) in enumerate(zip(panes, rects)):
        mod = pane.get("module", "M-BULLETS")
        fn = MODULES.get(mod)
        if pane.get("icons", spec.get("icons")) == "auto" and mod in ICON_MODULES:
            pane = _auto_icons(pane, mod)
        number = (i + 1) if ctx.index else None
        if number and not pane.get("header") and pane.get("title") and mod in ("M-BADGECOL", "M-CARD", "M-HEADCOL"):
            pane = {**pane, "title": f"{number}. {pane['title']}"}
        if pane.get("header"):
            r = pane_header(ctx, r, pane["header"], number, pane.get("status"), dark=pane.get("header_dark", False),
                            style=pane.get("header_style"))
        if fn is None:
            ctx.warn(f"部品 {mod} は compose 未対応（custom で描画すること）")
            continue
        try:
            _draw_filled(ctx, fn, r, pane)
        except Exception as e:   # 1 部品の失敗でスライド全体を止めない
            ctx.warn(f"{mod} の描画でエラー: {e}")
        if pane.get("focus"):   # 並びの中で 1 つだけ強調（人の資料の並びの 8% が 1 つだけ色・枠を変える）
            import figures      # 枠は濃い紫 2.25pt（明るい紫の枠は灰の面の外では目立たない）
            _rect(slide, r[0] - 0.04, r[1] - 0.04, r[2] + 0.08, r[3] + 0.08, fill=None, line=figures.EMPH["line"],
                  line_w=figures.EMPH["w"])
    vertical = grid in VERTICAL
    for i, t in enumerate(gaps_types):
        if t and i + 1 < len(rects):
            draw_connector(ctx, t, rects[i], rects[i + 1], vertical)
    if used and "R-CONV" in (spec.get("connectors") or []):
        _arrow_v(slide, x + w / 2, y + h + 0.08, 0.32)
    _finish(ctx, spec, devices, rects, (x, y, w, h))
    _LAST_ICONS[:] = ctx.icons
    _CUR["rich_base"] = None
    _CUR["strength"] = None
    return ctx.warnings


_LAST_ICONS = []   # 直前に描いたスライドのアイコン（build_deck がレポートとアイコン一覧に使う）

# 文字の列を横に並べる部品。情報量が違うのに同じ幅にすると、少ない列が空き、多い列が詰まる（均等割りに逃げない）
_TEXT_COLS = {"M-CARD", "M-HEADCOL", "M-BADGECOL", "M-BULLETS", "M-CHECKLIST", "M-NUMLIST"}


def _pane_chars(p):
    n = len(str(p.get("title", ""))) + len(str(p.get("header", ""))) + len(str(p.get("text", "")))
    return n + sum(len(str(t.get("text", t) if isinstance(t, dict) else t)) for t in p.get("items") or [])


def _content_ratios(spec, panes, grid):
    """横に並ぶ文字の列の幅を、文字量に比例させる（0.75〜1.4 倍）。差が小さい（1.3 倍未満）なら均等のまま。
    spec["equal"]: true で均等に固定。"""
    if spec.get("equal") or len(panes) < 2 or grid not in HORIZONTAL or grid in ("G-23", "G-32", "G-SIDE", "G-3W", "G-IO"):
        return None
    if not all(p.get("module", "M-BULLETS") in _TEXT_COLS for p in panes):
        return None
    if any(p.get("number") or p.get("head") for p in panes):     # 番号付きの強み・並列のカードは同じ幅
        return None
    lens = [max(_pane_chars(p), 1) for p in panes]
    if max(lens) / min(lens) < 1.3:
        return None
    mean = sum(lens) / len(lens)
    return [min(1.4, max(0.75, n / mean)) for n in lens]

# アイコンを入れられる部品と、項目の入れ物のキー（None は部品そのものに 1 つ）。
# 人の資料でアイコンが多い部品: フロー 57%・アイコン格子 82%・バッジ列 90%・循環 77%・構成図 57%
ICON_MODULES = {"M-CARD": None, "M-HEADCOL": None, "M-BADGECOL": None, "M-CARDSTACK": "cards", "M-ICONGRID": "items",
                "M-DIAGRAM-FLOW": "nodes", "M-CYCLE": "nodes", "M-ACTORS": "actors", "M-ORG-ACTORS": "actors",
                "M-HUB": "nodes", "M-HUBSPOKE": "nodes", "M-EQUATION": "terms", "M-ICONROW": "items", "M-ARCH": "zones"}


def _auto_icons(pane, mod):
    """icons: "auto" の部品で、icon の無い項目に検索の指定（icon_q か、なければラベル）を入れる。"""
    key = ICON_MODULES.get(mod)
    pane = dict(pane)
    if key is None:
        sp = icon_spec(pane, True, "title")
        if sp:
            pane["icon"] = sp
        return pane

    def fill(items):
        out = []
        for it in items or []:
            it = {"label": it} if isinstance(it, str) else dict(it)
            sp = icon_spec(it, True)
            if sp:
                it["icon"] = sp
            out.append(it)
        return out
    if mod == "M-ARCH":     # 帯の中の要素
        pane["zones"] = [{**z, "items": fill(z.get("items"))} for z in pane.get("zones") or []]
    elif mod == "M-HUB":
        pane["nodes"] = fill(pane.get("nodes"))
        if isinstance(pane.get("center"), dict):
            pane["center"] = fill([pane["center"]])[0]
    else:
        pane[key] = fill(pane.get(key))
    return pane


def _finish(ctx, spec, devices, rects, box):
    """本文を描いた後に重ねる表現（範囲・スタンプ・注釈・強調枠・凡例）。"""
    for dv in devices:
        if dv["type"] == "V-SCOPE":
            draw_scope(ctx, dv, rects)
        elif dv["type"] == "V-STAMP":
            if stamp_allowed(spec):
                draw_stamp(ctx, dv, box)
            else:
                ctx.warn("STAMP: 成果物・画面の見本ではないページの「イメージ」札は描かない（試算・仮置きは source に前提を書く）")
        elif dv["type"] in DEVICES_EXT:
            try:
                DEVICES_EXT[dv["type"]](ctx, dv, box)
            except Exception as e:
                ctx.warn(f"{dv['type']} の描画でエラー: {e!r}")
    if spec.get("illustrative") and stamp_allowed(spec) and not any(dv["type"] == "V-STAMP" for dv in devices):
        draw_stamp(ctx, {"text": "イメージ"}, box)


DEVICES_EXT = {}   # artifacts.py が V-ANNOT / V-HILITE / V-LEGEND を登録する


# ═══════════════════════════════════════════════════════════════
#  スライド・デッキ
# ═══════════════════════════════════════════════════════════════
def new_presentation():
    from pptx import Presentation
    prs = Presentation(_h.TEMPLATE_PATH)
    while len(prs.slides) > 0:
        rid = prs.slides._sldIdLst[0].get(qn("r:id"))
        prs.part.drop_rel(rid)
        del prs.slides._sldIdLst[0]
    return prs


def msg_reserve(spec):
    """メッセージの右端に空ける幅。体制図の凡例を置くページで、メッセージを手前で折り返して重ねない。
    「イメージ」の札はタイトルの行の右端に置くので、メッセージの幅は削らない。"""
    mods = {p.get("module") for p in spec.get("panes") or []}
    if "M-ORG" in mods:
        return 4.0
    return 0.0


def new_content_slide(prs, title, message=None, breadcrumb=None, layout=None, reserve=0.0):
    slide = prs.slides.add_slide(prs.slide_layouts[_h.layout_idx(layout or "content")])
    _h.clear_placeholders(slide)
    if breadcrumb:
        _h.add_breadcrumb(slide, _h.tidy_ja(breadcrumb))
    _h.add_title(slide, title)
    # タイトル枠とメッセージ行（MSG_Y）が重なるテーマへの対応（1 行タイトル前提、28pt に 0.53in 必要）
    msg_y = _h.MSG_Y
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0 and ph.top is not None and ph.height is not None:
            top_in, h_in = ph.top / 914400, ph.height / 914400
            if top_in + h_in > _h.MSG_Y - 0.01:
                left, width, top = ph.left, ph.width, ph.top   # 継承値を先に読む（高さだけ設定すると幅が 0 になる）
                if top_in + 0.53 <= _h.MSG_Y + 0.01:     # 枠を 1 行分に詰める
                    new_h = _h.MSG_Y - 0.01 - top_in
                else:                                     # 詰められない → メッセージ行を下げる
                    new_h = max(h_in, 0.53)
                    msg_y = top_in + new_h
                ph.left, ph.top, ph.width, ph.height = left, top, width, Inches(new_h)
    body_top = _h.CY
    if message:
        two = _h.message_lines(message, reserve) >= 2
        if _h.message_lines(message, reserve) > 2:
            _MSG_WARN[slide.slide_id] = "メッセージが 2 行（約 80 字）を超える"
        elif two and _h.message_tail(message, reserve) < 10:
            _MSG_WARN[slide.slide_id] = (f"メッセージの 2 行目が約 {_h.message_tail(message, reserve)} 字だけになる"
                                         "（言い回しを詰めて 1 行に収めるか、2 行目まで書く）")
        _h.add_message_line(slide, message, reserve=reserve)
        body_top = _h.CY_2LINE if two else _h.CY
        if msg_y != _h.MSG_Y:
            ml = slide.shapes[len(slide.shapes) - 1]
            ml.top = Inches(msg_y)
            if not two:
                ml.height = Inches(max(0.36, _h.CY - 0.02 - msg_y))
            body_top += msg_y - _h.MSG_Y
    _BODY_TOP[slide.slide_id] = body_top
    _h.set_footer(slide)
    return slide


_BODY_TOP = {}    # slide_id → 本文の開始位置（メッセージが 2 行なら下がる）
_MSG_WARN = {}


def body_box(slide):
    """そのスライドの本文領域 (x, y, w, h)。"""
    top = _BODY_TOP.get(slide.slide_id, _h.CY)
    return (_h.ML, top, _h.CW, _h.BY - top)


def _ph_size(slide, idx, default):
    """プレースホルダーの文字の大きさ（pt）。スライド → レイアウト → マスターの順に sz を探す。"""
    for src in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        for ph in src.placeholders:
            if ph.placeholder_format.idx == idx or (idx == 0 and ph.placeholder_format.type == 1):
                for e in ph._element.iter():
                    if e.tag.endswith(("}defRPr", "}rPr")) and e.get("sz"):
                        return int(e.get("sz")) / 100
    return default


def _fill_placeholders(slide, values, notes, color):
    """values: role → text。notes: layout_notes[...].placeholders。未使用 ph は半角スペースで埋める。"""
    roles = {int(k.split("_")[1]): v.get("role") for k, v in (notes or {}).items()}
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx in (20, 21):
            continue
        role = roles.get(idx) or ("title" if idx == 0 else None)
        text = values.get(role) if role else None
        tf = ph.text_frame
        tf.clear()
        if text and role in ("title", "subtitle", "description") and "\n" not in text:
            # 2 行になる表紙・中扉のタイトルは、語の切れ目で改行する（「プロジェ／クト」のように折り返さない）
            text = _wrap_safe(_h.tidy_ja(text), max(ph.width / 914400 - 0.2, 1.0), _ph_size(slide, idx, 36 if role == "title" else 14))
        for k, line in enumerate((text if text else " ").split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = _h.tidy_ja(line)
            r.font.name = _font()
            if role in ("date", "presenter", "body") and idx != 0:
                r.font.size = Pt(12)
            if role == "title":
                r.font.bold = True
                r._r.get_or_add_rPr().set("cap", "none")   # レイアウトの全大文字化を解除（表記を本文とそろえる）
            if color is not None:   # None ならレイアウト側の既定色（背景に合わせて設計済み）を使う
                r.font.color.rgb = color


def _clone_layout_ph(slide, idx):
    """日付など python-pptx がスライドに複製しない placeholder を、レイアウトから複製する。"""
    import copy
    if any(ph.placeholder_format.idx == idx for ph in slide.placeholders):
        return
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx == idx:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
            return


def add_cover(prs, outline):
    key = outline.get("cover_layout") or "cover"
    slide = prs.slides.add_slide(prs.slide_layouts[_h.layout_idx(key, _h.layout_idx("cover"))])
    _notes = _h.THEME.get("layout_notes", {})
    for k, v in ((_notes.get(key) or _notes.get("cover") or {}).get("placeholders", {})).items():
        if v.get("role") == "date" and outline.get("date"):
            _clone_layout_ph(slide, int(k.split("_")[1]))
    notes = _h.THEME.get("layout_notes", {})
    pn = (notes.get(key) or notes.get("cover") or {}).get("placeholders", {})
    bg = (notes.get(key) or notes.get("cover") or {}).get("background", "")
    dark = bg not in ("white", "light") if bg else _h.THEME.get("cover_text_color") != "BLACK"
    color = _h.WHITE if dark else _h.BLACK
    vals = {"title": outline.get("title", ""), "subtitle": outline.get("subtitle", ""),
            "date": outline.get("date", ""), "presenter": outline.get("author", ""), "body": outline.get("author", "")}
    _fill_placeholders(slide, vals, pn, None)
    return slide


def add_section(prs, title, number=None, description=None):
    slide = prs.slides.add_slide(prs.slide_layouts[_h.layout_idx("section")])
    sec = _h.THEME.get("layout_notes", {}).get("section", {})
    dark = sec.get("background", "dark") not in ("white", "light")
    vals = {"title": title, "section_number": f"{number:02d}" if isinstance(number, int) else (number or ""),
            "subtitle": description or "", "description": description or ""}
    _fill_placeholders(slide, vals, sec.get("placeholders", {}), None)
    _h.set_footer(slide)
    return slide


def add_statement(prs, text, kicker=None, support=None):
    """言い切りページ: 濃い紫の単色の面に、大きな一文だけを置く（物語の転換点に、デッキで最大 3 枚）。
    公式 Starter Pack の Key Message 型に当たる。グラデーションは使わない。"""
    slide = prs.slides.add_slide(prs.slide_layouts[_h.layout_idx("content")])
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    _rect(slide, 0, 0, _h.SLIDE_W if hasattr(_h, "SLIDE_W") else 13.333, 7.5, fill=_h.DARKEST_PURPLE)
    x, w = _h.ML + 0.4, 10.6
    y = 2.35
    if kicker:
        _textbox(slide, x, y - 0.55, w, 0.4, [kicker], 14, _h.LIGHT_PURPLE, bold=True, anchor="b", space_after=0)
    plain = _plain(text)
    size, ok = fit_size([plain], w, 2.3, (40, 36, 32, 28), 0)
    hh = est_height([plain], w, size, 0)
    _textbox(slide, x, y, w, hh, [text], size, _h.WHITE, bold=True, space_after=0, line_spacing=1.1)
    if support:
        _textbox(slide, x, y + hh + 0.3, w, 1.2, [support], 18, _h.LIGHTEST_PURPLE, space_after=0)
    _hline(slide, x, y - 0.12, 1.2, _h.CORE_PURPLE, 3.0)
    return slide


def add_agenda(prs, sections, title=None, current=None, message=None):
    """目次。current（0 始まり）を渡すと、その章だけ濃くした「目次＋現在地」になる（中扉の代わり）。
    message には章の結論文を入れる。"""
    lang_ja = _h.LANG == "ja"
    slide = new_content_slide(prs, title or ("目次" if lang_ja else "Agenda"), message, None)
    n = max(len(sections), 1)
    x, y, w = _h.ML + 0.6, body_box(slide)[1] - (0.1 if not message else -0.1), _h.CW - 1.2
    items = [s if isinstance(s, dict) else {"title": str(s)} for s in sections]
    has_desc = any(it.get("description") for it in items)
    rh = min(1.15 if has_desc else 0.95, (_h.BY - y - 0.2) / n)
    for i, it in enumerate(items):
        ry = y + i * rh
        on = current is None or i == current
        num = _rect(slide, x, ry + 0.1, 0.9, rh - 0.2, fill=_h.DARKEST_PURPLE if on else None,
                    line=None if on else _h.LIGHT_GRAY, line_w=0.75)
        _shape_text(num, f"{i + 1:02d}", 18, _h.WHITE if on else _h.MID_GRAY, True, inset=(0, 0, 0, 0), wrap=False)
        paras = [(it.get("title", ""), {"bold": True, "size": 18, "color": _h.DARKEST_PURPLE if on else _h.MID_GRAY,
                                         "space_after": 2})]
        if it.get("description"):     # 章の結論文（中扉を置かないデッキでは、ここで各章の結論を示す）
            paras.append((it["description"], {"size": 14, "color": _h.TEXT_BODY if on else _h.MID_GRAY}))
        _textbox(slide, x + 1.15, ry, w - 1.15, rh, paras, 18, anchor="m", space_after=2)
        _hline(slide, x + 1.15, ry + rh - 0.02, w - 1.15)
    return slide


def _mark_lang(prs):
    """日本語のデッキは、言語の指定が無い文字に lang="ja-JP" を付ける。指定が無いとマスターの既定（en-US）の
    行分割になり、「、」「ー」が行頭に来る（禁則処理が効かない）。"""
    if _h.LANG != "ja":
        return
    for slide in prs.slides:
        for tag in ("a:rPr", "a:endParaRPr"):
            for el in slide._element.iter(qn(tag)):
                if el.get("lang") is None:
                    el.set("lang", "ja-JP")
                    el.set("altLang", "en-US")


def build_deck(outline, out_path, custom=None, prs=None):
    """outline（v2）からデッキを生成して保存する。戻り値: {"slides": [{no, type, warnings}], "path": ...}

    custom: {スライド番号(1 始まり): fn(slide, slide_spec, prs)} — compose 未対応の本文を自前で描く。
    """
    out_path = _paths.check_path(out_path, write=True)
    custom = custom or {}
    prs = prs or new_presentation()
    _h.reset_slide_counter()
    report = []
    sections = [s.get("title", "") for s in outline.get("slides", []) if s.get("type") == "section"]
    sec_no = 0
    n_body = sum(1 for s in outline.get("slides", []) if s.get("type", "content") == "content")
    mode = (outline.get("settings") or {}).get("divider", "auto")
    # 用途: present（投影。本文 16pt）/ read（読み物。本文 14・12pt）
    _DECK["mode"] = outline.get("mode") or (outline.get("settings") or {}).get("mode") or "read"
    # 人の資料: 中扉は本文 13 枚に 1 枚。少ないデッキでは目次＋現在地を中扉の代わりにする
    agenda_div = mode == "agenda" or (mode == "auto" and sections and n_body / len(sections) < 8)
    # 章ラベル（パンくず）: 人の資料の 40% にある。中扉の title か本文の chapter を持ち回り、現在地だけを示す
    auto_bc = (outline.get("settings") or {}).get("breadcrumb", "auto") not in (False, "off", "none")
    chapter = None
    for no, s in enumerate(outline.get("slides", []), start=1):
        t = s.get("type", "content")
        warns = []
        if t == "section":
            chapter = s.get("title") or chapter
        elif t == "content" and s.get("chapter"):
            chapter = s["chapter"]
        bc = s.get("breadcrumb")
        if bc is None and auto_bc and chapter:
            bc = f"{chapter} > {s['crumb']}" if s.get("crumb") else chapter
        if t == "cover":
            add_cover(prs, outline)
        elif t == "agenda":
            add_agenda(prs, s.get("items") or sections, s.get("title"))
        elif t == "section":
            sec_no += 1
            if agenda_div and sections:
                add_agenda(prs, sections, s.get("agenda_title"), current=sec_no - 1, message=s.get("description"))
            else:
                add_section(prs, s.get("title", ""), s.get("number", sec_no), s.get("description"))
        elif t == "closing":
            _h.make_closing_slide(prs, s.get("title", "Thank You"), s.get("subtitle"))
        elif t == "statement":
            add_statement(prs, s.get("message") or s.get("title", ""), s.get("kicker"), s.get("support"))
        else:
            slide = new_content_slide(prs, s.get("title", ""), s.get("message"), bc or None, reserve=msg_reserve(s))
            if no in custom:
                custom[no](slide, s, prs)
            else:
                _DECK["body"] = True          # 本文の部品だけ、投影用の文字の写像（_eff）を掛ける
                _LAST_ICONS.clear()
                try:
                    warns = render_body(slide, s, prs)
                finally:
                    _DECK["body"] = False
            merge_text_into_boxes(slide)
        icons = list(_LAST_ICONS) if t == "content" else []
        _LAST_ICONS.clear()
        report.append({"no": no, "type": t, "title": s.get("title", ""), "warnings": warns, "icons": icons})
    _h.strip_sections(prs)
    _mark_lang(prs)
    _h.unify_fonts(prs, _h.LANG)      # 日本語は英数字・コード・グラフまで Meiryo UI、英語は Graphik
    prs.save(out_path)
    _icon_sheet(report, out_path)
    return {"path": out_path, "slides": report}


def merge_text_into_boxes(slide):
    """色の矩形の上に重ねたテキストボックスを、矩形の文字として書き直す（v6.8）。人の資料は色の箱の文字の 95% を
    箱に直接書く。テキストボックスの段落を矩形に移し、位置の差を矩形の余白（inset）にするので、見た目は変わらない。
    対象: 塗りか枠のある矩形（prst rect）で文字が空のもの × 塗り・線の無い回転していないテキストボックスで、その内側にあり、
    z 順で間に重なる図形が無く、矩形の中のテキストボックスがそれ 1 つのもの。戻り値: 書き直した数。"""
    import copy
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if os.environ.get("ACNPPTX_NO_MERGE"):          # 比べるとき用（合わせる前の形で出す）
        return 0
    E, tol = 914400, int(0.02 * 914400)
    shapes = [sh for sh in slide.shapes if not sh.is_placeholder]

    def box(sh):
        return sh.left, sh.top, sh.left + sh.width, sh.top + sh.height

    def overlap(a, b):
        return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]

    def inside(a, b):
        return a[0] >= b[0] - tol and a[1] >= b[1] - tol and a[2] <= b[2] + tol and a[3] <= b[3] + tol

    def plain_tb(sh):
        if sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX or not sh.has_text_frame or not sh.text_frame.text.strip():
            return False
        if (sh.rotation or 0) != 0:
            return False
        sp = sh._element.spPr
        if sp.find(qn("a:solidFill")) is not None or sp.find(qn("a:gradFill")) is not None:
            return False
        ln = sp.find(qn("a:ln"))
        return ln is None or ln.find(qn("a:solidFill")) is None

    def host_ok(sh):
        if sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or (sh.rotation or 0) != 0:
            return False
        g = sh._element.find(".//" + qn("a:prstGeom"))
        if g is None or g.get("prst") != "rect":
            return False
        if sh.has_text_frame and sh.text_frame.text.strip():
            return False
        sp = sh._element.spPr
        ln = sp.find(qn("a:ln"))
        return sp.find(qn("a:solidFill")) is not None or (ln is not None and ln.find(qn("a:solidFill")) is not None)
    done, gone = 0, set()
    for k, tb in enumerate(shapes):
        if not plain_tb(tb):
            continue
        b = box(tb)
        host = None
        for j in range(k - 1, -1, -1):                 # z 順で下にある図形を上から順に見る
            o = shapes[j]
            if id(o) in gone or not overlap(box(o), b):
                continue
            if host_ok(o) and inside(b, box(o)) and getattr(o, "_merged", False) is False:
                host = o
            break                                       # 最初に重なる図形が矩形でなければ合わせない
        if host is None:
            continue
        hb = box(host)
        area = lambda q: (q[2] - q[0]) * (q[3] - q[1])
        # 箱の中に文字の塊が複数あるとき（アイコンのラベル＋本文）は、いちばん大きい塊だけを箱の文字にする
        others = [t for t in shapes if t is not tb and id(t) not in gone and plain_tb(t) and inside(box(t), hb)
                  and shapes.index(t) > shapes.index(host)]
        if any(area(box(t)) > area(b) for t in others):
            continue
        # 列・行をまたぐ帯（現在の工程の列など）の中に本文の塊が複数あるなら、どれの箱でもないので合わせない
        if any(area(box(t)) >= area(hb) * 0.15 for t in others):
            continue
        src = tb._element.find(qn("p:txBody"))
        dst = host._element.find(qn("p:txBody"))
        new = copy.deepcopy(src)
        bp = new.find(qn("a:bodyPr"))
        ins = {k_: int(bp.get(k_, d_)) for k_, d_ in (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))}
        bp.set("lIns", str(max(0, b[0] - hb[0] + ins["lIns"])))
        bp.set("tIns", str(max(0, b[1] - hb[1] + ins["tIns"])))
        bp.set("rIns", str(max(0, hb[2] - b[2] + ins["rIns"])))
        bp.set("bIns", str(max(0, hb[3] - b[3] + ins["bIns"])))
        if bp.get("anchor") is None:
            bp.set("anchor", "t")
        for af in ("a:spAutoFit", "a:normAutofit", "a:noAutofit"):   # 箱の大きさは変えない
            for e in bp.findall(qn(af)):
                bp.remove(e)
        etree.SubElement(bp, qn("a:noAutofit"))
        # テキストボックスで縦の位置を中央・下に寄せていたときは、箱の余白で同じ位置になる（anchor はそのまま）
        if dst is not None:
            host._element.replace(dst, new)
        else:
            host._element.append(new)
        tb._element.getparent().remove(tb._element)
        gone.add(id(tb))
        host._merged = True
        done += 1
    return done


def _icon_sheet(report, out_path):
    """使ったアイコンの一覧（<出力名>_icons.json と _icons.png）。Claude が画像で見て、内容に合わない絵を icon で差し替える。"""
    rows = [(r["no"], q, a) for r in report for q, a in r.get("icons", [])]
    stem = os.path.splitext(out_path)[0]
    if not rows:
        for ext in ("_icons.json", "_icons.png"):
            if os.path.isfile(stem + ext):
                os.remove(stem + ext)
        return
    import json
    import asset_search
    with open(stem + "_icons.json", "w", encoding="utf-8") as f:
        json.dump([{"slide": n, "query": q, "asset": a} for n, q, a in rows], f, ensure_ascii=False, indent=1)
    try:
        from contact_sheet import make_sheet
        im = make_sheet([asset_search.asset_path(a) for _, _, a in rows],
                        [f"s{n} {a} {q[:16]}" for n, q, a in rows], cols=6, cell=160)
        im.save(stem + "_icons.png")
    except Exception as e:
        print("アイコン一覧の画像を作れなかった:", e)


# 実物・流れの部品と注釈（artifacts.py）を登録する
import artifacts as _artifacts  # noqa: E402
_artifacts.register(MODULES, DEVICES_EXT)
import diagrams as _diagrams  # noqa: E402  旧パターン N/Z/AA/AC/AE/AG/AK/AL/AM の部品
_diagrams.register(MODULES)
import org as _org  # noqa: E402  体制図
_org.register(MODULES)
import figures as _figures  # noqa: E402  アイコンで描く概念図（ハブ・式・アイコン行・アイコンのフロー）
_figures.register(MODULES)
import dense as _dense  # noqa: E402  情報量の多いページの部品（ゾーン・Why 行・行見出しマトリクス・ステップ・図＋番号解説）
_dense.register(MODULES)
import patterns as _patterns  # noqa: E402  人の提案書の型（実績一覧・列カード・五方向の器・成果物のイメージ）
MODULES.update(_patterns.MODULES)
SUPPORTED_MODULES = sorted(MODULES)


if __name__ == "__main__":
    import argparse
    import json
    ap = argparse.ArgumentParser()
    ap.add_argument("outline")
    ap.add_argument("out")
    ap.add_argument("--theme")
    a = ap.parse_args()
    with open(_paths.check_path(a.outline), encoding="utf-8") as f:
        ol = json.load(f)
    _h.set_lang(ol.get("language", "ja"))
    _h.load_theme(a.theme or ol.get("theme"))
    rep = build_deck(ol, a.out)
    for r in rep["slides"]:
        flag = "WARN" if r["warnings"] else "ok"
        print(f"{r['no']:>2} {r['type']:<8} {flag} {r['title'][:30]}")
        for w_ in r["warnings"]:
            print("     -", w_)
    print("saved", rep["path"])
