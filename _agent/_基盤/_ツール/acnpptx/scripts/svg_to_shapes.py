"""
svg_to_shapes.py — Claude が描いた簡単な SVG を、PowerPoint の「編集できるネイティブ図形」に変換して挿入する。

画像として貼るのではなく、要素ごとに DrawingML 図形を作り、全体を 1 つのグループにまとめる。

対応要素（これ以外はエラー）:
  svg / g / title / desc
  rect            → 矩形（rx は無視：ブランド規定で角丸禁止）
  circle, ellipse → 楕円
  line            → 直線コネクタ
  polyline        → 開いたフリーフォーム（custGeom）
  polygon, path   → フリーフォーム（custGeom）。path は M L H V C S Q T A Z（相対含む）。A は 3 次ベジェで近似
  text            → テキストボックス（text-anchor / font-size / font-weight / fill 対応、tspan は文字列のみ）
変換:
  viewBox を配置枠（x, y, w, h インチ）に等倍縮尺で収める（中央揃え）
  transform は translate / scale / matrix / rotate に対応（回転を含む rect・ellipse はフリーフォーム化）
  色はテーマのパレットに丸める（snap=True）。gradient・filter・image・use・clipPath・mask は使用不可

使い方:
    from svg_to_shapes import add_svg
    grp = add_svg(slide, svg_text, x=ML, y=CY, w=6.0, h=4.0, name="概念図")
CLI（検証用）:
    python svg_to_shapes.py input.svg out.pptx [--theme acn-technology-graphik]
"""

import math
import os
import re
import sys

from lxml import etree

from safexml import PARSER as _PARSER  # noqa: E402  外部由来の SVG を読むので実体展開・ネットワークを無効にしたパーサー
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

_SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)
import paths as _paths  # noqa: E402

SVG_NS = "http://www.w3.org/2000/svg"
ALLOWED = {"svg", "g", "title", "desc", "rect", "circle", "ellipse", "line", "polyline", "polygon", "path",
           "text", "tspan"}
EMU_PER_IN = 914400

_NAMED = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
          "gray": "808080", "grey": "808080", "purple": "800080", "orange": "FFA500", "yellow": "FFFF00",
          "lightgray": "D3D3D3", "lightgrey": "D3D3D3", "darkgray": "A9A9A9", "silver": "C0C0C0"}


class SvgError(ValueError):
    pass


# ── 色 ───────────────────────────────────────────────────────────
def _parse_color(v):
    if v is None:
        return None
    v = v.strip().lower()
    if v in ("none", "transparent", ""):
        return "none"
    if v.startswith("#"):
        h = v[1:]
        if len(h) == 3:
            h = "".join(c * 2 for c in h)
        if len(h) != 6:
            raise SvgError(f"色の形式が不正: {v}")
        return h.upper()
    m = re.match(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", v)
    if m:
        return "".join(f"{int(x):02X}" for x in m.groups())
    if v in _NAMED:
        return _NAMED[v]
    if v.startswith("url("):
        raise SvgError("gradient / pattern（url(...)）は使用不可。単色で描くこと")
    raise SvgError(f"未対応の色: {v}")


def theme_palette():
    """helpers の現在のテーマ色からパレット（hex のリスト）を作る。"""
    import helpers as h
    names = ["DARKEST_PURPLE", "DARK_PURPLE", "CORE_PURPLE", "LIGHT_PURPLE", "LIGHTEST_PURPLE", "BLACK",
             "TEXT_BODY", "MID_GRAY", "LIGHT_GRAY", "OFF_WHITE", "WHITE"]
    pal = []
    for n in names:
        c = getattr(h, n, None)
        if c is not None:
            pal.append(str(c).upper())
    for k, v in (getattr(h, "THEME", {}) or {}).get("tokens", {}).items():
        if isinstance(v, str) and v.startswith("#") and len(v) == 7:
            pal.append(v[1:].upper())
    return list(dict.fromkeys(pal))


def _snap(hexv, palette):
    if not palette or hexv in (None, "none"):
        return hexv
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))

    def d(p):
        pr, pg, pb = (int(p[i:i + 2], 16) for i in (0, 2, 4))
        rm = (r + pr) / 2   # 簡易な知覚距離（redmean）
        return (2 + rm / 256) * (r - pr) ** 2 + 4 * (g - pg) ** 2 + (2 + (255 - rm) / 256) * (b - pb) ** 2
    return min(palette, key=d)


# ── 変換行列 ─────────────────────────────────────────────────────
def _mul(m1, m2):
    a1, b1, c1, d1, e1, f1 = m1
    a2, b2, c2, d2, e2, f2 = m2
    return (a1 * a2 + c1 * b2, b1 * a2 + d1 * b2, a1 * c2 + c1 * d2, b1 * c2 + d1 * d2,
            a1 * e2 + c1 * f2 + e1, b1 * e2 + d1 * f2 + f1)


def _parse_transform(t):
    m = (1, 0, 0, 1, 0, 0)
    if not t:
        return m
    for name, args in re.findall(r"(\w+)\s*\(([^)]*)\)", t):
        a = [float(x) for x in re.split(r"[\s,]+", args.strip()) if x]
        if name == "translate":
            m = _mul(m, (1, 0, 0, 1, a[0], a[1] if len(a) > 1 else 0))
        elif name == "scale":
            m = _mul(m, (a[0], 0, 0, a[1] if len(a) > 1 else a[0], 0, 0))
        elif name == "matrix":
            m = _mul(m, tuple(a))
        elif name == "rotate":
            r = math.radians(a[0])
            rot = (math.cos(r), math.sin(r), -math.sin(r), math.cos(r), 0, 0)
            if len(a) == 3:
                m = _mul(m, (1, 0, 0, 1, a[1], a[2]))
                m = _mul(m, rot)
                m = _mul(m, (1, 0, 0, 1, -a[1], -a[2]))
            else:
                m = _mul(m, rot)
        else:
            raise SvgError(f"未対応の transform: {name}")
    return m


def _apply(m, x, y):
    a, b, c, d, e, f = m
    return a * x + c * y + e, b * x + d * y + f


def _has_rotation(m):
    return abs(m[1]) > 1e-9 or abs(m[2]) > 1e-9


# ── path 解析 ────────────────────────────────────────────────────
_TOK = re.compile(r"[MmLlHhVvCcSsQqTtAaZz]|[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?")


def _arc_to_beziers(x1, y1, rx, ry, phi, fa, fs, x2, y2):
    """SVG 円弧 → 3 次ベジェ列（W3C 実装ノートの方式）。"""
    if rx == 0 or ry == 0:
        return [("L", (x2, y2))]
    phi = math.radians(phi)
    cp, sp = math.cos(phi), math.sin(phi)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    rx, ry = abs(rx), abs(ry)
    lam = x1p ** 2 / rx ** 2 + y1p ** 2 / ry ** 2
    if lam > 1:
        rx, ry = rx * math.sqrt(lam), ry * math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    coef = math.sqrt(max(0, num / den)) if den else 0
    if fa == fs:
        coef = -coef
    cxp, cyp = coef * rx * y1p / ry, -coef * ry * x1p / rx
    cx = cp * cxp - sp * cyp + (x1 + x2) / 2
    cy = sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a
    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dt > 0:
        dt -= 2 * math.pi
    elif fs and dt < 0:
        dt += 2 * math.pi
    n = max(1, int(math.ceil(abs(dt) / (math.pi / 2))))
    seg = dt / n
    out = []
    k = 4 / 3 * math.tan(seg / 4)
    for i in range(n):
        a1, a2 = t1 + i * seg, t1 + (i + 1) * seg
        p1 = (math.cos(a1) - k * math.sin(a1), math.sin(a1) + k * math.cos(a1))
        p2 = (math.cos(a2) + k * math.sin(a2), math.sin(a2) - k * math.cos(a2))
        p3 = (math.cos(a2), math.sin(a2))
        pts = []
        for (ux, uy) in (p1, p2, p3):
            X, Y = rx * ux, ry * uy
            pts.append((cp * X - sp * Y + cx, sp * X + cp * Y + cy))
        out.append(("C", tuple(pts)))
    return out


def parse_path(d):
    """path の d 属性 → サブパスのリスト [[("M",(x,y)), ("L",(x,y)), ("C",(p1,p2,p3)), ("Q",(p1,p2)), ("Z",None)], ...]"""
    toks = _TOK.findall(d or "")
    i, cmd = 0, None
    cur = (0.0, 0.0)
    start = (0.0, 0.0)
    last_c2 = None
    last_q = None
    subs, sub = [], []

    def num():
        nonlocal i
        v = float(toks[i]); i += 1
        return v
    while i < len(toks):
        t = toks[i]
        if re.match(r"[A-Za-z]", t):
            cmd = t; i += 1
            if cmd in "Zz":
                sub.append(("Z", None)); cur = start; last_c2 = last_q = None
                continue
        elif cmd is None:
            raise SvgError("path の先頭がコマンドではない")
        rel = cmd.islower()
        C = cmd.upper()
        ox, oy = cur if rel else (0.0, 0.0)
        if C == "M":
            p = (num() + ox, num() + oy)
            if sub:
                subs.append(sub)
            sub = [("M", p)]; cur = start = p
            cmd = "l" if rel else "L"
            last_c2 = last_q = None
        elif C == "L":
            p = (num() + ox, num() + oy); sub.append(("L", p)); cur = p; last_c2 = last_q = None
        elif C == "H":
            x = num() + (cur[0] if rel else 0); cur = (x, cur[1]); sub.append(("L", cur)); last_c2 = last_q = None
        elif C == "V":
            y = num() + (cur[1] if rel else 0); cur = (cur[0], y); sub.append(("L", cur)); last_c2 = last_q = None
        elif C == "C":
            p1 = (num() + ox, num() + oy); p2 = (num() + ox, num() + oy); p3 = (num() + ox, num() + oy)
            sub.append(("C", (p1, p2, p3))); cur = p3; last_c2 = p2; last_q = None
        elif C == "S":
            p1 = (2 * cur[0] - last_c2[0], 2 * cur[1] - last_c2[1]) if last_c2 else cur
            p2 = (num() + ox, num() + oy); p3 = (num() + ox, num() + oy)
            sub.append(("C", (p1, p2, p3))); cur = p3; last_c2 = p2; last_q = None
        elif C == "Q":
            p1 = (num() + ox, num() + oy); p2 = (num() + ox, num() + oy)
            sub.append(("Q", (p1, p2))); cur = p2; last_q = p1; last_c2 = None
        elif C == "T":
            p1 = (2 * cur[0] - last_q[0], 2 * cur[1] - last_q[1]) if last_q else cur
            p2 = (num() + ox, num() + oy)
            sub.append(("Q", (p1, p2))); cur = p2; last_q = p1; last_c2 = None
        elif C == "A":
            rx, ry, phi, fa, fs = num(), num(), num(), int(num()), int(num())
            p = (num() + ox, num() + oy)
            sub.extend(_arc_to_beziers(cur[0], cur[1], rx, ry, phi, fa, fs, p[0], p[1]))
            cur = p; last_c2 = last_q = None
        else:
            raise SvgError(f"未対応の path コマンド: {cmd}")
    if sub:
        subs.append(sub)
    return subs


# ── 図形生成 ─────────────────────────────────────────────────────
class _Ctx:
    def __init__(self, scale, ox, oy, palette, font, warnings):
        self.scale, self.ox, self.oy = scale, ox, oy   # svg 単位 → EMU
        self.palette, self.font, self.warnings = palette, font, warnings

    def pt(self, m, x, y):
        X, Y = _apply(m, x, y)
        return self.ox + X * self.scale, self.oy + Y * self.scale

    def length(self, m, v):
        sx = math.hypot(m[0], m[1])
        sy = math.hypot(m[2], m[3])
        return v * self.scale * (sx + sy) / 2


def _style(el, inherited):
    st = dict(inherited)
    for k in ("fill", "stroke", "stroke-width", "opacity", "fill-opacity", "stroke-opacity",
              "stroke-dasharray", "font-size", "font-weight", "text-anchor", "font-family"):
        if el.get(k) is not None:
            st[k] = el.get(k)
    for part in (el.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            st[k.strip()] = v.strip()
    if "filter" in st or el.get("filter"):
        raise SvgError("filter（影・ぼかし）は使用不可")
    return st


def _set_alpha(color_fmt_xml_parent, alpha):
    if alpha is None or alpha >= 0.999:
        return
    clr = color_fmt_xml_parent.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))


def _apply_fill_line(shape, st, ctx, m, closed=True):
    op = float(st.get("opacity", 1))
    fill = _parse_color(st.get("fill", "black")) if closed else "none"
    stroke = _parse_color(st.get("stroke", "none"))
    if fill == "none" or fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(_snap(fill, ctx.palette))
        _set_alpha(shape.fill._xPr.find(qn("a:solidFill")), op * float(st.get("fill-opacity", 1)))
    ln = shape.line
    if stroke == "none" or stroke is None:
        ln.fill.background()
    else:
        ln.fill.solid()
        ln.fill.fore_color.rgb = RGBColor.from_string(_snap(stroke, ctx.palette))
        w = ctx.length(m, float(st.get("stroke-width", 1)))
        ln.width = Emu(max(int(w), 3175))   # 最小 0.25pt
        if st.get("stroke-dasharray") and st["stroke-dasharray"] != "none":
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ln.dash_style = MSO_LINE_DASH_STYLE.DASH
        lnxml = shape._element.spPr.find(qn("a:ln"))
        if lnxml is not None:
            _set_alpha(lnxml, op * float(st.get("stroke-opacity", 1)))


def _no_text_inset(shape):
    if not getattr(shape, "has_text_frame", False):
        return
    bp = shape.text_frame._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "rIns", "tIns", "bIns"):
        bp.set(k, "0")


def _custgeom_shape(shapes, subpaths_emu, closed_flags, name):
    """EMU 座標のサブパス群から custGeom 図形を作る。"""
    xs = [p[0] for sp in subpaths_emu for seg in sp for p in seg[1]]
    ys = [p[1] for sp in subpaths_emu for seg in sp for p in seg[1]]
    x0, y0 = min(xs), min(ys)
    w, h = max(max(xs) - x0, 1), max(max(ys) - y0, 1)
    shp = shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(int(x0)), Emu(int(y0)), Emu(int(w)), Emu(int(h)))
    shp.name = name
    spPr = shp._element.spPr
    prst = spPr.find(qn("a:prstGeom"))
    cust = etree.Element(qn("a:custGeom"))
    for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        etree.SubElement(cust, qn(tag))
    rect = etree.SubElement(cust, qn("a:rect")); rect.set("l", "l"); rect.set("t", "t"); rect.set("r", "r"); rect.set("b", "b")
    pl = etree.SubElement(cust, qn("a:pathLst"))
    for sp, closed in zip(subpaths_emu, closed_flags):
        path = etree.SubElement(pl, qn("a:path")); path.set("w", str(int(w))); path.set("h", str(int(h)))
        if not closed:
            path.set("fill", "none")
        for kind, pts in sp:
            tag = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo", "Q": "a:quadBezTo"}[kind]
            el = etree.SubElement(path, qn(tag))
            for (px, py) in pts:
                p = etree.SubElement(el, qn("a:pt"))
                p.set("x", str(int(round(px - x0)))); p.set("y", str(int(round(py - y0))))
        if closed:
            etree.SubElement(path, qn("a:close"))
    spPr.replace(prst, cust)
    _no_text_inset(shp)
    return shp


def _subpaths_to_emu(subs, ctx, m):
    out, closed = [], []
    for sub in subs:
        segs, cl = [], False
        for kind, v in sub:
            if kind == "Z":
                cl = True
                continue
            pts = [v] if kind in ("M", "L") else list(v)
            segs.append((kind, [ctx.pt(m, x, y) for (x, y) in pts]))
        if segs:
            out.append(segs); closed.append(cl)
    return out, closed


def _f(el, k, d=0.0):
    v = el.get(k)
    if v is None:
        return d
    return float(re.sub(r"[a-z%]+$", "", v))


def _render(el, shapes, ctx, m, st, counter):
    tag = etree.QName(el).localname
    if tag not in ALLOWED:
        raise SvgError(f"未対応の要素 <{tag}>。使えるのは {sorted(ALLOWED)}")
    m = _mul(m, _parse_transform(el.get("transform")))
    st = _style(el, st)
    counter[0] += 1
    nm = f"{tag}{counter[0]}"
    if tag in ("svg", "g"):
        for ch in el:
            if isinstance(ch.tag, str):
                _render(ch, shapes, ctx, m, st, counter)
        return
    if tag in ("title", "desc", "tspan"):
        return
    if tag == "rect":
        x, y, w, h = _f(el, "x"), _f(el, "y"), _f(el, "width"), _f(el, "height")
        if el.get("rx") or el.get("ry"):
            ctx.warnings.append(f"{nm}: rx/ry は無視（角丸はブランド規定で禁止）")
        if _has_rotation(m):
            subs = [[("M", (x, y)), ("L", (x + w, y)), ("L", (x + w, y + h)), ("L", (x, y + h)), ("Z", None)]]
            e, c = _subpaths_to_emu(subs, ctx, m)
            shp = _custgeom_shape(shapes, e, c, nm)
        else:
            (X0, Y0), (X1, Y1) = ctx.pt(m, x, y), ctx.pt(m, x + w, y + h)
            shp = shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(int(min(X0, X1))), Emu(int(min(Y0, Y1))),
                                   Emu(int(abs(X1 - X0))), Emu(int(abs(Y1 - Y0))))
            shp.name = nm
            _no_text_inset(shp)
        _apply_fill_line(shp, st, ctx, m)
    elif tag in ("circle", "ellipse"):
        cx, cy = _f(el, "cx"), _f(el, "cy")
        rx = _f(el, "r") if tag == "circle" else _f(el, "rx")
        ry = _f(el, "r") if tag == "circle" else _f(el, "ry")
        if _has_rotation(m):
            k = 0.5522847498
            subs = [[("M", (cx + rx, cy)),
                     ("C", ((cx + rx, cy + k * ry), (cx + k * rx, cy + ry), (cx, cy + ry))),
                     ("C", ((cx - k * rx, cy + ry), (cx - rx, cy + k * ry), (cx - rx, cy))),
                     ("C", ((cx - rx, cy - k * ry), (cx - k * rx, cy - ry), (cx, cy - ry))),
                     ("C", ((cx + k * rx, cy - ry), (cx + rx, cy - k * ry), (cx + rx, cy))), ("Z", None)]]
            e, c = _subpaths_to_emu(subs, ctx, m)
            shp = _custgeom_shape(shapes, e, c, nm)
        else:
            (X0, Y0), (X1, Y1) = ctx.pt(m, cx - rx, cy - ry), ctx.pt(m, cx + rx, cy + ry)
            shp = shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Emu(int(min(X0, X1))), Emu(int(min(Y0, Y1))),
                                   Emu(int(abs(X1 - X0))), Emu(int(abs(Y1 - Y0))))
            shp.name = nm
            _no_text_inset(shp)
        _apply_fill_line(shp, st, ctx, m)
    elif tag == "line":
        (X0, Y0), (X1, Y1) = ctx.pt(m, _f(el, "x1"), _f(el, "y1")), ctx.pt(m, _f(el, "x2"), _f(el, "y2"))
        shp = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(X0)), Emu(int(Y0)), Emu(int(X1)), Emu(int(Y1)))
        shp.name = nm
        st2 = dict(st); st2.setdefault("stroke", "black")
        if st2.get("stroke") == "none":
            st2["stroke"] = "black"
        stroke = _snap(_parse_color(st2["stroke"]), ctx.palette)
        shp.line.fill.solid(); shp.line.fill.fore_color.rgb = RGBColor.from_string(stroke)
        shp.line.width = Emu(max(int(ctx.length(m, float(st2.get("stroke-width", 1)))), 3175))
        if st2.get("stroke-dasharray") and st2["stroke-dasharray"] != "none":
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    elif tag in ("polyline", "polygon"):
        nums = [float(v) for v in re.split(r"[\s,]+", (el.get("points") or "").strip()) if v]
        pts = list(zip(nums[0::2], nums[1::2]))
        if len(pts) < 2:
            return
        sub = [("M", pts[0])] + [("L", p) for p in pts[1:]] + ([("Z", None)] if tag == "polygon" else [])
        e, c = _subpaths_to_emu([sub], ctx, m)
        shp = _custgeom_shape(shapes, e, c, nm)
        _apply_fill_line(shp, st, ctx, m, closed=(tag == "polygon"))
    elif tag == "path":
        subs = parse_path(el.get("d"))
        e, c = _subpaths_to_emu(subs, ctx, m)
        if not e:
            return
        shp = _custgeom_shape(shapes, e, c, nm)
        _apply_fill_line(shp, st, ctx, m, closed=any(c) or _parse_color(st.get("fill", "black")) != "none")
    elif tag == "text":
        txt = "".join(el.itertext()).strip()
        if not txt:
            return
        size_u = float(re.sub(r"[a-z]+$", "", str(st.get("font-size", 16))))
        size_emu = ctx.length(m, size_u)
        size_pt = max(8, round(size_emu / 12700))
        x, y = ctx.pt(m, _f(el, "x"), _f(el, "y"))
        anchor = st.get("text-anchor", "start")
        # 幅は文字数から概算（全角 1.0em、半角 0.55em）
        em = sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in txt)
        w = em * size_emu * 1.08 + 0.1 * EMU_PER_IN
        h = size_emu * 1.45
        left = {"middle": x - w / 2, "end": x - w}.get(anchor, x)
        tb = shapes.add_textbox(Emu(int(left)), Emu(int(y - size_emu * 1.05)), Emu(int(w)), Emu(int(h)))
        tb.name = nm
        tf = tb.text_frame
        tf.word_wrap = False
        _no_text_inset(tb)
        p = tf.paragraphs[0]
        p.alignment = {"middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}.get(anchor, PP_ALIGN.LEFT)
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size_pt)
        r.font.name = ctx.font
        r.font.bold = str(st.get("font-weight", "")).lower() in ("bold", "600", "700", "800", "900")
        col = _parse_color(st.get("fill", "black"))
        r.font.color.rgb = RGBColor.from_string(_snap(col if col != "none" else "000000", ctx.palette))


def add_svg(slide, svg, x, y, w, h, name="SVG図", snap=True, palette=None, font_name=None):
    """SVG（文字列またはファイルパス）を編集可能な図形グループとして slide に追加し、グループ図形を返す。

    x, y, w, h: 配置枠（インチ）。viewBox を等倍縮尺で枠内中央に収める。
    戻り値のグループには .svg_warnings（list[str]）が付く。
    """
    if isinstance(svg, str) and not svg.lstrip().startswith("<") and os.path.isfile(svg):
        with open(_paths.check_path(svg), encoding="utf-8") as f:
            svg = f.read()
    root = etree.fromstring(svg.encode("utf-8") if isinstance(svg, str) else svg, _PARSER)
    if etree.QName(root).localname != "svg":
        raise SvgError("ルート要素が <svg> ではない")
    vb = root.get("viewBox")
    if vb:
        vx, vy, vw, vh = [float(v) for v in re.split(r"[\s,]+", vb.strip())]
    else:
        vx, vy, vw, vh = 0, 0, _f(root, "width", 100), _f(root, "height", 100)
    bx, by, bw, bh = (v * EMU_PER_IN for v in (x, y, w, h))
    scale = min(bw / vw, bh / vh)
    ox = bx + (bw - vw * scale) / 2 - vx * scale
    oy = by + (bh - vh * scale) / 2 - vy * scale
    if palette is None and snap:
        try:
            palette = theme_palette()
        except Exception:
            palette = None
    if font_name is None:
        try:
            import helpers as _h
            font_name = _h.FONT
        except Exception:
            font_name = "Meiryo UI"
    warnings = []
    ctx = _Ctx(scale, ox, oy, palette if snap else None, font_name, warnings)
    grp = slide.shapes.add_group_shape()
    grp.name = name
    _render(root, grp.shapes, ctx, (1, 0, 0, 1, 0, 0), {}, [0])
    if len(grp.shapes) == 0:
        raise SvgError("描画できる要素がない")
    for shp in grp.shapes:   # 既定スタイル（accent1 参照）を外し、指定色だけにする
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
    grp.svg_warnings = warnings
    return grp


def lint_svg(svg_text):
    """使用可能要素・属性だけで書かれているかを検査し、問題のリストを返す（空なら OK）。"""
    issues = []
    try:
        root = etree.fromstring(svg_text.encode("utf-8"), _PARSER)
    except etree.XMLSyntaxError as e:
        return [f"XML として不正: {e}"]
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        tag = etree.QName(el).localname
        if tag not in ALLOWED:
            issues.append(f"未対応の要素 <{tag}>")
        for k in ("fill", "stroke"):
            v = el.get(k)
            if v and v.strip().startswith("url("):
                issues.append(f"<{tag}> {k}=url(...)（グラデーション等）は不可")
        if el.get("filter"):
            issues.append(f"<{tag}> filter は不可")
    return issues


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("svg")
    ap.add_argument("out")
    ap.add_argument("--theme")
    a = ap.parse_args()
    a.out = _paths.check_path(a.out, write=True)
    import helpers as _h
    if a.theme:
        _h.load_theme(a.theme)
    from pptx import Presentation
    prs = Presentation(_h.TEMPLATE_PATH)
    while len(prs.slides):
        rid = prs.slides._sldIdLst[0].get(qn("r:id")); prs.part.drop_rel(rid); del prs.slides._sldIdLst[0]
    s = prs.slides.add_slide(prs.slide_layouts[_h.layout_idx("content")])
    g = add_svg(s, a.svg, 0.42, 1.5, 12.5, 5.3)
    prs.save(a.out)
    print(f"saved {a.out}: {len(g.shapes)} shapes", *g.svg_warnings, sep="\n")
