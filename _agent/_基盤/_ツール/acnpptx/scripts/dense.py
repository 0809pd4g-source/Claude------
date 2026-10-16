"""
dense.py — 情報量が多いのに整理されて見やすいページの部品（v6.5）。

人の提案書（1 枚 400〜900 字）の作法を、寸法の決まった部品にした。
  リッチ本文 rich()   … ボックスの中身を「小見出し → 箇条書き 2 段 → 強調 1 句 → ⇒ 結論」のブロックの列で書く
  M-ZONE              … 見出し（帯か下線）＋面（灰か淡紫）＋リッチ本文。パッケージ概要・課題/解決の列の単位
  M-ROWHEAD（Why 型） … 番号＋課題（淡紫）◀ 記号 a/b/c ＋強みの要約（太字）｜縦線｜詳細。columns の見出し帯付き
  M-RHMATRIX          … 行見出し（灰）× 列見出し（下線・帯・矢羽）のマトリクス。セルはリッチ本文、主張する列は淡紫
  M-STEPS             … 下向きの五角形（番号＋工程名）＋説明ボックスを縦に並べる。右端にアイコンの流れ（任意）
  M-FIGNOTE           … 左に図（他の部品を入れ子）、右に番号の丸＋解説のリスト（図とリストで同じ番号）
  M-CARD（強み型）    … 番号タブ＋紫の見出しブロック＋灰の面の要点（太字）＋補足。下端の結論帯は closing
  M-ICONGRID（実績型）… アイコン（小）＋数字（太字）＋説明の格子。アイコンは文字の左で主役にしない

文字の段（人の資料の実測）: ゾーン見出し 16 → 小見出し 14 太字濃紫 → 本文 12 → 補足 11 → 注記 11。
隙間の 3 段階: セル・行 0.06 / カード・列 0.12〜0.16 / ゾーン 0.4〜1.0（三角を挟む）。
"""

import math

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as S
from pptx.enum.text import PP_ALIGN

import helpers as _h


def _c():
    import compose
    return compose


# 本文の大きさ → (小見出し, 補足)。13pt は使わない
TYPE = {16: (18, 14), 14: (16, 12), 12: (14, 11), 11: (14, 11)}
ICON = {"inline": 0.45, "grid": 0.55, "head": 0.7}     # アイコンの大きさ（人の資料の中央値 0.47in、p75 0.81in）


def body_sizes():
    """リッチ本文の大きさの候補（面を埋める最大を選ぶ。人の資料は面の広い提案書で 14〜16pt、密なページで 12・11pt）。
    11pt は密なページ（density high・本文 400 字以上）だけ。スライド内のゾーンでそろえる大きさがあればそれだけ。"""
    c = _c()
    if c._CUR.get("rich_base"):
        return (c._CUR["rich_base"],)
    return (16, 14, 12) + ((11,) if c._CUR.get("allow11") else ())


# ═══════════════════════════════════════════════════════════════
#  リッチ本文
# ═══════════════════════════════════════════════════════════════
def rich(body, bullet="•"):
    """body（文字列 / 文字列のリスト / ブロックのリスト）を、役割付きの段落 [(text, {"role":..})] にする。
    ブロック: {"h": 小見出し} {"p": 文} {"b": [文 | {"text":..,"sub":[..]}]} {"n": [文..]}（①②）
              {"concl": 結論}（⇒ 太字） {"note": 補足}（小さめ・灰）。文字列の要素は箇条書き 1 段目。
    文中の **語** は太字、==語== は太字＋濃い紫（色付きは 1 ブロック 1 句まで）。"""
    if body is None:
        return []
    if isinstance(body, dict):
        body = [body]
    if isinstance(body, str):
        return [(body, {"role": "body"})]
    out = []
    circled = "①②③④⑤⑥⑦⑧⑨⑩"
    for blk in body:
        if isinstance(blk, str):
            out.append((blk, {"role": "body", "bullet": bullet}))
            continue
        if not isinstance(blk, dict):
            continue
        if "h" in blk:
            out.append((blk["h"], {"role": "sub"}))
        if "p" in blk:
            out.append((blk["p"], {"role": "body"}))
        for it in blk.get("b") or []:
            t, subs = (it, []) if isinstance(it, str) else (it.get("text", ""), it.get("sub") or [])
            out.append((t, {"role": "body", "bullet": blk.get("mark", bullet)}))
            for sb in subs:
                out.append((sb, {"role": "aux", "bullet": "–", "level": 1}))
        for k, t in enumerate(blk.get("n") or []):
            out.append((f"{circled[k] if k < 10 else k + 1} {t}", {"role": "body"}))
        if "text" in blk and not any(k in blk for k in ("h", "p", "b", "n")):
            out.append((blk["text"], {"role": "body"}))
        if "concl" in blk:
            t = blk["concl"]
            out.append((t if t.startswith(("⇒", "→")) else "⇒ " + t, {"role": "concl"}))
        if "note" in blk:
            out.append((blk["note"], {"role": "aux", "color": _h.TEXT_SUB}))
    return out


def sized(paras, base, color=None, sub_color=None, align=None, space=4):
    """役割の段落に、本文の大きさ base からの大きさ・太字・色・間隔を付ける。"""
    sub, aux = TYPE.get(base, (base + 2, max(11, base - 1)))
    out = []
    for i, (t, o) in enumerate(paras):
        role = o.get("role", "body")
        oo = {k: v for k, v in o.items() if k != "role"}
        if role == "sub":
            oo.update(size=sub, bold=True, color=sub_color or _h.DARKEST_PURPLE, space_after=2)
            if i:
                oo["space_before"] = 4
        elif role == "aux":
            oo.setdefault("color", color or _h.TEXT_BODY)
            oo.update(size=aux, space_after=2)
        elif role == "concl":
            oo.update(size=base, bold=True, color=sub_color or _h.DARKEST_PURPLE, space_after=0)
        else:
            oo.setdefault("color", color or _h.TEXT_BODY)
            oo.update(size=base, space_after=space)
        if align is not None:
            oo["align"] = align
        out.append((t, oo))
    return out


def height(paras, w, base):
    """sized 済みの段落の高さ（段落ごとの大きさ・箇条書きの字下げを入れて見積もる）。"""
    c = _c()
    tot = 0.08
    for t, o in paras:
        sz = o.get("size", base)
        ww = w - 0.12 - (0.18 * (o.get("level", 0) + 1) if o.get("bullet") else 0)
        if o.get("bold"):
            ww /= 1.08
        tot += c.est_lines(t, ww + 0.12, sz) * c._eff(sz) * 1.22 / 72 + (o.get("space_after", 4) + o.get("space_before", 0)) / 72
    return tot


def fit(body_paras, w, h, sizes=None, **kw):
    """収まる最大の本文の大きさと、その大きさで sized した段落。収まらなければ最小で (…, False)。"""
    sizes = sizes or body_sizes()
    for s in sizes:
        p = sized(body_paras, s, **kw)
        if height(p, w, s) <= h:
            return s, p, True
    s = sizes[-1]
    return s, sized(body_paras, s, **kw), False


def draw_text(ctx, x, y, w, h, paras, base, anchor="t", inset=(0.06, 0.04, 0.06, 0.04)):
    c = _c()
    tb = c._textbox(ctx.slide, x, y, w, h, paras, base, anchor=anchor, space_after=4, emphasis=ctx.emphasis, inset=inset)
    for p, (_, o) in zip(tb.text_frame.paragraphs, paras):
        if o.get("space_before"):
            from pptx.util import Pt
            p.space_before = Pt(o["space_before"])
    return tb


def rich_box(ctx, x, y, w, h, body, anchor="t", warn_label="", sizes=None, **kw):
    """リッチ本文を枠に描き、使った高さを返す。"""
    paras = rich(body, kw.pop("bullet", "•"))
    if not paras:
        return 0
    base, sp, ok = fit(paras, w - 0.12, h - 0.08, sizes, **kw)
    if not ok:
        ctx.warn(f"{warn_label}: 文字量が枠を超える（要点を削るか、ゾーンを分ける）")
    draw_text(ctx, x, y, w, h, sp, base, anchor=anchor)
    return height(sp, w - 0.12, base)


def chars_of(body):
    return sum(len(t) for t, _ in rich(body))


# ═══════════════════════════════════════════════════════════════
#  共通の小物
# ═══════════════════════════════════════════════════════════════
def tri(ctx, cx, cy, vw, vh, direction="right", color=None):
    """向きの三角（灰）。vw×vh は見た目の幅×高さ。人の資料: 高さは指す先の箱の 65〜70%、文字は入れない。"""
    c = _c()
    rot = {"up": 0, "right": 90, "down": 180, "left": 270}[direction]
    w0, h0 = (vh, vw) if rot in (90, 270) else (vw, vh)    # 回転前の幅・高さ（回転で縦横が入れ替わる）
    t = c._rect(ctx.slide, cx - w0 / 2, cy - h0 / 2, w0, h0, fill=color or _h.LIGHT_GRAY, shape=S.ISOSCELES_TRIANGLE)
    t.rotation = rot
    return t


def num_circle(ctx, cx, cy, d, n, fill=None, size=12):
    c = _c()
    b = c._rect(ctx.slide, cx - d / 2, cy - d / 2, d, d, fill=fill or _h.DARK_PURPLE, shape=S.OVAL)
    c._shape_text(b, str(n), size, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
    return b


def zone_header(ctx, x, y, w, text, style="band", emph=False):
    """ゾーン見出し。band: 濃い紫の帯に白の太字 / rule: 太字の黒＋0.75pt の下線。高さを返す。"""
    c = _c()
    if style == "rule":
        size, ok = c.fit_size([text], w - 0.1, 0.42, (16, 14))
        c._textbox(ctx.slide, x, y, w, 0.4, [text], size, _h.BLACK, bold=True, anchor="b", align=PP_ALIGN.CENTER,
                   inset=(0.02, 0, 0.02, 0.03), space_after=0)
        c._hline(ctx.slide, x, y + 0.42, w, _h.BLACK, 0.75)
        return 0.5
    hh = Z()["hdr"]
    band = c._rect(ctx.slide, x, y, w, hh, fill=_h.CORE_PURPLE if emph else _h.DARK_PURPLE)
    size, ok = c.fit_size([text], w - 0.2, hh + 0.12, (16, 14, 12))     # 1 行なら帯の高さに収まる（見積もりの余白分を足す）
    if not ok:
        ctx.warn(f"ゾーン見出しが長い: {text[:16]}")
    c._shape_text(band, text, size, _h.WHITE, True, inset=(0.06, 0, 0.06, 0))
    return hh


def Z():
    """ゾーンの寸法（density max は見出しの帯 0.34・面の内側の余白 0.08。ほかは 0.4・0.12）。"""
    return {"hdr": 0.34, "pad": 0.08} if _c()._CUR.get("max") else {"hdr": 0.4, "pad": 0.12}


def _fill(kind):
    return {"gray": _h.OFF_WHITE, "purple": _h.LIGHTEST_PURPLE, "white": _h.WHITE, "none": None}.get(kind, _h.OFF_WHITE)


def icon_row(ctx, x, y, w, icons, isz=None):
    """小さなアイコン＋11pt のラベルを横に並べる（ゾーン内の図）。使った高さを返す。"""
    c = _c()
    n = len(icons)
    if not n:
        return 0
    isz = isz or ICON["grid"]
    cw = w / n
    lab_h = max(c.est_height([it.get("label", "")], cw - 0.06, 11, 0) for it in icons)
    for i, it in enumerate(icons):
        cx = x + cw * i + cw / 2
        if it.get("icon"):
            c._place_visual(ctx, it["icon"], (cx - isz / 2, y, isz, isz), color=_h.DARKEST_PURPLE)
        if it.get("label"):
            c._textbox(ctx.slide, cx - cw / 2, y + isz + 0.04, cw, lab_h, [it["label"]], 11, _h.TEXT_BODY,
                       align=PP_ALIGN.CENTER, space_after=0, inset=(0.02, 0, 0.02, 0))
    return isz + 0.06 + lab_h


# ═══════════════════════════════════════════════════════════════
#  M-ZONE
# ═══════════════════════════════════════════════════════════════
def m_zone(ctx, r, d):
    """{"title": "課題・ニーズ", "style": "band|rule", "fill": "gray|purple|white|none", "emph": false,
        "lead": "全社 Claude 導入の壁"（面の先頭の太字・任意）, "body": [リッチ本文],
        "icons": [{"icon": "q:shield", "label": "セキュリティ設定"}, ...]（面の上部に小さく・任意）,
        "sections": [{"title":..., "body": [...]}, ...]（ゾーンを縦に分ける・任意）}"""
    c = _c()
    x, y, w, h = r
    style = d.get("style", "band")
    secs = d.get("sections")
    if secs:          # 縦に分けたゾーン: 中身の量に比例した高さ
        need = [max(0.8, chars_of(sc.get("body")) / 30 * 0.2 + 0.6) for sc in secs]
        gap = 0.08
        avail = h - gap * (len(secs) - 1)
        yy = y
        for sc, nd in zip(secs, need):
            hh = avail * nd / sum(need)
            m_zone(ctx, (x, yy, w, hh), {**{k: v for k, v in d.items() if k not in ("sections", "_blk")}, **sc})
            yy += hh + gap
        return
    if d.get("title"):
        used = zone_header(ctx, x, y, w, d["title"], style, d.get("emph"))
        y, h = y + used + (0.0 if style == "band" else 0.06), h - used - (0.0 if style == "band" else 0.06)
    elif d.get("_hdr_pad"):          # 横に並ぶ他のゾーンに見出しがあるときは、面の上端をそろえる
        y, h = y + d["_hdr_pad"], h - d["_hdr_pad"]
    fill = _fill(d.get("fill", "gray" if style == "band" else "none"))
    if style == "summary" and d.get("fill") is None:
        fill = None
    if d.get("emph") and d.get("fill") is None:
        fill = _h.LIGHTEST_PURPLE
    if fill is not None:
        c._rect(ctx.slide, x, y, w, h, fill=fill)
    elif d.get("frame"):
        c._panel(ctx.slide, x, y, w, h)
    pad = Z()["pad"]
    iy = y + pad
    body = list(d.get("body") or ([d["text"]] if d.get("text") else []))
    if d.get("lead"):
        body = [{"h": d["lead"]}] + body
    if body and style != "summary" and d.get("anchor") is None:
        # 中身の塊（アイコン＋本文）が面の 9 割未満なら、塊を面の縦中央に置く（上に詰めて下を空けない）
        icon_h = (ICON["grid"] + 0.3 + 0.12) if d.get("icons") else 0
        base_, sp_, _ = fit(rich(body, d.get("bullet", "•")), w - 2 * pad + 0.08 - 0.12, y + h - iy - icon_h - pad / 2 - 0.08)
        blk = icon_h + height(sp_, w - 2 * pad + 0.08 - 0.12, base_)
        if d.get("_blk"):           # 並ぶゾーン: 中身の量が近いゾーンどうしは、いちばん多い塊に合わせて同じだけ下げる（上端がそろう）。
            hd = (Z()["hdr"] if style != "rule" else 0.56) if d.get("title") else (d.get("_hdr_pad") or 0)
            row = d["_blk"] - hd - pad - 0.3      # 量が大きく違う（8 割未満）ゾーンは自分の塊で縦中央に置く（上に詰めて下を空けない）
            if blk >= row * 0.8:
                blk = max(blk, row)
        if blk < (h - 2 * pad) * 0.9:
            iy += (h - 2 * pad - blk) / 2
    if d.get("icons"):
        iy += icon_row(ctx, x + pad, iy, w - 2 * pad, d["icons"]) + 0.12
    if body and style == "summary":       # 要旨（エグゼクティブサマリー）: 番号の箱＋灰の面の行（同じ高さ）
        summary_rows(ctx, (x, y, w, h), body)
    elif body:
        used = rich_box(ctx, x + pad - 0.04, iy, w - 2 * pad + 0.08, y + h - iy - pad / 2, body,
                        anchor=d.get("anchor", "t"), warn_label=f"M-ZONE「{d.get('title', '')[:10]}」",
                        bullet=d.get("bullet", "•"))
        if used < (y + h - iy) * 0.45 and not d.get("icons"):
            ctx.warn(f"M-ZONE「{d.get('title', '')[:10]}」: 面の下が空く（根拠・数字を足すか、面を小さくする）")


def summary_rows(ctx, r, body):
    """要旨の行: 左に番号の箱（濃い紫・白の数字）、右に灰の面（要点 16〜18 太字＋説明 14〜16）。行はすべて同じ高さ。
    項目は b / n の文字列（**語** は太字）か {"lead": 要点, "text": 説明}。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    items = []
    for blk in body:
        if isinstance(blk, str):
            items.append({"text": blk})
            continue
        for it in (blk.get("b") or []) + (blk.get("n") or []):
            items.append(it if isinstance(it, dict) else {"text": it})
        if blk.get("p"):
            items.append({"text": blk["p"]})
    n = max(len(items), 1)
    gap = 0.12
    nw = 0.55
    tw = w - nw - 0.06 - 0.3

    def paras(it, base):
        out = []
        if it.get("lead"):
            out.append((it["lead"], {"size": base + 2, "bold": True, "color": _h.DARKEST_PURPLE, "space_after": 2}))
        if it.get("text"):
            out.append((it["text"], {"size": base, "color": _h.TEXT_BODY, "space_after": 0}))
        return out
    avail = (h - gap * (n - 1)) / n
    base = 14
    for sz in (18, 16, 14):
        if all(height(paras(it, sz), tw, sz) <= avail - 0.16 for it in items):
            base = sz
            break
    need = max(height(paras(it, base), tw, base) for it in items) + 0.24
    rh = min(avail, max(need, 0.7))
    tot = rh * n + gap * (n - 1)
    yy = y + max(0, (h - tot) / 2)
    for k, it in enumerate(items):
        nb = c._rect(s, x, yy, nw, rh, fill=_h.DARK_PURPLE)
        c._shape_text(nb, str(k + 1), 20, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        c._rect(s, x + nw + 0.06, yy, w - nw - 0.06, rh, fill=_h.OFF_WHITE)
        ps = paras(it, base)
        if height(ps, tw, base) > rh:
            ctx.warn(f"M-ZONE(要旨): 項目 {k + 1} が長い（1 項目 2〜3 行）")
        draw_text(ctx, x + nw + 0.2, yy, tw + 0.1, rh, ps, base, anchor="m")
        yy += rh + gap
    if tot < h * 0.7:
        ctx.warn("M-ZONE(要旨): 下が空く（各項目に根拠・数字を足す。目安は 1 項目 2〜3 行）")


# ═══════════════════════════════════════════════════════════════
#  M-ROWHEAD（Why 型）
# ═══════════════════════════════════════════════════════════════
def zone_base(r, d):
    """ゾーンの本文が収まる最大の大きさ（描かずに見積もる）。横に並ぶゾーンの大きさをそろえるのに使う。"""
    x, y, w, h = r
    if d.get("sections"):
        secs = d["sections"]
        need = [max(0.8, chars_of(sc.get("body")) / 30 * 0.2 + 0.6) for sc in secs]
        avail = h - 0.08 * (len(secs) - 1)
        return min(zone_base((x, y, w, avail * nd / sum(need)), {**{k: v for k, v in d.items() if k != "sections"}, **sc})
                   for sc, nd in zip(secs, need))
    style = d.get("style", "band")
    if d.get("title"):
        used = Z()["hdr"] if style != "rule" else 0.5
        y, h = y + used + (0 if style == "band" else 0.06), h - used - (0 if style == "band" else 0.06)
    pad = Z()["pad"]
    iy = y + pad + ((ICON["grid"] + 0.3 + 0.12) if d.get("icons") else 0)
    body = list(d.get("body") or ([d["text"]] if d.get("text") else []))
    if d.get("lead"):
        body = [{"h": d["lead"]}] + body
    if not body:
        return 16
    return fit(rich(body, d.get("bullet", "•")), w - 2 * pad + 0.08 - 0.12, y + h - iy - pad / 2 - 0.08)[0]


def zone_need(r, d, base):
    """ゾーンを大きさ base で描いたときに要る高さ（見出し＋アイコン＋本文＋余白）。sections は枠いっぱいとみなす。"""
    x, y, w, h = r
    if d.get("sections") or not (d.get("body") or d.get("text") or d.get("lead")):
        return h
    style = d.get("style", "band")
    head = (Z()["hdr"] if style != "rule" else 0.56) if d.get("title") else (d.get("_hdr_pad") or 0)
    icons = (ICON["grid"] + 0.3 + 0.12) if d.get("icons") else 0
    body = list(d.get("body") or ([d["text"]] if d.get("text") else []))
    if d.get("lead"):
        body = [{"h": d["lead"]}] + body
    sp = sized(rich(body, d.get("bullet", "•")), base)
    return min(h, head + Z()["pad"] + icons + height(sp, w - 2 * Z()["pad"] - 0.04, base) + 0.3)


def is_why(d):
    return bool(d.get("columns")) or any(isinstance(rw, dict) and (rw.get("lead") or rw.get("badge") or rw.get("detail"))
                                         for rw in d.get("rows") or [])


def m_whyrows(ctx, r, d):
    """{"columns": ["貴社の課題", "アクセンチュアの強み"],
        "rows": [{"head": "競合に対する遅れの挽回", "sub": "（補足）", "badge": "a",
                  "lead": "業界の**豊富な支援経験**を生かし、先行者の轍を避ける",
                  "detail": ["…**要点**…", "…"]}, ...], "numbered": true}"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    rows = [rw for rw in d.get("rows") or [] if isinstance(rw, dict)]
    n = max(len(rows), 1)
    numbered = d.get("numbered", True)
    num_w = 0.4 if numbered else 0
    longest = max([sum(c._char_w(ch, 14) for ch in str(rw.get("head", ""))) for rw in rows] + [1.0])
    head_w = d.get("head_w") or min(2.5, max(1.6, longest / 2 + 0.35))
    has_lead = any(rw.get("lead") for rw in rows)
    has_badge = any(rw.get("badge") for rw in rows)
    tri_w, badge_w = 0.24, (0.36 if has_badge else 0)
    lead_w = d.get("lead_w") or (min(3.0, w * 0.24) if has_lead else 0)
    left_w = num_w + head_w
    dx0 = x + left_w + 0.1 + tri_w + 0.1        # 右側（記号・要約・詳細）の左端
    if d.get("columns"):
        cols = d["columns"]
        for (cx_, cw_), t in zip(((x, left_w), (dx0, x + w - dx0)), cols):
            zone_header(ctx, cx_, y, cw_, t, "band")
        y, h = y + 0.4 + 0.14, h - 0.54
    gap = 0.16
    detail_w = x + w - dx0 - badge_w - (0.1 if badge_w else 0) - lead_w - (0.25 if lead_w else 0)
    det_sizes = body_sizes()
    need = []
    for rw in rows:
        dp = sized(rich(rw.get("detail")), det_sizes[-1])
        lp = sized(rich(rw.get("lead")), 14)
        need.append(max(height(dp, detail_w - 0.12, det_sizes[-1]), height(lp, lead_w - 0.1, 14) if lead_w else 0,
                        0.7) + 0.14)
    avail = h - gap * (n - 1)
    heights = [avail * k / sum(need) for k in need]
    heights = [min(hh, nd * 1.7 + 0.2) for hh, nd in zip(heights, need)]
    # 詳細の文字は全行で同じ大きさ
    dsz = min(fit(rich(rw.get("detail")), detail_w - 0.12, hh - 0.1, (14, 12) + ((11,) if c._CUR.get("allow11") else ()))[0]
              for rw, hh in zip(rows, heights)) if rows else 12
    tot = sum(heights) + gap * (n - 1)
    ry = y + max(0, (h - tot) / 2) if tot < h * 0.9 else y
    for i, (rw, rh) in enumerate(zip(rows, heights)):
        if numbered:
            nb = c._rect(s, x, ry, num_w, rh, fill=_h.LIGHTEST_PURPLE)
            c._shape_text(nb, str(rw.get("number", i + 1)), 16, _h.DARKEST_PURPLE, True, inset=(0, 0, 0, 0), wrap=False)
        hb = c._rect(s, x + num_w + 0.06 if numbered else x, ry, head_w - (0.06 if numbered else 0), rh,
                     fill=_h.CORE_PURPLE if rw.get("emph") else _h.LIGHTEST_PURPLE)
        hp = [(rw.get("head", ""), {"bold": True, "color": _h.WHITE if rw.get("emph") else _h.DARKEST_PURPLE,
                                     "align": PP_ALIGN.CENTER, "space_after": 2})]
        if rw.get("sub"):
            hp.append((rw["sub"], {"size": 11, "color": _h.WHITE if rw.get("emph") else _h.TEXT_BODY, "align": PP_ALIGN.CENTER}))
        hs, ok = c.fit_size(hp, head_w - 0.2, rh - 0.1, (16, 14, 12))
        c._textbox(s, hb.left / 914400, ry, hb.width / 914400, rh, [(t, {**o, "size": o.get("size", hs)}) for t, o in hp],
                   hs, anchor="m", space_after=2, inset=(0.08, 0.04, 0.08, 0.04))
        tri(ctx, x + left_w + 0.1 + tri_w / 2, ry + rh / 2, tri_w, rh * 0.7, "left")
        bx = dx0
        if badge_w:
            if rw.get("badge"):
                b = c._rect(s, bx, ry + rh / 2 - 0.16, badge_w, 0.32, fill=_h.DARK_PURPLE)
                c._shape_text(b, str(rw["badge"]), 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
            bx += badge_w + 0.1
        if lead_w:
            lp = sized(rich(rw.get("lead")), 14, color=_h.DARKEST_PURPLE)
            lp = [(t, {**o, "bold": True}) for t, o in lp]
            ls = 14 if height(lp, lead_w - 0.1, 14) <= rh else 12
            if ls == 12:
                lp = [(t, {**o, "size": 12}) for t, o in lp]
            draw_text(ctx, bx, ry, lead_w, rh, lp, ls, anchor="m")
            bx += lead_w + 0.12
            c._connector(s, bx, ry + 0.08, bx, ry + rh - 0.08, _h.MID_GRAY, 0.75)
            bx += 0.13
        dp = sized(rich(rw.get("detail")), dsz)
        if not rw.get("detail"):
            ctx.warn(f"M-ROWHEAD(Why): 行「{rw.get('head', '')[:10]}」に詳細が無い")
        elif height(dp, detail_w - 0.12, dsz) > rh:
            ctx.warn(f"M-ROWHEAD(Why): 行「{rw.get('head', '')[:10]}」の詳細が長い")
        draw_text(ctx, bx, ry, x + w - bx, rh, dp, dsz, anchor="m")
        ry += rh + gap


# ═══════════════════════════════════════════════════════════════
#  M-RHMATRIX
# ═══════════════════════════════════════════════════════════════
def m_rhmatrix(ctx, r, d):
    """{"title": "全社AI変革"（列の上の帯・任意）, "columns": [{"label": "専門人材による AI 駆動 BPR", "sub": "構造変革の圧縮実行"}, ...],
        "col_style": "rule|band|chevron", "col_ratios": [1, 1], "emph_col": 1,
        "rows": [{"head": "変革対象", "cells": [リッチ本文 | {"icon": .., "text": ..}
                                                | {"body": リッチ本文, "status": "実施済", "done": true}, ...], "why": false}, ...],
        "groups": [{"label": "構想フェーズ", "span": 2}, ...]（列見出しの上の大きな工程の帯・任意。v6.8）,
        "scope": [{"cells": [[0, 0], [1, 0]] | "cols": [0], "label": "スコープ 1"}]（囲む枠＋札・任意。v6.8）}
    … 人の資料の 2 段の工程見出し: 上に工程のまとまり（濃い紫の五方向）、下に細かい工程（淡い紫の五方向）、
    その下にセル。終わった工程のセルは灰、状況の札（黒）をセルの左上に載せる。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    cols = [cl if isinstance(cl, dict) else {"label": str(cl)} for cl in d.get("columns") or []]
    rows = d.get("rows") or []
    nc = max(len(cols), 1)
    head_w = d.get("head_w") or 1.35
    cg, rg = 0.12, 0.06
    ratios = d.get("col_ratios") or [1] * nc
    cw_tot = w - head_w - cg * nc
    cws = [cw_tot * k / sum(ratios) for k in ratios]
    cxs = []
    xx = x + head_w + cg
    for cw in cws:
        cxs.append(xx)
        xx += cw + cg
    emph_col = d.get("emph_col")
    if d.get("title"):
        band = c._rect(s, cxs[0], y, x + w - cxs[0], 0.36, fill=_h.DARK_PURPLE)
        c._shape_text(band, d["title"], 16, _h.WHITE, True)
        y, h = y + 0.42, h - 0.42
    style = d.get("col_style", "rule")
    ch = 0.5 if style == "rule" else 0.44
    if d.get("groups"):             # 工程のまとまり（2 段の見出しの上の段）
        k0 = 0
        for g in d["groups"]:
            k1 = min(nc - 1, k0 + int(g.get("span", 1)) - 1)
            gx0, gx1 = cxs[k0], cxs[k1] + cws[k1]
            gb = c.homeplate(s, gx0, y, gx1 - gx0, 0.36, fill=_h.DARKEST_PURPLE)
            c._shape_text(gb, g.get("label", ""), 14, _h.WHITE, True)
            k0 = k1 + 1
        y, h = y + 0.42, h - 0.42
    for cl, cx, cw in zip(cols, cxs, cws):
        lab = cl.get("label", "") + (f"　— {cl['sub']}" if cl.get("sub") else "")
        if style == "rule":
            size, _ = c.fit_size([lab], cw - 0.1, ch, (14, 12))
            c._textbox(s, cx, y, cw, ch - 0.06, [lab], size, _h.TEXT_BODY, align=PP_ALIGN.CENTER, anchor="b",
                       space_after=0, inset=(0.02, 0, 0.02, 0.03))
            c._hline(s, cx, y + ch - 0.03, cw, _h.BLACK, 0.75)
        else:
            if style == "chevron":      # 列の幅いっぱいの五方向の矢印（列の隙間がそのまま矢印の隙間になる）
                b = c.homeplate(s, cx, y, cw, ch, fill=_h.LIGHTEST_PURPLE)
            else:
                b = c._rect(s, cx, y, cw, ch, fill=_h.DARK_PURPLE)
            size, _ = c.fit_size([lab], cw - 0.2, ch, (14, 12))
            c._shape_text(b, lab, size, _h.DARKEST_PURPLE if style == "chevron" else _h.WHITE, True)
    y, h = y + ch + 0.08, h - ch - 0.08
    n = max(len(rows), 1)

    def cell_paras(cell, base):
        if isinstance(cell, dict) and cell.get("icon") is not None and "text" in cell:
            return sized(rich(cell["text"]), base)
        if isinstance(cell, dict) and any(k in cell for k in ("body", "status", "done")):
            return sized(rich(cell.get("body") or cell.get("text") or ""), base,
                         color=_h.TEXT_SUB if cell.get("done") else None)
        return sized(rich(cell), base)
    sizes = body_sizes()
    if nc >= 5:                     # 列が多い格子は 14pt まで（大きい字で細い列に折り返すと、語の途中で切れやすい）
        sizes = tuple(z for z in sizes if z <= 14) or sizes
    need = []
    for rw in rows:
        nd = 0.45
        cells_ = (rw.get("cells") or [])[:nc]
        span = rw.get("why") and len(cells_) == 1
        for k, cell in enumerate(cells_):
            iw = (x + w - cxs[0] if span else cws[k]) - 0.24 - (ICON["inline"] + 0.12 if isinstance(cell, dict) and cell.get("icon") else 0)
            nd = max(nd, height(cell_paras(cell, sizes[-1]), iw, sizes[-1]) + 0.12)
        need.append(nd)
    avail = h - rg * (n - 1)
    heights = [avail * k / sum(need) for k in need] if need else []
    heights = [min(hh, nd * 1.9 + 0.2) for hh, nd in zip(heights, need)]
    # 全セルで同じ大きさ（収まる最大）
    base = sizes[-1]
    for sz in sizes:
        okk = True
        for rw, rh in zip(rows, heights):
            cells_ = (rw.get("cells") or [])[:nc]
            span = rw.get("why") and len(cells_) == 1
            for k, cell in enumerate(cells_):
                iw = (x + w - cxs[0] if span else cws[k]) - 0.24 - (ICON["inline"] + 0.12 if isinstance(cell, dict) and cell.get("icon") else 0)
                if height(cell_paras(cell, sz), iw, sz) > rh - 0.06:
                    okk = False
        if okk:
            base = sz
            break
    ry = y
    cell_box, cell_box_rows = {}, []
    for rw, rh in zip(rows, heights):
        why = rw.get("why")
        hb = c._rect(s, x, ry, head_w, rh, fill=_h.DARK_PURPLE if why else _h.MID_GRAY)
        hs, _ = c.fit_size([rw.get("head", "")], head_w - 0.16, rh, (16, 14, 12))
        c._shape_text(hb, rw.get("head", ""), hs, _h.WHITE, True, inset=(0.06, 0.02, 0.06, 0.02))
        # 行の中のセルは、いちばん多いセルに合わせて同じ高さから書き始める（見出しの上端をそろえる）
        cells_r = rw.get("cells") or []
        rmax = max([height(cell_paras(cl, base), cw_ - 0.24 - (ICON["inline"] + 0.16 if isinstance(cl, dict) and cl.get("icon") else 0), base)
                    for cl, cw_ in zip(cells_r, cws if not (why and len(cells_r) == 1) else [x + w - cxs[0]]) if cl is not None] + [0.3])
        top_pad = max(0.04, (rh - rmax) / 2)
        if any(isinstance(cl, dict) and cl.get("status") for cl in cells_r):
            top_pad = max(top_pad, 0.22)          # 状況の札の下から書く
        for k, (cx, cw) in enumerate(zip(cxs, cws)):
            cells = rw.get("cells") or []
            cell = cells[k] if k < len(cells) else None
            if why and len(cells) == 1 and k == 0:          # Why 行は 1 セルで全列をまたぐ
                cw = x + w - cx
            elif why and len(cells) == 1:
                continue
            done = isinstance(cell, dict) and cell.get("done")
            c._rect(s, cx, ry, cw, rh, fill=_h.LIGHT_GRAY if done else _h.LIGHTEST_PURPLE if k == emph_col else _h.OFF_WHITE)
            cell_box[(len(cell_box_rows), k)] = (cx, ry, cw, rh)
            if cell is None:
                continue
            tx, tw = cx + 0.08, cw - 0.16
            if isinstance(cell, dict) and cell.get("icon"):
                isz = ICON["inline"]
                c._place_visual(ctx, cell["icon"], (tx + 0.04, ry + (rh - isz) / 2, isz, isz), color=_h.DARKEST_PURPLE)
                tx, tw = tx + isz + 0.16, tw - isz - 0.16
            ps = cell_paras(cell, base)
            if height(ps, tw - 0.08, base) > rh:
                ctx.warn(f"M-RHMATRIX: 「{rw.get('head', '')[:8]}」×{k + 1} 列目が長い")
            draw_text(ctx, tx, ry + top_pad, tw, rh - top_pad, ps, base, anchor="t")
            if isinstance(cell, dict) and cell.get("status"):      # 状況の札（黒・セルの左上に載せる）
                st = str(cell["status"])
                tg = c._rect(s, cx - 0.04, ry - 0.04, 0.3 + sum(c._char_w(ch_, 11) for ch_ in st) * 1.1, 0.24, fill=_h.BLACK)
                c._shape_text(tg, st, 11, _h.WHITE, True, inset=(0.04, 0, 0.04, 0), wrap=False)
        cell_box_rows.append(ry)
        ry += rh + rg
    for sc in d.get("scope") or []:          # 範囲の枠（セル・列を囲む）＋札
        keys = [tuple(k_) for k_ in sc.get("cells") or []] or [(ri, ci) for ci in sc.get("cols") or [] for ri in range(len(rows))]
        bxs = [cell_box[k_] for k_ in keys if k_ in cell_box]
        if not bxs:
            continue
        fx0, fy0 = min(b_[0] for b_ in bxs) - 0.05, min(b_[1] for b_ in bxs) - 0.05
        fx1, fy1 = max(b_[0] + b_[2] for b_ in bxs) + 0.05, max(b_[1] + b_[3] for b_ in bxs) + 0.05
        c._rect(s, fx0, fy0, fx1 - fx0, fy1 - fy0, fill=None, line=_h.CORE_PURPLE, line_w=2.25)
        if sc.get("label"):
            lab = sc["label"]
            tw_ = 0.3 + sum(c._char_w(ch_, 12) for ch_ in lab) * 1.1
            tg = c._rect(s, fx1 - tw_, fy1, tw_, 0.28, fill=_h.CORE_PURPLE)
            c._shape_text(tg, lab, 12, _h.WHITE, True, inset=(0.03, 0, 0.03, 0), wrap=False)


# ═══════════════════════════════════════════════════════════════
#  M-STEPS
# ═══════════════════════════════════════════════════════════════
def m_steps(ctx, r, d):
    """{"steps": [{"label": "ツール変換", "title": "既存資産で土台を作る", "body": [リッチ本文], "icon": "q:server"}, ...],
        "icon_col": true（右端にアイコンの流れ）, "icon_labels": true}"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    steps = d.get("steps") or []
    n = max(len(steps), 1)
    gap = 0.1
    blk_w = d.get("label_w") or 1.5
    icol = 0.8 if d.get("icon_col") and any(st.get("icon") for st in steps) else 0
    bx = x + blk_w + 0.12
    bw = w - blk_w - 0.12 - (icol + 0.15 if icol else 0)
    bodies = [([{"h": st["title"]}] if st.get("title") else []) + list(st.get("body") or []) for st in steps]
    sizes = body_sizes()
    need = [height(sized(rich(b), sizes[-1]), bw - 0.3, sizes[-1]) + 0.2 for b in bodies]
    avail = h - gap * (n - 1)
    heights = [avail * k / sum(need) for k in need]
    base = min(fit(rich(b), bw - 0.3, hh - 0.12)[0] for b, hh in zip(bodies, heights))
    ry = y
    centers = []
    for i, (st, b, rh) in enumerate(zip(steps, bodies, heights)):
        blk = c._rect(s, x, ry, blk_w, rh, fill=_h.DARK_PURPLE, shape=S.FLOWCHART_OFFPAGE_CONNECTOR)
        c._shape_text(blk, st.get("label", ""), 16 if len(st.get("label", "")) <= 6 else 14, _h.WHITE, True,
                      anchor="m", inset=(0.1, 0.3, 0.1, rh * 0.2))
        tag = c._rect(s, x, ry, 0.32, 0.28, fill=_h.DARKEST_PURPLE, line=_h.WHITE, line_w=1.0)
        c._shape_text(tag, str(st.get("number", i + 1)), 12, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        c._rect(s, bx, ry, bw, rh, fill=_h.OFF_WHITE)
        ps = sized(rich(b), base)
        if height(ps, bw - 0.3, base) > rh:
            ctx.warn(f"M-STEPS: 工程「{st.get('label', '')[:8]}」の説明が長い")
        draw_text(ctx, bx + 0.12, ry, bw - 0.2, rh, ps, base, anchor="m")
        if icol:
            icon_shp = None
            isz = min(ICON["inline"], rh * 0.5)
            icx = x + w - icol / 2
            lab = st.get("icon_label") if d.get("icon_labels", True) else None
            ty = ry + (rh - isz - (0.24 if lab else 0)) / 2
            if st.get("icon"):
                if c._place_visual(ctx, st["icon"], (icx - isz / 2, ty, isz, isz), color=_h.DARKEST_PURPLE) is not None:
                    icon_shp = s.shapes[-1]
            if lab:
                c._textbox(s, icx - icol / 2, ty + isz + 0.02, icol, 0.24, [lab], 11, _h.DARK_PURPLE, bold=True,
                           align=PP_ALIGN.CENTER, space_after=0, inset=(0, 0, 0, 0))
            centers.append((icx, ty, ty + isz + (0.26 if lab else 0), icon_shp))
        ry += rh + gap
    for (cx_, t0, b0, sa), (_, t1, _, sb) in zip(centers, centers[1:]):
        # アイコンの間の細い矢印（1pt・短い直線）。アイコンとラベルの並びの順番を示すだけなので接続しない。
        # 上のアイコンのラベルの下から次のアイコンの上まで（線がラベルの文字に重ならない）
        if t1 - b0 > 0.2:
            c._connector(s, cx_, b0 + 0.05, cx_, t1 - 0.05, color=_h.TEXT_BODY, width=1.0, arrow=True)


# ═══════════════════════════════════════════════════════════════
#  M-FIGNOTE
# ═══════════════════════════════════════════════════════════════
def m_fignote(ctx, r, d):
    """{"figure": {"module": "M-HUB", ...}（左の図）, "ratio": 0.56,
        "notes": [{"title": "経営直轄で切り出す", "text": "…"}, ...],
        "markers": [{"n": 1, "x": 0.12, "y": 0.3}, ...]（図の枠に対する 0〜1 の位置。図の中に同じ番号の丸を置く）}"""
    c = _c()
    x, y, w, h = r
    fig = d.get("figure") or {}
    ratio = d.get("ratio", 0.56)
    fw = w * ratio - 0.15
    fn = c.MODULES.get(fig.get("module", ""))
    if fn:
        fr = (x, y, fw, h)
        if fig.get("header"):
            fr = c.pane_header(ctx, fr, fig["header"], style=fig.get("header_style"))
        fn(ctx, fr, fig)
    else:
        ctx.warn("M-FIGNOTE: figure に部品が無い")
    dd = 0.34
    for mk in d.get("markers") or []:
        num_circle(ctx, x + fw * mk.get("x", 0), y + h * mk.get("y", 0), dd, mk.get("n", 1), _h.CORE_PURPLE)
    notes = d.get("notes") or []
    nx = x + w * ratio + 0.15
    nw = x + w - nx
    tw = nw - dd - 0.15
    paras = [sized(rich(([{"h": nt["title"]}] if nt.get("title") else []) + list(nt.get("body") or ([nt["text"]] if nt.get("text") else []))), 12) for nt in notes]
    sizes = body_sizes()
    base = sizes[-1]
    for sz in sizes:
        ps = [sized(rich(([{"h": nt["title"]}] if nt.get("title") else []) + list(nt.get("body") or ([nt["text"]] if nt.get("text") else []))), sz)
              for nt in notes]
        if sum(height(p, tw, sz) for p in ps) + 0.15 * (len(ps) - 1) <= h:
            base, paras = sz, ps
            break
    hs = [height(p, tw, base) for p in paras]
    gap = min(0.5, (h - sum(hs)) / max(len(hs) - 1, 1)) if len(hs) > 1 else 0
    tot = sum(hs) + gap * (len(hs) - 1)
    ny = y + max(0, (h - tot) / 2)
    for k, (nt, p, hh) in enumerate(zip(notes, paras, hs)):
        num_circle(ctx, nx + dd / 2, ny + dd / 2 + 0.02, dd, nt.get("n", k + 1), _h.CORE_PURPLE)
        draw_text(ctx, nx + dd + 0.12, ny, tw, hh, p, base)
        ny += hh + gap


# ═══════════════════════════════════════════════════════════════
#  M-CARD（強み型）
# ═══════════════════════════════════════════════════════════════
def is_strength(d):
    return any(d.get(k) for k in ("head", "number", "lead"))


def _strength_frame(r, d):
    """強みカードの縦の配分（アイコン・番号タブ・見出しブロックの高さ、要点の面に使える高さ）。"""
    x, y, w, h = r
    isz = (ICON["head"] if w >= 1.6 else w * 0.42) if d.get("icon") else 0
    top = y + (isz + 0.1 if isz else 0) + (0.3 if d.get("number") else 0)
    hb_h = min(1.0, max(0.62, (y + h - top) * 0.26))
    rest = y + h - top - hb_h - 0.08
    return isz, hb_h, rest


def strength_sizes(r, d):
    """並ぶ強みカードをそろえるための値（要点の面に要る高さ・見出しの大きさ・補足の大きさ）。"""
    c = _c()
    x, y, w, h = r
    isz, hb_h, rest = _strength_frame(r, d)
    lead_h = rest * (0.55 if d.get("text") or d.get("body") else 1.0)
    if d.get("lead") and (d.get("text") or d.get("body")):
        lp0 = [(t, {**o, "bold": True}) for t, o in sized(rich(d["lead"]), 14)]
        lead_h = min(lead_h, max(0.8, height(lp0, w - 0.3, 14) + 0.35))
    head = d.get("head") or d.get("title") or ""
    hs, _ = c.fit_size([head], w - 0.24, hb_h - 0.08, (18, 16, 14))
    tb = 12
    if d.get("text") or d.get("body"):
        tb = fit(rich(d.get("body") or d["text"]), w - 0.12, rest - lead_h - 0.2, (14, 12, 11))[0]
    return {"lead": lead_h, "head": hs, "text": tb}


def m_strength(ctx, r, d):
    """{"number": "01", "head": "CXO との信頼関係", "lead": "技術提案にとどまらず**経営の意思決定**に関与する",
        "text": "数十年かけて築いた関係資産で、短期間では模倣できない", "icon": "q:handshake"}
    並ぶカードは、要点の面の高さ・見出しと補足の文字の大きさを並びでそろえる（compose が _CUR["strength"] に入れる）。"""
    c = _c()
    x, y, w, h = r
    s = ctx.slide
    uni = c._CUR.get("strength") or strength_sizes(r, d)
    top = y
    if d.get("icon"):
        isz = ICON["head"] if w >= 1.6 else w * 0.42
        c._place_visual(ctx, d["icon"], (x + (w - isz) / 2, top, isz, isz), color=_h.DARKEST_PURPLE)
        top += isz + 0.1
    tab_h = 0.3 if d.get("number") else 0
    if tab_h:
        tab = c._rect(s, x, top, 0.56, tab_h, fill=_h.DARK_PURPLE)
        c._shape_text(tab, str(d["number"]), 14, _h.WHITE, True, inset=(0, 0, 0, 0), wrap=False)
        top += tab_h
    head = d.get("head") or d.get("title") or ""
    hb_h = min(1.0, max(0.62, (y + h - top) * 0.26))
    hb = c._rect(s, x, top, w, hb_h, fill=_h.DARK_PURPLE)
    hs = uni["head"]
    c._shape_text(hb, head, hs, _h.WHITE, True, inset=(0.1, 0.03, 0.1, 0.03))
    top += hb_h + 0.08
    lead_h = uni["lead"]
    if d.get("lead"):
        c._rect(s, x, top, w, lead_h, fill=_h.OFF_WHITE)
        c._rect(s, x, top, 0.06, lead_h, fill=_h.DARK_PURPLE)
        lp = [(t, {**o, "bold": True}) for t, o in sized(rich(d["lead"]), 14, color=_h.DARKEST_PURPLE)]
        ls = 14 if height(lp, w - 0.3, 14) <= lead_h else 12
        lp = [(t, {**o, "size": ls}) for t, o in lp]
        if height(lp, w - 0.3, ls) > lead_h:
            ctx.warn(f"M-CARD(強み)「{head[:8]}」: 要点が長い")
        draw_text(ctx, x + 0.14, top, w - 0.2, lead_h, lp, ls, anchor="m")
        top += lead_h + 0.1
    if d.get("text") or d.get("body"):
        body = d.get("body") or d["text"]
        rich_box(ctx, x, top + 0.06, w, y + h - top - 0.06, body, anchor="t", warn_label=f"M-CARD(強み)「{head[:8]}」",
                 sizes=(uni["text"],), align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════
#  M-ICONGRID（実績型）
# ═══════════════════════════════════════════════════════════════
def is_stat(d):
    return any(isinstance(it, dict) and it.get("value") for it in d.get("items") or [])


def m_statgrid(ctx, r, d):
    """{"items": [{"icon": "q:flask", "value": "30+ years", "text": "業界の大手企業との協業年数"}, ...], "cols": 2}"""
    c = _c()
    x, y, w, h = r
    items = [it for it in d.get("items") or [] if isinstance(it, dict)]
    n = max(len(items), 1)
    cols = d.get("cols") or (1 if n <= 3 else 2 if n <= 8 else 3)
    rows = math.ceil(n / cols)
    gx, gy = 0.3, 0.12
    cw = (w - gx * (cols - 1)) / cols
    rh = (h - gy * (rows - 1)) / rows
    isz = min(ICON["grid"], rh * 0.6)
    tw = cw - isz - 0.2
    vs = 20 if all(len(str(it.get("value", ""))) <= 10 for it in items) else 16
    paras_of = lambda it, sz: [(str(it.get("value", "")), {"size": vs, "bold": True, "color": _h.DARK_PURPLE, "space_after": 0})] + \
        sized(rich(it.get("text") or it.get("label")), sz)
    base = 12
    for sz in (16, 14, 12):
        if all(height(paras_of(it, sz), tw, sz) <= rh * 0.8 for it in items):
            base = sz
            break
    # 行の高さは中身に合わせ（行の間を疎らにしない）、格子全体を縦中央に。数字の上端は全セルでそろえる
    ph = max(height(paras_of(it, base), tw, base) for it in items)
    rh2 = min(rh, max(ph + 0.35, isz + 0.3))
    gy2 = min(0.4, gy + (rh - rh2) * 0.5)
    tot = rh2 * rows + gy2 * (rows - 1)
    y0 = y + (h - tot) / 2
    for i, it in enumerate(items):
        rr, cc = divmod(i, cols) if d.get("order", "row") == "row" else (i % rows, i // rows)
        bx, by = x + cc * (cw + gx), y0 + rr * (rh2 + gy2)
        ps = paras_of(it, base)
        if it.get("icon"):
            c._place_visual(ctx, it["icon"], (bx, by + (rh2 - ph) / 2, isz, isz), color=_h.DARKEST_PURPLE)
        draw_text(ctx, bx + isz + 0.18, by + (rh2 - ph) / 2, tw, ph + 0.05, ps, base, anchor="t")
        if rr < rows - 1:
            c._hline(ctx.slide, bx, by + rh2 + gy2 / 2, cw, _h.LIGHT_GRAY, 0.75)


def register(modules):
    modules.update({"M-ZONE": m_zone, "M-RHMATRIX": m_rhmatrix, "M-STEPS": m_steps, "M-FIGNOTE": m_fignote,
                    "M-WHY": m_whyrows, "M-STATGRID": m_statgrid, "M-STRENGTH": m_strength})
