"""三者定例資料からONS月次報告のスライドの文字・表を抜き出す（赤字は [赤:...] で示す）。

使い方:
  python -I extract_slides.py <三者定例資料.pptx> <開始スライド番号> [終了スライド番号] > out.txt
- スライド番号はファイル上の通し番号（画面右下のページ番号は1つ小さいことがある）
- 図（件数推移・大分類表・未解決表）は画像なので抜き出せない。PowerPointで画像に書き出して目で読む
  （PowerShell: $pres.Slides.Item(N).Export("sN.png","PNG",1600,900)）
"""
import sys

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
src, start = sys.argv[1], int(sys.argv[2])
end = int(sys.argv[3]) if len(sys.argv) > 3 else 10**6
p = Presentation(src)
print("total", len(p.slides))


def is_red(run):
    c = run.font.color
    return c is not None and c.type is not None and getattr(c, "rgb", None) is not None and str(c.rgb) == "FF0000"


def walk(shapes, d=0):
    for s in shapes:
        if s.shape_type == 6:  # グループ
            walk(s.shapes, d + 1)
            continue
        if s.has_text_frame:
            for para in s.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t.strip():
                    reds = [r.text for r in para.runs if is_red(r)]
                    print("  " * d + t + ("  [赤:" + "|".join(reds) + "]" if reds else ""))
        if s.has_table:
            for row in s.table.rows:
                print("  " * d + " | ".join(c.text.replace("\n", "/") for c in row.cells))
        if s.shape_type == 13:
            print("  " * d + "[画像]", s.name)


for i, sl in enumerate(p.slides, 1):
    if start <= i <= end:
        print(f"===== slide {i}")
        walk(sl.shapes)
