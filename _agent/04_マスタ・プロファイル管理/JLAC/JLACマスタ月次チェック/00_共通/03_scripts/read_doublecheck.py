"""Excelで計算・保存したダブルチェック（make_doublecheck.py の出力）を読み取りだけで開き、照合の結果をJSONにする。保存はしない。
使い方: python 00_共通/03_scripts/read_doublecheck.py <ダブルチェック.xlsx> <出力.json> [--seconds <Excelでの計算秒数>]
  - 「照合」シート：E列（判定）が「不一致」の項目、「—」（比較できない）の数
  - 「今回」シート：行ごとの一致列（BG〜BK）が0の行（各20行まで。原本行・JLAC11コード）
  - 数式の計算結果が保存されていない（Excelで計算していない）ときは status を「未計算」にする
"""
import sys, json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
args = sys.argv[1:]
secs = None
if '--seconds' in args:
    i = args.index('--seconds'); secs = int(float(args[i + 1])); del args[i:i + 2]
src, dst = args[0], args[1]

wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
ws = wb['照合']
rows = list(ws.iter_rows(values_only=True))
title = rows[0][0] if rows and rows[0] else ''
HR = 6  # 見出しの行（make_doublecheck.py と同じ）
items = [r for r in rows[HR:] if r and r[1] not in (None, '') and r[0] not in (None, '')]
judged = [r[4] for r in items]
uncomputed = all(j is None for j in judged)
mismatch = [{'kind': r[0], 'item': r[1], 'excel': r[2], 'tool': r[3], 'memo': r[5]} for r in items if r[4] == '不一致']
dash = sum(1 for j in judged if j == '—')
match = sum(1 for j in judged if j == '一致')

ROWCOLS = {'BG': '新規判定', 'BH': 'KA-18', 'BI': '前月差分', 'BJ': 'KA-07', 'BK': 'KA-08'}
from openpyxl.utils import column_index_from_string as ci
idx = {k: ci(k) - 1 for k in ROWCOLS}
rowmis = {k: {'title': f'{k}列 {v}', 'count': 0, 'rows': []} for k, v in ROWCOLS.items()}
cur = wb['今回']
for r in cur.iter_rows(min_row=2, values_only=True):
    for k, j in idx.items():
        if j < len(r) and r[j] == 0:
            x = rowmis[k]; x['count'] += 1
            if len(x['rows']) < 20:
                x['rows'].append({'sheet': r[32], 'row': r[33], 'code': r[12]})
wb.close()
out = {
    'file': src, 'status': '未計算' if uncomputed else '計算済み', 'seconds': secs, 'title': title,
    'total': len(items), 'match': match, 'dash': dash, 'comparable': len(items) - dash,
    'mismatch': mismatch, 'rowMismatch': list(rowmis.values()),
}
json.dump(out, open(dst, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f"照合：{out['status']}　項目{out['total']}・一致{match}・不一致{len(mismatch)}・比較できない{dash}　行ごとの不一致 " + '、'.join(f"{x['title']} {x['count']}" for x in out['rowMismatch']))
