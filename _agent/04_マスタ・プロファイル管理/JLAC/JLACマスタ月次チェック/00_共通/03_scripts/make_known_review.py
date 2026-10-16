"""既知（前月も同じ指摘）としてコメントしなかった指摘を論点ごとにまとめ、
本田さんと「マスタの仕様か、誤りか」を判断するための確認資料（xlsx）を作る。毎月の月次実行で作る。

対象の行（2026-10-16 ユーザー決定）：
  ・使用中の行（適用終了日が99999999）はすべて対象
  ・適用終了している行は、前月の版でも同じ内容（32列すべて同じ＝前月も終了済みで変更なし）のものだけ外す。
    今月終了した行・内容が変わった行は対象に残す

使い方：
  python make_known_review.py <月次実行の明細.json または run_check.mjs --json の出力> <出力xlsx または出力フォルダ>
         [--target <今月の提出ファイル>] [--prev <前月公開版>] [--label <今月の表示名>]
  ・月次実行の明細.json なら、提出ファイル・前月公開版はjsonの files から読む
  ・出力にフォルダを渡すと「YYYYMMDD_JLACマスタ既知指摘の仕様確認_<提出版>_vN.xlsx」を新しい版で作る（上書きしない）
"""
import argparse
import csv
import datetime
import json
import os
import re
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHECKS = ('KA-10', 'KA-12', 'KA-23', 'KA-08')
OPEN_END = '99999999'
HEAD = PatternFill('solid', fgColor='DDEBF7')
ASK = PatternFill('solid', fgColor='FFF2CC')
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')


def s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def norm(v):
    v = unicodedata.normalize('NFKC', v)
    v = re.sub(r'[‐-―−－﹣]', '-', v)
    return re.sub(r'\s', '', v)


def strip_marks(v):
    return re.sub(r'[()\[\]・･\-‘’\'"“”、,]', '', norm(v)).casefold()


def read_table(path):
    """{シート名: (見出し, {行番号: 行})}。CSVはシート名を空にする。"""
    if path.lower().endswith('.csv'):
        raw = open(path, 'rb').read()
        for enc in ('utf-8-sig', 'cp932'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = list(csv.reader(text.splitlines()))
        return {'': (rows[0], {n: r for n, r in enumerate(rows[1:], start=2)})}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if rows and 'JLAC11コード' in [s(x) for x in rows[0]]:
            out[ws.title] = ([s(x) for x in rows[0]], {n: r for n, r in enumerate(rows[1:], start=2)})
    return out


def key32(row, ix):
    cols = list(ix.values())[:32]
    v = [s(row[i]) if i < len(row) else '' for i in cols]
    return tuple('00000000' if (h == '適用開始日' and x in ('0', '00000000')) else x for h, x in zip(list(ix)[:32], v))


def next_name(out_dir, tag):
    today = datetime.date.today().strftime('%Y%m%d')
    base = f'JLACマスタ既知指摘の仕様確認_{tag}'
    vs = [int(m.group(1)) for f in os.listdir(out_dir) for m in [re.match(rf'\d{{8}}_{re.escape(base)}_v(\d+)\.xlsx$', f)] if m]
    return os.path.join(out_dir, f'{today}_{base}_v{max(vs, default=0) + 1}.xlsx')


def add_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    last = len(header)
    for c in range(1, last + 1):
        cell = ws.cell(1, c)
        cell.font = BOLD
        cell.fill = ASK if c == last else HEAD
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            if cell.column == last:
                cell.fill = ASK
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('state')
    ap.add_argument('out')
    ap.add_argument('--target')
    ap.add_argument('--prev')
    ap.add_argument('--label')
    a = ap.parse_args()
    st = json.load(open(a.state, encoding='utf-8'))
    files = st.get('files', {})
    target, prev = a.target or files.get('target'), a.prev or files.get('prev')
    if not target or not prev:
        raise SystemExit('提出ファイルと前月公開版が分かりません（--target と --prev を指定してください）')
    label = a.label or st.get('label') or os.path.basename(target)
    tag = re.sub(r'[\s　-]', '', label).replace('JLACセンター', 'JLACセンター')

    tables = read_table(target)
    prev_tab = list(read_table(prev).values())[0]
    pix = {h: i for i, h in enumerate(prev_tab[0])}
    prev_keys = {key32(r, pix) for r in prev_tab[1].values()}

    def row_of(f):
        sheet = f.get('sheet', '')
        h, rows = tables.get(sheet) or next(iter(tables.values()))
        return {x: i for i, x in enumerate(h)}, rows.get(f['row'])

    def in_scope(f):
        ix, r = row_of(f)
        if r is None:
            return True
        if s(r[ix['適用終了日']]) == OPEN_END:
            return True
        return key32(r, ix) not in prev_keys  # 前月も同じ内容で終了済みの行だけ外す

    known = defaultdict(list)
    for f in st['findings']:
        if f['check'] in CHECKS and f.get('known') and f.get('level') == 'strict':
            known[f['check']].append(f)
    total = {k: len(known[k]) for k in CHECKS}
    act = {k: [f for f in known[k] if in_scope(f)] for k in CHECKS}
    ended_kept = {k: sum(1 for f in act[k] if s(row_of(f)[1][row_of(f)[0]['適用終了日']]) != OPEN_END) for k in CHECKS}

    def val(f, col):
        ix, r = row_of(f)
        return s(r[ix[col]]) if r is not None and col in ix else ''

    def code(f):
        return f"{f['row']}行目（{val(f, 'JLAC11コード')}）"

    def group_names(check, key_col, cols):
        groups = OrderedDict()
        for f in act[check]:
            combo = tuple(val(f, c) for c in cols)
            groups.setdefault(val(f, key_col), OrderedDict()).setdefault(combo, []).append(f)
        rows, breakdown = [], defaultdict(int)
        for key, combos in sorted(groups.items(), key=lambda x: -sum(len(v) for v in x[1].values())):
            dts = {c: sorted({val(f, 'データタイプ') for f in fs}) for c, fs in combos.items()}
            dt_list = [tuple(v) for v in dts.values()]
            text = ' '.join(''.join(c) for c in combos)
            if len(set(dt_list)) == len(dt_list) and all(len(v) == 1 for v in dt_list):
                cls = '定量・定性などの書き分け（データタイプが書き方ごとに違う）'
            elif re.search(r'定性|定量|空腹|随時|陽性コントロール比|インデックス', text):
                cls = '検査の種類の書き分け（名称に定性・定量・空腹時・随時などを含む）'
            else:
                cls = 'その他（要判断）'
            n_rows = sum(len(v) for v in combos.values())
            breakdown[cls] += n_rows
            listing = '\n'.join(f"・{'／'.join(c)}（{len(fs)}行、{'/'.join(dts[c])}）" for c, fs in combos.items())
            rows.append([key, n_rows, len(combos), listing, cls, '\n'.join(code(fs[0]) for fs in combos.values()), None])
        return rows, dict(breakdown), len(groups)

    r10, b10, n10 = group_names('KA-10', '大項目', ['FHIR項目名称', 'FHIR識別文字列', '略称'])
    r12, b12, n12 = group_names('KA-12', '略称', ['大項目', 'FHIR項目名称', 'FHIR識別文字列'])

    g23 = OrderedDict()
    for f in act['KA-23']:
        m = re.search(r'検査方法\(JLAC10-測定法\)「(.*?)」が、.*?測定法コード「(\w+)」の名称「(.*?)」', f['msg'])
        m2 = re.search(r'測定法コード「(\w+)」が、.*?見当たりません', f['msg'])
        if m:
            key = (m.group(2), m.group(1), m.group(3))
        elif m2:
            key = (m2.group(1), val(f, '検査方法(JLAC10-測定法)'), '（コード表にこのコードが無い）')
        else:
            key = ('?', '?', f['msg'][:60])
        g23.setdefault(key, []).append(f)
    r23, b23 = [], defaultdict(int)
    for (mc, mv, tv), fs in sorted(g23.items(), key=lambda x: -len(x[1])):
        if tv.startswith('（コード表に'):
            cls = 'コード表に無い測定法コード（999＝その他の測定法など）'
        elif norm(mv) == norm(tv):
            cls = '全角・半角・空白だけの違い'
        elif strip_marks(mv) == strip_marks(tv):
            cls = '括弧・記号だけの違い'
        else:
            cls = '名称が違う（要判断）'
        b23[cls] += len(fs)
        ex = '\n'.join(code(f) for f in fs[:3]) + ('\n…' if len(fs) > 3 else '')
        r23.append([mc, mv, tv, len(fs), cls, ex, None])

    r08, b08 = [], defaultdict(int)
    for f in sorted(act['KA-08'], key=lambda x: x['row']):
        m = re.search(r'測定法\(JLAC11\)「(.*?)」が、.*?「測定法名称」列の「(.*?)」', f['msg'])
        mv, lv = m.groups() if m else (f['msg'][:60], '?')
        if '_' in mv and norm(mv.split('_', 1)[1]) == norm(lv):
            cls = '「方法名_試薬名」の書き方（試薬名は一致）'
        elif norm(mv) == norm(lv):
            cls = '全角・半角・空白だけの違い'
        elif strip_marks(mv) == strip_marks(lv):
            cls = '全角・半角・記号（‘’など）・大文字小文字だけの違い'
        else:
            cls = '名称が違う（要判断）'
        b08[cls] += 1
        r08.append([f['row'], val(f, 'JLAC11コード'), val(f, '販売名称'), mv, lv, cls, None])

    def fmt(b):
        return '\n'.join(f'{k}：{v}行' for k, v in sorted(b.items(), key=lambda x: -x[1]))

    def scope(k):
        return f"{len(act[k])}（うち適用終了済み {ended_kept[k]}）"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '説明'
    ws.append([f'JLACマスタ 既知とした指摘の確認（{label}）'])
    ws.append([])
    ws.append(['目的', 'チェックツールが「既知（前月も同じ指摘）」としてコメントしなかった指摘を論点ごとにまとめ、「マスタの仕様（直さなくてよい）か、誤り（JLACセンターに確認する）か」をご判断いただく。既知は「前月の公開版にも同じ状態で載っていた」という意味で、JLACセンターに指摘して了承をもらったとは限らない（2026-10-13に確認）。'])
    ws.append(['対象の行', '使用中の行（適用終了日99999999）はすべて対象。適用終了している行は、前月の版でも同じ内容（32列すべて同じ）のものだけ外し、今月終了した行・内容が変わった行は残す。'])
    ws.append(['元にしたもの', f"今月：{os.path.basename(target)}／前月：{os.path.basename(prev)}／チェックツール：{st.get('tool', {}).get('version') or st.get('version', '')}"])
    ws.append(['ACNの見立てについて', '書き方の違いを機械的に分類したもの（推定）で、仕様かどうかの判断ではない。各シートの黄色の列に、仕様か誤りかをご記入ください。'])
    ws.append([])
    ws.append(['チェック', 'チェック内容', '既知（全体）', '対象の行', '論点の数', 'ACNの見立て（対象の行の内訳）', '判断いただきたいこと', '本田さんの判断'])
    ws.append(['KA-10', '大項目をキーにFHIR項目名称/FHIR識別文字列/略称が整合すること', total['KA-10'], scope('KA-10'), f'大項目 {n10}種類', fmt(b10),
               '同じ大項目で書き方が複数あるのは仕様か（とくに「その他（要判断）」）', None])
    ws.append(['KA-12', '略称をキーに大項目/FHIR項目名称/FHIR識別文字列が整合すること', total['KA-12'], scope('KA-12'), f'略称 {n12}種類', fmt(b12),
               '同じ略称で大項目・FHIR名が複数あるのは仕様か（KA-10と同じ論点）', None])
    ws.append(['KA-23', '検査方法(JLAC10-測定法)がJLAC10測定法コードと整合すること', total['KA-23'], scope('KA-23'), f'名称の組 {len(g23)}種類', fmt(b23),
               'マスタの検査方法の名称がJLAC10コード表と違うのは許容か。全角・半角や括弧だけの違いは指摘しなくてよいか', None])
    ws.append(['KA-08', '測定法(JLAC11)がJLAC11測定法コードと整合すること', total['KA-08'], scope('KA-08'), f'{len(act["KA-08"])}行', fmt(b08),
               '表記だけの違いは指摘しなくてよいか。名称が違う行はJLACセンターに確認するか', None])
    ws['A1'].font = Font(bold=True, size=13)
    for c in range(1, 9):
        ws.cell(8, c).font = BOLD
        ws.cell(8, c).fill = ASK if c == 8 else HEAD
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = WRAP
    for r in range(9, 13):
        ws.cell(r, 8).fill = ASK
    for col, w in zip('ABCDEFGH', (12, 30, 10, 14, 14, 46, 50, 30)):
        ws.column_dimensions[col].width = w

    add_sheet(wb, 'KA-10（大項目）', ['大項目', '対象の行数', '書き方の数', '書き方（FHIR項目名称／FHIR識別文字列／略称、行数、データタイプ）', 'ACNの見立て', '代表行（書き方ごとの先頭行）', '本田さんの判断（仕様／誤り）'], r10, (18, 8, 8, 60, 28, 30, 22))
    add_sheet(wb, 'KA-12（略称）', ['略称', '対象の行数', '書き方の数', '書き方（大項目／FHIR項目名称／FHIR識別文字列、行数、データタイプ）', 'ACNの見立て', '代表行（書き方ごとの先頭行）', '本田さんの判断（仕様／誤り）'], r12, (14, 8, 8, 60, 28, 30, 22))
    add_sheet(wb, 'KA-23（検査方法）', ['JLAC10測定法コード', 'マスタの検査方法(JLAC10-測定法)', 'JLAC10コード表の名称', '対象の行数', 'ACNの見立て', '代表行（先頭3行）', '本田さんの判断（仕様／誤り）'], r23, (10, 30, 30, 8, 22, 30, 22))
    add_sheet(wb, 'KA-08（測定法）', ['行', 'JLAC11コード', '販売名称', 'マスタの測定法(JLAC11)', '17桁コードリストの測定法名称', 'ACNの見立て', '本田さんの判断（仕様／誤り）'], r08, (8, 20, 28, 36, 30, 26, 22))

    out = next_name(a.out, tag) if os.path.isdir(a.out) else a.out
    if os.path.exists(out):
        raise SystemExit(f'出力先がすでにあります（上書きしません）：{out}')
    wb.save(out)
    print('OUT ' + out)
    print(json.dumps({'total': total, 'target': {k: len(v) for k, v in act.items()}, 'ended_kept': ended_kept,
                      'KA-10': b10, 'KA-12': b12, 'KA-23': dict(b23), 'KA-08': dict(b08)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
