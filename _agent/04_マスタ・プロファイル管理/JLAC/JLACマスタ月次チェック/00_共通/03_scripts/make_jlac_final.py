"""JLACマスタの完成版3点（ポータル公開用・基盤取り込み用・新のみ）を作る（内々の効率化用）。

決まり（2026-09版の現物 `20260915_JLACセンターからの回答受領後_9月版公開用マスタ準備.xlsx` の3シートで全行一致を確認、2026-10-16）：
  1 公開用（新旧混合） … 整備済みのマスタ（32列）そのまま
  2 取り込み用（新旧混合） … 1のうち JLAC11コードが空欄の行を、上から順に xyzxyzxyzxyz00001… で埋める（17桁・一意）。他の列は1と同じ
  3 新のみ（旧コード除外版） … 1のうち 適用終了日 >= 20260824 の行（99999999 を含む）。並びは1のまま
書式は公開版CSVと同じ（UTF-8・BOMなし・全項目を""で囲む・CRLF）。

使い方:
  python make_jlac_final.py --in <整備済みのマスタ（.csv または .xlsx）> [--sheet <シート名>] --ym 202610 --ymd 20261015 --out <出力フォルダ>
  python make_jlac_final.py --selftest <9月版公開用マスタ準備.xlsx>   … 9月の3シートを再現できるか

ファイル名：公開用は9月の公開CSVの名前（［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_YYYYMM_YYYYMMDD.csv）。
取り込み用・新のみの正式なファイル名は未確認のため、--name-tor / --name-new で指定する（既定は仮の名前）。
"""
import argparse
import csv
import os
import sys

CUTOFF = '20260824'  # 第6・第7改定の見直し前の旧コードを除く境目（KT 2026-10-15・9月版で確認）
DUMMY = 'xyzxyzxyzxyz{:05d}'
NAME_PUB = '［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_{ym}_{ymd}'
NAME_TOR = '【仮】共有項目JLACコードマスタ_取り込み用_{ym}_{ymd}'
NAME_NEW = '【仮】共有項目JLACコードマスタ_旧コード除外版_{ym}_{ymd}'


def cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def read_rows(path, sheet=None):
    if path.lower().endswith('.csv'):
        with open(path, encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f))
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    head = rows[0]
    n = len(head)
    while n and not head[n - 1]:
        n -= 1  # 見出しの無い右端の列（作業用の33列目など）は落とす
    return [r[:n] + [''] * (n - len(r[:n])) for r in rows]


def build(rows):
    head, data = rows[0], rows[1:]
    j, e = head.index('JLAC11コード'), head.index('適用終了日')
    tor, k = [], 0
    for r in data:
        r = list(r)
        if r[j] == '':
            k += 1
            r[j] = DUMMY.format(k)
        tor.append(r)
    new = [r for r in data if r[e] >= CUTOFF]
    return [head] + data, [head] + tor, [head] + new, k


def write(path, rows):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        csv.writer(f, quoting=csv.QUOTE_ALL, lineterminator='\r\n').writerows(rows)


def check(pub, tor, new, k):
    """作ったものの点検（NGがあれば送らない）"""
    head = pub[0]
    j, e = head.index('JLAC11コード'), head.index('適用終了日')
    msgs = []
    codes = [r[j] for r in tor[1:]]
    msgs.append(('OK' if all(len(c) == 17 for c in codes) else 'NG', '取り込み用：JLAC11コードがすべて17桁'))
    dups = len(codes) - len(set(codes))
    msgs.append(('情報', f'取り込み用：埋めた行 {k}件／JLAC11コードの重複 {dups}件（新旧混合のため同じコードが複数行あるのは通常。ダミーは一意）'))
    dummies = [c for c in codes if c.startswith('xyzxyzxyzxyz')]
    msgs.append(('OK' if len(dummies) == len(set(dummies)) == k else 'NG', 'ダミーが一意で、空欄の行数と同じ'))
    msgs.append(('OK' if all(r[e] >= CUTOFF for r in new[1:]) else 'NG', f'新のみ：適用終了日がすべて {CUTOFF} 以降'))
    msgs.append(('OK' if all(r[j] != '' for r in new[1:]) else '要確認', '新のみ：JLAC11コードの空欄なし'))
    msgs.append(('情報', f'行数：公開用 {len(pub) - 1}／取り込み用 {len(tor) - 1}／新のみ {len(new) - 1}'))
    return msgs


def selftest(path):
    import openpyxl
    names = openpyxl.load_workbook(path, read_only=True).sheetnames
    pick = lambda key: [n for n in names if key in n][0]
    pub0 = read_rows(path, pick('公開用'))
    tor0 = read_rows(path, pick('取り込み用'))
    new0 = read_rows(path, pick('新のみ'))
    pub, tor, new, k = build(pub0)
    ok = (pub == pub0, tor == tor0, new == new0)
    print(f'9月版の再現：公開用 {"一致" if ok[0] else "不一致"}／取り込み用 {"一致" if ok[1] else "不一致"}（埋めた行 {k}）／新のみ {"一致" if ok[2] else "不一致"}（{len(new) - 1}行）')
    return 0 if all(ok) else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='inp')
    ap.add_argument('--sheet')
    ap.add_argument('--ym'); ap.add_argument('--ymd'); ap.add_argument('--out')
    ap.add_argument('--name-tor', default=NAME_TOR); ap.add_argument('--name-new', default=NAME_NEW)
    ap.add_argument('--selftest')
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding='utf-8')
    if a.selftest:
        return selftest(a.selftest)
    pub, tor, new, k = build(read_rows(a.inp, a.sheet))
    os.makedirs(a.out, exist_ok=True)
    f = lambda t: os.path.join(a.out, t.format(ym=a.ym, ymd=a.ymd) + '.csv')
    write(f(NAME_PUB), pub); write(f(a.name_tor), tor); write(f(a.name_new), new)
    msgs = check(pub, tor, new, k)
    lines = ['# JLACマスタ 完成版3点の作成結果', ''] + [f'- **{s}** {m}' for s, m in msgs]
    open(os.path.join(a.out, '作成結果.md'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('\n'.join(lines))
    return 1 if any(s == 'NG' for s, _ in msgs) else 0


if __name__ == '__main__':
    sys.exit(main())
