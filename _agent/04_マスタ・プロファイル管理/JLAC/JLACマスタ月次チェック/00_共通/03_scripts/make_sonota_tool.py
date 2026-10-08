"""『その他の測定法』コード（ダミーコード）作成ツール（Excelひな形）を作る。

考え方は作成ツール「電子カルテ情報共有サービス対応JLACコード表作成ツール_20250127.xlsx」の
「ダミーコード生成用」シートと同じ（定常業務マニュアル JLACマスタチェック手順 v3 の7章）。
32列のマスタ（表示用単位2・XML用単位2あり）用。データは空で、使う人がマスタを貼る。

使い方：python make_sonota_tool.py <出力xlsx>
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

HEADER = ['区分名称', '予備区分名称', '救急フラグ', '生活習慣病フラグ', 'データ区分', '大項目', 'FHIR項目名称',
          'FHIR識別文字列', '略称', '販売名称', '材料(JLAC11)', '測定法(JLAC11)', 'JLAC11コード', '表示用単位',
          '表示用単位2', 'XML用単位', 'XML用単位2', '材料(JLAC10)', '検査方法(JLAC10-測定法)', 'JLAC10コード',
          '基準値フラグ(下限)', '基準値フラグ(上限)', '基準値フラグ(判定)', 'データタイプ', '値の下限', '値の上限',
          '数値型の場合の形式', 'コード型の場合の値リスト', 'コード型のOID', '並び順', '適用開始日', '適用終了日']
assert len(HEADER) == 32
MAX_ROWS = 20000      # 貼れるマスタの行数（見出しを除く）
MAX_OUT = 3000        # 結果シートの行数
SONOTA = "その他の測定法"

S_MASTER = "①マスタ貼り付け"
S_PREV = "②前月公開版（任意）"
S_WORK = "③作業（触らない）"
S_OUT = "④結果"
q = lambda s: f"'{s}'"
ix = {h: i for i, h in enumerate(HEADER)}
col = lambda h: CL(ix[h] + 1)
M, T, ST, EN = col("JLAC11コード"), col("JLAC10コード"), col("適用開始日"), col("適用終了日")
LAST = MAX_ROWS + 1

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


def head(ws, values):
    ws.append(values)
    for c in range(1, len(values) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def build(out):
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "使い方"

    # ①マスタ貼り付け：見出しだけ置く（貼るときは見出しごとA1に貼る）
    wm = wb.create_sheet(S_MASTER)
    head(wm, HEADER)

    # ②前月公開版：JLAC11コードの列だけ
    wp = wb.create_sheet(S_PREV)
    head(wp, ["JLAC11コード（前月公開版のJLAC11コード列をここに貼る）"])
    wp.column_dimensions["A"].width = 40

    # ③作業
    ww = wb.create_sheet(S_WORK)
    work_head = ["元のJLAC11コード", "000版JLAC11コード", "000版の重複(A)", "マスタに既存の000", "抽出対象(○)",
                 "候補の連番", "同じ000版に有効な行(99999999)", "前月公開版にあった",
                 "999版JLAC10コード", "単位（表示用単位〜XML用単位2）", "FHIR識別文字列",
                 "元の行でJLAC10が違う", "元の行で単位が違う", "元の行でFHIR識別文字列が違う"]
    head(ww, work_head)
    mref = q(S_MASTER)
    for n in range(2, LAST + 1):
        ww.append([
            f'=IF({mref}!{M}{n}="","",{mref}!{M}{n}&"")',
            f'=IF(A{n}="","",IF(MID(A{n},13,3)="000","",LEFT(A{n},12)&"000"&RIGHT(A{n},2)))',
            f'=IF(B{n}="","",IF(COUNTIF($B$2:B{n},B{n})>1,"A",""))',
            f'=IF(B{n}="","",IF(COUNTIF({mref}!$M$2:$M${LAST},B{n})>0,"既存",""))',
            f'=IF(AND(B{n}<>"",C{n}="",D{n}=""),"○","")',
            f'=IF(E{n}="○",COUNTIF($E$2:E{n},"○"),"")',
            f'=IF(E{n}<>"○","",IF(COUNTIFS($B$2:$B${LAST},B{n},{mref}!${EN}$2:${EN}${LAST},"99999999")>0,"あり","なし"))',
            f'=IF(E{n}<>"○","",IF(COUNTA({q(S_PREV)}!$A:$A)<=1,"（未貼付）",IF(COUNTIF({q(S_PREV)}!$A:$A,A{n})>0,"あり","なし")))',
            f'=IF(B{n}="","",LEFT({mref}!{T}{n}&"",12)&"999"&RIGHT({mref}!{T}{n}&"",2))',
            f'=IF(B{n}="","",{mref}!N{n}&"|"&{mref}!O{n}&"|"&{mref}!P{n}&"|"&{mref}!Q{n})',
            f'=IF(B{n}="","",{mref}!H{n}&"")',
            f'=IF(E{n}<>"○","",IF(COUNTIFS($B$2:$B${LAST},B{n},$I$2:$I${LAST},"<>"&I{n})>0,"違う",""))',
            f'=IF(E{n}<>"○","",IF(COUNTIFS($B$2:$B${LAST},B{n},$J$2:$J${LAST},"<>"&J{n})>0,"違う",""))',
            f'=IF(E{n}<>"○","",IF(COUNTIFS($B$2:$B${LAST},B{n},$K$2:$K${LAST},"<>"&K{n})>0,"違う",""))',
        ])
    for i in range(1, len(work_head) + 1):
        ww.column_dimensions[CL(i)].width = 20

    # ④結果
    wo = wb.create_sheet(S_OUT)
    extra = ["【確認】元にしたJLAC11コード（先頭行）", "【確認】要確認事項"]
    head(wo, ["（作業用）行"] + HEADER + [None] + extra)
    wref = q(S_WORK)
    for r in range(2, MAX_OUT + 2):
        a = f"$A{r}"
        idx = lambda c: f"INDEX({mref}!{c}:{c},{a})"
        line = [f'=IFERROR(MATCH(ROW()-1,{wref}!$F:$F,0),"")']
        for h in HEADER:
            c = col(h)
            if h in ("測定法(JLAC11)", "検査方法(JLAC10-測定法)"):
                f = f'"{SONOTA}"'
            elif h == "販売名称":
                f = '""'
            elif h == "JLAC11コード":
                f = f"INDEX({wref}!B:B,{a})"
            elif h == "JLAC10コード":
                f = f'IF({idx(T)}&""="","",LEFT({idx(T)}&"",12)&"999"&RIGHT({idx(T)}&"",2))'
            elif h == "適用終了日":
                f = f'IF(AND(INDEX({wref}!G:G,{a})="あり",{idx(EN)}&""<>"99999999"),"99999999",{idx(EN)}&"")'
            else:
                f = f'IF({idx(c)}&""="","",{idx(c)})'
            line.append(f'=IF({a}="","",{f})')
        line.append(None)
        line.append(f'=IF({a}="","",INDEX({wref}!A:A,{a}))')
        line.append(
            f'=IF({a}="","",'
            f'IF(INDEX({wref}!G:G,{a})="なし","元の行がすべて終了済み：作るかどうか要判断。","")'
            f'&IF(AND(INDEX({wref}!G:G,{a})="あり",{idx(EN)}&""<>"99999999"),"先頭行は終了済み→有効な行があるため終了日を99999999にした（KA-34）。","")'
            f'&IF(OR({idx(ST)}&""="0",{idx(ST)}&""="00000000"),"開始日が00000000（KA-M1：新規は00000000以外）。","")'
            f'&IF(INDEX({wref}!H:H,{a})="あり","元の行は前月公開版からある：開始日を元の行のまま写すか、今月の15日にするか要判断。","")'
            f'&IF(INDEX({wref}!L:L,{a})="違う","元の行でJLAC10コードが違う（先頭行の999版を使用。どの行のJLAC10にするか要判断）。","")'
            f'&IF(INDEX({wref}!M:M,{a})="違う","元の行で単位（表示用単位・XML用単位・単位2）が違う（先頭行を使用）。","")'
            f'&IF(INDEX({wref}!N:N,{a})="違う","元の行でFHIR識別文字列が違う（先頭行を使用）。",""))'
        )
        wo.append(line)
    wo.column_dimensions["A"].hidden = True
    for i in range(2, len(HEADER) + 2):
        wo.column_dimensions[CL(i)].width = 14
    wo.column_dimensions[CL(ix["JLAC11コード"] + 2)].width = 20
    nh = len(HEADER) + 3
    for i in (nh, nh + 1):
        wo.cell(row=1, column=i).fill = NOTE_FILL
    wo.column_dimensions[CL(nh)].width = 22
    wo.column_dimensions[CL(nh + 1)].width = 70

    # 使い方
    layout_ok = (f'=IF(AND({mref}!{M}1="JLAC11コード",{mref}!{T}1="JLAC10コード",{mref}!O1="表示用単位2",'
                 f'{mref}!Q1="XML用単位2",{mref}!{ST}1="適用開始日",{mref}!{EN}1="適用終了日"),"OK",'
                 f'"NG：列の並びが違う（32列のマスタを見出しごとA1に貼る）")')
    n_rows = f'=COUNTA({mref}!A:A)-1'
    rows_ok = f'=IF(COUNTA({mref}!A:A)-1>{MAX_ROWS},"NG：{MAX_ROWS:,}行を超えている（ツールを作り直す）","OK")'
    n_cand = f'=MAX({wref}!F:F)'
    cand_ok = f'=IF(MAX({wref}!F:F)>{MAX_OUT},"NG：結果シートの行数（{MAX_OUT:,}行）を超えている","OK")'
    n_existing = f'=COUNTIF({mref}!${M}$2:${M}${LAST},"????????????000??")'
    lines = [
        ["『その他の測定法』コード（ダミーコード）作成ツール"],
        ["JLACマスタの各コードについて、測定法を「その他の測定法」（JLAC11の13〜15桁=000／JLAC10=999）にしたコードのうち、マスタにまだ無いものを作る。ダミーコード＝その他の測定法コード。"],
        [],
        ["使い方"],
        ["1", f"今月のマスタ（32列）を全件、見出しごと「{S_MASTER}」のA1に値で貼る（33列目にコメント列があっても可）"],
        ["2", f"（任意）前月公開版のJLAC11コード列を「{S_PREV}」のA1に見出しごと貼る。元のコードが前月からあるのに000版が無いものに印が付く"],
        ["3", "下の「確認」がすべてOKであることを見る（計算に数十秒かかることがある）"],
        ["4", f"「{S_OUT}」に出た行が、追加する000コードの候補。マスタに入れるときは値で貼り、右側の黄色い【確認】列は除く"],
        ["5", "【確認】要確認事項に書かれた行は、本田さんと相談して開始日・終了日などを決める"],
        [],
        ["確認", "結果"],
        ["列の並び", layout_ok],
        ["マスタの行数", n_rows],
        ["行数の上限", rows_ok],
        ["マスタにすでにある000コード（JLAC11の13〜15桁が000の行）", n_existing],
        ["新しく作る000コードの候補", n_cand],
        ["結果シートの上限", cand_ok],
        [],
        ["作り方（各列の決まり方）"],
        ["・", "JLAC11コードの13〜15桁を000、JLAC10コードの13〜15桁を999に置き換える"],
        ["・", "同じ000コードにまとまる行のうち、いちばん上の行の名称・FHIR識別文字列・単位・並び順・適用開始日などを写す（9月公開版の000コード1,179件もこの決め方と一致）"],
        ["・", "測定法(JLAC11)・検査方法(JLAC10-測定法)は「その他の測定法」、販売名称は空欄"],
        ["・", "マスタにすでにある000コードは作らない"],
        ["・", "適用終了日：同じ000コードにまとまる行に有効な行（99999999）があれば99999999（KA-34）。それ以外は先頭行を写す"],
        ["・", "同じ000コードにまとまる行で、JLAC10コード・単位・FHIR識別文字列が違うときは【確認】要確認事項に書く（先頭行の値を使っている。9月公開版の000コードでも、JLAC10コード（999版）が先頭行以外の行から取られていたものが61件あった）"],
        [],
        ["元にしたもの"],
        ["・", "本田さんの作成ツール「電子カルテ情報共有サービス対応JLACコード表作成ツール_20250127.xlsx」の「ダミーコード生成用」シート。元のツールは30列の旧マスタ用で、32列のマスタにそのまま貼るとJLAC10の999版が材料(JLAC10)から作られるため、このツールを作った"],
        ["・", "手順の正本：定常業務マニュアル「JLACマスタチェック手順」7章。ツールの作り直し：00_共通/03_scripts/make_sonota_tool.py"],
        ["・", f"上限：マスタ{MAX_ROWS:,}行・結果{MAX_OUT:,}行。超えるときはスクリプトの MAX_ROWS・MAX_OUT を変えて作り直す"],
    ]
    for l in lines:
        guide.append(l)
    guide["A1"].font = Font(bold=True, size=13)
    for r in (4, 11, 19, 27):
        guide.cell(row=r, column=1).font = BOLD
    for r in range(12, 18):
        guide.cell(row=r, column=2).fill = NOTE_FILL
        guide.cell(row=r, column=2).font = BOLD
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


if __name__ == "__main__":
    build(sys.argv[1])
    print("ok")
