"""確認結果のダブルチェック用Excelを作る。元データを値で貼り込み、判定はすべてExcel関数で計算する（ツールのコードとは独立した検算）。
使い方: python 00_共通/03_scripts/make_doublecheck.py <ACN確認用xlsx> <出力xlsx> [<YYYYMM提出分フォルダ>]（省略時は最新の月フォルダ。入力ファイル名・新規開始行・公開日は下の定数を毎月差し替え）
前月の基準は前月公開CSV（全件）。新旧混合リストの提出にも対応（JLAC11空欄の行は全項目をつないだキーで前月と突き合わせる）。
10/1提出版（前月の基準＝9月FIXの新のみシート）で使った版は make_doublecheck_20261001提出版用.py
"""
import sys, csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

sys.stdout.reconfigure(encoding='utf-8')
BASE = Path(__file__).resolve().parent.parent  # 00_共通
ROOT = BASE.parent  # JLACマスタ月次チェック
MONTH = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else sorted(d for d in ROOT.iterdir() if d.is_dir() and d.name[:6].isdigit() and d.name.endswith('提出分'))[-1]
IN = MONTH / '01_input'
NEW = IN / '1006再提出［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_202609_20260915.xlsx'
PREV = IN / '［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_202609_20260915.csv'
J10 = IN / '140jlac10_1.xlsx'
J11 = IN / 'jlac11_1_1.0.xlsx'
LIST11 = ROOT.parent / '20260930_JLACマスター新旧対応表確認' / '01_input' / 'jlac11_3_1.1b.xlsx'
NEW_SHEETS = ['1006提出_（新旧コード混合リスト）']
NEWROW, START = 13436, 20261015
LABEL = 'JLACセンター 2026-10-06再提出版（新旧混合リスト）'
NOTE_COL = 33  # JLACセンターの連絡用の列（見出しに「コメント」）。無い月は None
ACN, OUT = Path(sys.argv[1]), Path(sys.argv[2])


def rows(path, sheet, min_row=1, ncol=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = [list(r if ncol is None else (list(r) + [None] * ncol)[:ncol]) for r in wb[sheet].iter_rows(min_row=min_row, values_only=True)]
    wb.close()
    return out


def pad(v, n):
    s = '' if v is None else str(v).strip()
    return s.zfill(n) if s.isdigit() else s


wb = openpyxl.Workbook()
chk = wb.active
chk.title = '照合'
bold = Font(bold=True)

# ---- 今回（メイン＋ABO） ----
cur = wb.create_sheet('今回')
hdr = rows(NEW, NEW_SHEETS[0], 1, 32)[0]
cur.append(hdr + ['シート', '原本行'])
notes = {}
for sh in NEW_SHEETS:
    for i, r in enumerate(rows(NEW, sh, 1, 40)[1:], start=2):
        if any(v not in (None, '') for v in r[:32]):
            cur.append(r[:32] + [sh, i])
            if NOTE_COL and r[NOTE_COL - 1] not in (None, ''):
                notes[cur.max_row] = r[NOTE_COL - 1]
N = cur.max_row  # 最終行
R = lambda c: f"'今回'!${c}$2:${c}${N}"

# ---- 前月 ----
prev = wb.create_sheet('前月')
with open(PREV, encoding='utf-8-sig', newline='') as fp:
    pr = [(r + [''] * 32)[:32] for r in csv.reader(fp)]
for r in pr:
    if any(v != '' for v in r):
        prev.append([v if v != '' else None for v in r])
PN = prev.max_row
KEY = lambda r: '&"|"&'.join(f'${L(c)}{r}' for c in range(1, 33))
prev['AG1'] = 'キー：全項目'; prev['AH1'] = '今回に同じコード（空欄行は同じキー）の行数'; prev['AI1'] = '前月内の同じキーの行数（空欄行）'; prev['AJ1'] = '空欄行：今回で減った分（同じキーの行数で按分）'
for r in range(2, PN + 1):
    prev[f'AG{r}'] = '=' + KEY(r)
    prev[f'AH{r}'] = f"=IF($M{r}&\"\"<>\"\",COUNTIF('今回'!$M$2:$M${N},$M{r}),COUNTIF('今回'!$BS$2:$BS${N},$AG{r}))"
    prev[f'AI{r}'] = f'=IF($M{r}&""="",COUNTIF($AG$2:$AG${PN},$AG{r}),"")'
    prev[f'AJ{r}'] = f'=IF($M{r}&""="",MAX(0,AI{r}-AH{r})/AI{r},0)'
P = lambda c: f"'前月'!${c}$2:${c}${PN}"

# ---- コード表（値のまま。コードの桁そろえのみ） ----
s = wb.create_sheet('JLAC11材料'); s.append(['材料コード', '材料名'])
for r in rows(J11, '材料コード', 2, 8):
    if r[2] not in (None, ''): s.append([pad(r[2], 3), r[3]])
s = wb.create_sheet('JLAC11単位'); s.append(['結果単位コード', '結果単位名(1)', '結果単位(2)'])
for r in rows(J11, '結果単位コード', 2, 6):
    if r[0] not in (None, ''): s.append([str(r[0]).strip(), r[1], r[2]])
s = wb.create_sheet('JLAC11_17桁'); s.append(['キー（測定物|識別|測定法）', '測定法名称'])
for r in rows(LIST11, 'JLAC11_17桁コードリスト', 2, 44):
    if r[3] not in (None, ''): s.append([f'{str(r[3]).strip()}|{pad(r[5], 4)}|{pad(r[9], 3)}', r[10]])
s = wb.create_sheet('JLAC10分析物'); s.append(['分析物コード', '分析物名'])
for r in rows(J10, '分析物コード', 6, 10):
    if r[3] not in (None, ''): s.append([str(r[3]).strip(), r[4]])
s = wb.create_sheet('JLAC10材料'); s.append(['材料コード', '材料名'])
for r in rows(J10, '材料コード', 10, 12):
    if r[6] not in (None, '') and str(r[6]).strip().isdigit(): s.append([pad(r[6], 3), r[7]])
s = wb.create_sheet('JLAC10測定法'); s.append(['測定法コード', '測定法名', '測定法名(2)'])
for r in rows(J10, '測定法コード', 6, 10):
    if r[3] not in (None, '') and str(r[3]).strip().isdigit(): s.append([pad(r[3], 3), r[4], r[5]])

# ---- ツールの結果（突き合わせ用） ----
tool_new, tool_ids, tool_rec, tool_code = {}, {}, {}, {}
awb = openpyxl.load_workbook(ACN, read_only=True, data_only=True)
for name in awb.sheetnames:
    if name.startswith('JLACセンター向けコメント_確認用'):
        if name != 'JLACセンター向けコメント_確認用' and len(NEW_SHEETS) < 2:
            continue
        src = NEW_SHEETS[0] if name == 'JLACセンター向けコメント_確認用' else NEW_SHEETS[1]
        for i, r in enumerate(awb[name].iter_rows(min_row=2, values_only=True), start=2):
            tool_new[(src, i)] = r[33]
            tool_ids[(src, i)] = r[34] or ''
            tool_rec[(src, i)] = r[35] or ''  # 記録のみ（既知・共有対象外など）
            tool_code[(src, i)] = r[12]
tsum = {}
# ツールv5以降、手順書由来の項目は「その他の指摘事項」シートに表示IDで出るため内部IDへ戻す
ALIAS = {'その他1（削除行）': 'STEP1-DEL', 'その他2（項目の変更）': 'STEP1-DIFF', 'その他3（表記の変更）': 'STEP1-DIFF表記',
         'その他4（適用終了日）': 'STEP1-END', 'その他5（新規行の適用開始日）': 'STEP2-START'}
for _sh in ('チェック結果一覧', 'その他の指摘事項'):
    if _sh not in awb.sheetnames:
        continue
    _rows = list(awb[_sh].iter_rows(values_only=True))
    _h = list(_rows[0])
    _ci, _ki = _h.index('コメント対象 行数'), _h.index('既知（コメントなし）件数')
    for r in _rows[1:]:
        tsum[ALIAS.get(r[0], r[0])] = (int(r[_ci] or 0), int(r[_ki] or 0))
awb.close()
# ツールの行ごとの結果から数える（指摘の件数は1行に複数出ることがあるため、行数で突き合わせる）
both = {k: f"{tool_ids[k]} {tool_rec[k]}" for k in tool_new}
t_new = sum(1 for v in tool_new.values() if v == '新規')
t_new_from = sum(1 for (sh, i), v in tool_new.items() if sh == NEW_SHEETS[0] and i >= NEWROW and v == '新規')
t_rows_from = sum(1 for (sh, i) in tool_new if sh == NEW_SHEETS[0] and i >= NEWROW)
t_first_new = min((i for (sh, i), v in tool_new.items() if v == '新規'), default='—')
t_blank_exist = sum(1 for k, v in tool_new.items() if v == '既存' and tool_code[k] in (None, ''))
t_note = sum(1 for k in tool_new if 'KA-26(連絡用の列)' in tool_rec[k])
t_diff = sum(1 for k in tool_new if any(x in both[k] for x in ('KA-M2', 'KA-M3', 'その他2', 'その他3', 'その他4')))
t_k08 = lambda kind: sum(1 for k, v in tool_new.items() if v == kind and 'KA-08' in tool_ids[k])

# ---- 今回シートの補助列（すべてExcel関数） ----
def norm(x):
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(ASC({x})," ",""),"‐","-"),"−","-"),"―","-"),"–","-")'

helpers = [
    ('AI', '前月行（コード、空欄行は全項目キーで突き合わせ）', lambda r: f"=IF($M{r}&\"\"<>\"\",IFERROR(MATCH($M{r},'前月'!$M$1:$M${PN},0),\"\"),IFERROR(MATCH($BS{r},'前月'!$AG$1:$AG${PN},0),\"\"))"),
    ('AJ', '新規判定', lambda r: f'=IF(AI{r}="","新規","既存")'),
    ('AK', 'KA-18 開始日が0', lambda r: f'=IF($AE{r}&""="0",1,0)'),
    ('AL', '前月と違う列数（開始・終了日以外）', lambda r: f"=IF(AI{r}=\"\",\"\",SUMPRODUCT(--(($A{r}:$AD{r}&\"\")<>(INDEX('前月'!$A$1:$AD${PN},AI{r},0)&\"\"))))"),
    ('AM', 'KA-M2 開始日変更', lambda r: f"=IF(AI{r}=\"\",\"\",IF(VALUE($AE{r}&\"\")<>VALUE(INDEX('前月'!$AE$1:$AE${PN},AI{r})&\"\"),1,0))"),
    ('AN', '終了日変更', lambda r: f"=IF(AI{r}=\"\",\"\",IF(VALUE($AF{r}&\"\")<>VALUE(INDEX('前月'!$AF$1:$AF${PN},AI{r})&\"\"),1,0))"),
    ('AO', 'JLAC11材料名（表）', lambda r: f"=IF(LEN($M{r}&\"\")<>17,\"\",IFERROR(VLOOKUP(MID($M{r},10,3),'JLAC11材料'!$A:$B,2,FALSE)&\"\",\"#コード無し\"))"),
    ('AP', 'KA-07 材料一致', lambda r: f'=IF(AO{r}="","",IF(EXACT(TRIM($K{r}),TRIM(SUBSTITUTE(AO{r},"　",""))),1,0))'),
    ('AQ', 'JLAC11測定法名（17桁リスト）', lambda r: f"=IF(LEN($M{r}&\"\")<>17,\"\",IFERROR(VLOOKUP(LEFT($M{r},5)&\"|\"&MID($M{r},6,4)&\"|\"&MID($M{r},13,3),'JLAC11_17桁'!$A:$B,2,FALSE)&\"\",\"\"))"),
    ('AR', 'KA-08 測定法一致', lambda r: f'=IF(OR(AQ{r}="",MID($M{r},13,3)="000"),"",IF({norm(f"$L{r}")}={norm(f"AQ{r}")},1,0))'),
    ('AS', 'KA-06 単位一致', lambda r: f"=IF(LEN($M{r}&\"\")<>17,\"\",IFERROR(IF(OR(EXACT($N{r},TRIM(VLOOKUP(MID($M{r},16,2),'JLAC11単位'!$A:$C,2,FALSE)&\"\")),ISNUMBER(FIND(\",\"&$N{r}&\",\",\",\"&SUBSTITUTE(SUBSTITUTE(VLOOKUP(MID($M{r},16,2),'JLAC11単位'!$A:$C,3,FALSE)&\"\",\" \",\"\"),\",,\",\",\")&\",\"))),1,0),0))"),
    ('AT', 'KA-33 000の相方の数', lambda r: f'=IF(MID($M{r},13,3)="000",COUNTIFS({R("M")},LEFT($M{r},12)&"???"&RIGHT($M{r},2))-1,"")'),
    ('AU', 'KA-31 同じJLAC11の数', lambda r: f'=COUNTIF({R("M")},$M{r})'),
    ('AV', 'KA-10 不整合', lambda r: f'=IF(TRIM($F{r})="","",IF(COUNTIFS({R("F")},$F{r},{R("BL")},"<>"&$BL{r})>0,1,0))'),
    ('AW', 'KA-11 不整合', lambda r: f'=IF(TRIM($H{r})="","",IF(COUNTIFS({R("H")},$H{r},{R("BM")},"<>"&$BM{r})>0,1,0))'),
    ('AX', 'KA-12 不整合', lambda r: f'=IF(TRIM($I{r})="","",IF(COUNTIFS({R("I")},$I{r},{R("BN")},"<>"&$BN{r})>0,1,0))'),
    ('AY', 'KA-17 不整合', lambda r: f'=IF(TRIM($G{r})="","",IF(COUNTIFS({R("G")},$G{r},{R("AD")},"<>"&$AD{r})>0,1,0))'),
    ('AZ', 'KA-32 不整合', lambda r: f'=IF(TRIM($T{r})="","",IF(COUNTIFS({R("T")},$T{r},{R("H")},"<>"&$H{r})>0,1,0))'),
    ('BA', 'KA-21 分析物コードあり', lambda r: f"=IF(COUNTIF('JLAC10分析物'!$A:$A,LEFT($T{r},5))>0,1,0)"),
    ('BB', 'KA-22 材料(JLAC10)一致', lambda r: f"=IF(EXACT(TRIM($R{r}),TRIM(SUBSTITUTE(IFERROR(VLOOKUP(MID($T{r},10,3),'JLAC10材料'!$A:$B,2,FALSE)&\"\",\"#\"),\"　\",\"\"))),1,0)"),
    ('BC', 'KA-23 測定法(JLAC10)一致', lambda r: (lambda v1, v2, s: f'=IFERROR(IF(OR(EXACT({s},{v1}),EXACT({s},{v1}&{v2}),EXACT({s},{v1}&"("&{v2}&")"),EXACT({s},{v1}&"（"&{v2}&"）"),EXACT({s},{v1}&"["&{v2}&"]"),EXACT({s},{v1}&" "&{v2}),AND({v2}<>"",EXACT({s},{v2}))),1,0),0)')(
        f"TRIM(SUBSTITUTE(VLOOKUP(MID($T{r},13,3),'JLAC10測定法'!$A:$C,2,FALSE)&\"\",\"　\",\"\"))",
        f"TRIM(SUBSTITUTE(VLOOKUP(MID($T{r},13,3),'JLAC10測定法'!$A:$C,3,FALSE)&\"\",\"　\",\"\"))", f'TRIM($S{r})')),
    ('BD', 'KA-34 不整合', lambda r: f'=IF(AND(MID($M{r},13,3)="000",$AF{r}&""<>"99999999"),IF(COUNTIFS({R("M")},LEFT($M{r},12)&"???"&RIGHT($M{r},2),{R("AF")},"<>99999999")-1>0,1,0),"")'),
    ('BE', 'ツール：新規/既存', None),
    ('BF', 'ツール：コメント対象ID', None),
    ('BG', '新規判定の一致', lambda r: f'=IF(AJ{r}=BE{r},1,0)'),
    ('BH', 'KA-18の一致', lambda r: f'=IF(AK{r}=IF(ISNUMBER(SEARCH("KA-18",BF{r}&" "&BP{r})),1,0),1,0)'),
    ('BI', '前月差分の一致', lambda r: '=IF(BO{0}=IF('.format(r) + '+'.join(f'ISNUMBER(SEARCH("{x}",BF{r}&" "&BP{r}))' for x in ('KA-M2', 'KA-M3', 'その他2', 'その他3', 'その他4')) + '>0,1,0),1,0)'),
    ('BJ', 'KA-07の一致', lambda r: f'=IF((AP{r}=0)=ISNUMBER(SEARCH("KA-07",BF{r}&" "&BP{r})),1,0)'),
    ('BK', 'KA-08の一致', lambda r: f'=IF(OR(AND(AJ{r}="新規",AR{r}=0),AND(AJ{r}="既存",AR{r}=0,OR(BQ{r}=1,BT{r}=1)))=ISNUMBER(SEARCH("KA-08",BF{r})),1,0)'),
    ('BL', 'キー：FHIR項目名称|識別|略称', lambda r: f'=$G{r}&"|"&$H{r}&"|"&$I{r}'),
    ('BM', 'キー：大項目|FHIR項目名称|略称', lambda r: f'=$F{r}&"|"&$G{r}&"|"&$I{r}'),
    ('BN', 'キー：大項目|FHIR項目名称|識別', lambda r: f'=$F{r}&"|"&$G{r}&"|"&$H{r}'),
    ('BO', '前月と差分あり', lambda r: f'=IF(AI{r}="",0,IF(AL{r}+AM{r}+AN{r}>0,1,0))'),
    ('BP', 'ツール：記録のみのチェックID（既知・共有対象外など）', None),
    ('BQ', 'KA-08 前月の測定法での一致', lambda r: f'=IF(OR(AI{r}="",AQ{r}="",MID($M{r},13,3)="000"),"",IF({norm(f"INDEX('前月'!$L$1:$L${PN},AI{r})&\"\"")}={norm(f"AQ{r}")},1,0))'),
    ('BR', '連絡用の列（提出ファイルの33列目）', None),
    ('BS', 'キー：全項目', lambda r: '=' + KEY(r)),
    ('BT', '測定法(JLAC11)が前月から変更', lambda r: f"=IF(AI{r}=\"\",\"\",IF(EXACT($L{r}&\"\",INDEX('前月'!$L$1:$L${PN},AI{r})&\"\"),0,1))"),
]
for col, title, f in helpers:
    cur[f'{col}1'] = title
    cur[f'{col}1'].font = bold
for r in range(2, N + 1):
    key = (cur[f'AG{r}'].value, cur[f'AH{r}'].value)
    for col, title, f in helpers:
        if col == 'BE': cur[f'BE{r}'] = tool_new.get(key, '')
        elif col == 'BF': cur[f'BF{r}'] = tool_ids.get(key, '')
        elif col == 'BP': cur[f'BP{r}'] = tool_rec.get(key, '')
        elif col == 'BR': cur[f'BR{r}'] = notes.get(r)
        else: cur[f'{col}{r}'] = f(r)
for c in range(1, 73):
    cur.cell(row=1, column=c).font = bold
cur.freeze_panes = 'N2'
cur.auto_filter.ref = f'A1:BT{N}'

# ---- 照合シート ----
chk['A1'] = f'確認結果のダブルチェック（{LABEL}）'; chk['A1'].font = Font(bold=True, size=14)
chk['A2'] = '元データ（今回提出・前月公開CSV・コード表）を値のまま貼り込み、判定はすべてExcel関数で計算した（ツールのコードは使っていない）。D列はツールの結果。E列がすべて「一致」なら、ツールの確認結果はExcelでの独立した検算と同じ。'
chk['A3'] = '不一致がある場合：「今回」シートの補助列（AI〜BT）でフィルタして該当行を確認する。行ごとの一致列（BG〜BK）が0の行が食い違い。'
chk['A4'] = '目視確認（推奨）：指摘のある種類ごとに「今回」シートで2〜3行を選び、原本の値・コード表を自分の目で確認する。'
for c in ('A2', 'A3', 'A4'):
    chk[c].alignment = Alignment(wrap_text=True, vertical='top')
    chk.merge_cells(f'{c}:F{c[1:]}')
    chk.row_dimensions[int(c[1:])].height = 36
head = ['区分', '確認項目', 'Excelでの再計算', 'ツールの結果', '判定', '計算方法・メモ']
HR = 6
for i, h in enumerate(head, start=1):
    chk.cell(row=HR, column=i, value=h).font = bold
    chk.cell(row=HR, column=i).fill = PatternFill('solid', fgColor='DDEBF7')
T = lambda k, known=False: tsum[k][0] + (tsum[k][1] if known else 0)
cnt = lambda c, crit: f'=COUNTIF({R(c)},{crit})'
items = [
    ('A 件数の突き合わせ', '前月（公開版・全件）の行数', f"=COUNTA({P('A')})", '—', '前月公開CSV'),
    ('A 件数の突き合わせ', '今回提出の行数', f'=COUNTA({R("A")})', '—', ''),
    ('A 件数の突き合わせ', '新規行の数（前月に同じJLAC11コードが無い行）', cnt('AJ', '"新規"'), t_new, ''),
    ('A 件数の突き合わせ', '前月にあって今回に無い行（削除行。コードあり）', f"=SUMPRODUCT(({P('M')}&\"\"<>\"\")*({P('AH')}=0))", '—', '今回に同じJLAC11コードが無い前月の行'),
    ('A 件数の突き合わせ', '前月にあって今回に無い行（削除行。JLAC11空欄）', f"=SUM({P('AJ')})", '—', '全項目が同じ行の数を前月と今回で比べ、減った分（前月シートAJ列）'),
    ('A 件数の突き合わせ', '前月にあって今回に無い行（削除行の合計）', f'=C10+C11', T('STEP1-DEL', True), 'ツール：その他1（削除行）'),
    ('A 件数の突き合わせ', '前月でJLAC11が空欄の行', f"=SUMPRODUCT(--({P('M')}&\"\"=\"\"))", '—', 'JLAC10のみの見直し前行'),
    ('A 件数の突き合わせ', '今回のJLAC11空欄の行のうち、前月に全項目が同じ行がある行', f'=COUNTIFS({R("M")},"",{R("AI")},">0")', t_blank_exist, 'ツール：空欄で「既存」と判定した行'),
    ('A 件数の突き合わせ', '最初の新規行の原本行番号', f"=INDEX({R('AH')},MATCH(\"新規\",{R('AJ')},0))", t_first_new, f'JLACセンター連絡：{NEWROW}行目以降が新規'),
    ('A 件数の突き合わせ', f'{NEWROW}行目以降の行数', f'=COUNTIFS({R("AG")},"{NEW_SHEETS[0]}",{R("AH")},">={NEWROW}")', t_rows_from, ''),
    ('A 件数の突き合わせ', f'{NEWROW}行目以降で新規の行数', f'=COUNTIFS({R("AG")},"{NEW_SHEETS[0]}",{R("AH")},">={NEWROW}",{R("AJ")},"新規")', t_new_from, '差がある場合は、同じコードが前月にある行（コード重複など）'),
    ('A 件数の突き合わせ', f'新規行のうち適用開始日={START}', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},{START})', t_new - T('STEP2-START', True), f'JLACセンター連絡：適用開始日{START}'),
    ('A 件数の突き合わせ', '連絡用の列（33列目）に記載のある行', f'=COUNTA({R("BR")})', t_note, 'ツール：KA-26(連絡用の列)として記録のみ'),
    ('B 指摘の検算', 'KA-18 適用開始日が0の行', f'=SUM({R("AK")})', T('KA-18', True), ''),
    ('B 指摘の検算', '前月と値が違う行（STEP1・KA-M）', f'=SUM({R("BO")})', t_diff, 'ツール：行ごとのチェックIDにKA-M2・KA-M3・その他2〜4がある行（連携済みを含む）'),
    ('B 指摘の検算', '　うち適用開始日の変更（KA-M2）', f'=SUM({R("AM")})', T('KA-M2', True), ''),
    ('B 指摘の検算', '　うち適用終了日の変更', f'=SUM({R("AN")})', T('STEP1-END', True) + T('KA-M3', True), ''),
    ('B 指摘の検算', 'KA-07 材料(JLAC11)がコード表と不一致', cnt('AP', 0), T('KA-07', True), ''),
    ('B 指摘の検算', 'KA-08 測定法(JLAC11)が17桁リストと不一致（新規行）', f'=COUNTIFS({R("AJ")},"新規",{R("AR")},0)', t_k08('新規'), 'ASCで全角→半角、ハイフン類・空白をそろえて比較'),
    ('B 指摘の検算', 'KA-08 既存行で今回不一致、かつ前月は一致または測定法の値が変わった行', f'=COUNTIFS({R("AJ")},"既存",{R("AR")},0,{R("BQ")},1)+COUNTIFS({R("AJ")},"既存",{R("AR")},0,{R("BQ")},0,{R("BT")},1)', t_k08('既存'), 'ツールは指摘の文面（値を含む）が前月と同じなら既知とするため、前月も不一致でも値が変われば新しい指摘になる（BQ・BT列）'),
    ('B 指摘の検算', '行ごとの不一致：新規判定', cnt('BG', 0), 0, '0なら全行でツールと一致（ツール側はコメント対象と記録のみの両方のIDで照合）'),
    ('B 指摘の検算', '行ごとの不一致：KA-18', cnt('BH', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：前月差分', cnt('BI', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：KA-07', cnt('BJ', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：KA-08', cnt('BK', 0), 0, ''),
    ('C 見逃しの確認', 'KA-01 予備区分名称・基準値フラグに値がある', f'=SUMPRODUCT(--(LEN({R("B")}&{R("U")}&{R("V")}&{R("W")})>0))', T('KA-01', True), ''),
    ('C 見逃しの確認', 'KA-02 JLAC11コードが17桁でない', f'=SUMPRODUCT(--(LEN({R("M")}&"")<>17))', T('KA-02', True), ''),
    ('C 見逃しの確認', 'KA-03 救急・生活習慣病フラグが0/1以外', f'=SUMPRODUCT(--((({R("C")}&"")<>"0")*(({R("C")}&"")<>"1")))+SUMPRODUCT(--((({R("D")}&"")<>"0")*(({R("D")}&"")<>"1")))', T('KA-03', True), ''),
    ('C 見逃しの確認', 'KA-04 データ区分が1/2以外', f'=SUMPRODUCT(--((({R("E")}&"")<>"1")*(({R("E")}&"")<>"2")))', T('KA-04', True), ''),
    ('C 見逃しの確認', 'KA-06 表示用単位が結果単位コードと不一致', cnt('AS', 0), T('KA-06', True), ''),
    ('C 見逃しの確認', 'KA-10 大項目キーの不整合（既知を含む）', cnt('AV', 1), T('KA-10', True), '空腹時/随時など設計上の区分で、前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-11 FHIR識別文字列キーの不整合', cnt('AW', 1), T('KA-11', True), ''),
    ('C 見逃しの確認', 'KA-12 略称キーの不整合（既知を含む）', cnt('AX', 1), T('KA-12', True), '前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-13 データタイプがPQ/CD以外', f'=SUMPRODUCT(--(({R("X")}<>"PQ")*({R("X")}<>"CD")))', T('KA-13', True), ''),
    ('C 見逃しの確認', 'KA-14 CDで値の下限・上限が空', f'=SUMPRODUCT(({R("X")}="CD")*(((TRIM({R("Y")}&"")="")+(TRIM({R("Z")}&"")=""))>0))', T('KA-14', True), ''),
    ('C 見逃しの確認', 'KA-15 CDで値リストが空', f'=SUMPRODUCT(({R("X")}="CD")*(TRIM({R("AB")}&"")=""))', T('KA-15', True), ''),
    ('C 見逃しの確認', 'KA-16 CDでOIDが空', f'=SUMPRODUCT(({R("X")}="CD")*(TRIM({R("AC")}&"")=""))', T('KA-16', True), ''),
    ('C 見逃しの確認', 'KA-17 同じFHIR項目名称で並び順が違う', cnt('AY', 1), T('KA-17', True), ''),
    ('C 見逃しの確認', 'KA-19 適用終了日が8桁の99999999/実在日付でない', f'=SUMPRODUCT(--(LEN({R("AF")}&"")<>8))+SUMPRODUCT((({R("AF")}&"")<>"99999999")*ISERROR(DATEVALUE(TEXT({R("AF")},"0000-00-00"))))', T('KA-19', True), ''),
    ('C 見逃しの確認', 'KA-20 JLAC10コードが17桁でない', f'=SUMPRODUCT(--(LEN({R("T")}&"")<>17))', T('KA-20', True), ''),
    ('C 見逃しの確認', 'KA-21 分析物コードがJLAC10に無い', cnt('BA', 0), T('KA-21', True), ''),
    ('C 見逃しの確認', 'KA-22 材料(JLAC10)がコード表と不一致', cnt('BB', 0), T('KA-22', True), ''),
    ('C 見逃しの確認', 'KA-23 検査方法(JLAC10)がコード表と不一致（既知を含む）', cnt('BC', 0), T('KA-23', True), '前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-25 Excelエラー値', f'=SUMPRODUCT(--ISNUMBER(MATCH(\'今回\'!$A$2:$AF${N}&"",{{"#N/A","#NAME?","#REF!","#VALUE!","#DIV/0!","#NUM!","#NULL!"}},0)))', T('KA-25', True), ''),
    ('C 見逃しの確認', 'KA-27 区分名称が空', f'=SUMPRODUCT(--(TRIM({R("A")}&"")=""))', T('KA-27', True), ''),
    ('C 見逃しの確認', 'KA-28 大項目が空', f'=SUMPRODUCT(--(TRIM({R("F")}&"")=""))', T('KA-28', True), ''),
    ('C 見逃しの確認', 'KA-29 FHIR項目名称が空', f'=SUMPRODUCT(--(TRIM({R("G")}&"")=""))', T('KA-29', True), ''),
    ('C 見逃しの確認', 'KA-30 FHIR識別文字列が空', f'=SUMPRODUCT(--(TRIM({R("H")}&"")=""))', T('KA-30', True), ''),
    ('C 見逃しの確認', 'KA-31 JLAC11コードの重複（行数）', f'=COUNTIF({R("AU")},">1")', T('KA-31', True), 'メイン＋ABOを合算'),
    ('C 見逃しの確認', 'KA-32 同じJLAC10で識別文字列が違う', cnt('AZ', 1), T('KA-32', True), ''),
    ('C 見逃しの確認', 'KA-K ケルビン記号（U+212A）', f'=SUMPRODUCT(--ISNUMBER(FIND(_xlfn.UNICHAR(8490),\'今回\'!$A$2:$AF${N})))', T('KA-K', True), ''),
    ('C 見逃しの確認', 'KA-33 000コードに相方が無い', cnt('AT', 0), T('KA-33', True), ''),
    ('C 見逃しの確認', 'KA-34 終了した測定法の000版が終了している', cnt('BD', 1), T('KA-34', True), ''),
    ('C 見逃しの確認', 'KA-M1 新規で開始日0または終了日≠99999999', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},0)+COUNTIFS({R("AJ")},"新規",{R("AF")},"<>99999999")', T('KA-M1', True), ''),
    ('C 見逃しの確認', 'KA-M3 終了日が変わり99999999になった', f'=COUNTIFS({R("AN")},1,{R("AF")},99999999)', T('KA-M3', True), ''),
    ('C 見逃しの確認', f'Step2 新規で適用開始日≠{START}', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},"<>{START}")', T('STEP2-START', True), ''),
    ('C 見逃しの確認', 'KA-26 33列目以降の値（連絡用の列を除く）', '目視', '—', '原本シートの連絡用の列（見出しに「コメント」）より右が空であることを目視で確認'),
    ('C 見逃しの確認', 'KA-DUMMY ダミーコード', f'=SUMPRODUCT(--(ISNUMBER(SEARCH("XXX",{R("M")}))+ISNUMBER(SEARCH("ZZZ",{R("M")}))+ISNUMBER(SEARCH("999999",{R("M")}))+ISNUMBER(SEARCH("仮",{R("M")}))>0))', T('KA-DUMMY', True), '目安（判定条件が未定義）'),
]
for i, (k, name, f, t, memo) in enumerate(items, start=HR + 1):
    chk.cell(row=i, column=1, value=k)
    chk.cell(row=i, column=2, value=name)
    chk.cell(row=i, column=3, value=f)
    chk.cell(row=i, column=4, value=t)
    chk.cell(row=i, column=5, value=f'=IF(D{i}="—","—",IF(C{i}=D{i},"一致","不一致"))')
    chk.cell(row=i, column=6, value=memo)
last = HR + len(items)
chk.cell(row=last + 2, column=2, value='一致しなかった項目の数').font = bold
chk.cell(row=last + 2, column=3, value=f'=COUNTIF(E{HR + 1}:E{last},"不一致")').font = bold
for col, w in zip('ABCDEF', (16, 52, 16, 14, 9, 60)):
    chk.column_dimensions[col].width = w
chk.freeze_panes = f'A{HR + 1}'
chk.auto_filter.ref = f'A{HR}:F{last}'
wb.save(OUT)
print('saved', OUT, 'rows', N, 'items', len(items))
