"""確認結果のダブルチェック用Excelを作る。元データを値で貼り込み、判定はすべてExcel関数で計算する（ツールのコードとは独立した検算）。
使い方: python 00_共通/03_scripts/make_doublecheck.py <ACN確認用xlsx> <出力xlsx> [<YYYYMM提出分フォルダ>]（省略時は最新の月フォルダ。入力ファイル名・新規開始行・公開日は下の定数を毎月差し替え）
"""
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

sys.stdout.reconfigure(encoding='utf-8')
BASE = Path(__file__).resolve().parent.parent  # 00_共通
ROOT = BASE.parent  # JLACマスタ月次チェック
MONTH = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else sorted(d for d in ROOT.iterdir() if d.is_dir() and d.name[:6].isdigit() and d.name.endswith('提出分'))[-1]
IN = MONTH / '01_input'
NEW = IN / '［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_202609_20260915からの更新_20261001.xlsx'
FIX = IN / '20260915_JLACセンターからの回答受領後_9月版公開用マスタ準備.xlsx'
J10 = IN / '140jlac10_1.xlsx'
J11 = IN / 'jlac11_1_1.0.xlsx'
LIST11 = ROOT.parent / '20260930_JLACマスター新旧対応表確認' / '01_input' / 'jlac11_3_1.1b.xlsx'
NEW_SHEETS = ['1001提出_共有項目JLACコード', '免疫血液学的検査ABO、Rh']
NEWROW, START = 8367, 20261015
ACN, OUT = Path(sys.argv[1]), Path(sys.argv[2])


def rows(path, sheet, min_row=1, ncol=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = [list(r if ncol is None else (list(r) + [None] * ncol)[:ncol]) for r in wb[sheet].iter_rows(min_row=min_row, values_only=True)]
    wb.close()
    return out


def pad(v, n):
    s = '' if v is None else str(v).strip()
    return s.zfill(n) if s.isdigit() else s


wb = openpyxl.Workbook()
chk = wb.active
chk.title = '照合'
bold = Font(bold=True)

# ---- 今回（メイン＋ABO） ----
cur = wb.create_sheet('今回')
hdr = rows(NEW, NEW_SHEETS[0], 1, 32)[0]
cur.append(hdr + ['シート', '原本行'])
for sh in NEW_SHEETS:
    for i, r in enumerate(rows(NEW, sh, 1, 32)[1:], start=2):
        if any(v not in (None, '') for v in r):
            cur.append(r + [sh, i])
N = cur.max_row  # 最終行
R = lambda c: f"'今回'!${c}$2:${c}${N}"

# ---- 前月 ----
prev = wb.create_sheet('前月_新のみ')
pr = rows(FIX, '202609_JLACマスタ(新のみ)', 1, 32)
for r in pr:
    if any(v not in (None, '') for v in r):
        prev.append(r)
PN = prev.max_row
mixed = wb.create_sheet('前月_新旧混合')
mixed.append(['JLAC11コード'])
for r in rows(FIX, '202609_JLACマスタ(新旧混合)公開用', 2, 32):
    if any(v not in (None, '') for v in r):
        mixed.append([r[12] if r[12] not in (None, '') else None])
MN = mixed.max_row

# ---- コード表（値のまま。コードの桁そろえのみ） ----
s = wb.create_sheet('JLAC11材料'); s.append(['材料コード', '材料名'])
for r in rows(J11, '材料コード', 2, 8):
    if r[2] not in (None, ''): s.append([pad(r[2], 3), r[3]])
s = wb.create_sheet('JLAC11単位'); s.append(['結果単位コード', '結果単位名(1)', '結果単位(2)'])
for r in rows(J11, '結果単位コード', 2, 6):
    if r[0] not in (None, ''): s.append([str(r[0]).strip(), r[1], r[2]])
s = wb.create_sheet('JLAC11_17桁'); s.append(['キー（測定物|識別|測定法）', '測定法名称'])
for r in rows(LIST11, 'JLAC11_17桁コードリスト', 2, 44):
    if r[3] not in (None, ''): s.append([f'{str(r[3]).strip()}|{pad(r[5], 4)}|{pad(r[9], 3)}', r[10]])
s = wb.create_sheet('JLAC10分析物'); s.append(['分析物コード', '分析物名'])
for r in rows(J10, '分析物コード', 6, 10):
    if r[3] not in (None, ''): s.append([str(r[3]).strip(), r[4]])
s = wb.create_sheet('JLAC10材料'); s.append(['材料コード', '材料名'])
for r in rows(J10, '材料コード', 10, 12):
    if r[6] not in (None, '') and str(r[6]).strip().isdigit(): s.append([pad(r[6], 3), r[7]])
s = wb.create_sheet('JLAC10測定法'); s.append(['測定法コード', '測定法名', '測定法名(2)'])
for r in rows(J10, '測定法コード', 6, 10):
    if r[3] not in (None, '') and str(r[3]).strip().isdigit(): s.append([pad(r[3], 3), r[4], r[5]])

# ---- ツールの結果（突き合わせ用） ----
tool_new, tool_ids, tool_rec = {}, {}, {}
awb = openpyxl.load_workbook(ACN, read_only=True, data_only=True)
for name in awb.sheetnames:
    if name.startswith('JLACセンター向けコメント_確認用'):
        src = NEW_SHEETS[0] if name == 'JLACセンター向けコメント_確認用' else NEW_SHEETS[1]
        for i, r in enumerate(awb[name].iter_rows(min_row=2, values_only=True), start=2):
            tool_new[(src, i)] = r[33]
            tool_ids[(src, i)] = r[34] or ''
            tool_rec[(src, i)] = r[35] or ''  # 記録のみ（既知・共有対象外など）
tsum = {}
# ツールv5以降、手順書由来の項目は「その他の指摘事項」シートに表示IDで出るため内部IDへ戻す
ALIAS = {'その他1（削除行）': 'STEP1-DEL', 'その他2（項目の変更）': 'STEP1-DIFF', 'その他3（表記の変更）': 'STEP1-DIFF表記',
         'その他4（適用終了日）': 'STEP1-END', 'その他5（新規行の適用開始日）': 'STEP2-START'}
for _sh in ('チェック結果一覧', 'その他の指摘事項'):
    if _sh not in awb.sheetnames:
        continue
    _rows = list(awb[_sh].iter_rows(values_only=True))
    _h = list(_rows[0])
    _ci, _ki = _h.index('コメント対象 行数'), _h.index('既知（コメントなし）件数')
    for r in _rows[1:]:
        tsum[ALIAS.get(r[0], r[0])] = (int(r[_ci] or 0), int(r[_ki] or 0))
import re
notice = next((str(r[1]) for r in awb['送付前確認事項'].iter_rows(values_only=True) if r[1] and 'JLAC11空欄' in str(r[1])), '')
m = re.search(r'JLAC11コードあり(\d+)件・JLAC11空欄(\d+)行', notice)
tool_old, tool_blank = (int(m.group(1)), int(m.group(2))) if m else ('（表示なし）', '（表示なし）')
awb.close()

# ---- 今回シートの補助列（すべてExcel関数） ----
def norm(x):
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(ASC({x})," ",""),"‐","-"),"−","-"),"―","-"),"–","-")'

helpers = [
    ('AI', '前月行', lambda r: f"=IFERROR(MATCH($M{r},'前月_新のみ'!$M$1:$M${PN},0),\"\")"),
    ('AJ', '新規判定', lambda r: f'=IF(AI{r}="","新規","既存")'),
    ('AK', 'KA-18 開始日が0', lambda r: f'=IF($AE{r}&""="0",1,0)'),
    ('AL', '前月と違う列数（開始・終了日以外）', lambda r: f"=IF(AI{r}=\"\",\"\",SUMPRODUCT(--(($A{r}:$AD{r}&\"\")<>(INDEX('前月_新のみ'!$A$1:$AD${PN},AI{r},0)&\"\"))))"),
    ('AM', 'KA-M2 開始日変更', lambda r: f"=IF(AI{r}=\"\",\"\",IF(VALUE($AE{r}&\"\")<>VALUE(INDEX('前月_新のみ'!$AE$1:$AE${PN},AI{r})&\"\"),1,0))"),
    ('AN', '終了日変更', lambda r: f"=IF(AI{r}=\"\",\"\",IF(VALUE($AF{r}&\"\")<>VALUE(INDEX('前月_新のみ'!$AF$1:$AF${PN},AI{r})&\"\"),1,0))"),
    ('AO', 'JLAC11材料名（表）', lambda r: f"=IFERROR(VLOOKUP(MID($M{r},10,3),'JLAC11材料'!$A:$B,2,FALSE)&\"\",\"#コード無し\")"),
    ('AP', 'KA-07 材料一致', lambda r: f'=IF(EXACT(TRIM($K{r}),TRIM(SUBSTITUTE(AO{r},"　",""))),1,0)'),
    ('AQ', 'JLAC11測定法名（17桁リスト）', lambda r: f"=IFERROR(VLOOKUP(LEFT($M{r},5)&\"|\"&MID($M{r},6,4)&\"|\"&MID($M{r},13,3),'JLAC11_17桁'!$A:$B,2,FALSE)&\"\",\"\")"),
    ('AR', 'KA-08 測定法一致', lambda r: f'=IF(OR(AQ{r}="",MID($M{r},13,3)="000"),"",IF({norm(f"$L{r}")}={norm(f"AQ{r}")},1,0))'),
    ('AS', 'KA-06 単位一致', lambda r: f"=IFERROR(IF(OR(EXACT($N{r},TRIM(VLOOKUP(MID($M{r},16,2),'JLAC11単位'!$A:$C,2,FALSE)&\"\")),ISNUMBER(FIND(\",\"&$N{r}&\",\",\",\"&SUBSTITUTE(SUBSTITUTE(VLOOKUP(MID($M{r},16,2),'JLAC11単位'!$A:$C,3,FALSE)&\"\",\" \",\"\"),\",,\",\",\")&\",\"))),1,0),0)"),
    ('AT', 'KA-33 000の相方の数', lambda r: f'=IF(MID($M{r},13,3)="000",COUNTIFS({R("M")},LEFT($M{r},12)&"???"&RIGHT($M{r},2))-1,"")'),
    ('AU', 'KA-31 同じJLAC11の数', lambda r: f'=COUNTIF({R("M")},$M{r})'),
    ('AV', 'KA-10 不整合', lambda r: f'=IF(TRIM($F{r})="","",IF(COUNTIFS({R("F")},$F{r},{R("BL")},"<>"&$BL{r})>0,1,0))'),
    ('AW', 'KA-11 不整合', lambda r: f'=IF(TRIM($H{r})="","",IF(COUNTIFS({R("H")},$H{r},{R("BM")},"<>"&$BM{r})>0,1,0))'),
    ('AX', 'KA-12 不整合', lambda r: f'=IF(TRIM($I{r})="","",IF(COUNTIFS({R("I")},$I{r},{R("BN")},"<>"&$BN{r})>0,1,0))'),
    ('AY', 'KA-17 不整合', lambda r: f'=IF(TRIM($G{r})="","",IF(COUNTIFS({R("G")},$G{r},{R("AD")},"<>"&$AD{r})>0,1,0))'),
    ('AZ', 'KA-32 不整合', lambda r: f'=IF(TRIM($T{r})="","",IF(COUNTIFS({R("T")},$T{r},{R("H")},"<>"&$H{r})>0,1,0))'),
    ('BA', 'KA-21 分析物コードあり', lambda r: f"=IF(COUNTIF('JLAC10分析物'!$A:$A,LEFT($T{r},5))>0,1,0)"),
    ('BB', 'KA-22 材料(JLAC10)一致', lambda r: f"=IF(EXACT(TRIM($R{r}),TRIM(SUBSTITUTE(IFERROR(VLOOKUP(MID($T{r},10,3),'JLAC10材料'!$A:$B,2,FALSE)&\"\",\"#\"),\"　\",\"\"))),1,0)"),
    ('BC', 'KA-23 測定法(JLAC10)一致', lambda r: (lambda v1, v2, s: f'=IFERROR(IF(OR(EXACT({s},{v1}),EXACT({s},{v1}&{v2}),EXACT({s},{v1}&"("&{v2}&")"),EXACT({s},{v1}&"（"&{v2}&"）"),EXACT({s},{v1}&"["&{v2}&"]"),EXACT({s},{v1}&" "&{v2}),AND({v2}<>"",EXACT({s},{v2}))),1,0),0)')(
        f"TRIM(SUBSTITUTE(VLOOKUP(MID($T{r},13,3),'JLAC10測定法'!$A:$C,2,FALSE)&\"\",\"　\",\"\"))",
        f"TRIM(SUBSTITUTE(VLOOKUP(MID($T{r},13,3),'JLAC10測定法'!$A:$C,3,FALSE)&\"\",\"　\",\"\"))", f'TRIM($S{r})')),
    ('BD', 'KA-34 不整合', lambda r: f'=IF(AND(MID($M{r},13,3)="000",$AF{r}&""<>"99999999"),IF(COUNTIFS({R("M")},LEFT($M{r},12)&"???"&RIGHT($M{r},2),{R("AF")},"<>99999999")-1>0,1,0),"")'),
    ('BE', 'ツール：新規/既存', None),
    ('BF', 'ツール：コメント対象ID', None),
    ('BG', '新規判定の一致', lambda r: f'=IF(AJ{r}=BE{r},1,0)'),
    ('BH', 'KA-18の一致', lambda r: f'=IF(AK{r}=IF(ISNUMBER(SEARCH("KA-18",BF{r}&" "&BP{r})),1,0),1,0)'),
    ('BI', '前月差分の一致', lambda r: f'=IF(BO{r}=IF(ISNUMBER(SEARCH("STEP1",BF{r}&" "&BP{r}))+ISNUMBER(SEARCH("KA-M",BF{r}&" "&BP{r}))+ISNUMBER(SEARCH("その他",BF{r}&" "&BP{r}))>0,1,0),1,0)'),
    ('BJ', 'KA-07の一致', lambda r: f'=IF((AP{r}=0)=ISNUMBER(SEARCH("KA-07",BF{r}&" "&BP{r})),1,0)'),
    ('BK', 'KA-08の一致', lambda r: f'=IF(AND(AJ{r}="新規",AR{r}=0)=ISNUMBER(SEARCH("KA-08",BF{r})),1,0)'),
    ('BL', 'キー：FHIR項目名称|識別|略称', lambda r: f'=$G{r}&"|"&$H{r}&"|"&$I{r}'),
    ('BM', 'キー：大項目|FHIR項目名称|略称', lambda r: f'=$F{r}&"|"&$G{r}&"|"&$I{r}'),
    ('BN', 'キー：大項目|FHIR項目名称|識別', lambda r: f'=$F{r}&"|"&$G{r}&"|"&$H{r}'),
    ('BO', '前月と差分あり', lambda r: f'=IF(AI{r}="",0,IF(AL{r}+AM{r}+AN{r}>0,1,0))'),
    ('BP', 'ツール：記録のみのチェックID（既知・共有対象外など）', None),
]
for col, title, f in helpers:
    cur[f'{col}1'] = title
    cur[f'{col}1'].font = bold
for r in range(2, N + 1):
    key = (cur[f'AG{r}'].value, cur[f'AH{r}'].value)
    for col, title, f in helpers:
        if col == 'BE': cur[f'BE{r}'] = tool_new.get(key, '')
        elif col == 'BF': cur[f'BF{r}'] = tool_ids.get(key, '')
        elif col == 'BP': cur[f'BP{r}'] = tool_rec.get(key, '')
        else: cur[f'{col}{r}'] = f(r)
for c in range(1, 69):
    cur.cell(row=1, column=c).font = bold
cur.freeze_panes = 'N2'
cur.auto_filter.ref = f'A1:BP{N}'

# ---- 照合シート ----
chk['A1'] = '確認結果のダブルチェック（JLACセンター 2026-10-01提出版）'; chk['A1'].font = Font(bold=True, size=14)
chk['A2'] = '元データ（今回提出・前月FIX・コード表）を値のまま貼り込み、判定はすべてExcel関数で計算した（ツールのコードは使っていない）。D列はツールの結果。E列がすべて「一致」なら、ツールの確認結果はExcelでの独立した検算と同じ。'
chk['A3'] = '不一致がある場合：「今回」シートの補助列（AI〜BO）でフィルタして該当行を確認する。行ごとの一致列（BG〜BK）が0の行が食い違い。'
chk['A4'] = '目視確認（推奨）：指摘のある種類ごとに「今回」シートで2〜3行を選び、原本の値・コード表を自分の目で確認する。'
for c in ('A2', 'A3', 'A4'):
    chk[c].alignment = Alignment(wrap_text=True, vertical='top')
    chk.merge_cells(f'{c}:F{c[1:]}')
    chk.row_dimensions[int(c[1:])].height = 36
head = ['区分', '確認項目', 'Excelでの再計算', 'ツールの結果', '判定', '計算方法・メモ']
HR = 6
for i, h in enumerate(head, start=1):
    chk.cell(row=HR, column=i, value=h).font = bold
    chk.cell(row=HR, column=i).fill = PatternFill('solid', fgColor='DDEBF7')
T = lambda k, known=False: tsum[k][0] + (tsum[k][1] if known else 0)
cnt = lambda c, crit: f'=COUNTIF({R(c)},{crit})'
items = [
    ('A 件数の突き合わせ', '前月FIX（新のみ）の行数', f"=COUNTA('前月_新のみ'!$M$2:$M${PN})", '—', ''),
    ('A 件数の突き合わせ', '今回提出の行数（メイン＋ABO）', f'=COUNTA({R("M")})', '—', 'メインとABOシートの合計'),
    ('A 件数の突き合わせ', '新規行の数（前月FIXに同じJLAC11コードが無い行）', cnt('AJ', '"新規"'), 75, ''),
    ('A 件数の突き合わせ', '整合式：前月 −（今回 − 新規）', f'=C7-(C8-C9)', 0, '0なら、行の抜け・二重がない'),
    ('A 件数の突き合わせ', '前月FIX（新のみ）にあって今回に無いコード', f"=SUMPRODUCT(--(COUNTIF({R('M')},'前月_新のみ'!$M$2:$M${PN})=0))", 0, '行の削除'),
    ('A 件数の突き合わせ', '前月の新旧混合にあって、新のみにも今回にも無いJLAC11コード（旧コード）', f"=SUMPRODUCT((('前月_新旧混合'!$A$2:$A${MN}&\"\")<>\"\")*(COUNTIF('前月_新のみ'!$M$2:$M${PN},'前月_新旧混合'!$A$2:$A${MN}&\"\")=0)*(COUNTIF({R('M')},'前月_新旧混合'!$A$2:$A${MN}&\"\")=0)/COUNTIF('前月_新旧混合'!$A$2:$A${MN},'前月_新旧混合'!$A$2:$A${MN}&\"\"))", tool_old, '送付前確認事項（旧コードの抜け）'),
    ('A 件数の突き合わせ', '前月の新旧混合でJLAC11が空欄の行（今回の提出には無い）', f"=SUMPRODUCT(--(('前月_新旧混合'!$A$2:$A${MN}&\"\")=\"\"))", tool_blank, 'JLAC10のみの見直し前行'),
    ('A 件数の突き合わせ', '最初の新規行の原本行番号', f"=INDEX({R('AH')},MATCH(\"新規\",{R('AJ')},0))", NEWROW, 'JLACセンター連絡：8367行目以降が新規'),
    ('A 件数の突き合わせ', f'{NEWROW}行目以降（メイン）の行数', f'=COUNTIFS({R("AG")},"1001*",{R("AH")},">={NEWROW}")', 75, ''),
    ('A 件数の突き合わせ', f'{NEWROW}行目以降（メイン）で新規の行数', f'=COUNTIFS({R("AG")},"1001*",{R("AH")},">={NEWROW}",{R("AJ")},"新規")', 75, ''),
    ('A 件数の突き合わせ', f'新規行のうち適用開始日={START}', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},{START})', 75, 'JLACセンター連絡：適用開始日20261015'),
    ('B 指摘の検算', 'KA-18 適用開始日が0の行', f'=SUM({R("AK")})', T('KA-18', True), '共有対象外として記録のみの行を含む'),
    ('B 指摘の検算', '前月と値が違う行（STEP1・KA-M）', f'=SUM({R("BO")})', T('STEP1-DIFF', True) + T('STEP1-DIFF表記', True) + T('STEP1-END', True) + T('KA-M2', True) + T('KA-M3', True), '表記だけの変更42行＋ABO 2行（共有対象外として記録のみ）'),
    ('B 指摘の検算', '　うち適用開始日の変更（KA-M2）', f'=SUM({R("AM")})', T('KA-M2'), ''),
    ('B 指摘の検算', '　うち適用終了日の変更', f'=SUM({R("AN")})', T('STEP1-END') + T('KA-M3'), ''),
    ('B 指摘の検算', 'KA-07 材料(JLAC11)がコード表と不一致', cnt('AP', 0), T('KA-07', True), ''),
    ('B 指摘の検算', 'KA-08 測定法(JLAC11)が17桁リストと不一致（新規行）', f'=COUNTIFS({R("AJ")},"新規",{R("AR")},0)', T('KA-08'), 'ASCで全角→半角、ハイフン類・空白をそろえて比較'),
    ('B 指摘の検算', '行ごとの不一致：新規判定', cnt('BG', 0), 0, '0なら全行でツールと一致（ツール側はコメント対象と記録のみの両方のIDで照合）'),
    ('B 指摘の検算', '行ごとの不一致：KA-18', cnt('BH', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：前月差分', cnt('BI', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：KA-07', cnt('BJ', 0), 0, ''),
    ('B 指摘の検算', '行ごとの不一致：KA-08', cnt('BK', 0), 0, ''),
    ('C 見逃しの確認', 'KA-01 予備区分名称・基準値フラグに値がある', f'=SUMPRODUCT(--(LEN({R("B")}&{R("U")}&{R("V")}&{R("W")})>0))', T('KA-01', True), ''),
    ('C 見逃しの確認', 'KA-02 JLAC11コードが17桁でない', f'=SUMPRODUCT(--(LEN({R("M")}&"")<>17))', T('KA-02', True), ''),
    ('C 見逃しの確認', 'KA-03 救急・生活習慣病フラグが0/1以外', f'=SUMPRODUCT(--((({R("C")}&"")<>"0")*(({R("C")}&"")<>"1")))+SUMPRODUCT(--((({R("D")}&"")<>"0")*(({R("D")}&"")<>"1")))', T('KA-03', True), ''),
    ('C 見逃しの確認', 'KA-04 データ区分が1/2以外', f'=SUMPRODUCT(--((({R("E")}&"")<>"1")*(({R("E")}&"")<>"2")))', T('KA-04', True), ''),
    ('C 見逃しの確認', 'KA-06 表示用単位が結果単位コードと不一致', cnt('AS', 0), T('KA-06', True), ''),
    ('C 見逃しの確認', 'KA-10 大項目キーの不整合（既知を含む）', cnt('AV', 1), T('KA-10', True), '空腹時/随時など設計上の区分で、前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-11 FHIR識別文字列キーの不整合', cnt('AW', 1), T('KA-11', True), ''),
    ('C 見逃しの確認', 'KA-12 略称キーの不整合（既知を含む）', cnt('AX', 1), T('KA-12', True), '前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-13 データタイプがPQ/CD以外', f'=SUMPRODUCT(--(({R("X")}<>"PQ")*({R("X")}<>"CD")))', T('KA-13', True), ''),
    ('C 見逃しの確認', 'KA-14 CDで値の下限・上限が空', f'=SUMPRODUCT(({R("X")}="CD")*(((TRIM({R("Y")}&"")="")+(TRIM({R("Z")}&"")=""))>0))', T('KA-14', True), ''),
    ('C 見逃しの確認', 'KA-15 CDで値リストが空', f'=SUMPRODUCT(({R("X")}="CD")*(TRIM({R("AB")}&"")=""))', T('KA-15', True), ''),
    ('C 見逃しの確認', 'KA-16 CDでOIDが空', f'=SUMPRODUCT(({R("X")}="CD")*(TRIM({R("AC")}&"")=""))', T('KA-16', True), ''),
    ('C 見逃しの確認', 'KA-17 同じFHIR項目名称で並び順が違う', cnt('AY', 1), T('KA-17', True), ''),
    ('C 見逃しの確認', 'KA-19 適用終了日が8桁の99999999/実在日付でない', f'=SUMPRODUCT(--(LEN({R("AF")}&"")<>8))+SUMPRODUCT((({R("AF")}&"")<>"99999999")*ISERROR(DATEVALUE(TEXT({R("AF")},"0000-00-00"))))', T('KA-19', True), ''),
    ('C 見逃しの確認', 'KA-20 JLAC10コードが17桁でない', f'=SUMPRODUCT(--(LEN({R("T")}&"")<>17))', T('KA-20', True), ''),
    ('C 見逃しの確認', 'KA-21 分析物コードがJLAC10に無い', cnt('BA', 0), T('KA-21', True), ''),
    ('C 見逃しの確認', 'KA-22 材料(JLAC10)がコード表と不一致', cnt('BB', 0), T('KA-22', True), ''),
    ('C 見逃しの確認', 'KA-23 検査方法(JLAC10)がコード表と不一致（既知を含む）', cnt('BC', 0), T('KA-23', True), '前月FIXでも同じ（既知）'),
    ('C 見逃しの確認', 'KA-25 Excelエラー値', f'=SUMPRODUCT(--ISNUMBER(MATCH(\'今回\'!$A$2:$AF${N}&"",{{"#N/A","#NAME?","#REF!","#VALUE!","#DIV/0!","#NUM!","#NULL!"}},0)))', T('KA-25', True), ''),
    ('C 見逃しの確認', 'KA-27 区分名称が空', f'=SUMPRODUCT(--(TRIM({R("A")}&"")=""))', T('KA-27', True), ''),
    ('C 見逃しの確認', 'KA-28 大項目が空', f'=SUMPRODUCT(--(TRIM({R("F")}&"")=""))', T('KA-28', True), ''),
    ('C 見逃しの確認', 'KA-29 FHIR項目名称が空', f'=SUMPRODUCT(--(TRIM({R("G")}&"")=""))', T('KA-29', True), ''),
    ('C 見逃しの確認', 'KA-30 FHIR識別文字列が空', f'=SUMPRODUCT(--(TRIM({R("H")}&"")=""))', T('KA-30', True), ''),
    ('C 見逃しの確認', 'KA-31 JLAC11コードの重複（行数）', f'=COUNTIF({R("AU")},">1")', T('KA-31', True), 'メイン＋ABOを合算'),
    ('C 見逃しの確認', 'KA-32 同じJLAC10で識別文字列が違う', cnt('AZ', 1), T('KA-32', True), ''),
    ('C 見逃しの確認', 'KA-K ケルビン記号（U+212A）', f'=SUMPRODUCT(--ISNUMBER(FIND(_xlfn.UNICHAR(8490),\'今回\'!$A$2:$AF${N})))', T('KA-K', True), ''),
    ('C 見逃しの確認', 'KA-33 000コードに相方が無い', cnt('AT', 0), T('KA-33', True), ''),
    ('C 見逃しの確認', 'KA-34 終了した測定法の000版が終了している', cnt('BD', 1), T('KA-34', True), ''),
    ('C 見逃しの確認', 'KA-M1 新規で開始日0または終了日≠99999999', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},0)+COUNTIFS({R("AJ")},"新規",{R("AF")},"<>99999999")', T('KA-M1', True), ''),
    ('C 見逃しの確認', 'KA-M3 終了日が変わり99999999になった', f'=COUNTIFS({R("AN")},1,{R("AF")},99999999)', T('KA-M3', True), ''),
    ('C 見逃しの確認', f'Step2 新規で適用開始日≠{START}', f'=COUNTIFS({R("AJ")},"新規",{R("AE")},"<>{START}")', T('STEP2-START', True), ''),
    ('C 見逃しの確認', 'KA-26 33列目以降の値', '目視', '—', '原本シートのAG列以降が空であることを目視で確認'),
    ('C 見逃しの確認', 'KA-DUMMY ダミーコード', f'=SUMPRODUCT(--(ISNUMBER(SEARCH("XXX",{R("M")}))+ISNUMBER(SEARCH("ZZZ",{R("M")}))+ISNUMBER(SEARCH("999999",{R("M")}))+ISNUMBER(SEARCH("仮",{R("M")}))>0))', T('KA-DUMMY', True), '目安（判定条件が未定義）'),
]
for i, (k, name, f, t, memo) in enumerate(items, start=HR + 1):
    chk.cell(row=i, column=1, value=k)
    chk.cell(row=i, column=2, value=name)
    chk.cell(row=i, column=3, value=f)
    chk.cell(row=i, column=4, value=t)
    chk.cell(row=i, column=5, value=f'=IF(D{i}="—","—",IF(C{i}=D{i},"一致","不一致"))')
    chk.cell(row=i, column=6, value=memo)
last = HR + len(items)
chk.cell(row=last + 2, column=2, value='一致しなかった項目の数').font = bold
chk.cell(row=last + 2, column=3, value=f'=COUNTIF(E{HR + 1}:E{last},"不一致")').font = bold
for col, w in zip('ABCDEF', (16, 52, 16, 14, 9, 60)):
    chk.column_dimensions[col].width = w
chk.freeze_panes = f'A{HR + 1}'
chk.auto_filter.ref = f'A{HR}:F{last}'
wb.save(OUT)
print('saved', OUT, 'rows', N, 'items', len(items))
