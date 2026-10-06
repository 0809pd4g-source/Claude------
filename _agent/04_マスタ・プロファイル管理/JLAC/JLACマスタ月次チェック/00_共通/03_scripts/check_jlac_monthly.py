"""【凍結・参照実装】JLACマスタ月次チェック（KA-01〜KA-M3）のPython版。今後は更新しない。
ロジックの正本は 02_output の最新版チェックツールHTML（jlac-core）。実行は run_check.mjs を使う。
この版で作ったExcelはopenpyxlの保存で外部データ接続が欠落し、Excelで修復を求められる。

コメント付与の方針（定常業務マニュアル JLACマスタチェック手順 Step0〜2 に準拠）:
  - 新規行（前月FIXに無いJLAC11コード）: 全チェック結果をコメント
  - 既存行: 前月FIXとの差分、および前月FIXでは通っていたチェックの失敗（再発）のみコメント
  - 前月FIXでも同じ失敗だったもの（既知）はコメントせず件数のみ集計
"""
import sys, re, json, unicodedata, datetime
from collections import defaultdict, Counter
from pathlib import Path
from copy import copy
import openpyxl
from openpyxl.styles import Alignment

sys.stdout.reconfigure(encoding='utf-8')
BASE = Path(__file__).resolve().parent.parent
IN = BASE / '01_input'

# ===== 設定（毎月差し替え） =====
NEW_FILE = IN / '［令和8年9月末頃以降本番環境適用開始］共有項目JLACコードマスタ_202609_20260915からの更新_20261001.xlsx'
NEW_SHEETS = ['1001提出_共有項目JLACコード', '免疫血液学的検査ABO、Rh']
PREV_FILE = IN / '20260915_JLACセンターからの回答受領後_9月版公開用マスタ準備.xlsx'
PREV_SHEET_NEWONLY = '202609_JLACマスタ(新のみ)'
PREV_SHEET_MIXED = '202609_JLACマスタ(新旧混合)公開用'
PREV_LABEL = '202609版'
FORM_FILE = IN / '【電カル共有】JLACマスタ_JLACセンター20260901提出版_確認事項一覧_20260909.xlsx'
FORM_SHEET = 'JLACセンター向けコメント_確認用'
FORM_PREV_SUBMIT_SHEET = '_令和8年9月末頃以降本番環境適用開始_共有項目JLACコード'
J10_FILE = IN / '140jlac10_1.xlsx'
J11_FILE = IN / 'jlac11_1_1.0.xlsx'
J11_LIST_FILE = BASE.parent / '20260930_JLACマスター新旧対応表確認' / '01_input' / 'jlac11_3_1.1b.xlsx'
EXPECTED_START = '20261015'
EXPECTED_END = '20261014'
OUT_XLSX = BASE / '02_output' / '20261003_JLACマスタ確認事項一覧_JLACセンター20261001提出版_v1.xlsx'
OUT_JSON = Path(sys.argv[1]) if len(sys.argv) > 1 else None
# ================================

NCOL = 32
COMMENT_HDR = 'JLACセンター向けコメント'
(KUBUN, YOBI, KYUKYU, SEIKATSU, DATAKUBUN, DAIKOMOKU, FHIRNAME, FHIRID, RYAKU, HANBAI,
 MAT11, METH11, JLAC11, UNIT, UNIT2, XUNIT, XUNIT2, MAT10, METH10, JLAC10,
 FLG_L, FLG_H, FLG_J, DTYPE, VMIN, VMAX, NUMFMT, CODELIST, OID, ORDER, START, END) = range(NCOL)
EXCEL_ERRORS = ('#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A', '#SPILL!', '#CALC!', '#GETTING_DATA')


def sv(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


HYPHENS = re.compile('[‐-―−－﹣]')


def nz(x):
    """全角半角・ハイフン類・空白の差を吸収した比較用文字列"""
    return HYPHENS.sub('-', unicodedata.normalize('NFKC', x)).replace(' ', '').replace('　', '')


def date_norm(x):
    return x.zfill(8) if x.isdigit() else x


def code_pad(v, n):
    s = sv(v).strip()
    return s.zfill(n) if s.isdigit() else s


class Row:
    def __init__(self, sheet, xrow, vals, extra):
        self.sheet, self.xrow, self.raw, self.extra = sheet, xrow, vals, extra
        self.s = [sv(v) for v in vals]
        self.code = self.s[JLAC11].strip()


def load(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    header, rows = None, []
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        r = list(r)
        if i == 1:
            header = [sv(v) for v in r[:NCOL]]
            continue
        if not any(v not in (None, '') for v in r):
            continue
        vals = (r + [None] * NCOL)[:NCOL]
        extra = [j for j, v in enumerate(r[NCOL:], start=NCOL + 1) if v not in (None, '')]
        rows.append(Row(sheet, i, vals, extra))
    wb.close()
    return header, rows


# ---------- コード表 ----------
def load_tables():
    t = {}
    wb = openpyxl.load_workbook(J11_FILE, read_only=True, data_only=True)
    t['mat11'] = {code_pad(r[2], 3): sv(r[3]).strip(' 　') for r in wb['材料コード'].iter_rows(min_row=2, values_only=True) if r[2] not in (None, '')}
    t['unit11'] = {}
    for r in wb['結果単位コード'].iter_rows(min_row=2, values_only=True):
        if r[0] in (None, ''):
            continue
        names = {sv(r[1]).strip()} | {x.strip() for x in sv(r[2]).split(',') if x.strip()}
        t['unit11'][sv(r[0]).strip()] = (sv(r[1]).strip(), names - {''})
    wb.close()

    wb = openpyxl.load_workbook(J10_FILE, read_only=True, data_only=True)
    t['ana10'] = {}
    for r in wb['分析物コード'].iter_rows(min_row=6, values_only=True):
        if r[3] not in (None, ''):
            t['ana10'][sv(r[3]).strip()] = (sv(r[4]).strip(), sv(r[8]).strip())
    t['mat10'] = {}
    for r in wb['材料コード'].iter_rows(min_row=10, values_only=True):
        if len(r) > 7 and r[6] not in (None, '') and sv(r[6]).strip().isdigit():
            t['mat10'][code_pad(r[6], 3)] = sv(r[7]).strip(' 　')
    t['meth10'] = {}
    for r in wb['測定法コード'].iter_rows(min_row=6, values_only=True):
        if r[3] not in (None, '') and sv(r[3]).strip().isdigit():
            n1, n2 = sv(r[4]).strip(' 　'), sv(r[5]).strip(' 　')
            t['meth10'][code_pad(r[3], 3)] = (n1, n2)
    wb.close()

    t['list11'] = {}
    if J11_LIST_FILE.exists():
        wb = openpyxl.load_workbook(J11_LIST_FILE, read_only=True, data_only=True)
        for r in wb['JLAC11_17桁コードリスト'].iter_rows(min_row=2, values_only=True):
            if r[3] in (None, ''):
                continue
            key = (sv(r[3]).strip(), code_pad(r[5], 4), code_pad(r[9], 3))
            mats = {code_pad(r[7], 3)} | {code_pad(r[i], 3) for i in range(13, 29, 2) if r[i] not in (None, '')}
            units = {sv(r[11]).strip()} | {sv(r[i]).strip() for i in range(29, 39, 2) if r[i] not in (None, '')}
            d = t['list11'].setdefault(key, {'meth': sv(r[10]).strip(), 'ana': sv(r[4]).strip(), 'mats': set(), 'units': set()})
            d['mats'] |= mats
            d['units'] |= units
        wb.close()
    return t


# ---------- 行単位チェック ----------
def row_checks(r, H, T):
    """[(check_id, level, message)] level: strict=コメント対象 / ref=参考（報告のみ）"""
    s, out = r.s, []
    add = lambda cid, msg, level='strict': out.append((cid, level, msg))

    for c in (YOBI, FLG_L, FLG_H, FLG_J):
        if s[c] != '':
            add('KA-01', f'{H[c]}は空欄である必要があります。「{s[c]}」が設定されているため、空欄としてください')
    code = s[JLAC11]
    if code.strip() == '':
        add('KA-02', 'JLAC11コードが空欄です。設定をお願いします')
    elif len(code) != 17:
        add('KA-02', f'JLAC11コードは17桁である必要があります（現在{len(code)}桁：「{code}」）。ご確認ください')
    for c in (KYUKYU, SEIKATSU):
        if s[c] not in ('0', '1'):
            add('KA-03', f'{H[c]}は0または1である必要があります（現在「{s[c]}」）。ご確認ください')
    if s[DATAKUBUN] not in ('1', '2'):
        add('KA-04', f'データ区分は1または2である必要があります（現在「{s[DATAKUBUN]}」）。ご確認ください')

    if len(code) == 17:
        ana, ide, mat, meth, unit = code[0:5], code[5:9], code[9:12], code[12:15], code[15:17]
        # KA-06 表示用単位 × JLAC11結果単位コード
        u = T['unit11'].get(unit)
        if u is None:
            add('KA-06', f'JLAC11コードの結果単位コード「{unit}」がJLAC11結果単位コード表に存在しません。ご確認ください')
        elif s[UNIT] not in u[1]:
            add('KA-06', f'表示用単位「{s[UNIT]}」が、JLAC11結果単位コード「{unit}」の単位（{"／".join(sorted(u[1]))}）と一致しません（大文字・小文字を区別）。ご確認ください')
        # KA-07 材料(JLAC11) × JLAC11材料コード
        m = T['mat11'].get(mat)
        if m is None:
            add('KA-07', f'JLAC11コードの材料コード「{mat}」がJLAC11材料コード表に存在しません。ご確認ください')
        elif s[MAT11].strip() != m:
            if nz(s[MAT11]) == nz(m) and re.search('[（）]', s[MAT11]):
                add('KA-07', f'材料(JLAC11)について、他のレコードと合わせて、「()」は半角としてください。\n{s[MAT11]}→{m}')
            elif nz(s[MAT11]) == nz(m):
                add('KA-07', f'材料(JLAC11)について、JLAC11材料コード表の表記に合わせてください。\n{s[MAT11]}→{m}')
            else:
                add('KA-07', f'材料(JLAC11)「{s[MAT11]}」が、JLAC11材料コード「{mat}」の名称「{m}」と一致しません。ご確認ください')
        # KA-08 測定法(JLAC11) × JLAC11コードリスト（参考：1.1b版）
        if T['list11'] and meth != '000':
            e = T['list11'].get((ana, ide, meth))
            if e is None:
                add('KA-08', f'JLAC11コードリスト(1.1b)に測定物「{ana}」識別「{ide}」測定法「{meth}」の組合せが見当たらない', 'ref')
            else:
                if nz(s[METH11]) != nz(e['meth']):
                    add('KA-08', f'測定法(JLAC11)「{s[METH11]}」が、JLAC11コード表（17桁コードリスト）の測定法名称「{e["meth"]}」と表記が異なります。ご確認ください')
                if mat not in e['mats'] or unit not in e['units']:
                    add('KA-08', f'コードリスト(1.1b)の材料/単位の組合せに材料「{mat}」単位「{unit}」が無い', 'ref')

    if s[DTYPE] not in ('PQ', 'CD'):
        add('KA-13', f'データタイプはPQまたはCDである必要があります（現在「{s[DTYPE]}」）。ご確認ください')
    if s[DTYPE] == 'CD':
        if s[VMIN].strip() == '' or s[VMAX].strip() == '':
            add('KA-14', 'データタイプがCDの場合、値の下限・値の上限の設定が必要です。ご確認ください')
        if s[CODELIST].strip() == '':
            add('KA-15', 'データタイプがCDの場合、コード型の場合の値リストの設定が必要です。ご確認ください')
        if s[OID].strip() == '':
            add('KA-16', 'データタイプがCDの場合、コード型のOIDの設定が必要です。ご確認ください')

    st, en = s[START], s[END]
    if not valid_date8(st, '00000000'):
        if st == '0':
            add('KA-18', '適用開始日は8桁である必要があります。そのため、0→00000000に修正ください')
        else:
            add('KA-18', f'適用開始日は8桁で「00000000」または実在する日付である必要があります（現在「{st}」）。ご確認ください')
    if not valid_date8(en, '99999999'):
        add('KA-19', f'適用終了日は8桁で「99999999」または実在する日付である必要があります（現在「{en}」）。ご確認ください')

    c10 = s[JLAC10]
    if len(c10) != 17:
        add('KA-20', f'JLAC10コードは17桁である必要があります（現在{len(c10)}桁：「{c10}」）。ご確認ください')
    else:
        a10, m10, me10 = c10[0:5], c10[9:12], c10[12:15]
        if a10 not in T['ana10']:
            add('KA-21', f'JLAC10コードの分析物コード「{a10}」がJLAC10分析物コード表に存在しません。ご確認ください')
        mm = T['mat10'].get(m10)
        if mm is None:
            add('KA-22', f'JLAC10コードの材料コード「{m10}」がJLAC10材料コード表に存在しません。ご確認ください')
        elif s[MAT10].strip() != mm:
            add('KA-22', f'材料(JLAC10)「{s[MAT10]}」が、JLAC10コードの材料コード「{m10}」の名称「{mm}」と一致しません。ご確認ください')
        me = T['meth10'].get(me10)
        if me is None:
            add('KA-23', f'JLAC10コードの測定法コード「{me10}」がJLAC10測定法コード表に存在しません。ご確認ください')
        else:
            n1, n2 = me
            cands = {n1, n1 + n2, f'{n1}({n2})', f'{n1}（{n2}）', f'{n1}[{n2}]', f'{n1}［{n2}］', f'{n1} {n2}', n2} - {''}
            if s[METH10].strip() not in cands:
                add('KA-23', f'検査方法(JLAC10-測定法)「{s[METH10]}」が、JLAC10コードの測定法コード「{me10}」の名称「{n1}{("／" + n2) if n2 else ""}」と一致しません。ご確認ください')

    for c in range(NCOL):
        if s[c].strip() in EXCEL_ERRORS:
            add('KA-25', f'{H[c]}にExcelのエラー値「{s[c]}」が含まれています。ご確認ください')
    if r.extra:
        add('KA-26', f'33列目以降（{",".join(map(str, r.extra))}列目）に値が入力されています。32列以内としてください')
    for c, cid in ((KUBUN, 'KA-27'), (DAIKOMOKU, 'KA-28'), (FHIRNAME, 'KA-29'), (FHIRID, 'KA-30')):
        if s[c].strip() == '':
            add(cid, f'{H[c]}が空欄です。設定をお願いします')
    if code and (not re.fullmatch(r'[0-9A-Z]{17}', code) or re.search(r'仮|X{3,}|Z{3,}|9{6,}', code)):
        add('KA-DUMMY', f'JLAC11コード「{code}」はダミー（仮）コードの可能性があります。正式なコードかご確認ください')
    for c in range(NCOL):
        if 'K' in s[c]:
            add('KA-K', f'{H[c]}にケルビン記号「K」(U+212A)が使われています。半角英字「K」に修正ください')
    return out


def valid_date8(x, special):
    if len(x) != 8 or not x.isdigit():
        return False
    if x == special:
        return True
    try:
        datetime.datetime.strptime(x, '%Y%m%d')
        return True
    except ValueError:
        return False


# ---------- グループチェック ----------
def combos(rows, keyf, valf):
    g = defaultdict(lambda: defaultdict(list))
    for r in rows:
        k = keyf(r)
        if k.strip():
            g[k][valf(r)].append(r)
    return g


def group_findings(new_rows, base_rows):
    """[(row, check_id, message, known)]"""
    out = []
    specs = [
        ('KA-10', '大項目', lambda r: r.s[DAIKOMOKU], lambda r: (r.s[FHIRNAME], r.s[FHIRID], r.s[RYAKU]), 'FHIR項目名称/FHIR識別文字列/略称'),
        ('KA-11', 'FHIR識別文字列', lambda r: r.s[FHIRID], lambda r: (r.s[DAIKOMOKU], r.s[FHIRNAME], r.s[RYAKU]), '大項目/FHIR項目名称/略称'),
        ('KA-12', '略称', lambda r: r.s[RYAKU], lambda r: (r.s[DAIKOMOKU], r.s[FHIRNAME], r.s[FHIRID]), '大項目/FHIR項目名称/FHIR識別文字列'),
        ('KA-17', 'FHIR項目名称', lambda r: r.s[FHIRNAME], lambda r: (r.s[ORDER],), '並び順'),
        ('KA-32', 'JLAC10コード', lambda r: r.s[JLAC10], lambda r: (r.s[FHIRID],), 'FHIR識別文字列'),
    ]
    for cid, kname, keyf, valf, vname in specs:
        g_new, g_base = combos(new_rows, keyf, valf), combos(base_rows, keyf, valf)
        for k, g in g_new.items():
            if len(g) < 2:
                continue
            base_set = set(g_base.get(k, {}).keys())
            newcomers = set(g) - base_set
            shown = '、'.join('／'.join(v) for v in sorted(g))
            msg = f'同一の{kname}「{k}」に対し、{vname}が複数存在します（{shown}）。統一していただけますと幸いです'
            if cid == 'KA-17':
                msg = f'同一のFHIR項目名称「{k}」で並び順が複数（{shown}）存在します。並び順を統一してください'
            if cid == 'KA-32':
                msg = f'同一のJLAC10コード「{k}」に異なるFHIR識別文字列（{shown}）が設定されています。ご確認ください'
            for v, rows in g.items():
                # 前月から有る組合せ側には付けず、今回新たに生じた組合せ側にのみ付ける
                known = not newcomers or v not in newcomers
                for r in rows:
                    out.append((r, cid, msg, known))

    # KA-31 JLAC11コード重複
    base_cnt = Counter(r.code for r in base_rows)
    by_code = defaultdict(list)
    for r in new_rows:
        by_code[r.code].append(r)
    for code, rows in by_code.items():
        if code and len(rows) > 1:
            where = '、'.join(f'{r.sheet} {r.xrow}行目' for r in rows)
            msg = f'同一のJLAC11コード「{code}」が{len(rows)}行存在します（{where}）。正しい行のみ残していただけますと幸いです'
            for r in rows:
                out.append((r, 'KA-31', msg, base_cnt[code] > 1))

    # KA-33 / KA-34 その他の測定法(000)
    def sets(rows):
        spec = defaultdict(list)
        others = {}
        for r in rows:
            c = r.code
            if len(c) != 17:
                continue
            key = c[:12] + c[15:]
            if c[12:15] == '000':
                others[key] = r
            else:
                spec[key].append(r)
        return spec, others
    spec_n, oth_n = sets(new_rows)
    spec_b, oth_b = sets(base_rows)
    for key, r in oth_n.items():
        if key not in spec_n:
            msg = '本レコードは『その他の測定法』コード（JLAC11の13〜15桁=000）ですが、測定法が決まったコードが同時に存在しません。ご確認ください'
            out.append((r, 'KA-33', msg, key in oth_b and key not in spec_b))
        ended = [x for x in spec_n.get(key, []) if x.s[END] != '99999999']
        if ended and r.s[END] != '99999999':
            names = '、'.join(x.code for x in ended)
            msg = f'測定法コード（{names}）に適用終了日が設定されているため、本『その他の測定法』コードの適用終了日は99999999である必要があります（現在「{r.s[END]}」）。ご確認ください'
            b = oth_b.get(key)
            known = b is not None and b.s[END] != '99999999' and any(x.s[END] != '99999999' for x in spec_b.get(key, []))
            out.append((r, 'KA-34', msg, known))
    return out


# ---------- 前月比較 ----------
def compare_prev(r, p, H):
    """既存行の差分チェック。[(check_id, message)]"""
    out = []
    ps, ns = p.s, r.s
    if date_norm(ps[START]) != date_norm(ns[START]):
        out.append(('KA-M2', f'{PREV_LABEL}から適用開始日が変更されています（{ps[START]}→{ns[START]}）。既存コードの適用開始日は変更しない想定のため、ご確認ください'))
    if date_norm(ps[END]) != date_norm(ns[END]):
        if ns[END] == '99999999':
            out.append(('KA-M3', f'{PREV_LABEL}から適用終了日が変更されていますが（{ps[END]}→{ns[END]}）、変更後が99999999になっています。ご確認ください'))
        elif ps[END] == '99999999':
            if ns[END] != EXPECTED_END:
                out.append(('STEP1-END', f'適用終了日が{ns[END]}に設定されています。今回の切替日（{EXPECTED_START}）の前日（{EXPECTED_END}）を想定しておりましたが、意図した日付かご確認ください'))
            else:
                out.append(('STEP1-END-OK', f'適用終了日 99999999→{ns[END]}（想定どおり）'))
        else:
            out.append(('STEP1-END', f'{PREV_LABEL}から適用終了日が変更されています（{ps[END]}→{ns[END]}）。適用終了日のバックデート・再設定にあたるため、意図した変更かご確認ください'))
    diffs = defaultdict(list)
    for c in range(NCOL):
        if c not in (START, END) and ps[c] != ns[c]:
            diffs[(ps[c], ns[c])].append(H[c])
    for (old, new), cols in diffs.items():
        cols = '・'.join(cols)
        if nz(old) == nz(new):
            out.append(('STEP1-DIFF表記', f'{PREV_LABEL}から{cols}の表記が変更されています（「{old}」→「{new}」）。全角・記号の半角統一による変更と認識しておりますが、相違ないかご確認ください'))
        else:
            out.append(('STEP1-DIFF', f'{PREV_LABEL}から{cols}が変更されています（「{old}」→「{new}」）。意図した変更かご確認ください'))
    return out


def main():
    T = load_tables()
    H = None
    new_rows, by_sheet = [], {}
    for sh in NEW_SHEETS:
        h, rows = load(NEW_FILE, sh)
        H = H or h
        by_sheet[sh] = rows
        new_rows += rows
    _, base_rows = load(PREV_FILE, PREV_SHEET_NEWONLY)
    _, mixed_rows = load(PREV_FILE, PREV_SHEET_MIXED)
    _, prev_submit = load(FORM_FILE, FORM_PREV_SUBMIT_SHEET)

    base_map = defaultdict(list)
    for p in base_rows:
        base_map[p.code].append(p)
    base_rc = {}
    for p in base_rows:
        base_rc.setdefault(p.code, set()).update((cid, msg) for cid, lv, msg in row_checks(p, H, T))

    comments = defaultdict(list)   # id(row) -> [msg]
    findings = []                  # dict for report
    def rec(r, cid, level, msg, known, status):
        findings.append({'sheet': r.sheet, 'row': r.xrow, 'code': r.code, 'check': cid, 'level': level,
                         'msg': msg, 'known': known, 'status': status})

    new_codes = []
    for r in new_rows:
        is_new = r.code not in base_map
        if is_new:
            new_codes.append(r)
        for cid, level, msg in row_checks(r, H, T):
            known = (not is_new) and (cid, msg) in base_rc.get(r.code, set())
            status = '既存行' if not is_new else '新規行'
            rec(r, cid, level, msg, known, status)
            if level == 'strict' and not known:
                comments[id(r)].append(msg)
        if is_new:
            if date_norm(r.s[START]) == '00000000' or r.s[END] != '99999999':
                msg = f'新規追加のJLAC11コードは、適用開始日≠00000000 かつ 適用終了日=99999999 である必要があります（現在 適用開始日「{r.s[START]}」・適用終了日「{r.s[END]}」）。ご確認ください'
                rec(r, 'KA-M1', 'strict', msg, False, '新規行')
                comments[id(r)].append(msg)
            if r.s[START] != EXPECTED_START and date_norm(r.s[START]) != '00000000':
                msg = f'新規追加行の適用開始日が「{r.s[START]}」となっています。今回の公開日（{EXPECTED_START}）を想定しておりましたが、意図した日付かご確認ください'
                rec(r, 'STEP2-START', 'strict', msg, False, '新規行')
                comments[id(r)].append(msg)
        else:
            p = base_map[r.code][0]
            for cid, msg in compare_prev(r, p, H):
                level = 'info' if cid == 'STEP1-END-OK' else 'strict'
                rec(r, cid, level, msg, False, '既存行')
                if level == 'strict':
                    comments[id(r)].append(msg)

    for r, cid, msg, known in group_findings(new_rows, base_rows):
        rec(r, cid, 'strict', msg, known, '新規行' if r.code not in base_map else '既存行')
        if not known:
            comments[id(r)].append(msg)

    # ---- 参考チェック（コメントしない） ----
    base_xunits = {p.s[XUNIT] for p in base_rows} | {p.s[XUNIT2] for p in base_rows}
    base_unit10 = {(p.s[JLAC10][15:17], p.s[UNIT]) for p in base_rows}
    base_ana_dai = defaultdict(set)
    for p in base_rows:
        base_ana_dai[p.code[:5]].add(p.s[DAIKOMOKU])
    for r in new_codes:
        for c in (XUNIT, XUNIT2):
            if r.s[c] and r.s[c] not in base_xunits:
                rec(r, 'KA-09', 'ref', f'{H[c]}「{r.s[c]}」は{PREV_LABEL}に無い単位（UCUM掲載を目視確認）', False, '新規行')
        if (r.s[JLAC10][15:17], r.s[UNIT]) not in base_unit10:
            rec(r, 'KA-24', 'ref', f'JLAC10結果識別「{r.s[JLAC10][15:17]}」×表示用単位「{r.s[UNIT]}」は{PREV_LABEL}に無い組合せ', False, '新規行')
        if r.code[:5] in base_ana_dai and r.s[DAIKOMOKU] not in base_ana_dai[r.code[:5]]:
            rec(r, 'KA-05', 'ref', f'JLAC11測定物「{r.code[:5]}」の大項目が既存（{"／".join(sorted(base_ana_dai[r.code[:5]]))}）と異なる「{r.s[DAIKOMOKU]}」', False, '新規行')

    # ---- 範囲（削除）分析 ----
    new_set = {r.code for r in new_rows}
    base_set = set(base_map)
    mixed_set = {r.code for r in mixed_rows}
    prev_submit_set = {r.code for r in prev_submit}
    deleted_from_newonly = sorted(base_set - new_set)
    old_only = mixed_set - base_set
    old_missing = old_only - new_set
    mixed_dup = sum(1 for c, n in Counter(r.code for r in mixed_rows).items() if n > 1)
    newonly_vs_mixed_diff = 0
    mixed_map = defaultdict(list)
    for m in mixed_rows:
        mixed_map[m.code].append(m)
    for p in base_rows:
        if not any(m.s == p.s for m in mixed_map.get(p.code, [])):
            newonly_vs_mixed_diff += 1
    scope = {
        'new_rows': {sh: len(rs) for sh, rs in by_sheet.items()},
        'base_newonly_rows': len(base_rows), 'base_mixed_rows': len(mixed_rows), 'prev_submit_rows': len(prev_submit),
        'new_codes': len(new_codes),
        'deleted_from_newonly': len(deleted_from_newonly),
        'deleted_from_newonly_moved_to_abo': sum(1 for c in deleted_from_newonly if c in {r.code for r in by_sheet[NEW_SHEETS[1]]}),
        'old_only_codes_in_mixed': len(old_only), 'old_only_missing_in_new': len(old_missing),
        'old_missing_end_dist': Counter(mixed_map[c][0].s[END] for c in old_missing).most_common(10),
        'mixed_dup_codes': mixed_dup, 'newonly_rows_differs_from_mixed': newonly_vs_mixed_diff,
        'prev_submit_minus_mixed': len(prev_submit_set - mixed_set), 'mixed_minus_prev_submit': len(mixed_set - prev_submit_set),
        'prev_submit_contains_old': len(prev_submit_set & old_only),
        'abo_in_base': sum(1 for r in by_sheet[NEW_SHEETS[1]] if r.code in base_map),
        'new_code_start': Counter(r.s[START] for r in new_codes).most_common(),
        'new_code_end': Counter(r.s[END] for r in new_codes).most_common(),
        'new_code_sheets': Counter(r.sheet for r in new_codes).most_common(),
        'list11_loaded': bool(T['list11']),
        'new_material_pairs': Counter((r.s[MAT10], r.s[MAT11]) for r in new_codes).most_common(),
    }

    write_output(by_sheet, comments, H)
    n_comment_rows = sum(1 for rs in by_sheet.values() for r in rs if comments.get(id(r)))
    summary = {'scope': scope, 'n_comment_rows': n_comment_rows,
               'comment_rows_by_sheet': {sh: sum(1 for r in rs if comments.get(id(r))) for sh, rs in by_sheet.items()},
               'findings': findings,
               'new_rows': [{'sheet': r.sheet, 'row': r.xrow, 'code': r.code, 'dai': r.s[DAIKOMOKU], 'fhir': r.s[FHIRNAME],
                             'mat11': r.s[MAT11], 'meth11': r.s[METH11], 'mat10': r.s[MAT10], 'jlac10': r.s[JLAC10],
                             'unit': r.s[UNIT], 'start': r.s[START], 'end': r.s[END]} for r in new_codes]}
    if OUT_JSON:
        OUT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=list), encoding='utf-8')
    cnt = Counter((f['check'], f['level'], f['status'], f['known']) for f in findings)
    for k in sorted(cnt):
        print(k, cnt[k])
    print('comment rows:', summary['comment_rows_by_sheet'])
    print(json.dumps(scope, ensure_ascii=False, default=list, indent=1))
    print('出力:', OUT_XLSX)


def write_output(by_sheet, comments, H):
    wb = openpyxl.load_workbook(NEW_FILE)
    fwb = openpyxl.load_workbook(FORM_FILE, read_only=False)
    fws = fwb[FORM_SHEET]
    widths = {k: v.width for k, v in fws.column_dimensions.items() if v.width}
    fwb.close()
    idx = 0
    for sh in NEW_SHEETS:
        rows = by_sheet[sh]
        if sh != NEW_SHEETS[0] and not any(comments.get(id(r)) for r in rows):
            continue
        name = FORM_SHEET if sh == NEW_SHEETS[0] else f'{FORM_SHEET}_{sh[-6:]}'[:31]
        src = wb[sh]
        ws = wb.create_sheet(name, idx)
        idx += 1
        for c in range(1, NCOL + 1):
            sc = src.cell(row=1, column=c)
            dc = ws.cell(row=1, column=c, value=sc.value)
            dc.font, dc.fill, dc.alignment, dc.border = copy(sc.font), copy(sc.fill), copy(sc.alignment), copy(sc.border)
        hc = ws.cell(row=1, column=NCOL + 1, value=COMMENT_HDR)
        sc = src.cell(row=1, column=1)
        hc.fill, hc.border = copy(sc.fill), copy(sc.border)
        for c in range(1, NCOL + 2):
            f = copy(ws.cell(row=1, column=c).font if c <= NCOL else sc.font)
            f.b = True
            ws.cell(row=1, column=c).font = f
        out_r = 2
        for r in rows:
            for c in range(1, NCOL + 1):
                sc = src.cell(row=r.xrow, column=c)
                dc = ws.cell(row=out_r, column=c, value=sc.value)
                dc.number_format = sc.number_format
            msgs = list(dict.fromkeys(comments.get(id(r), [])))
            if msgs:
                text = msgs[0] if len(msgs) == 1 else '\n'.join('・' + m for m in msgs)
                cc = ws.cell(row=out_r, column=NCOL + 1, value=text)
                cc.alignment = Alignment(wrap_text=True, vertical='top')
            out_r += 1
        for k, w in widths.items():
            ws.column_dimensions[k].width = w
        ws.freeze_panes = 'N2'
        ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(NCOL + 1)}{out_r - 1}'
    wb.create_sheet('JLACセンター原本➡', idx)
    wb.active = 0
    for w in wb.worksheets:
        w.sheet_view.tabSelected = (w.title == FORM_SHEET)
    OUT_XLSX.parent.mkdir(exist_ok=True)
    wb.save(OUT_XLSX)


if __name__ == '__main__':
    main()
