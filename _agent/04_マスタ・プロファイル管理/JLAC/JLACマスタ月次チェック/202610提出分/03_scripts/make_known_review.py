"""10月版で「既知」としてコメントしなかった指摘のうち、使用中の行（適用終了日99999999）の分を
論点ごとにまとめ、本田さんと「仕様か誤りか」を判断するための確認資料を作る。

使い方：python make_known_review.py <findings.json（run_check.mjs --json の出力）> <10月版マスタxlsx> <出力xlsx>
"""
import json
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHECKS = ('KA-10', 'KA-12', 'KA-23', 'KA-08')
HEAD = PatternFill('solid', fgColor='DDEBF7')
ASK = PatternFill('solid', fgColor='FFF2CC')
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')


def s(v):
    return '' if v is None else str(v)


def norm(v):
    v = unicodedata.normalize('NFKC', v)
    v = re.sub(r'[‐-―−－﹣]', '-', v)
    return re.sub(r'\s', '', v)


def strip_marks(v):
    return re.sub(r'[()\[\]・･\-‘’\'"“”、,]', '', norm(v)).casefold()


def load(fj, fx):
    findings = json.load(open(fj, encoding='utf-8'))['findings']
    ws = openpyxl.load_workbook(fx, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    data = {n: r for n, r in enumerate(rows[1:], start=2)}
    return findings, data, ix


def add_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    last = len(header)
    for c in range(1, last + 1):
        cell = ws.cell(1, c)
        cell.font = BOLD
        cell.fill = ASK if c == last else HEAD
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            if cell.column == last:
                cell.fill = ASK
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def main(fj, fx, out):
    findings, data, ix = load(fj, fx)

    def active(n):
        return s(data[n][ix['適用終了日']]).replace('.0', '') == '99999999'

    def code(n):
        return f"{n}行目（{s(data[n][ix['JLAC11コード']])}）"

    known = defaultdict(list)
    for f in findings:
        if f['check'] in CHECKS and f['known'] and f['level'] == 'strict':
            known[f['check']].append(f)
    total = {k: len(known[k]) for k in CHECKS}
    act = {k: [f for f in known[k] if active(f['row'])] for k in CHECKS}

    def group_names(check, key_col, cols):
        groups = OrderedDict()
        for f in act[check]:
            r = data[f['row']]
            combo = tuple(s(r[ix[c]]) for c in cols)
            groups.setdefault(s(r[ix[key_col]]), OrderedDict()).setdefault(combo, []).append(f['row'])
        rows, breakdown = [], defaultdict(int)
        for key, combos in sorted(groups.items(), key=lambda x: -sum(len(v) for v in x[1].values())):
            dts = {c: sorted({s(data[n][ix['データタイプ']]) for n in ns}) for c, ns in combos.items()}
            dt_list = [tuple(v) for v in dts.values()]
            text = ' '.join(''.join(c) for c in combos)
            if len(set(dt_list)) == len(dt_list) and all(len(v) == 1 for v in dt_list):
                cls = '定量・定性などの書き分け（データタイプが書き方ごとに違う）'
            elif re.search(r'定性|定量|空腹|随時|陽性コントロール比|インデックス', text):
                cls = '検査の種類の書き分け（名称に定性・定量・空腹時・随時などを含む）'
            else:
                cls = 'その他（要判断）'
            n_rows = sum(len(v) for v in combos.values())
            breakdown[cls] += n_rows
            listing = '\n'.join(f"・{'／'.join(c)}（{len(ns)}行、{'/'.join(dts[c])}）" for c, ns in combos.items())
            examples = '\n'.join(code(ns[0]) for ns in combos.values())
            rows.append([key, n_rows, len(combos), listing, cls, examples, None])
        return rows, dict(breakdown), len(groups)

    r10, b10, n10 = group_names('KA-10', '大項目', ['FHIR項目名称', 'FHIR識別文字列', '略称'])
    r12, b12, n12 = group_names('KA-12', '略称', ['大項目', 'FHIR項目名称', 'FHIR識別文字列'])

    g23 = OrderedDict()
    for f in act['KA-23']:
        m = re.search(r'検査方法\(JLAC10-測定法\)「(.*?)」が、.*?測定法コード「(\w+)」の名称「(.*?)」', f['msg'])
        m2 = re.search(r'測定法コード「(\w+)」が、.*?見当たりません', f['msg'])
        if m:
            key = (m.group(2), m.group(1), m.group(3))
        elif m2:
            key = (m2.group(1), s(data[f['row']][ix['検査方法(JLAC10-測定法)']]), '（コード表にこのコードが無い）')
        else:
            key = ('?', '?', f['msg'][:60])
        g23.setdefault(key, []).append(f['row'])
    r23, b23 = [], defaultdict(int)
    for (mc, mv, tv), ns in sorted(g23.items(), key=lambda x: -len(x[1])):
        if tv.startswith('（コード表に'):
            cls = 'コード表に無い測定法コード（999＝その他の測定法）'
        elif norm(mv) == norm(tv):
            cls = '全角・半角・空白だけの違い'
        elif strip_marks(mv) == strip_marks(tv):
            cls = '括弧・記号だけの違い'
        else:
            cls = '名称が違う（要判断）'
        b23[cls] += len(ns)
        ex = '\n'.join(code(n) for n in ns[:3]) + ('\n…' if len(ns) > 3 else '')
        r23.append([mc, mv, tv, len(ns), cls, ex, None])

    r08, b08 = [], defaultdict(int)
    for f in sorted(act['KA-08'], key=lambda x: x['row']):
        m = re.search(r'測定法\(JLAC11\)「(.*?)」が、.*?「測定法名称」列の「(.*?)」', f['msg'])
        mv, lv = m.groups() if m else (f['msg'][:60], '?')
        if '_' in mv and norm(mv.split('_', 1)[1]) == norm(lv):
            cls = '「方法名_試薬名」の書き方（試薬名は一致）'
        elif norm(mv) == norm(lv):
            cls = '全角・半角・空白だけの違い'
        elif strip_marks(mv) == strip_marks(lv):
            cls = '全角・半角・記号（‘’など）・大文字小文字だけの違い'
        else:
            cls = '名称が違う（要判断）'
        b08[cls] += 1
        r = data[f['row']]
        r08.append([f['row'], s(r[ix['JLAC11コード']]), s(r[ix['販売名称']]), mv, lv, cls, None])

    def fmt(b):
        return '\n'.join(f'{k}：{v}行' for k, v in sorted(b.items(), key=lambda x: -x[1]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '説明'
    ws.append(['JLACマスタ 10月版で「既知」とした指摘の確認（使用中の行のみ）'])
    ws.append([])
    ws.append(['目的', '10月版のチェックで「前月も同じ＝了承済み」としてコメントしなかった指摘は、8月・9月にJLACセンターへ同じ論点で指摘したことがなく、了承をもらったものではなかった。使用中の行（適用終了日99999999）に絞り、論点ごとに「マスタの仕様（直さなくてよい）か、誤り（JLACセンターに確認する）か」をご判断いただきたい。適用終了しているコードは対象外。'])
    ws.append(['元にしたもの', '10月版：1006再提出（新旧混合リスト）。チェック：JLACマスタチェックツール最新版（v20。既知の件数は10月版の確認時〔v12〕と同じ）。8月・9月のコメント：20260810_202608版JLACマスタ_JLACセンター向けコメント／【電カル共有】JLACマスタ_JLACセンター20260901提出版_確認事項一覧_20260909'])
    ws.append(['ACNの見立てについて', '書き方の違いを機械的に分類したもの（推定）で、仕様かどうかの判断ではない。各シートの黄色の列に、仕様か誤りかをご記入ください。'])
    ws.append([])
    ws.append(['チェック', 'チェック内容', '既知（全体）', '使用中の行', '論点の数', 'ACNの見立て（使用中の行の内訳）', '判断いただきたいこと', '本田さんの判断'])
    ws.append(['KA-10', '大項目をキーにFHIR項目名称/FHIR識別文字列/略称が整合すること', total['KA-10'], len(act['KA-10']), f'大項目 {n10}種類', fmt(b10),
               '同じ大項目で書き方が複数あるのは仕様か（とくに「その他（要判断）」）。判定の基準をFHIR識別文字列・FHIR項目名称に変える方向（R012）でよいか', None])
    ws.append(['KA-12', '略称をキーに大項目/FHIR項目名称/FHIR識別文字列が整合すること', total['KA-12'], len(act['KA-12']), f'略称 {n12}種類', fmt(b12),
               '同じ略称で大項目・FHIR名が複数あるのは仕様か（KA-10と同じ論点）', None])
    ws.append(['KA-23', '検査方法(JLAC10-測定法)がJLAC10測定法コードと整合すること', total['KA-23'], len(act['KA-23']), f'名称の組 {len(g23)}種類', fmt(b23),
               'マスタの検査方法の名称がJLAC10コード表と違うのは許容か。全角・半角や括弧だけの違いは指摘しなくてよいか', None])
    ws.append(['KA-08', '測定法(JLAC11)がJLAC11測定法コードと整合すること', total['KA-08'], len(act['KA-08']), f'{len(act["KA-08"])}行', fmt(b08),
               '「方法名_試薬名」の書き方（例：IFCC対応法_N-アッセイ…）はマスタの命名ルールか。名称が違う行はJLACセンターに確認するか', None])
    ws['A1'].font = Font(bold=True, size=13)
    for c in range(1, 9):
        ws.cell(7, c).font = BOLD
        ws.cell(7, c).fill = ASK if c == 8 else HEAD
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = WRAP
    for r in range(8, 12):
        ws.cell(r, 8).fill = ASK
    for col, w in zip('ABCDEFGH', (12, 30, 10, 10, 14, 46, 50, 30)):
        ws.column_dimensions[col].width = w

    add_sheet(wb, 'KA-10（大項目）', ['大項目', '使用中の行数', '書き方の数', '書き方（FHIR項目名称／FHIR識別文字列／略称、行数、データタイプ）', 'ACNの見立て', '代表行（書き方ごとの先頭行）', '本田さんの判断（仕様／誤り）'], r10, (18, 8, 8, 60, 28, 30, 22))
    add_sheet(wb, 'KA-12（略称）', ['略称', '使用中の行数', '書き方の数', '書き方（大項目／FHIR項目名称／FHIR識別文字列、行数、データタイプ）', 'ACNの見立て', '代表行（書き方ごとの先頭行）', '本田さんの判断（仕様／誤り）'], r12, (14, 8, 8, 60, 28, 30, 22))
    add_sheet(wb, 'KA-23（検査方法）', ['JLAC10測定法コード', 'マスタの検査方法(JLAC10-測定法)', 'JLAC10コード表の名称', '使用中の行数', 'ACNの見立て', '代表行（先頭3行）', '本田さんの判断（仕様／誤り）'], r23, (10, 30, 30, 8, 22, 30, 22))
    add_sheet(wb, 'KA-08（測定法）', ['行', 'JLAC11コード', '販売名称', 'マスタの測定法(JLAC11)', '17桁コードリストの測定法名称', 'ACNの見立て', '本田さんの判断（仕様／誤り）'], r08, (8, 20, 28, 36, 30, 26, 22))
    wb.save(out)
    print(json.dumps({'total': total, 'active': {k: len(v) for k, v in act.items()}, 'groups': {'KA-10': n10, 'KA-12': n12, 'KA-23': len(g23)},
                      'KA-10': b10, 'KA-12': b12, 'KA-23': dict(b23), 'KA-08': dict(b08)}, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main(*sys.argv[1:4])
