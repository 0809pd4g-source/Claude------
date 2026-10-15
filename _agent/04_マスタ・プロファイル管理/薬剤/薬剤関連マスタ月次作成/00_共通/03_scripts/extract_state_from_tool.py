"""作成ツール（Excel）から、自動作成（generate_yakuzai_master.py）が使う「基準データ」を書き出す。最初の1回だけ使う。

書き出すもの（CSV・UTF-8）：
  薬価基準.csv          … 「薬価基準収載医薬品コード(…)」シートのA〜E列（シートの並びのまま。成分名は頭9桁で最初の行から引くため並びが意味を持つ）
  一般名処方マスタ.csv  … 「一般名処方マスタ(…)」シートのA〜E列（並びのまま）
  例外コード表.csv      … 「一般名処方マスタ（例外コード表）」シートのB〜E列（並びのまま）
  YJ-list.csv          … 「YJ-list」シートのA〜F列（検証用：その月のYJコードリスト）
  前回GCM.csv          … 「STEP7_前回分との整合チェック」のG列＝前回の行のA〜E列（検証用：前月の公開版）

使い方: python -I extract_state_from_tool.py <作成ツール.xlsx> <出力フォルダ>
"""
import csv
import os
import sys

import openpyxl


def s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def sheet(wb, prefix):
    hit = [n for n in wb.sheetnames if n.startswith(prefix)]
    if len(hit) != 1:
        raise SystemExit(f'シートが特定できない: {prefix} → {hit}')
    return wb[hit[0]], hit[0]


def write(path, head, rows):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(head)
        w.writerows(rows)


def main():
    tool, out = sys.argv[1], sys.argv[2]
    os.makedirs(out, exist_ok=True)
    wb = openpyxl.load_workbook(tool, read_only=True, data_only=True)
    log = []

    ws, name = sheet(wb, '薬価基準収載医薬品コード(')
    rows = [[s(c) for c in r[:5]] for r in ws.iter_rows(min_row=2, max_col=5, values_only=True) if r[1]]
    write(os.path.join(out, '薬価基準.csv'), ['区分', '薬価基準収載医薬品コード', '成分名', '規格', 'コード追加日'], rows)
    log.append(f'薬価基準.csv  {len(rows)}行  （{name}）')

    ws, name = sheet(wb, '一般名処方マスタ(')
    rows = [[s(c) for c in r[:5]] for r in ws.iter_rows(min_row=2, max_col=5, values_only=True) if r[1]]
    write(os.path.join(out, '一般名処方マスタ.csv'), ['区分', '一般名コード', '一般名処方の標準的な記載', '成分名', '規格'], rows)
    log.append(f'一般名処方マスタ.csv  {len(rows)}行  （{name}）')

    ws = wb['一般名処方マスタ（例外コード表）']
    rows = [[s(c) for c in r[1:5]] for r in ws.iter_rows(min_row=3, max_col=5, values_only=True) if r[2]]
    write(os.path.join(out, '例外コード表.csv'), ['区分', '一般名コード', '一般名処方の標準的な記載', '成分名'], rows)
    log.append(f'例外コード表.csv  {len(rows)}行')

    ws = wb['YJ-list']
    rows = [[s(c) for c in r[:6]] for r in ws.iter_rows(min_row=2, max_col=6, values_only=True) if r[0]]
    write(os.path.join(out, 'YJ-list.csv'), ['ＹＪコード', '医薬品名', '会社名', 'リスト登録年月日', 'リスト除外年月日', '変更区分'], rows)
    log.append(f'YJ-list.csv  {len(rows)}行')

    ws = wb['STEP7_前回分との整合チェック']
    rows = [[s(c) for c in r[:5]] for r in ws.iter_rows(min_row=3, max_col=7, values_only=True) if r[0] and r[6] == '前回']
    write(os.path.join(out, '前回GCM.csv'), ['剤形・規格・銘柄不明コード', '薬剤成分名(剤形・規格・銘柄不明)', '(空欄)', '登録年月日', '除外年月日'], rows)
    log.append(f'前回GCM.csv  {len(rows)}行')

    open(os.path.join(out, '抽出元.txt'), 'w', encoding='utf-8').write(os.path.basename(tool) + '\n' + '\n'.join(log) + '\n')
    sys.stdout.reconfigure(encoding='utf-8')
    print('\n'.join(log))


if __name__ == '__main__':
    main()
