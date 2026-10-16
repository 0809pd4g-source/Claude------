# -*- coding: utf-8 -*-
"""令和9年度案件（担当8件）の要求事項の詳細化：説明スライドのたたき台を作る。

土台：01_input/開発課題検討会_20261013_厚労省側.pptx（厚労省フォーマット。原本は変更しない）
出力：02_output/20261016_令和9年度案件_要求事項詳細化_説明スライド_v1.pptx
中身の正本：02_output/ の各案件の詳細化案（MD）
"""
import copy
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "01_input" / "開発課題検討会_20261013_厚労省側.pptx"
OUT = ROOT / "02_output" / "20261016_令和9年度案件_要求事項詳細化_説明スライド_v1.pptx"

NAVY = RGBColor(0x1F, 0x49, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)
TITLE_BASE = "X．令和9年度案件の要求事項の詳細化について"
TOTAL = 9
LEFT, WIDTH = 39.0, 702.0
BOTTOM = 506.0  # 注記の上まで
TABLE_STYLE_GRID = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
NOTE_RED = "※赤字：案件台帳（10/13時点）の記載からの追記箇所（案）"

# ---------------------------------------------------------------- 中身
# 項目は (文, 赤字か)。表は {"after": ラベル名, "cols": [...], "widths": [...], "rows": [[...]]}
CASES = [
    {
        "id": "011", "name": "健診PDF帳票の表示件数の拡張",
        "lead": "健診PDFの表示範囲を、XML形式と同じ「過去5年分」にそろえるため、対象・数え方・レイアウトの考え方を具体化した。",
        "sections": [
            ("背景", [
                ("健診結果のPDF帳票は直近5回分までを表示している。一方、資格確認端末で「電子カルテ情報共有サービスを利用する」設定の場合、XML形式では過去5年間のすべての健診データ（同一年内の複数回受診を含む）を取得できる（技術解説書 v2.1.0）。", True),
                ("健診結果報告書のオン資への保存期間は5年間と公開している（利活用WG）。", False),
            ]),
            ("要求事項（詳細化案）", [
                ("健診結果のPDF帳票の表示上限を5回分から5年分に変更する。", False),
                ("対象は「利用する」設定の場合の、医療機関向け・マイナポータル向けのPDF帳票とする。「利用しない」設定は現行どおり。", True),
                ("「5年分」はXML形式と同じく、受診日が過去5年以内の健診をすべて表示する（回数上限なし）。", True),
                ("6回分以上の場合は、現行の5回分の表を1単位として新しい順に改ページすることを想定する（レイアウトは要件定義で作成）。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("優先度について、10/5開発課題検討会での説明（中）と一覧（低）のどちらとするか、ご確認いただきたい。", False),
                ("対象を医療機関向け・マイナポータル向けの両方とし、「利用しない」設定を対象外とする方向でよろしいでしょうか。", False),
                ("保存期間の延長（案件178）とは連動させず、表示範囲は5年で固定する方向でよろしいでしょうか。", False),
            ]),
        ],
        "table": {"after": "背景", "cols": ["取得の形式", "現行", "本案件の対応後"], "widths": [140, 280, 282],
                  "rows": [["XML形式", "過去5年間のすべての健診データ", "変更なし"],
                           ["PDF形式（健診結果報告書）", "直近5回分", "過去5年間のすべての健診データ"]]},
    },
    {
        "id": "207", "name": "自己情報提供不可フラグ設定した患者に対する患者サマリーの登録時のエラーのアドオンチェック",
        "lead": "8/3にご議論いただいた方針に沿って、対象のフラグ・更新削除の扱い・エラーの文言案を具体化した。",
        "notes": ["※CISS1116E：住所・郵便番号が不開示のため診療情報提供書を登録・取得できないときのエラー（技術解説書 v2.1.0）"],
        "sections": [
            ("背景", [
                ("自己情報提供不可フラグは、DV被害者等がマイナンバーカードを相手方のもとに置いてきた場合等に、再交付までマイナポータルでの閲覧を制限するフラグである。当該患者はマイナポータルで患者サマリーを閲覧できない。", False),
                ("8/3開発課題検討会にて、当該患者の患者サマリーは登録エラーとし、臨床情報は現行どおり登録可とする方針となった。", False),
            ]),
            ("要求事項（詳細化案）", [
                ("患者サマリー登録時、登録対象患者に自己情報提供不可フラグが設定されている場合は登録エラーを返却する。", False),
                ("対象は自己情報提供不可フラグのみとする（不開示該当フラグの患者はマイナポータルを閲覧できるため対象外）。", True),
                ("更新も同様に登録エラーとし、削除は受け付ける。フラグ設定前に登録された患者サマリーは削除・変更しない。", True),
                ("エラー時は、住所・郵便番号の不開示エラー（CISS1116E）と同様に、①登録できない理由 ②療養計画書を紙媒体等で交付することを画面に表示する。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("エラーメッセージの文言を「この患者はマイナポータルで患者サマリーを閲覧できない状態のため、患者サマリーを登録できません。療養計画書は紙媒体等で交付してください。」としてよろしいでしょうか。", False),
                ("文言には「DV」「自己情報提供不可」等、フラグの種類が分かる表現を用いない方向でよろしいでしょうか。", False),
                ("医療機関等への影響確認を「あり（紙の療養計画書への切り替えの周知）」とする方向でよろしいでしょうか。", False),
            ]),
        ],
    },
    {
        "id": "211", "name": "標準マスタにおける1つのコードに対するdisplay値変更時における経過措置期間の設定機能の追加",
        "lead": "方式（②〜④）は見積もり結果を踏まえて決定するが、方式によらず決められる対象・判定の条件を具体化した。",
        "notes": ["※display値：コードに対応する名称。codeとdisplayの整合性チェック：登録された名称が標準マスタの名称と一致するかの確認（技術解説書 v2.1.0 表35）"],
        "sections": [
            ("背景", [
                ("中央基盤のマスタチェックは、標準マスタで終了日が最新のレコードのdisplay値と一致しない場合はエラーとしている。一部のJ-FAGY・JLACコードでdisplay値のみが変更となり、旧display値で登録するとエラーとなった。", False),
                ("9/28開発課題検討会にて、パターン②〜④のいずれかを採用する方向とし、運用負荷と改修費用を踏まえて方式を決定することとなった（見積もりは10/19目途）。", False),
            ]),
            ("要求事項（詳細化案）", [
                ("1つのコードに対してdisplay値のみの変更があった場合、経過措置期間内であれば旧display値でもワーニングを返却して登録できるようにする。", False),
                ("対象は、codeとdisplayの整合性チェックの対象となる標準マスタとする（健診マスタ・医療機関等マスタは対象外）。", True),
                ("チェック対象日付が「旧display値の適用終了日＋経過措置日数」以前であれば受け付ける。ワーニングでは新しいdisplay値と期限を示す。", True),
                ("設定方式の候補は②マスタのバージョンごと／③マスタごと／④全マスタ統一とする（①レコードごとは削除）。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("方式・日数の初期値・日数の設定の運用は見積もり結果を踏まえて決定し、それ以外は上記で確定する方向でよろしいでしょうか。", False),
                ("医療機関等マスタ（診療情報提供書の宛先）を対象外とする方向でよろしいでしょうか。", False),
            ]),
        ],
    },
    {
        "id": "216", "name": "サポートするプロファイルバージョンの終了(令和9年度)",
        "lead": "案件192に続き、CLINS1.12.0を令和9年12月リリース（10/5の説明のとおり）でシステム的に利用終了とする方向で、終了のしかたを案件192とそろえて具体化した。",
        "notes": ["※CLINS：電子カルテ情報共有サービスのFHIR実装ガイド（JP-CLINS）。1.12.0・1.13.0はその版。医療機関はデータを送るときに準拠する版を記載する"],
        "sections": [
            ("背景", [
                ("CLINS1.13.0より、感染症において患者共有フラグをFHIRで記述できるようになった。全国的な運用開始版の技術解説書では、CLINS1.13.0版に対応することと記述する予定である。", False),
                ("案件192「サポートするプロファイルバージョンの終了(令和9年3末)」にて、CLINS1.12.0はモデルベンダの改修スケジュールを考慮し、令和9年度の改修に合わせて利用終了とすることとした。モデル医療機関へは2026年9月に周知済み。", True),
            ]),
            ("要求事項（詳細化案）", [
                ("CLINS1.12.0をシステム的に利用終了とする。", False),
                ("令和9年12月リリース以降、CLINS1.12.0の形式で作成したデータの登録（新規・更新）はエラーとする。エラーの返し方は案件192と同じとする。", True),
                ("利用終了前にCLINS1.12.0で登録されたデータは引き続き閲覧でき、削除も受け付ける。更新はCLINS1.13.0以降でデータ全体を送信してもらうことを想定する（要件整理で現行の作りを確認のうえ確定）。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("リリース時期を令和9年12月と明記し、終了のしかたを案件192とそろえる方向でよろしいでしょうか。", False),
                ("全国の医療機関・ベンダーへは、技術解説書（全国的な運用開始版）等で、終了の時期とエラーの内容を周知する方向でよろしいでしょうか。", False),
            ]),
        ],
        "table": {"after": "背景", "cols": ["案件", "利用終了とするバージョン", "リリース"], "widths": [140, 340, 222],
                  "rows": [["192（令和8年度）", "CLINS 1.10.0・1.11.0、健診 1.5.0", "令和8年度3月"],
                           ["216（本案件）", "CLINS 1.12.0", "令和9年12月"]]},
    },
    {
        "id": "218", "name": "検査・感染症の検査結果に対する注意事項の追記",
        "lead": "10/5にご説明した方針（「検出せず」は「0」で登録し、閲覧時に注意書きを表示する）に沿って、表示の条件と文言案を具体化した。",
        "sections": [
            ("背景", [
                ("定量検査の結果は、数値か「検査結果を登録できない理由」を設定する仕様であり、「検出せず」は登録できない理由として連携するよう案内している。この場合、当該結果は他医療機関へ共有されない。", False),
                ("10/5開発課題検討会にて、「検出せず」は「0」として登録する方針とし、閲覧時の注意書きを整えたうえで、技術解説書・FAQ等で周知する方針を説明した。", True),
            ]),
            ("要求事項（詳細化案）", [
                ("PDF帳票（他医療機関向け臨床情報・マイナポータル向け臨床情報・マイナポータル向け患者サマリー）と資格確認端末の臨床情報詳細画面の注意書き欄に、結果値「0」には「検出せず」が含まれる場合がある旨を追記する。", False),
                ("注意書きは、検査情報・感染症情報の欄がある場合に、結果値によらず常に表示する。", True),
                ("「0」として登録する対象は「検出せず」の場合とし、「定量下限未満」は現行どおり登録できない理由を設定する。周知は本案件のリリース以降に厚労省にて行う。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("注意書きの文言を、以下としてよろしいでしょうか。", False),
                ("　他医療機関向け：「検査情報・感染症情報の結果値『0』には、検出せず（検出限界未満）の結果が含まれる場合があります。」", False),
                ("　マイナポータル向け：「検査の結果が『0』となっている場合、検査で検出されなかった（検出せず）ことを表している場合があります。」", False),
                ("注意書きを常に表示し、対象の帳票・画面を上記の4つとする方向でよろしいでしょうか。", False),
            ]),
        ],
    },
    {
        "id": "219", "name": "医療機関の承継が発生した際における患者サマリーのマイナポでの処方情報・検査情報の表示対象範囲の正規化",
        "lead": "9/28に決定した整理結果を要求事項に反映し、承継の判定方法と対象範囲を具体化した。",
        "sections": [
            ("背景", [
                ("患者サマリーの処方・検査情報は、患者サマリーに記述された医療機関コードを基に紐づけて表示している。承継で医療機関コードが変わると、承継後に登録した患者サマリーには承継前の処方・検査情報が表示されない。", False),
                ("9/28開発課題検討会にて、指導日から30日間の処方・検査情報を承継の前後を問わず表示すること、処方・検査情報ごとの登録医療機関名は表示しないことを決定した。", False),
            ]),
            ("要求事項（詳細化案）", [
                ("承継後に登録された患者サマリーについて、承継元医療機関として登録した処方・検査情報も含めて、指導日から30日間分を抽出・表示する。", False),
                ("承継の有無は、総合ポータルサイトの承継の届出により判定する（案件141と同様）。届出のない吸収合併等は承継として扱わない。", True),
                ("統合により承継元が複数ある場合、承継が複数回行われている場合も、紐づくすべての承継元を含める。", True),
                ("対象は患者サマリーの臨床情報閲覧機能（PDF・XML）とする。技術解説書（患者サマリー編）の「同一医療機関が…」の記載はあわせて改訂する。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("承継の判定を、案件141と同じく登録時点の承継の届出で行う方向でよろしいでしょうか。", False),
                ("本案件のリリース前に登録された患者サマリーは対象外とする方向でよろしいでしょうか。", False),
            ]),
        ],
    },
    {
        "id": "220", "name": "承継申請に基づく医療機関コードの変換処理における旧医療機関発行した患者サマリーの挙動の変更",
        "lead": "案件141の改修により、承継前に指導した患者サマリーを更新すると指導医療機関名が変わる点の対応である。9/28の方針と合わせて要否をご判断いただきたい。",
        "sections": [
            ("背景", [
                ("案件141にて、承継の届出が行われている場合は承継先の医療機関名称を表示する改修を行った。その結果、承継元で登録した患者サマリーを承継後に更新すると、指導医療機関が承継先と表示される。", False),
                ("一方、9/28開発課題検討会にて、医療機関名を承継しない案件205は改修しない方針として取り下げた。", False),
            ]),
            ("要求事項（詳細化案）", [
                ("指導日が承継日より前の患者サマリーの更新の場合は、医療機関名・医療機関コードを承継先に変換する処理の対象外とする。", False),
                ("承継日は総合ポータルサイトの届出の承継日、指導日は患者サマリーに記録された指導日とする。指導日が承継日以後の場合は現行どおり変換する。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("案件205の取り下げと同じ考え方で本案件を見送るか、更新の前後で名称が変わらないよう対応するか、ご判断いただきたい。", False),
                ("対応する場合、承継後に承継前の指導日の患者サマリーを新規に登録した場合も、変換の対象外とする方向でよろしいでしょうか。", False),
            ]),
        ],
        "table": {"after": "要求事項（詳細化案）", "cols": ["例：A病院がB病院に承継（承継日4/1）", "指導医療機関の表示（現行 → 本案件の対応後）", "処方・検査情報（案件219の対応後）"],
                  "widths": [250, 250, 202],
                  "rows": [["①A病院が3/20に指導・登録（更新なし）", "A病院（変更なし）", "A病院の分"],
                           ["②①をB病院が4/10に更新", "B病院 → A病院", "A病院の分"],
                           ["③B病院が4/15に指導・登録", "B病院（変更なし）", "A病院・B病院の分"]]},
    },
    {
        "id": "223", "name": "電子署名の危殆化対応",
        "lead": "TSAの署名方式をRSA-2048からECDSAへ移行する対応を、2段階（2027年9月・2028年3月）で進める。背景は案件台帳の記載から差し替える案（確認中）とし、将来の方式変更への備えを追加した。",
        "note_red": "※赤字：案件台帳（10/13時点）の記載からの追記箇所（案）。【検討中】は、台帳の背景・目的から差し替える案（確認中）",
        "notes": ["※TSA：文書に「この時刻に存在した」ことの証明（タイムスタンプ）を付ける仕組み。EE証明書：TSAが署名に使う証明書。ECDSA：RSAより強度の高い署名方式",
                  "※アーカイブタイムスタンプ：長期保存する文書の署名の有効性を延ばすために、新たに付けるタイムスタンプ"],
        "sections": [
            ("背景（案）", [
                ("【検討中】電子カルテ情報共有サービスのTSA（Time Stamp Authority）は、RSA-2048を用いてタイムスタンプ署名を行っている。現行のEE証明書は2027年10月に有効期限を迎える。", True),
                ("【検討中】RSA-2048は2030年までが移行完遂期間とされており、より高いセキュリティ強度を有する署名方式への移行が必要となる。", True),
            ]),
            ("要求事項（詳細化案）", [
                ("①TSAのEE証明書を、RSA鍵を用いた現行方式のまま更新する（2027年9月リリース予定）。", False),
                ("②署名検証機能及びタイムスタンプ付与・検証機能を、新たな署名方式（ECDSA-256/384を想定）に対応させる（2028年3月リリース予定）。", False),
                ("③将来の暗号方式の変更に備え、署名アルゴリズムを差し替えやすい構成とする。", True),
                ("④②より前にRSA-2048のタイムスタンプが付与された文書について、2030年以降も署名の有効性を確認できるよう、新たな署名方式によるアーカイブタイムスタンプの付与の要否を要件定義にて整理する。", True),
            ]),
            ("ご確認いただきたい事項", [
                ("耐量子計算機暗号への移行（9/28開発課題検討会のアジェンダ「電子署名の暗号方式変更」）は本案件の対象外とし、本案件では③の備えにとどめる方向でよろしいでしょうか。", False),
                ("取得した文書の署名検証を実装している医療機関・ベンダーに、②のリリースまでに新方式への対応を周知する方向でよろしいでしょうか。", False),
            ]),
        ],
    },
]

OVERVIEW_ROWS = [
    # ID, 区分, 案件名称, 優先度, 詳細化の概要, ご確認いただきたい事項（要点）
    ["011", "健診文書", "健診PDF帳票の表示件数の拡張", "低※", "XML形式と同じ「過去5年分」にそろえる。対象・数え方・改ページの考え方", "優先度（中／低）、対象の帳票"],
    ["207", "患者サマリー", "自己情報提供不可フラグ設定した患者に対する患者サマリーの登録時のエラーのアドオンチェック", "高", "対象のフラグ、更新・削除の扱い、エラーの文言案", "エラーの文言、影響確認"],
    ["211", "標準マスタ", "標準マスタにおける1つのコードに対するdisplay値変更時における経過措置期間の設定機能の追加", "中", "対象マスタ、旧display値を受け付ける条件、方式の候補を②〜④に", "見積もり後に決める範囲"],
    ["216", "全体", "サポートするプロファイルバージョンの終了(令和9年度)", "高", "令和9年12月リリース、終了のしかたを案件192とそろえる", "リリース時期、周知"],
    ["218", "臨床情報：感染症・検査", "検査・感染症の検査結果に対する注意事項の追記", "高", "注意書きの表示の条件、対象の帳票・画面、文言案", "注意書きの文言"],
    ["219", "患者サマリー", "医療機関の承継が発生した際における患者サマリーのマイナポでの処方情報・検査情報の表示対象範囲の正規化", "高", "9/28の決定事項の反映、承継の判定方法、対象範囲", "承継の判定の時点"],
    ["220", "患者サマリー", "承継申請に基づく医療機関コードの変換処理における旧医療機関発行した患者サマリーの挙動の変更", "低", "変換の対象外とする条件、219と合わせた見え方", "本案件の要否"],
    ["223", "全体", "電子署名の危殆化対応", "高", "背景・目的の改め、2段階のリリース、将来の方式変更への備え", "耐量子への移行の扱い、周知"],
]


# ---------------------------------------------------------------- 部品
def chars_per_line(font_pt, width_pt):
    return max(10, int((width_pt - 14) / font_pt))


def est_lines(text, font_pt, width_pt):
    n = 0
    for c in text:
        n += 1 if ord(c) > 0x2E7F else 0.55
    return max(1, math.ceil(n / chars_per_line(font_pt, width_pt)))


def set_bullet(paragraph, char="•"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Pt(11).emu))
    pPr.set("indent", str(-Pt(11).emu))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def add_text(slide, x, y, w, h, items, size, bullet=True, name=None):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, (text, red) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if bullet:
            set_bullet(p)
        p.space_after = Pt(1)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = RED if red else BLACK
    return tb


def add_label(slide, x, y, text, w=160):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(18))
    s.name = "ラベル：" + text
    s.fill.solid()
    s.fill.fore_color.rgb = NAVY
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_top = tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return s


def add_table(slide, x, y, widths, header, rows, size=9, name=None, center_cols=()):
    nrows, ncols = len(rows) + 1, len(header)
    line_h = size * 1.35
    heights = []
    for r in [header] + rows:
        lines = max(est_lines(str(v), size, widths[i]) for i, v in enumerate(r))
        heights.append(lines * line_h + 6)
    gf = slide.shapes.add_table(nrows, ncols, Pt(x), Pt(y), Pt(sum(widths)), Pt(sum(heights)))
    if name:
        gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(e)
    sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = TABLE_STYLE_GRID
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for i, wv in enumerate(widths):
        tbl.columns[i].width = Pt(wv)
    for ri, r in enumerate([header] + rows):
        tbl.rows[ri].height = Pt(heights[ri])
        for ci, v in enumerate(r):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Pt(3)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if ri == 0 or ci in center_cols:
                p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(v)
            run.font.size = Pt(size)
            run.font.bold = ri == 0
            run.font.color.rgb = WHITE if ri == 0 else BLACK
    return gf, sum(heights)


def new_slide(prs, layout, title):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type not in (1, 3):  # TITLE, CENTER_TITLE
            ph._element.getparent().remove(ph._element)
    s.shapes.title.text_frame.text = title
    return s


def section_height(items, size):
    line_h = size * 1.3
    lines = sum(est_lines(t, size, WIDTH - 12) for t, _ in items)
    return lines * line_h + len(items) * 1.5 + 3


def remove_unused_masters(prs):
    """使っていないスライドマスタ・レイアウトを外す（土台に残る Accenture のフッター等を持ち込まない）。"""
    used = {s.slide_layout.part.partname for s in prs.slides}
    lst = prs.slide_masters._sldMasterIdLst
    for master, sldMasterId in list(zip(prs.slide_masters, lst)):
        if not any(l.part.partname in used for l in master.slide_layouts):
            prs.part.drop_rel(sldMasterId.rId)
            lst.remove(sldMasterId)
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if layout.part.partname not in used:
                master.slide_layouts.remove(layout)


# ---------------------------------------------------------------- 組み立て
def build():
    prs = Presentation(str(SRC))
    layout = prs.slides[4].slide_layout  # 「2. 血液型…」と同じレイアウト（タイトルのプレースホルダー付き）

    # 土台のスライドはすべて外す（原本ファイルは変更しない）
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)

    # 1枚目：一覧
    s = new_slide(prs, layout, f"{TITLE_BASE}（1/{TOTAL}）")
    add_text(s, LEFT, 58, WIDTH, 34, [
        ("令和9年度案件のうち8件について、案件台帳の要求事項を詳細化した（案）。各案件の詳細は次ページ以降のとおり。", False),
        ("各案件の「ご確認いただきたい事項」について、ご確認いただきたい。", False)], 12, name="リード")
    add_table(s, 12, 100, [32, 70, 228, 40, 230, 156],
              ["ID", "区分", "案件名称", "優先度", "詳細化の概要", "ご確認いただきたい事項（要点）"],
              OVERVIEW_ROWS, size=9, name="一覧表", center_cols=(0, 3))
    note = add_text(s, 12, 470, 756, 30, [
        ("※優先度は10/5開発課題検討会資料の「8. 令和9年度案件について 一覧」による。011は一覧では「低」、同日の口頭説明では「中」（要確認）", False),
        ("※詳細化案の文言はいずれも（案）。案件台帳への反映は、ご確認の結果を踏まえて行う", False)], 9, bullet=False, name="注記")

    # 2枚目以降：案件ごと
    for idx, c in enumerate(CASES, start=2):
        title = f"{TITLE_BASE}（{idx}/{TOTAL}）"
        s = new_slide(prs, layout, title)
        # 案件名とリード
        add_text(s, LEFT, 57, WIDTH, 20, [(f"案件ID{c['id']}「{c['name']}」", False)], 12, bullet=False, name="案件名")
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
        lead_items = [(c["lead"], False)]
        size = 12
        # 高さを見積もって、入らなければ11ptに下げる（パターン集：本文12pt、詰まるページは11pt）
        def total_h(sz):
            h = 76 + section_height(lead_items, sz) + 4
            for label, items in c["sections"]:
                h += 20 + section_height(items, sz) + 6
                if c.get("table") and c["table"]["after"] == label:
                    t = c["table"]
                    h += sum((max(est_lines(str(v), 9, t["widths"][i]) for i, v in enumerate(r)) * 9 * 1.35 + 6)
                             for r in [t["cols"]] + t["rows"]) + 6
            return h
        notes = c.get("notes", [])
        if total_h(size) > BOTTOM - 12 * len(notes):
            size = 11
        y = 76
        lh = section_height(lead_items, size)
        add_text(s, LEFT, y, WIDTH, lh, lead_items, size, name="リード")
        y += lh + 4
        for label, items in c["sections"]:
            add_label(s, LEFT, y, label, w=175 if len(label) > 6 else 130)
            y += 20
            h = section_height(items, size)
            add_text(s, LEFT, y, WIDTH, h, items, size, name="本文：" + label)
            y += h + 6
            t = c.get("table")
            if t and t["after"] == label:
                _, th = add_table(s, LEFT, y, t["widths"], t["cols"], t["rows"], size=9, name="表：" + label)
                y += th + 6
        has_red = any(red for _, items in c["sections"] for _, red in items)
        if has_red:
            add_text(s, LEFT, 512, WIDTH, 14, [(c.get("note_red", NOTE_RED), False)], 9, bullet=False, name="注記：赤字")
        if notes:
            ny = (512 if has_red else 524) - 12 * len(notes)
            add_text(s, LEFT, ny, WIDTH, 12 * len(notes), [(n, False) for n in notes], 9, bullet=False, name="注記：用語")
        c["_end_y"] = y
        c["_notes_y"] = (512 if has_red else 524) - 12 * len(notes)
        c["_size"] = size

    remove_unused_masters(prs)
    prs.save(str(OUT))
    for c in CASES:
        print(c["id"], "本文", c["_size"], "pt", "下端", round(c["_end_y"], 1), "注記の上端", c["_notes_y"])
    print("masters", len(prs.slide_masters), "layouts", sum(len(m.slide_layouts) for m in prs.slide_masters))
    print("saved", OUT)


if __name__ == "__main__":
    build()
