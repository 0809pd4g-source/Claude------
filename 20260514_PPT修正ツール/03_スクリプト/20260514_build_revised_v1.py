# -*- coding: utf-8 -*-
import os, shutil
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

WORKDIR    = r"C:\Users\masato.nakazawa\Downloads\Claudeプロジェクト"
INPUT_NAME = "20260430_RESTfulヒアリング結果共有_FINDEX.pptx"
OUTDIR_NAME= "20260430_RESTfulヒアリング結果共有_FINDEX"

INPUT    = os.path.join(WORKDIR, "input", INPUT_NAME)
OUTDIR   = os.path.join(WORKDIR, OUTDIR_NAME)
ORIGINAL = os.path.join(OUTDIR, INPUT_NAME)
MODIFIED = os.path.join(OUTDIR, INPUT_NAME.replace(".pptx", "_修正後.pptx"))

os.makedirs(OUTDIR, exist_ok=True)
if os.path.exists(INPUT):
    shutil.move(INPUT, ORIGINAL)

PAREN_TRANS = str.maketrans({"（": "(", "）": ")"})
# Shape 0 (ph_idx=0) がスライドタイトル — 修正対象外
TITLE_TEXT = "1．RestfulAPI"

REPLACEMENTS = {
    0: [
        # 承認不要: 誤字 (変換ミス)
        ("自系列で", "時系列で"),
        # 承認不要 + 承認必須#1: 脱字補完 + 語彙修正
        ("ピンポイント検索によ通信不可の削減", "ピンポイント検索による通信回数の削減"),
        # 承認不要: 文法誤り
        ("芳しくなくと聞いている", "芳しくなかったと聞いている"),
        # 承認不要: 表記ゆれ (タイトル Shape は TITLE_TEXT 除外で守られる)
        ("RestfulAPI", "RESTful API"),
        # 承認必須 #2: 冗長性排除 (「ことで」重複)
        ("ことで、データをピンポイントに取得することで、",
         "ことで、データをピンポイントに取得し、"),
        # 承認必須 #3: きれいな日本語化
        ("臨床情報としてカバーできるとなお良い。",
         "臨床情報としてカバーされることが望ましい。"),
        # 承認必須 #4: 冗長性排除
        ("非常に大変なため、", "非常に困難なため、"),
        # 承認必須 #5: きれいな日本語化
        ("OSSを利用するのが良いのではないか。", "OSSを利用することが望ましい。"),
        # 承認必須 #6: きれいな日本語化
        ("公開するのが良い。", "公開することが望ましい。"),
    ],
}

# 重複除去
for k in REPLACEMENTS:
    seen, uniq = set(), []
    for o, n in REPLACEMENTS[k]:
        if (o, n) not in seen:
            seen.add((o, n)); uniq.append((o, n))
    REPLACEMENTS[k] = uniq


def _is_title_para(p):
    return TITLE_TEXT in "".join(r.text for r in p.runs)


def convert_parens_in_slide(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            if _is_title_para(p):
                continue
            for run in p.runs:
                if "（" in run.text or "）" in run.text:
                    run.text = run.text.translate(PAREN_TRANS)


def replace_in_paragraph(p, old, new):
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return 0
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return 1
    runs = list(p.runs)
    runs[0].text = full.replace(old, new, 1)
    for r in runs[1:]:
        r.text = ""
    return 1


def replace_in_slide(slide, replacements):
    misses = []
    for old, new in replacements:
        n = 0
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                if _is_title_para(p):
                    continue
                n += replace_in_paragraph(p, old, new)
        if n == 0:
            misses.append(old[:60])
    return misses


def duplicate_slide(prs, src_idx):
    src = prs.slides[src_idx]
    new_slide = prs.slides.add_slide(src.slide_layout)
    spTree = new_slide.shapes._spTree
    for child in list(spTree):
        tag = child.tag.split("}")[-1]
        if tag in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            spTree.remove(child)
    for child in src.shapes._spTree:
        tag = child.tag.split("}")[-1]
        if tag in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            spTree.append(deepcopy(child))
    return new_slide


def move_slide(prs, from_idx, to_idx):
    sldIdLst = prs.slides._sldIdLst
    target = list(sldIdLst)[from_idx]
    sldIdLst.remove(target)
    sldIdLst.insert(to_idx, target)


def add_revised_badge(prs, slide):
    w, h = Inches(1.2), Inches(0.45)
    x = prs.slide_width - w - Inches(0.25)
    y = Inches(0.10)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "修正後"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = "Meiryo UI"


# ===== Build =====
prs = Presentation(ORIGINAL)
n = len(prs.slides)
dups = [(i, duplicate_slide(prs, i)) for i in range(n)]
for orig_idx, dup in dups:
    convert_parens_in_slide(dup)
    misses = replace_in_slide(dup, REPLACEMENTS.get(orig_idx, []))
    add_revised_badge(prs, dup)
    for m in misses:
        print(f"[MISS] S{orig_idx+1}: {m}...")

# 交互配置: [S1..Sn, S1'..Sn'] -> [S1, S1', S2, S2', ...]
for k in range(n):
    move_slide(prs, n + k, 2 * k + 1)

prs.save(MODIFIED)
print(f"saved: {MODIFIED}")
