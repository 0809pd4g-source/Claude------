# -*- coding: utf-8 -*-
import os, shutil
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

WORKDIR     = r"C:\Users\masato.nakazawa\Downloads\Claudeプロジェクト"
INPUT_NAME  = "20260518_1330_開発進捗連絡会議_議事要旨(案).pptx"
OUTDIR_NAME = "20260518_1330_開発進捗連絡会議_議事要旨(案)"

INPUT    = os.path.join(WORKDIR, "input", INPUT_NAME)
OUTDIR   = os.path.join(WORKDIR, OUTDIR_NAME)
ORIGINAL = os.path.join(OUTDIR, INPUT_NAME)
MODIFIED = os.path.join(OUTDIR, INPUT_NAME.replace(".pptx", "_修正後.pptx"))

os.makedirs(OUTDIR, exist_ok=True)
if os.path.exists(INPUT):
    shutil.move(INPUT, ORIGINAL)

PAREN_TRANS = str.maketrans({"（": "(", "）": ")"})

REPLACEMENTS = {
    2: [  # Slide 3 (index 2)
        # 承認必須 #1: 冗長性排除 — #995 ToDo 内「コストについて」重複
        ("富士通側の運用保守コストについて厚労省へ連携する。",
         "富士通側の運用保守コストを厚労省へ連携する。"),
    ],
}

# 重複除去
for k in REPLACEMENTS:
    seen, uniq = set(), []
    for o, n in REPLACEMENTS[k]:
        if (o, n) not in seen:
            seen.add((o, n)); uniq.append((o, n))
    REPLACEMENTS[k] = uniq


def _is_title_shape(shape):
    """ph_idx=0 のタイトルプレースホルダーを除外"""
    return shape.is_placeholder and shape.placeholder_format.idx == 0


def _paragraphs_of_slide(slide):
    """スライド内の全段落をイテレート (テキストボックス + テーブルセル)"""
    for shape in slide.shapes:
        if _is_title_shape(shape):
            continue
        if shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        elif shape.shape_type == 19:  # TABLE
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def convert_parens_in_slide(slide):
    for p in _paragraphs_of_slide(slide):
        for run in p.runs:
            if "（" in run.text or "）" in run.text:
                run.text = run.text.translate(PAREN_TRANS)


def replace_in_paragraph(p, old, new):
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return 0
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return 1
    # 複数 run にまたがる場合: 先頭 run にまとめる
    runs = list(p.runs)
    runs[0].text = full.replace(old, new, 1)
    for r in runs[1:]:
        r.text = ""
    return 1


def replace_in_slide(slide, replacements):
    misses = []
    for old, new in replacements:
        n = sum(replace_in_paragraph(p, old, new)
                for p in _paragraphs_of_slide(slide))
        if n == 0:
            misses.append(old[:60])
    return misses


def fix_case_connector(slide):
    """承認不要 (助詞抜け): #1023 段落の「に変更する場合」+「5/18」の間に「と」を挿入。
    run 単位で外科的に修正し、前後 run のフォント書式を保持する。"""
    for shape in slide.shapes:
        if shape.shape_type != 19:
            continue
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    runs = list(p.runs)
                    for i in range(len(runs) - 1):
                        if (runs[i].text == "に変更する場合"
                                and runs[i + 1].text == "5/18"):
                            runs[i].text = "に変更する場合と"
                            return  # 1 箇所のみ


def duplicate_slide(prs, src_idx):
    src = prs.slides[src_idx]
    new_slide = prs.slides.add_slide(src.slide_layout)
    spTree = new_slide.shapes._spTree
    for child in list(spTree):
        tag = child.tag.split("}")[-1]
        if tag in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            spTree.remove(child)
    for child in src.shapes._spTree:
        tag = child.tag.split("}")[-1]
        if tag in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            spTree.append(deepcopy(child))
    return new_slide


def move_slide(prs, from_idx, to_idx):
    sldIdLst = prs.slides._sldIdLst
    target = list(sldIdLst)[from_idx]
    sldIdLst.remove(target)
    sldIdLst.insert(to_idx, target)


def add_revised_badge(prs, slide):
    w, h = Inches(1.2), Inches(0.45)
    x = prs.slide_width - w - Inches(0.25)
    y = Inches(0.10)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "修正後"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = "Meiryo UI"


# ===== Build =====
prs = Presentation(ORIGINAL)
n = len(prs.slides)
dups = [(i, duplicate_slide(prs, i)) for i in range(n)]

for orig_idx, dup in dups:
    convert_parens_in_slide(dup)
    if orig_idx == 2:  # Slide 3 のみ特殊処理
        fix_case_connector(dup)
    misses = replace_in_slide(dup, REPLACEMENTS.get(orig_idx, []))
    add_revised_badge(prs, dup)
    for m in misses:
        print(f"[MISS] S{orig_idx + 1}: {m}...")

# 交互配置: [S1..Sn, S1'..Sn'] -> [S1,S1', S2,S2', S3,S3']
for k in range(n):
    move_slide(prs, n + k, 2 * k + 1)

prs.save(MODIFIED)
print(f"saved: {MODIFIED}")
