#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
旅行しおり 冊子ジェネレーター（汎用ビルダー）
====================================================
旅行データを分離したデータファイル(trip_*.py)と組み合わせて使う。

Usage:
    python trip_egypt_greece.py          # Egypt/Greece版を生成
    # または
    from build_booklet import build
    from trip_egypt_greece import TRIP, OUTPUT_PATH
    build(TRIP, OUTPUT_PATH)

Trip config の構造 → trip_egypt_greece.py を参照。
新しい旅行は trip_egypt_greece.py をコピーして内容を書き換えるだけ。
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ═══════════════════════════════════════════════
# カラーパレット（固定）
# ═══════════════════════════════════════════════
_AMBER700 = RGBColor(0xB4, 0x53, 0x09)
_AMBER600 = RGBColor(0xD9, 0x77, 0x06)
_AMBER_BG  = "FEF3C7"
_AMBER_BG2 = "FFFBEB"
_SKY700   = RGBColor(0x03, 0x69, 0xA1)
_SKY600   = RGBColor(0x02, 0x84, 0xC7)
_SKY_BG   = "E0F2FE"
_SKY_BG2  = "F0F9FF"
_EMRLD700 = RGBColor(0x04, 0x78, 0x57)
_EMRLD_BG = "D1FAE5"
_ROSE600  = RGBColor(0xE1, 0x1D, 0x48)
_ROSE_BG  = "FFE4E6"
_SL800    = RGBColor(0x1E, 0x29, 0x3B)
_SL700    = RGBColor(0x33, 0x41, 0x55)
_SL600    = RGBColor(0x47, 0x56, 0x69)
_SL400    = RGBColor(0x94, 0xA3, 0xB8)
_SL100    = "F1F5F9"
_WHITE    = "FFFFFF"

# A5 usable width: 148 - 13*2 = 122mm
_PW   = 122   # page usable width (mm)
_TW   = 22    # schedule time col (mm)
_AW   = _PW - _TW  # schedule activity col

def _twip(mm):
    return int(mm * 56.693)

# ═══════════════════════════════════════════════
# テーマ解決（"a" = アンバー系, "b" = スカイ系）
# ═══════════════════════════════════════════════
def _th(theme):
    """theme -> (header_color, bg_light, bg_dark, bar_hex)"""
    if theme == "b":
        return _SKY700, _SKY_BG2, _SKY_BG, "0369A1"
    return _AMBER700, _AMBER_BG2, _AMBER_BG, "B45309"

def _box_colors(box_type, theme="a"):
    """box type -> (bg_hex, title_color)"""
    t = box_type or "tip"
    if t == "warning":
        return _ROSE_BG, _ROSE600
    if t == "hotel":
        return _SKY_BG2, _SKY700
    if t == "green":
        return _EMRLD_BG, _EMRLD700
    if t == "note":
        return _SL100, _SL700
    # "tip" / "info" / default → theme color
    _, bg, _, _ = _th(theme)
    hdr = _SKY700 if theme == "b" else _AMBER700
    return bg, hdr

# ═══════════════════════════════════════════════
# XML ヘルパー
# ═══════════════════════════════════════════════
def _shade(cell, fill_hex):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill_hex)
    tcPr.append(s)

def _cw(cell, mm):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcW')): tcPr.remove(old)
    w = OxmlElement('w:tcW')
    w.set(qn('w:w'), str(_twip(mm))); w.set(qn('w:type'), 'dxa')
    tcPr.append(w)

def _no_border(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
    b = OxmlElement('w:tcBorders')
    for s in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'none'); b.append(e)
    tcPr.append(b)

def _cm(cell, t=1.5, bo=1.5, l=2.5, r=2.5):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcMar')): tcPr.remove(old)
    m = OxmlElement('w:tcMar')
    for side, mm in [('top',t),('bottom',bo),('left',l),('right',r)]:
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(_twip(mm))); e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)

def _sp(para, bef=0, aft=0):
    pPr = para._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:spacing')): pPr.remove(old)
    s = OxmlElement('w:spacing')
    s.set(qn('w:before'), str(int(bef*20))); s.set(qn('w:after'), str(int(aft*20)))
    pPr.append(s)

def _indent(para, left_mm):
    pPr = para._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:ind')): pPr.remove(old)
    i = OxmlElement('w:ind'); i.set(qn('w:left'), str(_twip(left_mm)))
    pPr.append(i)

def _run(para, text, bold=False, italic=False, pt=10, color=None):
    r = para.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(pt); r.font.name = 'Yu Gothic'
    if color: r.font.color.rgb = color
    return r

def _para(doc, text='', bold=False, italic=False, pt=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, bef=0, aft=2):
    para = doc.add_paragraph(); para.alignment = align; _sp(para, bef, aft)
    if text: _run(para, text, bold=bold, italic=italic, pt=pt, color=color)
    return para

def _pb(doc):
    """改ページ段落"""
    para = doc.add_paragraph(); _sp(para, 0, 0)
    para.add_run().add_break(WD_BREAK.PAGE)
    return para

def _para_border(para, color_hex="D97706", size=6):
    pPr = para._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:pBdr')): pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(size)); b.set(qn('w:color'), color_hex)
    pBdr.append(b); pPr.append(pBdr)

# ═══════════════════════════════════════════════
# 共通UIパーツ
# ═══════════════════════════════════════════════

def _sp_para(doc, aft=4):
    """スペーサー段落"""
    p = doc.add_paragraph(); _sp(p, 0, aft); return p

def _chapter_header(doc, num, title, theme="a"):
    hdr_rgb, _, bg, bar = _th(theme)
    tbl = doc.add_table(rows=1, cols=2); tbl.style = 'Table Grid'
    c0 = tbl.cell(0,0); c1 = tbl.cell(0,1)
    # 番号セル
    _cw(c0, 18); _shade(c0, bar); _no_border(c0); _cm(c0, 3, 3, 4, 4)
    p0 = c0.paragraphs[0]; _sp(p0, 0, 0); p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p0, str(num), bold=True, pt=18, color=RGBColor(0xFF, 0xFF, 0xFF))
    # タイトルセル
    _cw(c1, _PW-18); _shade(c1, bg); _no_border(c1); _cm(c1, 3, 3, 5, 4)
    p1 = c1.paragraphs[0]; _sp(p1, 0, 0)
    _run(p1, title, bold=True, pt=14, color=hdr_rgb)
    _sp_para(doc, 6)

def _section_bar(doc, text, theme="a", bef=8):
    hdr, _, bg, _ = _th(theme)
    _sp_para(doc, bef)
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    c = tbl.cell(0,0); _cw(c, _PW); _shade(c, bg); _no_border(c); _cm(c, 3, 3, 4, 4)
    p = c.paragraphs[0]; _sp(p, 0, 0)
    _run(p, text, bold=True, pt=11, color=hdr)
    _sp_para(doc, 3)

def _info_box(doc, title, lines, bg, title_color, aft=6):
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    c = tbl.cell(0,0); _cw(c, _PW); _shade(c, bg); _no_border(c); _cm(c, 2.5, 2.5, 4, 4)
    if title:
        tp = c.paragraphs[0]; _sp(tp, 0, 2)
        _run(tp, title, bold=True, pt=9.5, color=title_color)
        for line in lines:
            lp = c.add_paragraph(); _sp(lp, 0, 1)
            _run(lp, line, pt=9, color=_SL700)
    else:
        first = True
        for line in lines:
            lp = c.paragraphs[0] if first else c.add_paragraph(); first = False
            _sp(lp, 0, 1); _run(lp, line, pt=9, color=_SL700)
    _sp_para(doc, aft)

def _render_box(doc, box, theme="a"):
    """box dict -> render as info_box or hotel_box"""
    btype = box.get("type", "tip")
    if btype == "hotel":
        _hotel_box(doc, box["hotel"], box.get("note", ""))
        return
    bg, clr = _box_colors(btype, theme)
    _info_box(doc, box.get("title",""), box.get("lines",[]), bg, clr, aft=box.get("aft",6))

def _hotel_box(doc, hotel, nights=""):
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    c = tbl.cell(0,0); _cw(c, _PW); _shade(c, _SKY_BG2); _no_border(c); _cm(c, 2, 2, 4, 4)
    p = c.paragraphs[0]; _sp(p, 0, 0)
    _run(p, "🏨 宿泊: ", bold=True, pt=9.5, color=_SKY700)
    _run(p, hotel, pt=9.5, color=_SL800)
    if nights: _run(p, f"  ({nights})", pt=8.5, color=_SL400)
    _sp_para(doc, 8)

def _add_image(doc, path, caption="", width_mm=118):
    """画像を挿入（中央揃え）。ファイルが見つからない場合はスキップ。
    path: プロジェクトルートからの相対パスまたは絶対パス
    """
    from pathlib import Path as _Path
    p = _Path(path)
    if not p.exists():
        print(f"  ⚠️  画像が見つかりません（スキップ）: {path}")
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _sp(para, 0, 2)
    para.add_run().add_picture(str(p), width=Mm(width_mm))
    if caption:
        cp = doc.add_paragraph(); _sp(cp, 0, 4)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cp, caption, italic=True, pt=8, color=_SL400)
    else:
        _sp_para(doc, 4)

def _day_header(doc, n, date, location, theme="a", badges=""):
    hdr, bg, _, _ = _th(theme)
    tbl = doc.add_table(rows=1, cols=1); tbl.style = 'Table Grid'
    c = tbl.cell(0,0); _cw(c, _PW); _shade(c, bg); _no_border(c); _cm(c, 3, 3, 5, 5)
    p1 = c.paragraphs[0]; _sp(p1, 0, 2)
    _run(p1, f"DAY {n}  ", bold=True, pt=16, color=hdr)
    _run(p1, date, pt=9.5, color=_SL600)
    if badges: _run(p1, f"  {badges}", pt=8.5, color=_SL400)
    p2 = c.add_paragraph(); _sp(p2, 2, 0)
    _run(p2, location, bold=True, pt=12, color=_SL800)
    _sp_para(doc, 3)

def _sched_table(doc, rows):
    tbl = doc.add_table(rows=0, cols=2); tbl.style = 'Table Grid'
    for item in rows:
        time = item[0]; act = item[1]; note = item[2] if len(item)>2 else ''
        row = tbl.add_row()
        tc = row.cells[0]; ac = row.cells[1]
        _cw(tc, _TW); _shade(tc, _SL100); _cm(tc, 1.5, 1.5, 2, 2)
        pt = tc.paragraphs[0]; pt.alignment = WD_ALIGN_PARAGRAPH.CENTER; _sp(pt, 0, 0)
        _run(pt, time, bold=True, pt=8.5, color=_SL600)
        _cw(ac, _AW); _cm(ac, 1.5, 1.5, 3, 2)
        ap = ac.paragraphs[0]; _sp(ap, 0, 1)
        _run(ap, act, pt=9.5, color=_SL800)
        if note:
            np = ac.add_paragraph(); _sp(np, 1, 0)
            _run(np, note, italic=True, pt=8, color=_SL400)
    _sp_para(doc, 4)

def _phrase_table(doc, rows):
    tbl = doc.add_table(rows=0, cols=3); tbl.style = 'Table Grid'
    for item in rows:
        jp = item[0]; lc = item[1]; rd = item[2] if len(item)>2 else ''
        row = tbl.add_row()
        for cell, val, w, clr in zip(row.cells, [jp,lc,rd], [35,47,40], [_SL700,_AMBER700,_SKY700]):
            _cw(cell, w); _cm(cell, 1.5, 1.5, 2, 2)
            p = cell.paragraphs[0]; _sp(p, 0, 0); _run(p, val, pt=8.5, color=clr)
    _sp_para(doc, 6)

def _check_table(doc, items, cols=3):
    iw = _PW // cols
    for i in range(0, len(items), cols):
        chunk = items[i:i+cols]
        tbl = doc.add_table(rows=1, cols=cols); tbl.style = 'Table Grid'
        for j, cell in enumerate(tbl.row_cells(0)):
            _cw(cell, iw); _no_border(cell); _cm(cell, 1, 1, 2, 2)
            p = cell.paragraphs[0]; _sp(p, 0, 0)
            if j < len(chunk): _run(p, f"□ {chunk[j]}", pt=9, color=_SL700)
        _sp_para(doc, 1)

def _flight_table(doc, flights):
    tbl = doc.add_table(rows=0, cols=4); tbl.style = 'Table Grid'
    # ヘッダー行
    hrow = tbl.add_row()
    for cell, h, w in zip(hrow.cells, ['便名','出発','到着','区間'], [28,23,23,48]):
        _cw(cell, w); _shade(cell, "334155"); _no_border(cell); _cm(cell, 1.5, 1.5, 2, 2)
        p = cell.paragraphs[0]; _sp(p, 0, 0)
        _run(p, h, bold=True, pt=8.5, color=RGBColor(0xFF,0xFF,0xFF))
    for flt in flights:
        row = tbl.add_row()
        for cell, val, w in zip(row.cells, [flt[0],flt[1],flt[2],flt[3]], [28,23,23,48]):
            _cw(cell, w); _cm(cell, 1.5, 1.5, 2, 2)
            p = cell.paragraphs[0]; _sp(p, 0, 0); _run(p, val, pt=8.5, color=_SL800)
    _sp_para(doc, 6)

def _hotel_table(doc, hotels):
    tbl = doc.add_table(rows=0, cols=2); tbl.style = 'Table Grid'
    for day, hotel in hotels:
        row = tbl.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        _cw(c0, 22); _cw(c1, 100); _shade(c0, _SL100)
        _cm(c0, 1.5, 1.5, 2, 2); _cm(c1, 1.5, 1.5, 3, 2)
        p0 = c0.paragraphs[0]; _sp(p0, 0, 0)
        _run(p0, day, bold=True, pt=8.5, color=_SL600)
        p1 = c1.paragraphs[0]; _sp(p1, 0, 0)
        _run(p1, hotel, pt=9, color=_SL800)
    _sp_para(doc, 6)

# ═══════════════════════════════════════════════
# セクションビルダー（データ受取型）
# ═══════════════════════════════════════════════

def _build_cover(doc, meta):
    for _ in range(4): _sp_para(doc, 4)

    # 国旗
    fp = _para(doc, meta.get("flags",""), pt=20, align=WD_ALIGN_PARAGRAPH.CENTER, bef=0, aft=8)

    # タイトル
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER; _sp(tp, 0, 4)
    for part in meta["title_parts"]:  # [{"text":"...", "color":"a"/"b"}]
        color = _AMBER700 if part["color"]=="a" else _SKY700
        _run(tp, part["text"], bold=True, pt=22, color=color)

    # 罫線
    hr = _para(doc, "", bef=0, aft=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para_border(hr, color_hex="D97706", size=6)

    # 日付・サブタイトル
    _para(doc, meta["dates"], pt=11, color=_SL600, align=WD_ALIGN_PARAGRAPH.CENTER, bef=0, aft=3)
    _para(doc, meta.get("subtitle",""), pt=10, color=_SL600, align=WD_ALIGN_PARAGRAPH.CENTER, bef=0, aft=12)

    # カバー写真（任意）
    if meta.get("image"):
        _add_image(doc, meta["image"], caption=meta.get("image_caption",""), width_mm=110)

    # ルート表
    tbl = doc.add_table(rows=0, cols=1); tbl.style = 'Table Grid'
    for icon, text, theme in meta["route"]:
        _, bg, bg2, _ = _th(theme)
        row = tbl.add_row()
        c = row.cells[0]; _cw(c, _PW); _shade(c, bg2); _no_border(c); _cm(c, 2.5, 2.5, 5, 4)
        p = c.paragraphs[0]; _sp(p, 0, 0)
        _run(p, f"{icon}  ", pt=10); _run(p, text, pt=10, color=_SL700)
    _sp_para(doc, 8)

    _para(doc, meta.get("note",""), italic=True, pt=9, color=_SL400,
          align=WD_ALIGN_PARAGRAPH.CENTER, bef=0, aft=0)
    _pb(doc)


def _build_toc(doc, toc_entries, meta):
    """toc_entries: [{"label": "...", "indent": 0/4}]"""
    _para(doc, "目次", bold=True, pt=16, color=_AMBER700, bef=4, aft=6)
    for e in toc_entries:
        para = doc.add_paragraph(); _sp(para, 0, 2)
        if e.get("indent"): _indent(para, e["indent"])
        _run(para, e["label"], bold=(e.get("indent",0)==0),
             pt=9 if e.get("indent") else 10,
             color=_SL600 if e.get("indent") else _SL800)
    _pb(doc)


def _build_overview(doc, overview):
    _chapter_header(doc, "◎", overview.get("chapter_title","旅行概要"), theme="a")
    _info_box(doc, overview.get("summary_title","旅行概要"),
              overview.get("summary",[]), _AMBER_BG2, _AMBER700)
    _section_bar(doc, "✈ フライト一覧", theme="a")
    _flight_table(doc, overview["flights"])
    _section_bar(doc, "🏨 宿泊施設一覧", theme="a")
    _hotel_table(doc, overview["hotels"])
    _pb(doc)


def _build_itinerary(doc, days, part_num=1, part_title="日程詳細"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for day in days:
        _pb(doc)  # 各DAYを必ず新ページから
        theme = day.get("theme", "a")
        _day_header(doc, day["n"], day["date"], day["location"],
                    theme=theme, badges=day.get("badges",""))
        if day.get("image"):
            _add_image(doc, day["image"], caption=day.get("image_caption",""))
        _sched_table(doc, day["schedule"])
        for box in day.get("boxes", []):
            _render_box(doc, box, theme=theme)
        if day.get("hotel"):
            _hotel_box(doc, day["hotel"], day.get("hotel_note",""))
    _pb(doc)


def _build_area_guide(doc, areas, part_num=2, part_title="エリアガイド"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for area in areas:
        theme = area.get("theme", "a")
        _section_bar(doc, f"{area.get('flag','')} {area['name']}", theme=theme)
        if area.get("overview"):
            hdr_clr, _, _, _ = _th(theme)
            _info_box(doc, area.get("overview_title","エリア概要"),
                      area["overview"], _AMBER_BG2 if theme=="a" else _SKY_BG2,
                      hdr_clr)
        for spot in area.get("spots",[]):
            _, bg, _, _ = _th(theme)
            hdr_clr, _, _, _ = _th(theme)
            _info_box(doc, f"📍 {spot['name']}", spot["details"], bg, hdr_clr, aft=3)
        if area.get("practical"):
            _para(doc, area.get("practical_title","実用情報"),
                  bold=True, pt=10, color=_AMBER700 if theme=="a" else _SKY700, bef=4, aft=2)
            for line in area["practical"]:
                _para(doc, f"• {line}", pt=9, color=_SL700, bef=0, aft=2)
        _sp_para(doc, 4)
    _pb(doc)


def _build_history(doc, history, part_num=3, part_title="歴史・文化入門"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for chapter in history:
        theme = chapter.get("theme","a")
        _section_bar(doc, f"{chapter.get('flag','')} {chapter['title']}", theme=theme)
        if chapter.get("overview"):
            _, bg, _, _ = _th(theme)
            hdr, _, _, _ = _th(theme)
            _info_box(doc, chapter.get("overview_title","概要"), chapter["overview"], bg, hdr)
        for sec in chapter.get("sections",[]):
            _, bg, _, _ = _th(theme)
            hdr, _, _, _ = _th(theme)
            _info_box(doc, sec["title"], sec["content"], bg, hdr, aft=4)
    _pb(doc)


def _build_food(doc, food, part_num=4, part_title="食事・お土産ガイド"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for section in food.get("sections",[]):
        theme = section.get("theme","a")
        _section_bar(doc, f"{section.get('flag','')} {section['title']}", theme=theme)
        for item in section.get("items",[]):
            _, bg, _, _ = _th(theme)
            hdr, _, _, _ = _th(theme)
            _info_box(doc, item["name"], item["details"], bg, hdr, aft=3)
    for warn in food.get("warnings",[]):
        _info_box(doc, warn["title"], warn["lines"], _ROSE_BG, _ROSE600)
    for extra in food.get("extra_sections",[]):
        theme = extra.get("theme","a")
        _section_bar(doc, f"{extra.get('flag','')} {extra['title']}", theme=theme)
        for box in extra.get("boxes",[]):
            _render_box(doc, box, theme=theme)
    _pb(doc)


def _build_phrases(doc, phrases, part_num=5, part_title="現地語フレーズ集"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for lang in phrases:
        theme = lang.get("theme","a")
        _section_bar(doc, f"{lang.get('flag','')} {lang['language']}", theme=theme)
        if lang.get("intro"):
            _, bg, _, _ = _th(theme); hdr, _, _, _ = _th(theme)
            _info_box(doc, lang.get("intro_title","発音のコツ"), lang["intro"], bg, hdr)
        for sec in lang.get("sections",[]):
            clr = _AMBER700 if theme=="a" else _SKY700
            _para(doc, sec["title"], bold=True, pt=10, color=clr, bef=4, aft=2)
            _phrase_table(doc, sec["rows"])
    _pb(doc)


def _build_practical(doc, practical, part_num=6, part_title="実用情報"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for sec in practical:
        theme = sec.get("theme","a")
        _section_bar(doc, f"{sec.get('icon','')} {sec['title']}", theme=theme)
        for box in sec.get("boxes",[]):
            _render_box(doc, box, theme=theme)
        if sec.get("table"):
            # Simple 2-col table
            tbl = doc.add_table(rows=0, cols=2); tbl.style = 'Table Grid'
            for name, val in sec["table"]:
                row = tbl.add_row()
                c0, c1 = row.cells[0], row.cells[1]
                _cw(c0, 55); _cw(c1, 67); _shade(c0, _SL100)
                _cm(c0, 1.5, 1.5, 2, 2); _cm(c1, 1.5, 1.5, 2, 2)
                p0 = c0.paragraphs[0]; _sp(p0, 0, 0)
                _run(p0, name, bold=True, pt=8.5, color=_SL700)
                p1 = c1.paragraphs[0]; _sp(p1, 0, 0)
                _run(p1, val, pt=8.5, color=_SL800)
            _sp_para(doc, 6)
    _pb(doc)


def _build_packing(doc, packing, part_num=7, part_title="持ち物チェックリスト"):
    _chapter_header(doc, part_num, part_title, theme="a")
    for cat_name, items in packing.get("categories",{}).items():
        _section_bar(doc, cat_name, theme="a")
        _check_table(doc, items, cols=3)
        _sp_para(doc, 4)
    if packing.get("print_checklist"):
        _section_bar(doc, "🖨 出発前・印刷物チェック", theme="a")
        _check_table(doc, packing["print_checklist"], cols=2)
    _sp_para(doc, 12)
    ep = _para(doc, packing.get("ending","Have a wonderful trip!"),
               bold=True, pt=14, color=_AMBER600, align=WD_ALIGN_PARAGRAPH.CENTER, bef=0, aft=0)


# ═══════════════════════════════════════════════
# メインエントリーポイント
# ═══════════════════════════════════════════════

def build(config, output_path):
    """
    config: TRIP dict from trip_*.py
    output_path: str or Path
    """
    print("📘 冊子ビルド開始...")
    doc = Document()

    sec = doc.sections[0]
    sec.page_width  = Mm(148); sec.page_height = Mm(210)
    sec.left_margin = Mm(13);  sec.right_margin  = Mm(13)
    sec.top_margin  = Mm(15);  sec.bottom_margin = Mm(15)

    doc.styles['Normal'].font.name = 'Yu Gothic'
    doc.styles['Normal'].font.size = Pt(9.5)

    meta     = config["meta"]
    overview = config["overview"]
    days     = config["days"]
    areas    = config.get("areas", [])
    history  = config.get("history", [])
    food     = config.get("food", {})
    phrases  = config.get("phrases", [])
    practical= config.get("practical", [])
    packing  = config.get("packing", {})

    print("  表紙..."); _build_cover(doc, meta)
    print("  目次...");  _build_toc(doc, config.get("toc",[]), meta)
    print("  旅行概要..."); _build_overview(doc, overview)
    print("  旅程 (DAY 1-15)..."); _build_itinerary(doc, days, 1, "日程詳細")
    print("  エリアガイド..."); _build_area_guide(doc, areas, 2, "エリアガイド")
    print("  歴史..."); _build_history(doc, history, 3, "歴史・文化入門")
    print("  食事・お土産..."); _build_food(doc, food, 4, "食事・お土産ガイド")
    print("  フレーズ集..."); _build_phrases(doc, phrases, 5, "現地語フレーズ集")
    print("  実用情報..."); _build_practical(doc, practical, 6, "実用情報")
    print("  持ち物リスト..."); _build_packing(doc, packing, 7, "持ち物チェックリスト")

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    size = out.stat().st_size // 1024
    print(f"\n✅ 完成: {out}")
    print(f"   ファイルサイズ: {size} KB")
    return out
